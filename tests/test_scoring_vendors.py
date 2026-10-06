from __future__ import annotations

import json
import sqlite3
from copy import deepcopy
from io import BytesIO

import pandas as pd
from pptx import Presentation

from src.modules.scoring import calculate_scoring
from src.modules import scoring_jobs
from src.modules.scoring_vendors import normalize_scoring_vendor_result, scoring_vendor_name
from src.modules.scoring_views import build_scoring_views, scoring_coverage_notes
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_api import scoring_api


def _presentation_text(presentation: Presentation) -> str:
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.has_chart:
                chart = shape.chart
                parts.extend(str(series.name) for series in chart.series if series.name)
                for plot in chart.plots:
                    parts.extend(
                        str(category.label)
                        for category in plot.categories
                        if getattr(category, 'label', None) is not None
                    )
    return '\n'.join(parts)


def test_scoring_vendor_name_removes_known_operator_prefixes_and_keeps_vendor_underscores():
    operators = ['3', 'EE', 'VF_UK']

    assert scoring_vendor_name('3_Ericsson', operators) == 'Ericsson'
    assert scoring_vendor_name('EE_Ericsson', operators) == 'Ericsson'
    assert scoring_vendor_name('VF_UK_Huawei', operators) == 'Huawei'
    assert scoring_vendor_name('Mixed_Vendor', operators) == 'Mixed_Vendor'
    assert scoring_vendor_name('VF_SA_Ericsson', ['Vodafone UK', 'VF_SA']) == 'Ericsson'


def test_scoring_engine_groups_prefixed_vendor_rows_and_gap_by_pure_vendor():
    configuration = scoring_configuration()
    configuration['metrics'] = [
        deepcopy(metric) for metric in configuration['metrics'] if metric['code'] == 'K1'
    ]
    source = pd.DataFrame([
        {'Campaign': '2026-Q2', 'Operator': operator, 'Vendor': vendor,
         'G_Level_1': 'Drive', 'G_Level_2': 'City', 'Session_Type': 'CALL',
         'Call_Status': status}
        for operator, vendor, status in (
            ('EE', 'EE_Ericsson', 'Completed'),
            ('3', '3_Ericsson', 'Failed'),
            ('VF_UK', 'VF_UK_Huawei', 'Failed'),
            ('EE', 'EE', 'Completed'),
            ('O2', 'O2', 'Failed'),
        )
    ])

    result = calculate_scoring(
        {'voice': source}, ['Operator', 'Vendor'], baseline_operator='EE',
        configuration=configuration,
    )

    assert {(row['operator'], row['vendor']) for row in result['scoring']} == {
        ('EE', 'Ericsson'), ('3', 'Ericsson'), ('VF_UK', 'Huawei'),
        ('EE', 'All'), ('O2', 'All'),
    }
    gap_contexts = {(row['operator'], row['vendor']) for row in result['gap']}
    assert gap_contexts == {('3', 'Ericsson'), ('VF_UK', 'Huawei'), ('O2', 'All')}


def _vendor_reference_result():
    configuration = scoring_configuration()
    configuration['metrics'] = [
        deepcopy(metric) for metric in configuration['metrics'] if metric['code'] == 'K1'
    ]
    observations = [
        # Exact baseline Ericsson should win over EE's All row in this context.
        ('EE', 'EE_Ericsson', '2026-Q2', 'North', 'City', 'Completed'),
        ('EE', 'EE', '2026-Q2', 'North', 'City', 'Failed'),
        ('O2', 'O2_Ericsson', '2026-Q2', 'North', 'City', 'Failed'),
        # Huawei uses the same context's All baseline.
        ('O2', 'O2_Huawei', '2026-Q2', 'North', 'City', 'Failed'),
        # Do not borrow an All baseline across a different campaign, region, or environment.
        ('O2', 'O2_Ericsson', '2026-Q3', 'North', 'City', 'Failed'),
        ('O2', 'O2_Nokia', '2026-Q2', 'South', 'City', 'Failed'),
        ('O2', 'O2_Ericsson', '2026-Q2', 'North', 'Connectionroad', 'Failed'),
    ]
    source = pd.DataFrame([
        {'Campaign': campaign, 'Operator': operator, 'Vendor': vendor, 'Region': region,
         'G_Level_1': 'Drive', 'G_Level_2': environment, 'Session_Type': 'CALL',
         'Call_Status': status}
        for operator, vendor, campaign, region, environment, status in observations
    ])
    levels = ['Operator', 'Vendor', 'Campaign', 'Region']
    result = calculate_scoring(
        {'voice': source}, levels, baseline_operator='EE', configuration=configuration,
    )
    job = {
        'levels': levels, 'aggregation_levels': levels, 'aggregation_contract_version': 2,
        'baseline_operator': 'EE', 'configuration': configuration,
    }
    return result, job


def test_all_vendor_reference_prefers_exact_and_stays_within_selected_contexts():
    result, job = _vendor_reference_result()
    gaps = [row for row in result['gap'] if row['operator'] == 'O2']

    assert {
        (row['vendor'], row['campaign'], row['region'], row['environment']) for row in gaps
    } == {
        ('Ericsson', '2026-Q2', 'North', 'DriveCity'),
        ('Huawei', '2026-Q2', 'North', 'DriveCity'),
    }
    exact_gap = next(row for row in gaps if row['vendor'] == 'Ericsson')
    exact_baseline = next(
        row for row in result['scoring']
        if row['operator'] == 'EE' and row['vendor'] == 'Ericsson'
        and row['campaign'] == '2026-Q2' and row['region'] == 'North'
        and row['environment'] == 'DriveCity'
    )
    assert exact_gap['baseline_points'] == exact_baseline['weighted_points']

    views = build_scoring_views(job, result)
    scalar = next(
        table for table in views['gap_summary_tables']
        if table['context'].get('vendor') == 'Huawei'
        and table['context'].get('campaign') == '2026-Q2'
        and table['context'].get('region') == 'North'
        and table['context']['environment'] == 'DriveCity'
    )
    scalar_row = next(row for row in scalar['rows'] if row.get('kpi_code') == 'K1')
    assert scalar_row['gaps']['O2'] is not None

    matrix = next(
        table for table in views['hierarchy_gap_tables']
        if table['context']['environment'] == 'DriveCity'
    )
    leaves = {}
    for column in matrix['hierarchy_columns']:
        path = {entry['level']: entry.get('value') for entry in column['path']}
        if path.get('Operator') == 'O2':
            leaves[(path.get('Vendor'), path.get('Campaign'), path.get('Region'))] = column['id']
    hierarchy_row = next(row for row in matrix['rows'] if row['kpi_code'] == 'K1')
    assert hierarchy_row['gaps'][leaves[('Ericsson', '2026-Q2', 'North')]] is not None
    assert hierarchy_row['gaps'][leaves[('Huawei', '2026-Q2', 'North')]] is not None
    assert hierarchy_row['gaps'][leaves[('Ericsson', '2026-Q3', 'North')]] is None
    assert hierarchy_row['gaps'][leaves[('Nokia', '2026-Q2', 'South')]] is None
    road = next(
        table for table in views['hierarchy_gap_tables']
        if table['context']['environment'] == 'DriveConnectionroad'
    )
    road_leaf = next(
        column['id'] for column in road['hierarchy_columns']
        if {entry['level']: entry.get('value') for entry in column['path']}.get('Operator') == 'O2'
    )
    road_row = next(row for row in road['rows'] if row['kpi_code'] == 'K1')
    assert road_row['gaps'][road_leaf] is None


def test_vendor_filter_selects_all_matching_operator_prefixes_and_outputs_legacy_labels(scoring_api):
    repository = scoring_api['repository']
    repository.replace_operator_mapping_groups([
        {'canonical': 'Three UK', 'aliases': ['3'], 'color': '#AABBCC'},
        {'canonical': 'Vodafone UK', 'aliases': ['VF_UK'], 'color': '#DDEEFF'},
    ])
    expected_operators = {'3', 'EE', 'VF_UK'}
    source_rows = [
        {'Operator': '3', 'Operator_Vendor': '3_Ericsson', 'Region': 'North', 'City': 'Leeds',
         'Campaign': '2026-Q2', 'score': 1},
        {'Operator': 'EE', 'Operator_Vendor': 'EE_Ericsson', 'Region': 'North', 'City': 'Leeds',
         'Campaign': '2026-Q2', 'score': 1},
        {'Operator': 'VF_UK', 'Operator_Vendor': 'VF_UK_Huawei', 'Region': 'North', 'City': 'Leeds',
         'Campaign': '2026-Q2', 'score': 1},
        {'Operator': 'VF_UK', 'Operator_Vendor': 'VF_UK', 'Region': 'North', 'City': 'Leeds',
         'Campaign': '2026-Q2', 'score': 1},
        {'Operator': 'O2', 'Operator_Vendor': 'O2', 'Region': 'North', 'City': 'Leeds',
         'Campaign': '2026-Q2', 'score': 1},
        {'Operator': 'Nokia Test', 'Operator_Vendor': 'Nokia', 'Region': 'North', 'City': 'Leeds',
         'Campaign': '2026-Q2', 'score': 1},
    ]
    for dataset_id in scoring_api['complete_dataset_ids']:
        frame = pd.DataFrame(source_rows)
        frame['Vendor'] = [
            'Ericsson', 'Ericsson', 'Huawei', 'Vodafone UK - All', 'O2 - All', 'Nokia',
        ]
        repository.replace_dataset_rows(dataset_id, frame)
        repository.replace_cdr_catalogue(
            dataset_id, vendors=['3_Ericsson', 'EE_Ericsson', 'VF_UK_Huawei', 'O2', 'Nokia'],
            vendors_only=['Ericsson', 'Huawei', 'Vodafone UK - All', 'O2 - All', 'Nokia'],
            regions=['North'], cities=['Leeds'], campaigns=['2026-Q2'],
            operators=['3', 'EE', 'VF_UK', 'O2'],
        )

    response = scoring_api['client'].post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'nr_mode': 'NSA',
        'aggregation_levels': ['Operator'],
        'context_filters': {'Vendor': ['Ericsson', 'Huawei']},
    })
    assert response.status_code == 200, response.text
    job = response.json()['job']
    assert job['context_filters']['Vendor'] == ['Ericsson', 'Huawei']
    completed = scoring_jobs.run_scoring_job(repository, job['id'])
    assert completed['status'] == 'completed', completed.get('error')
    frames = scoring_api['calls'][0]['frames']
    source_frames = frames.values() if isinstance(frames, dict) else [item[-1] for item in frames]
    actual_operators = {operator for frame in source_frames for operator in frame['Operator']}
    assert actual_operators == expected_operators

    alias_filter_response = scoring_api['client'].post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'nr_mode': 'NSA',
        'aggregation_levels': ['Operator'],
        'context_filters': {'Vendor': ['VF_UK - All Vendors']},
    })
    assert alias_filter_response.status_code == 200, alias_filter_response.text
    alias_filter_job = alias_filter_response.json()['job']
    alias_filter = scoring_jobs.run_scoring_job(repository, alias_filter_job['id'])
    assert alias_filter['status'] == 'completed', alias_filter.get('error')
    alias_frames = scoring_api['calls'][1]['frames']
    alias_source_frames = alias_frames.values() if isinstance(alias_frames, dict) else [item[-1] for item in alias_frames]
    assert {operator for frame in alias_source_frames for operator in frame['Operator']} == {'VF_UK'}

    operator_scoped_response = scoring_api['client'].post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'nr_mode': 'NSA',
        'aggregation_levels': ['Operator'],
        'context_filters': {'Vendor': ['Ericsson', 'Huawei'], 'Operator': ['3', 'EE', 'VF_UK']},
    })
    assert operator_scoped_response.status_code == 200, operator_scoped_response.text
    operator_scoped_job = operator_scoped_response.json()['job']
    operator_scoped = scoring_jobs.run_scoring_job(repository, operator_scoped_job['id'])
    assert operator_scoped['status'] == 'completed', operator_scoped.get('last_error')
    scoped_frames = scoring_api['calls'][2]['frames']
    scoped_source_frames = scoped_frames.values() if isinstance(scoped_frames, dict) else [item[-1] for item in scoped_frames]
    scoped_operators = {operator for frame in scoped_source_frames for operator in frame['Operator']}
    assert scoped_operators == {'3', 'EE', 'VF_UK'}

    # A legacy result can still contain prefixed identities. API views and PPT
    # output should normalize copies while the saved result keeps its original data.
    saved = scoring_jobs.get_scoring_job(repository, job['id'], include_result=True)['result']
    for row, vendor in zip(saved['scoring'], ['3_Ericsson', 'EE_Ericsson', 'VF_UK_Huawei', 'O2']):
        row['vendor'] = vendor
    saved['charts'][0]['vendor'] = 'EE'
    saved['gap'][0]['vendor'] = 'O2'
    with sqlite3.connect(repository.db_path) as connection:
        connection.execute(
            'UPDATE scoring_jobs SET result_json = ? WHERE id = ?',
            (json.dumps(saved, ensure_ascii=False), job['id']),
        )

    legacy_snapshot = deepcopy(saved)
    api_result = scoring_api['client'].get(f'/api/scoring/jobs/{job["id"]}')
    assert api_result.status_code == 200, api_result.text
    body = api_result.json()
    assert {row['vendor'] for row in body['scoring']} == {'Ericsson', 'Huawei', 'All'}
    assert {row['vendor'] for row in body['charts']} == {'All'}
    assert {row['vendor'] for row in body['gap']} == {'All'}
    assert '3_Ericsson' not in json.dumps(body)
    assert 'VF_UK_Huawei' not in json.dumps(body)

    ppt = scoring_api['client'].get(f'/scoring/jobs/{job["id"]}/export/ppt')
    assert ppt.status_code == 200, ppt.text
    deck = Presentation(BytesIO(ppt.content))
    deck_text = _presentation_text(deck)
    assert 'Ericsson' in deck_text and 'Huawei' in deck_text and 'All' in deck_text
    assert '3_Ericsson' not in deck_text and 'VF_UK_Huawei' not in deck_text

    persisted = scoring_jobs.get_scoring_job(repository, job['id'], include_result=True)['result']
    assert persisted == legacy_snapshot


def test_vendor_result_normalization_returns_copy_and_respects_underscore_operator():
    legacy = {
        'scoring': [{'Operator': 'VF_UK', 'Vendor': 'VF_UK_Huawei'}],
        'charts': [{'operator': '3', 'vendor': '3_Ericsson'}],
        'gap': [{'operator': 'EE', 'vendor': 'EE_Mixed_Vendor'}],
    }

    normalized = normalize_scoring_vendor_result(legacy)

    assert normalized['scoring'][0]['Vendor'] == 'Huawei'
    assert normalized['charts'][0]['vendor'] == 'Ericsson'
    assert normalized['gap'][0]['vendor'] == 'Mixed_Vendor'
    assert legacy['scoring'][0]['Vendor'] == 'VF_UK_Huawei'


def test_legacy_all_vendor_baseline_warning_is_refreshed_for_api_and_ppt(scoring_api):
    repository = scoring_api['repository']
    response = scoring_api['client'].post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'nr_mode': 'NSA',
        'aggregation_levels': ['Operator', 'Vendor'],
    })
    assert response.status_code == 200, response.text
    job = response.json()['job']
    completed = scoring_jobs.run_scoring_job(repository, job['id'])
    assert completed['status'] == 'completed', completed.get('error')

    saved = scoring_jobs.get_scoring_job(repository, job['id'], include_result=True)['result']
    vendor_by_operator = {
        'EE': 'All', 'O2': 'Ericsson', 'Three UK': 'Nokia', 'Vodafone UK': 'Huawei',
    }
    for row in saved['scoring']:
        row['vendor'] = vendor_by_operator[row['operator']]
    for row in saved['charts']:
        row['vendor'] = 'All'
    for row in saved['gap']:
        row['vendor'] = 'Ericsson'
    legacy_warning = 'Baseline EE is unavailable for one or more comparison groups.'
    saved['warnings'] = [legacy_warning]
    with sqlite3.connect(repository.db_path) as connection:
        connection.execute(
            'UPDATE scoring_jobs SET result_json = ? WHERE id = ?',
            (json.dumps(saved, ensure_ascii=False), job['id']),
        )
    saved_snapshot = deepcopy(saved)

    api_result = scoring_api['client'].get(f'/api/scoring/jobs/{job["id"]}')
    assert api_result.status_code == 200, api_result.text
    assert legacy_warning not in api_result.json()['warnings']
    assert any(
        row['vendor'] == 'All' for row in api_result.json()['scoring']
    )

    ppt = scoring_api['client'].get(f'/scoring/jobs/{job["id"]}/export/ppt')
    assert ppt.status_code == 200, ppt.text
    deck = Presentation(BytesIO(ppt.content))
    assert all(
        legacy_warning not in slide.notes_slide.notes_text_frame.text
        for slide in deck.slides
    )
    persisted = scoring_jobs.get_scoring_job(repository, job['id'], include_result=True)['result']
    assert persisted == saved_snapshot


def test_coverage_notes_name_excluded_kpis_for_each_vendor_campaign_context():
    configuration = scoring_configuration()
    totals = [
        {'environment': 'Combined', 'operator': 'VF_UK', 'vendor': 'Samsung', 'campaign': '2026-Q1',
         'region': 'North', 'city': 'Leeds', 'category': 'Overall', 'complete_coverage': False,
         'available_points': 933.2645, 'max_points': 1000.0},
        {'environment': 'Combined', 'operator': 'VF_UK', 'vendor': 'Samsung', 'campaign': '2026-Q2',
         'region': 'North', 'city': 'Leeds', 'category': 'Overall', 'complete_coverage': False,
         'available_points': 990.9, 'max_points': 1000.0},
        {'environment': 'Combined', 'operator': 'EE', 'vendor': 'Samsung', 'campaign': '2026-Q1',
         'region': 'North', 'city': 'Leeds', 'category': 'Overall', 'complete_coverage': True,
         'available_points': 1000.0, 'max_points': 1000.0},
    ]
    scoring_rows = [
        {'environment': 'DriveCity', 'operator': 'VF_UK', 'vendor': 'Samsung', 'campaign': '2026-Q1',
         'region': 'North', 'city': 'Leeds', 'kpi_code': 'K7', 'category': 'Classic Calls',
         'kpi': 'Call Setup Time > 10 s [%]', 'weighted_points': None, 'max_points': 56.875},
        {'environment': 'DriveCity', 'operator': 'VF_UK', 'vendor': 'Samsung', 'campaign': '2026-Q1',
         'region': 'North', 'city': 'Leeds', 'kpi_code': 'K10', 'category': 'Classic Calls',
         'kpi': 'Call Setup Time', 'weighted_points': None, 'max_points': 9.8605},
        {'environment': 'DriveCity', 'operator': 'VF_UK', 'vendor': 'Samsung', 'campaign': '2026-Q2',
         'region': 'North', 'city': 'Leeds', 'kpi_code': 'K11', 'category': 'Classic Calls',
         'kpi': 'Call Setup Success Ratio', 'weighted_points': None, 'max_points': 56.875},
    ]
    notes = scoring_coverage_notes({
        'configuration': configuration, 'totals': totals, 'scoring': scoring_rows,
    })

    combined_notes = notes['Combined']
    q1_note = next(note for note in combined_notes if '/ 2026-Q1:' in note)
    q2_note = next(note for note in combined_notes if '/ 2026-Q2:' in note)
    assert q1_note.startswith('VF_UK / Samsung / North / Leeds / 2026-Q1: maximum achievable scoring 933.2645 of 1000 points.')
    # Short notes: the first missing KPIs by environment, then how many more.
    assert 'KPIs without valid measurements: Call Setup Time > 10 s [%], Call Setup Time,' in q1_note
    assert '(DriveCity)' in q1_note and 'more.' in q1_note
    assert 'maximum achievable scoring 990.9 of 1000 points.' in q2_note
    assert 'Call Setup Success Ratio' in q2_note
    assert 'K7' not in q1_note and 'K10' not in q1_note and 'K11' not in q2_note
    assert len(combined_notes) == 2


def test_coverage_notes_group_series_with_the_same_missing_contributions():
    configuration = scoring_configuration()
    totals = [
        {'environment': 'Combined', 'operator': operator, 'vendor': 'All', 'campaign': '2026-Q1', 'region': None,
         'city': city, 'category': 'Overall', 'complete_coverage': False, 'available_points': 650.0, 'max_points': 1000.0}
        for operator in ('EE', 'VF_UK', 'O2') for city in ('Leeds', 'York')
    ]
    notes = scoring_coverage_notes({'configuration': configuration, 'totals': totals, 'scoring': []})['Combined']
    assert len(notes) == 1
    assert notes[0].startswith('EE / All / Leeds / 2026-Q1, EE / All / York / 2026-Q1, VF_UK / All / Leeds / 2026-Q1 and 3 more series:')
