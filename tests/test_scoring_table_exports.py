"""Displayed table modes preserve KPI evidence, aggregates and chart inputs."""
import csv
from copy import deepcopy
from io import BytesIO, StringIO

import pytest
from pptx import Presentation

from src.modules.scoring_exports import (
    _hierarchy_chart_color, _hierarchy_score_tables, _score_tables, _summary_extreme_columns,
    _summary_extreme_operators,
    _table_for_mode, export_scoring_csv, export_scoring_powerpoint,
)
from src.modules.scoring_views import build_scoring_views
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_exports import TEMPLATE, _result, _comparison_matrices, _nested_shapes


def _job():
    return {'aggregation_levels': ['Operator'], 'baseline_operator': 'EE',
            'configuration': scoring_configuration(), 'campaigns': ['UK_Q2_2026'], 'nr_mode': 'NSA'}


def _assert_row_matches_category_fill(table, row_index, expected_fill):
    first_fill = str(table.cell(row_index, 0).fill.fore_color.rgb).upper()
    assert first_fill == expected_fill
    assert all(str(table.cell(row_index, column).fill.fore_color.rgb).upper() == first_fill
               for column in range(len(table.columns)))


@pytest.mark.parametrize('mode', ['expanded', 'summary'])
def test_display_csv_has_category_points_and_mean_gap_without_double_counting(mode):
    job = _job()
    result = _result(missing={('O2 UK', 'K1')})
    rows = list(csv.DictReader(StringIO(export_scoring_csv(
        job, result, 'scoring', mode, environment='DriveCity',
    ))))
    matrix = build_scoring_views(job, result)['score_tables'][0]
    operator_rows = [row for row in rows if row['operator'] == 'O2 UK']
    category_rows = [row for row in operator_rows if row['row_type'] == 'category']
    assert len(category_rows) == 7
    assert all(row['kpi_value'] == '' for row in category_rows)
    assert {row['row_type'] for row in operator_rows} == ({'kpi', 'category', 'total'} if mode == 'expanded' else {'category', 'total'})
    total = next(row for row in operator_rows if row['row_type'] == 'total')
    valid_gaps = [row['gaps']['O2 UK'] for row in matrix['rows'] if row['gaps']['O2 UK'] is not None]
    assert float(total['gap_points']) == pytest.approx(sum(valid_gaps) / len(valid_gaps))
    assert float(total['score_points']) == pytest.approx(matrix['total']['values']['O2 UK']['points'])
    assert sum(float(row['score_points']) for row in category_rows) == pytest.approx(float(total['score_points']))


def test_summary_csv_gap_excludes_reference_and_keeps_missing_comparisons():
    result = _result(missing={('EE', 'K1')})
    rows = list(csv.DictReader(StringIO(export_scoring_csv(
        _job(), result, 'gap', 'summary', environment='DriveCity',
    ))))
    assert all(row['operator'] != 'EE' for row in rows)
    assert {row['row_type'] for row in rows} == {'category', 'total'}
    assert all(row['reference_operator'] == 'EE' for row in rows)


def test_csv_exposes_saved_global_raw_value_without_changing_category_or_total_rows():
    job = _job()
    result = _result()
    source = next(row for row in result['scoring'] if row['kpi_code'] == 'K1' and row['operator'] == 'O2 UK')
    result['global_kpis'] = [{
        **{key: source[key] for key in ('campaign', 'region', 'operator', 'dataset_type', 'kpi_code', 'kpi', 'category', 'kpi_type')},
        'value': 1234.5, 'sample_count': 100, 'environment': 'All Environments',
        'complete_coverage': True, 'missing_environments': [],
    }]
    rows = list(csv.DictReader(StringIO(export_scoring_csv(
        job, result, 'scoring', 'expanded', environment='all',
    ))))
    global_row = next(row for row in rows if row['environment'] == 'Combined'
                      and row['operator'] == 'O2 UK' and row['kpi_code'] == 'K1')
    assert global_row['kpi_value'] == '1234.5'
    assert global_row['score_points'] != global_row['kpi_value']
    assert all(row['kpi_value'] == '' for row in rows if row['row_type'] in {'category', 'total'})


def test_ppt_table_modes_include_bold_category_totals_and_keep_identical_charts():
    result = _result(missing={('O2 UK', 'K1')})
    presentations = [Presentation(BytesIO(export_scoring_powerpoint(
        _job(), result, TEMPLATE, table_mode=mode, environment='DriveCity',
    )))
                     for mode in ('expanded', 'summary')]
    expanded, summary = [_comparison_matrices(presentation) for presentation in presentations]
    expanded = expanded[1]
    summary = summary[0]
    assert [slide.shapes.title.text.split('\n')[0] for slide in presentations[0].slides
            if slide.shapes.title and slide.shapes.title.text.startswith('Scoring Tables')] == [
                'Scoring Tables — Summary', 'Scoring Tables — Breakdown',
            ]
    assert len(expanded.rows) == 32 + 7 + 2
    assert len(summary.rows) == 7 + 2
    for table in (summary, expanded):
        assert str(table.cell(0, 0).fill.fore_color.rgb) == '455B65'
        category_fill = 'E6F0F7' if table is summary else 'DCE5E9'
        category_cells = [row.cells[0] for row in list(table.rows)[1:-1]
                          if row.cells[0].text and row.cells[0].text != 'TOTAL']
        assert category_cells
        assert all(str(cell.fill.fore_color.rgb) == category_fill for cell in category_cells)
        if table is expanded:
            _assert_row_matches_category_fill(table, len(table.rows) - 1, category_fill)
        else:
            assert str(table.cell(len(table.rows) - 1, 0).fill.fore_color.rgb) == 'D8DFE4'
    assert all(summary.cell(index, 5).text_frame.paragraphs[0].font.bold for index in range(1, 8))
    assert all(summary.cell(index, 1).text.endswith(' total') for index in range(1, 8))
    assert expanded.cell(len(expanded.rows) - 1, 5).text == summary.cell(len(summary.rows) - 1, 4).text
    chart_values = [[list(series.values) for slide in presentation.slides for shape in _nested_shapes(slide.shapes)
                     if shape.has_chart for series in shape.chart.series] for presentation in presentations]
    assert chart_values[0] == chart_values[1]
    all_gap_slides = [slide for slide in presentations[1].slides
                      if slide.shapes.title.text.split('\n')[0] == 'GAP Analysis — All vs EE']
    table = next(shape.table for shape in all_gap_slides[0].shapes if shape.has_table)
    assert len(table.rows) == 1 + len(scoring_configuration()['metrics'])
    gap_kpis = [table.cell(row, 1).text for row in range(1, len(table.rows))]
    assert all(not kpi.casefold().endswith(' total') for kpi in gap_kpis)
    assert all(kpi.casefold() != 'average kpi gap' for kpi in gap_kpis)


def test_seven_category_tints_are_distinct_and_invalid_export_mode_is_rejected():
    colors = [_hierarchy_chart_color('#F7931E', index, 7) for index in range(7)]
    assert len(set(colors)) == 7
    assert colors[0] == '#F7931E'
    with pytest.raises(ValueError, match='Table mode'):
        export_scoring_powerpoint(_job(), _result(), TEMPLATE, table_mode='unsupported')


@pytest.mark.parametrize('layout', ['end', 'adjacent'])
def test_gap_placement_excludes_reference_and_preserves_score_and_gap_values(layout):
    job = _job()
    presentation = Presentation(BytesIO(export_scoring_powerpoint(
        job, _result(), TEMPLATE, gap_layout=layout, environment='DriveCity',
    )))
    slide = next(slide for slide in presentation.slides if slide.shapes.title.text.split('\n')[0] == 'Scoring Tables — Breakdown')
    table = next(shape.table for shape in slide.shapes if shape.has_table)
    if layout == 'end':
        headers = [cell.text for cell in table.rows[0].cells][5:]
        assert headers[:4] == ['EE', 'O2 UK', 'Three UK', 'Vodafone UK']
        assert len(headers) == 7
        assert all('GAP' in text and 'GAP EE' not in text for text in headers[4:])
    else:
        headers = [cell.text for cell in table.rows[2].cells][5:]
        assert headers == ['Score', 'Score', 'GAP', 'Score', 'GAP', 'Score', 'GAP']
        values = [cell.text for cell in table.rows[3].cells][5:]
        assert values[0] == '58.79'
        assert values[2] != '0.00'
    with pytest.raises(ValueError, match='GAP layout'):
        export_scoring_powerpoint(job, _result(), TEMPLATE, gap_layout='invalid', environment='DriveCity')


def _gap_slide_contents(presentation):
    contents = []
    for slide in presentation.slides:
        title = slide.shapes.title.text.split('\n')[0]
        if not title.startswith('GAP Analysis'):
            continue
        tables = [
            tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
            for shape in slide.shapes if shape.has_table
        ]
        contents.append((title, tuple(tables)))
    return contents


def _scoring_table_score_cells(table, aggregation_levels, layout, operator_count):
    if len(aggregation_levels) == 1 and layout == 'end':
        start_row = 1
        score_columns = list(range(4, 4 + operator_count))
    else:
        kind_row = len(aggregation_levels) + 1
        start_row = kind_row + 1
        score_columns = [
            index for index in range(4, len(table.columns))
            if table.cell(kind_row, index).text == 'Score'
        ]
    return [
        tuple(table.cell(row, column).text for column in score_columns)
        for row in range(start_row, len(table.rows))
    ]


@pytest.mark.parametrize('levels', [['Operator'], ['Operator', 'Region', 'Campaign']])
@pytest.mark.parametrize('layout', ['end', 'adjacent'])
def test_hiding_scoring_table_gaps_preserves_scores_and_gap_analysis(levels, layout):
    job = {**_job(), 'aggregation_levels': levels}
    result = _result()
    with_gaps = Presentation(BytesIO(export_scoring_powerpoint(
        job, result, TEMPLATE, gap_layout=layout, environment='DriveCity', show_gap_values=True,
    )))
    without_gaps = Presentation(BytesIO(export_scoring_powerpoint(
        job, result, TEMPLATE, gap_layout=layout, environment='DriveCity', show_gap_values=False,
    )))

    scoring_slide_with = next(slide for slide in with_gaps.slides
                              if slide.shapes.title.text.split('\n')[0] == 'Scoring Tables — Breakdown')
    scoring_slide_without = next(slide for slide in without_gaps.slides
                                 if slide.shapes.title.text.split('\n')[0] == 'Scoring Tables — Breakdown')
    table_with = next(shape.table for shape in scoring_slide_with.shapes if shape.has_table)
    table_without = next(shape.table for shape in scoring_slide_without.shapes if shape.has_table)
    operator_count = len({row['operator'] for row in result['scoring']})
    assert _scoring_table_score_cells(table_with, levels, layout, operator_count) == \
        _scoring_table_score_cells(table_without, levels, layout, operator_count)
    assert len(table_without.columns) < len(table_with.columns)
    header_row_count = 1 if len(levels) == 1 and layout == 'end' else len(levels) + 2
    assert all(
        'GAP' not in table_without.cell(row, column).text
        for row in range(header_row_count)
        for column in range(len(table_without.columns))
    )

    gap_slides_with = _gap_slide_contents(with_gaps)
    gap_slides_without = _gap_slide_contents(without_gaps)
    assert gap_slides_with
    assert gap_slides_without == gap_slides_with


def _multi_environment_result():
    result = _result(operators=('EE', 'O2 UK'), warnings=[])
    configuration = result['configuration']
    metrics = {metric['code']: metric for metric in configuration['metrics']}
    road_rows = deepcopy(result['scoring'])
    for row in road_rows:
        row['environment'] = 'DriveConnectionroad'
        row['max_points'] = metrics[row['kpi_code']]['contexts']['DriveConnectionroad']['max_points']
        row['weighted_points'] = row['score'] * row['max_points']
    result['scoring'].extend(road_rows)
    return result


@pytest.mark.parametrize('levels', [['Operator'], ['Operator', 'Region', 'Campaign']])
def test_all_environment_ppt_finishes_each_full_block_aggregate_first(levels):
    job = {**_job(), 'aggregation_levels': levels}
    result = _multi_environment_result()
    result.update({'aggregation_contract_version': 2, 'aggregation_levels': levels})
    presentation = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, table_mode='summary')))
    titles = [slide.shapes.title.text.split('\n')[0] for slide in presentation.slides]
    for environment in ('All Environments', 'Drive - City', 'Drive - Connecting Roads'):
        assert environment in titles
    assert titles.count('Scoring Tables — Summary') == 3
    assert titles.count('Scoring Tables — Breakdown') == 3
    assert titles.count('Best Network Scoring per Service') == 3
    assert titles.count('Scoring per Category') == 3
    for slide in presentation.slides:
        text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        if slide.shapes.title:
            assert all('Environment:' not in paragraph.text
                       for paragraph in slide.shapes.title.text_frame.paragraphs)
        if slide.shapes.title and slide.shapes.title.text.startswith('Scoring Tables'):
            assert any(environment in text for environment in (
                'All Environments', 'DriveCity', 'DriveConnectionroad',
            ))
    selected = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, environment='DriveCity')))
    assert [slide.shapes.title.text.split('\n')[0] for slide in selected.slides].count('Scoring Tables — Summary') == 1
    assert [slide.shapes.title.text.split('\n')[0] for slide in selected.slides].count('Scoring Tables — Breakdown') == 1
    assert all('DriveCity' in slide.shapes.title.text for slide in list(selected.slides)[2:])


def test_csv_environment_selection_filters_rows_and_rejects_unavailable_environment():
    rows = list(csv.DictReader(StringIO(export_scoring_csv(
        _job(), _multi_environment_result(), 'scoring', 'summary', environment='DriveConnectionroad',
    ))))
    assert {row['environment'] for row in rows} == {'DriveConnectionroad'}
    with pytest.raises(ValueError, match='not available'):
        export_scoring_powerpoint(_job(), _result(), TEMPLATE, environment='Walk')
    with pytest.raises(ValueError, match='not available'):
        export_scoring_csv(_job(), _result(), 'gap', 'summary', environment='Unknown')


def test_summary_extremes_keep_ties_and_ignore_missing_or_non_finite_values():
    best, worst = _summary_extreme_operators({
        'best-a': {'points': 10}, 'best-b': {'points': 10.0000000001},
        'middle': {'points': 7}, 'worst': {'points': 2}, 'missing': {'points': None},
        'invalid': {'points': float('nan')},
    })
    assert best == {'best-a', 'best-b'}
    assert worst == {'worst'}
    assert _summary_extreme_operators({'a': {'points': 4}, 'b': {'points': 4}}) == (set(), set())
    assert _summary_extreme_operators({'a': {'points': 4}, 'b': {'points': None}}) == (set(), set())


def test_hierarchy_summary_extremes_compare_operators_within_the_same_context():
    columns = [
        {'id': 'north-a', 'path': [{'level': 'Operator', 'value': 'A'}, {'level': 'Region', 'value': 'North'}]},
        {'id': 'north-b', 'path': [{'level': 'Operator', 'value': 'B'}, {'level': 'Region', 'value': 'North'}]},
        {'id': 'south-a', 'path': [{'level': 'Operator', 'value': 'A'}, {'level': 'Region', 'value': 'South'}]},
        {'id': 'south-b', 'path': [{'level': 'Operator', 'value': 'B'}, {'level': 'Region', 'value': 'South'}]},
    ]
    values = {'north-a': {'points': 100}, 'north-b': {'points': 90},
              'south-a': {'points': 10}, 'south-b': {'points': 20}}
    best, worst = _summary_extreme_columns(values, columns)
    assert best == {'north-a', 'south-b'}
    assert worst == {'north-b', 'south-a'}


@pytest.mark.parametrize('hierarchy', [False, True])
def test_summary_tables_preserve_category_fill_extreme_highlights_and_legend(hierarchy):
    operators = ('EE', 'O2 UK', 'Three UK', 'Vodafone UK')
    scores = {'EE': .6, 'O2 UK': .9, 'Three UK': .9, 'Vodafone UK': .3}
    job = _job()
    if hierarchy:
        job['aggregation_levels'] = ['Operator', 'Region', 'Campaign']
    result = _result(operators=operators, peer_score=lambda operator, index, metric: scores[operator])
    if hierarchy:
        result.update({'aggregation_contract_version': 2, 'aggregation_levels': job['aggregation_levels']})
    views = build_scoring_views(job, result)
    if hierarchy:
        matrix = next(item for item in views['hierarchy_score_tables']
                      if item['context'].get('environment') == 'DriveCity')
    else:
        matrix = next(item for item in views['score_tables']
                      if item['context'].get('environment') == 'DriveCity')
    matrix = _table_for_mode(matrix, 'summary')
    presentation = Presentation()
    if hierarchy:
        _hierarchy_score_tables(presentation, [matrix], views.get('threshold_legend', []),
                                show_gap_values=False, title='Scoring Tables — Summary')
    else:
        _score_tables(presentation, [matrix], views.get('threshold_legend', []),
                      show_gap_values=False, title='Scoring Tables — Summary')
    slide = presentation.slides[0]
    table = next(shape.table for shape in slide.shapes if shape.has_table and len(shape.table.columns) > 5)
    labels = [cell.text for shape in slide.shapes if shape.has_table and len(shape.table.columns) == 1
              for row in shape.table.rows for cell in row.cells]
    assert labels == ['Best operator', 'Worst operator']
    all_slide_table_text = [cell.text for shape in slide.shapes if shape.has_table
                            for row in shape.table.rows for cell in row.cells]
    assert not any(band in all_slide_table_text for band in ('Low', 'Medium', 'High', 'UltraHigh', 'Unavailable'))

    if hierarchy:
        score_columns = [index for index in range(5, len(table.columns))
                         if table.cell(len(matrix['hierarchy_levels']) + 1, index).text == 'Score']
        first_metric_row = len(matrix['hierarchy_levels']) + 2
        category_rows = [row_index for row_index in range(first_metric_row, len(table.rows) - 1)
                         if table.cell(row_index, 1).text.endswith(' total')]
    else:
        score_columns = list(range(4, 4 + len(matrix['operators'])))
        category_rows = [row_index for row_index in range(1, len(table.rows) - 1)
                         if table.cell(row_index, 1).text.endswith(' total')]
    assert category_rows
    assert str(table.cell(0, 0).fill.fore_color.rgb) == str(table.cell(0, 1).fill.fore_color.rgb) == '455B65'
    for row_index in category_rows:
        assert str(table.cell(row_index, 0).fill.fore_color.rgb) == 'E6F0F7'
        assert str(table.cell(row_index, 1).fill.fore_color.rgb) == 'E3E6E7'
    assert str(table.cell(len(table.rows) - 1, 0).fill.fore_color.rgb) == 'D8DFE4'
    for row_index in [*category_rows, len(table.rows) - 1]:
        colors = [str(table.cell(row_index, column).fill.fore_color.rgb) for column in score_columns]
        assert colors.count('C6EFCE') == 2  # Tied best operators.
        assert colors.count('FFC7CE') == 1



@pytest.mark.parametrize('mode', ['expanded', 'summary'])
@pytest.mark.parametrize(('layout', 'show_gap'), [('end', True), ('adjacent', True), ('end', False)])
def test_ppt_kpi_type_column_colors_and_category_background_aggregates(mode, layout, show_gap):
    presentation = Presentation(BytesIO(export_scoring_powerpoint(
        _job(), _result(), TEMPLATE, table_mode=mode, gap_layout=layout,
        show_gap_values=show_gap, environment='DriveCity',
    )))
    checked_types = set()
    for slide in presentation.slides:
        title = slide.shapes.title.text.split('\n')[0]
        if not title.startswith('Scoring Tables —') and not title.startswith('GAP Analysis'):
            continue
        table = next((shape.table for shape in slide.shapes if shape.has_table
                      and any(cell.text == 'Type of KPI' for cell in shape.table.rows[0].cells)), None)
        if table is None:
            continue
        type_column = next(i for i, cell in enumerate(table.rows[0].cells) if cell.text == 'Type of KPI')
        if title.startswith('Scoring Tables —'):
            assert type_column == 2
        for row in table.rows:
            cell = row.cells[type_column]
            if cell.text in {'Reliable', 'Diff'}:
                checked_types.add(cell.text)
                expected = 'D8EFCA' if cell.text == 'Reliable' else 'FFF2CC'
                assert str(cell.fill.fore_color.rgb) == expected
            elif row.cells[0].text.upper() in {'TOTAL', 'TOTALS'} or row.cells[1].text.endswith(' total'):
                assert cell.text == ''
                if title.startswith('Scoring Tables —'):
                    expected = 'E6F0F7' if title.endswith('Summary') else 'DCE5E9'
                    assert str(cell.fill.fore_color.rgb) == expected
    if mode == 'expanded':
        assert checked_types == {'Reliable', 'Diff'}
    else:
        assert checked_types == {'Reliable', 'Diff'}


def test_ppt_score_numeric_font_uses_available_space_without_exceeding_cell_height():
    result = _result()
    for mode in ['expanded', 'summary']:
        presentation = Presentation(BytesIO(export_scoring_powerpoint(
            _job(), result, TEMPLATE, table_mode=mode, gap_layout='adjacent', environment='DriveCity',
        )))
        desired = 'Scoring Tables — Summary' if mode == 'summary' else 'Scoring Tables — Breakdown'
        table = next(shape.table for slide in presentation.slides
                     if slide.shapes.title.text.split('\n')[0] == desired
                     for shape in slide.shapes if shape.has_table)
        font_sizes = []
        for row in list(table.rows)[3:]:
            for cell in list(row.cells)[5:]:
                size = cell.text_frame.paragraphs[0].font.size.pt
                assert size <= row.height.pt
                font_sizes.append(size)
        if mode == 'summary':
            assert max(font_sizes) > 7.5
