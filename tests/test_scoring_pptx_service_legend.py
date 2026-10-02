from pptx.util import Inches

from tests.test_scoring_exports import _export, _result


def test_service_key_is_editable_and_below_each_best_network_chart():
    presentation = _export(_result())
    service_groups = [
        shape for slide in presentation.slides for shape in slide.shapes
        if shape.name == 'Scoring Best Network Service Chart'
    ]
    assert service_groups
    for group in service_groups:
        chart = next(shape for shape in group.shapes if shape.has_chart)
        heading = next(shape for shape in group.shapes
                       if shape.name == 'Scoring Chart Service Key Heading')
        key = next(shape for shape in group.shapes
                   if shape.name == 'Scoring Chart Service Shade Key')
        assert heading.text == 'Service types'
        assert heading.top >= chart.top + chart.height
        assert key.top >= heading.top + heading.height
        assert key.left == chart.left
        assert key.width == chart.width
        assert key.top + key.height < Inches(7.03)
        cells = list(key.table.rows[0].cells)
        assert [cell.text for cell in cells] == ['Data', 'Voice']
        assert [str(cell.fill.fore_color.rgb) for cell in cells] == ['555555', 'C4C4C4']
        assert [str(cell.text_frame.paragraphs[0].font.color.rgb) for cell in cells] == [
            'FFFFFF', '17232D',
        ]
