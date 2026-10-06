from pathlib import Path
import zipfile

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

from src.modules.network_insights_export import export_network_insights_powerpoint, export_network_insights_word, overview_table

TEMPLATE = Path(__file__).resolve().parents[1] / 'assets/ppt-templates/Template_CDR_analysis.pptx'


def _analysis(comparison=False):
    rows = [{
        'operator': f'3 · Edinburgh · LTE · 2026-Q{1 + index % 2}', 'samples': 5751,
        'rsrp_mean': -89.1, 'rsrp_median': -90.4, 'rsrp_p10': -104.8,
        'low_coverage_share': 2.4, 'sinr_mean': 7.3, 'sinr_median': 6.9,
        'sinr_p10': -0.8, 'high_interference_share': 13.8, 'observed_enodebs': 140, 'observed_cells': 512,
        'deltas': {'rsrp_median': 1.2, 'low_coverage_share': -0.3, 'sinr_median': 0.2, 'high_interference_share': -1.1},
    } for index in range(25)]
    analysis = {'rf_rows': rows, 'maps': {}, 'spectrum': {}}
    if comparison:
        analysis['comparison'] = {'latest': '2026-Q2', 'previous': '2026-Q1'}
    return analysis


def test_network_ppt_preserves_dashboard_template_and_fits_group_cells(tmp_path):
    output = tmp_path / 'network.pptx'
    source = _analysis()
    export_network_insights_powerpoint(output, source, {}, {'group_label': 'Operator → City → Technology → Campaign'}, lambda *_args: b'', template=TEMPLATE)
    deck, reference = Presentation(output), Presentation(TEMPLATE)
    assert (deck.slide_width, deck.slide_height) == (reference.slide_width, reference.slide_height)
    assert deck.slides[0].slide_layout.name == 'Title Page'
    assert deck.slides[-1].slide_layout.name == 'Black logo end slide'
    assert len(deck.slide_layouts) == len(reference.slide_layouts)
    tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 1).text == "Samples"]
    assert len(tables) == 3
    columns, values = overview_table(source)
    actual = []
    for table in tables:
        assert [cell.text for cell in table.rows[0].cells] == columns
        assert table.columns[0].width > table.columns[1].width
        for row in table.rows:
            assert all(cell.text_frame.word_wrap is False for cell in row.cells)
            assert all(len(cell.text_frame.paragraphs) == 1 for cell in row.cells)
        actual.extend([[cell.text for cell in row.cells] for row in list(table.rows)[1:]])
    assert actual == values
    with zipfile.ZipFile(output) as generated, zipfile.ZipFile(TEMPLATE) as original:
        assert generated.read('ppt/theme/theme1.xml') == original.read('ppt/theme/theme1.xml')
        assert generated.read('ppt/slideMasters/slideMaster1.xml') == original.read('ppt/slideMasters/slideMaster1.xml')


def test_wide_network_ppt_repeats_group_without_losing_comparison_metrics(tmp_path):
    output = tmp_path / 'wide-network.pptx'
    source = _analysis(comparison=True)
    export_network_insights_powerpoint(output, source, {}, {}, lambda *_args: b'', template=TEMPLATE)
    tables = [shape.table for slide in Presentation(output).slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 1).text != "Observed eNodeBs"]
    columns, _rows = overview_table(source)
    assert {cell.text for table in tables for cell in table.rows[0].cells} == set(columns)
    assert all(table.cell(0, 0).text == 'Group' for table in tables)
    assert all(run.font.size.pt >= 8 for table in tables for row in table.rows for cell in row.cells for paragraph in cell.text_frame.paragraphs for run in paragraph.runs)
    for column in columns[1:]:
        assert sum(sum(cell.text == column for cell in table.rows[0].cells) for table in tables) == 3


def test_network_word_uses_landscape_fixed_single_line_tables(tmp_path):
    for comparison in [False, True]:
        output = tmp_path / f'network-{comparison}.docx'
        source = _analysis(comparison)
        export_network_insights_word(output, source, {}, {}, lambda *_args: b'')
        document = Document(output)
        section = document.sections[0]
        assert section.page_width > section.page_height
        columns, values = overview_table(source)
        # The Overview cards come first; the RF Quality Overview table follows them.
        table = next(table for table in document.tables if [cell.text for cell in table.rows[0].cells] == columns)
        assert table.autofit is False
        assert table.columns[0].width > table.columns[1].width
        assert sum(column.width for column in table.columns) <= section.page_width - section.left_margin - section.right_margin + 9144
        assert [cell.text for cell in table.rows[0].cells] == columns
        assert [[cell.text for cell in row.cells] for row in list(table.rows)[1:]] == values
        assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
        for row in table.rows:
            assert row._tr.trPr.find(qn('w:cantSplit')) is not None
            for cell in row.cells:
                assert cell._tc.tcPr.find(qn('w:noWrap')) is not None
                assert len(cell.paragraphs) == 1 and '\n' not in cell.text
                assert all(run.font.size.pt >= 8 for run in cell.paragraphs[0].runs)


def test_network_summaries_include_all_grouped_deployment_views_and_exclude_full_inventory(client, tmp_path, monkeypatch):
    import pandas as pd
    import src.DashboardAnalytic as app_module
    from src.modules import network_insights as ni, network_insights_export as export
    from src.modules.repository import local_now_iso

    client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=False)
    repo = app_module.repository

    def ready(name, kind, frame):
        source = tmp_path / name
        source.write_text('source', encoding='utf-8')
        identifier, _ = repo.add_dataset(name, str(source), 'admin')
        repo.replace_dataset_rows(identifier, frame)
        repo.update_dataset_profile(identifier, status='ready', progress=100, dataset_kind=kind,
                                    nr_mode='NSA' if kind == 'data' else None, row_count=len(frame),
                                    column_count=len(frame.columns), processed_at=local_now_iso())
        return identifier

    data_id = ready('CDR.xlsx', 'data', pd.DataFrame({'Operator': ['VF'] * 3, 'Campaign': ['2026-Q1'] * 3,
                                                     'LTE_PCell_RSRP_Avg': [-90, -95, -100], 'LTE_PCell_SINR_Avg': [10, 12, 8]}))
    ready('mapping.xlsx', 'mapping_vodafone', pd.DataFrame({'Site_ID': ['S1'], 'Cell ID': ['1'], 'eMOCNScenario': ['NNS']}))
    ready('clusters.geojson', 'clusters', pd.DataFrame({'Cluster_Field': ['Cluster']}))
    (tmp_path / 'clusters.geojson').write_text('{"type":"FeatureCollection","features":[{"type":"Feature","properties":{"Name":"Cluster A"},"geometry":{"type":"Polygon","coordinates":[[[0,0],[1,0],[1,1],[0,1],[0,0]]]}}]}', encoding='utf-8')
    captured = []

    def writer(destination, _analysis, deployment, _selection, _render):
        captured.append(deployment)
        destination.write_bytes(b'export')
        return destination

    monkeypatch.setattr(export, 'export_network_insights_word', writer)
    monkeypatch.setattr(export, 'export_network_insights_powerpoint', writer)
    for kind in ['word', 'powerpoint']:
        app_module.write_network_insights_summary({'datasets': {'data': [data_id]}}, kind, tmp_path / f'{kind}.output')
    for deployment in captured:
        assert [group['group'] for group in deployment['groupings']] == list(ni.DEPLOYMENT_GROUPINGS)
        tables = export.deployment_tables(deployment)
        assert len(tables) == len(ni.DEPLOYMENT_GROUPINGS)
        assert {columns[0] for _title, columns, _rows in tables} == set(ni.DEPLOYMENT_GROUPINGS.values())
        assert all('Full Site' not in title for title, _columns, _rows in tables)
        assert deployment['cluster_inventories']['rows'] == [{'operator': 'Vodafone', 'sites': 1, 'cells': 1}]


def test_cluster_sites_sources_are_both_present_in_ppt_and_word(tmp_path):
    analysis = _analysis()
    deployment = {'cluster_inventories': {'rows': [{'operator': 'Vodafone', 'sites': 37201, 'cells': 516898},
                                                 {'operator': 'Three', 'sites': 16107, 'cells': 140127}]}}
    ppt = tmp_path / 'density.pptx'
    word = tmp_path / 'density.docx'
    export_network_insights_powerpoint(ppt, analysis, deployment, {}, lambda *_args: b'', template=TEMPLATE)
    export_network_insights_word(word, analysis, deployment, {}, lambda *_args: b'')
    ppt_tables = [shape.table for slide in Presentation(ppt).slides for shape in slide.shapes if shape.has_table]
    word_tables = Document(word).tables
    for tables in [ppt_tables, word_tables]:
        observed = [table for table in tables if table.cell(0, 1).text == 'Observed eNodeBs']
        inventory = [table for table in tables if table.cell(0, 1).text == 'Inventory sites']
        assert sum(len(table.rows) - 1 for table in observed) == 25
        assert len(inventory) == 1
        assert [[cell.text for cell in row.cells] for row in list(inventory[0].rows)[1:]] == [
            ['Vodafone', '37,201', '516,898'], ['Three', '16,107', '140,127']]
