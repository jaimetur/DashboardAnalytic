from __future__ import annotations

import csv
import io
import json
import math
from pathlib import Path

import pandas as pd
import pytest

import src.DashboardAnalytic as app_module
from src.modules import network_insights as ni
from src.modules.cdr_reporting import _explicit_bucket_labels, _osm_map_tile_geometry, _status_chart_categories
from src.modules.repository import local_now_iso


TEMPLATE_PATH = Path(__file__).resolve().parent / 'fixtures' / 'rf-quality-template.csv'


def _login(client) -> None:
    response = client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=False)
    assert response.status_code == 303


def _add_ready_dataset(tmp_path: Path, name: str, kind: str, rows: pd.DataFrame, nr_mode: str | None = 'NSA') -> int:
    repository = app_module.repository
    source = tmp_path / name
    source.write_text('test source', encoding='utf-8')
    dataset_id, _created = repository.add_dataset(name, str(source), 'admin')
    repository.replace_dataset_rows(dataset_id, rows)
    repository.update_dataset_profile(
        dataset_id, status='ready', progress=100, dataset_kind=kind, nr_mode=nr_mode,
        row_count=len(rows), column_count=len(rows.columns), processed_at=local_now_iso(),
    )
    return dataset_id


def _add_ready_polygon(tmp_path, kind, name):
    filename = f'{kind}.geojson'
    identifier = _add_ready_dataset(tmp_path, filename, kind, pd.DataFrame({'Name': [name]}), None)
    (tmp_path / filename).write_text(json.dumps({
        'type': 'FeatureCollection', 'features': [{'type': 'Feature', 'properties': {'Name': name},
            'geometry': {'type': 'Polygon', 'coordinates': [[[-2, 53], [0, 53], [0, 55], [-2, 55], [-2, 53]]]}}],
    }), encoding='utf-8')
    return identifier


def _data_rows(campaign: str, rsrp_offset: float = 0.0) -> pd.DataFrame:
    count = 12
    return pd.DataFrame({
        'Operator': ['EE'] * 6 + ['Vodafone UK'] * 6,
        'Campaign': [campaign] * count,
        'Region': ['North'] * count,
        'City': ['Leeds'] * 6 + ['York'] * 6,
        'Test_Start_Latitude': [53.80 + index * 0.0001 for index in range(count)],
        'Test_Start_Longitude': [-1.55 + index * 0.0001 for index in range(count)],
        'LTE_PCell_RSRP_Avg': [-85 + rsrp_offset, -95, -105, -115, -118, -90, -80, -82, -84, -86, -112, -999],
        'LTE_PCell_SINR_Avg': [15, 10, 4, -2, -3, 22, 25, 18, 12, 8, -1, 6],
        'NR_PCell_RSRP_Avg': [None] * count,
        'NR_PCell_SINR_Avg': [None] * count,
        'LAC_CID_xARFCN': ['[LTE E-UTRA 20, 100, 154067457, 6300]'] * 6 + ['[LTE E-UTRA 3, 200, 30001922, 1300]->[LTE E-UTRA 7, 200, 30001923, 2850]'] * 6,
        'LTE_PCC_EARFCN': [6300] * 6 + [1300] * 6,
        'LTE_DL_PCell_Bandwidth': [10] * 6 + [20] * 6,
        'NR_DL_PCell_Band': [''] * count,
    })


def test_radio_identifiers_and_bands_are_parsed() -> None:
    assert ni.lte_enodeb_from_eci('154067457') == 601826
    assert ni.lte_enodeb_from_eci(12345) is None
    assert ni.lte_band_for_earfcn('6300') == 'B20'
    assert ni.lte_band_for_earfcn('1300') == 'B3'
    assert ni.lte_band_for_earfcn('x') == ''
    assert ni.band_class('B20') == 'Low'
    assert ni.band_class('B3') == 'Mid'
    assert ni.band_class('n78') == 'High (TDD)'
    assert ni.nr_band_label('NR BAND 78->NR BAND 1') == 'n78'
    assert ni.enodeb_label('L 601983') == '601983'
    cells = ni.parse_cell_trace('[LTE E-UTRA 7, 12, 30001923, 2850]->[UMTS 1, 3, 456, 10700]')
    assert cells[0] == {'technology': 'LTE', 'band': 'B7', 'cell': '30001923', 'channel': '2850'}
    assert cells[1]['technology'] == 'UMTS'


def test_samples_are_normalised_and_summarised_per_operator() -> None:
    frame = _data_rows('Q1 2026')
    columns = ni.resolve_source_columns(frame.columns, 'data')
    samples = ni.normalise_samples(frame, 'data', columns)

    assert samples['kind'].unique().tolist() == ['Data']
    assert math.isnan(samples.loc[11, 'lte_rsrp'])  # -999 is a placeholder
    assert samples.loc[0, 'enodebs'] == ('601826',)
    assert samples.loc[6, 'cells'] == ('30001922', '30001923')
    assert samples.loc[0, 'lte_band'] == 'B20'

    rows = ni.rf_summary(samples, 'lte', None, -110, 0)
    ee = next(row for row in rows if row['operator'] == 'EE')
    assert ee['samples'] == 6
    assert ee['low_coverage_share'] == pytest.approx(33.3)
    assert ee['high_interference_share'] == pytest.approx(33.3)
    assert ee['observed_enodebs'] == 1
    assert sum(item['share'] for item in ee['rsrp_classes']) == pytest.approx(100, abs=0.2)

    grouped = ni.rf_summary(samples, 'lte', 'city', -110, 0)
    assert {(row['operator'], row['group']) for row in grouped} == {('EE', 'Leeds'), ('Vodafone UK', 'York')}

    observed = ni.observed_spectrum(samples)
    assert {(row['operator'], row['band'], row['band_class']) for row in observed} == {('EE', 'B20', 'Low'), ('Vodafone UK', 'B3', 'Mid')}
    assert next(row for row in observed if row['band'] == 'B3')['typical_bandwidth_mhz'] == 20


def test_grid_cells_rank_weak_areas_and_build_map_payload() -> None:
    samples = pd.DataFrame({
        'latitude': [51.5] * 4 + [51.6] * 4,
        'longitude': [-0.1] * 4 + [-0.2] * 4,
        'lte_rsrp': [-115, -118, -112, -90, -80, -85, -82, -84],
        'city': ['London'] * 8,
        'region': ['South'] * 8,
    })
    cells, grid = ni.grid_cells(samples, 'lte_rsrp', -110, 250, 3)
    assert grid == 250
    assert len(cells) == 2
    ranked = ni.hotspots(cells)
    assert len(ranked) == 1
    assert ranked[0]['bad_share'] == 75
    assert ranked[0]['city'] == 'London'

    payload = ni.map_payload(cells, ni.RSRP_CLASSES, 'Coverage', 'dBm', _osm_map_tile_geometry)
    assert payload['type'] == 'map'
    assert [item['label'] for item in payload['legend']['items']][0] == 'Excellent (≥ -80 dBm)'
    assert sum(len(series['points']) for series in payload['series']) == 2
    assert ni.map_payload(pd.DataFrame(), ni.RSRP_CLASSES, 'Coverage', 'dBm', None)['type'] == 'empty'


def test_spectrum_holdings_are_validated_parsed_and_summarised() -> None:
    holdings = ni.parse_spectrum_csv('Operator\tBand\tDuplex\tBand Class\tBandwidth MHz\tNotes\nEE\tB20\t\t\t20\t2x10\nEE\tn78\t\t\t80\t\nO2\tB3\tFDD\tMid\t40\t')
    assert holdings[0] == {'operator': 'EE', 'band': 'B20', 'duplex': 'FDD', 'band_class': 'Low', 'bandwidth_mhz': 20.0, 'notes': '2x10'}
    assert holdings[1]['duplex'] == 'TDD' and holdings[1]['band_class'] == 'High (TDD)'
    assert ni.parse_spectrum_csv(ni.spectrum_holdings_csv(holdings)) == holdings
    assert ni.parse_spectrum_csv('Operator;Band;Bandwidth MHz\nEE;B1;30') == [
        {'operator': 'EE', 'band': 'B1', 'duplex': 'FDD', 'band_class': 'Mid', 'bandwidth_mhz': 30.0, 'notes': ''},
    ]
    assert ni.parse_spectrum_csv('  ') == []
    summary = {row['operator']: row for row in ni.licensed_spectrum_summary(holdings)}
    assert summary['EE'] == {'operator': 'EE', 'Low': 20.0, 'Mid': 0.0, 'High (TDD)': 80.0, 'total': 100.0}
    with pytest.raises(ValueError, match='bandwidth'):
        ni.normalise_spectrum_holdings([{'operator': 'EE', 'band': 'B20', 'bandwidth_mhz': 'wide'}])
    with pytest.raises(ValueError, match='band class'):
        ni.normalise_spectrum_holdings([{'operator': 'EE', 'band': 'B99', 'bandwidth_mhz': 10}])


def test_threshold_and_negative_bucket_labels_follow_the_configuration() -> None:
    assert _explicit_bucket_labels([1, 5, 20]) == ['<1', '1-5', '5-20', '20+']
    assert _explicit_bucket_labels([-110, -100, -90, -80]) == ['< -110', '-110 to -100', '-100 to -90', '-90 to -80', '-80+']
    frame = pd.DataFrame({'value': [-115, -100, 1.2, 2.0]})
    _result, labels, _colours = _status_chart_categories(frame, 'value', quality=True, threshold=-110)
    assert labels == ('< -110', '≥ -110')
    _result, labels, _colours = _status_chart_categories(frame, 'value', quality=True, threshold=1.6)
    assert labels == ('< 1.6', '≥ 1.6')


def test_bundled_rf_quality_template_is_valid() -> None:
    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    assert len(catalogue) == 116
    from src.modules.cdr_reporting import catalogue_csv, parse_catalog_csv
    restored = parse_catalog_csv(catalogue_csv(catalogue), 'nsa')
    assert [(entry.cdr_source, entry.chart_type, entry.kpi, entry.filters) for entry in restored] == [
        (entry.cdr_source, entry.chart_type, entry.kpi, entry.filters) for entry in catalogue
    ]
    from pptx import Presentation
    from src.modules.cdr_reporting import _named_slide_layout
    presentation = Presentation(str(Path(__file__).resolve().parents[1] / 'assets' / 'ppt-templates' / 'Template_CDR_analysis.pptx'))
    for layout_name in {entry.layout for entry in catalogue}:
        assert _named_slide_layout(presentation, layout_name) is not None, layout_name
    chart_types = {entry.chart_type for entry in catalogue}
    assert {'Distribution Stacked Vertical Bars', 'Threshold Stacked Vertical Bars', 'Average Vertical Bars'} <= chart_types
    maps = [entry for entry in catalogue if entry.chart_type == 'Map']
    assert len(maps) == 16
    assert {entry.source_kind for entry in maps} == {'data', 'voice', 'speech', 'all'}
    assert all(len(entry.kpi.split(' vs ')) == 3 for entry in maps)
    for chart_type in ('Distribution Stacked Vertical Bars', 'Threshold Stacked Vertical Bars'):
        nr_entries = [entry for entry in catalogue if entry.chart_type == chart_type and entry.chart_title.startswith('NR ')]
        assert len(nr_entries) == 8
        assert {entry.source_kind for entry in nr_entries} == {'data', 'voice', 'speech', 'all'}
        assert all((entry.kpi.startswith('NR_') or '_NR_' in entry.kpi) and entry.exclude_null_empty for entry in nr_entries)
        assert all(f'{entry.kpi} >=' in entry.filters and f'{entry.kpi} <=' in entry.filters for entry in nr_entries)
    assert catalogue[-1].layout == 'Black logo end slide'
    cdf_slides = {entry.slide for entry in catalogue if entry.chart_type == 'CDF Line'}
    assert len(cdf_slides) == 8
    for slide in cdf_slides:
        entries = [entry for entry in catalogue if entry.slide == slide]
        assert len(entries) == 4
        assert all(entry.layout == 'Title + 2 rows + 2 columns + comments right' for entry in entries)
        assert [entry.chart_type for entry in entries[:2]] == ['CDF Line', 'CDF Line']
        assert all(entry.chart_type == 'Histogram Line' for entry in entries[2:])
        assert all(entry.grouping_rows == 'Operator' and entry.grouping_columns == 'Campaign' for entry in entries[2:])


@pytest.mark.parametrize('field,values,edges', [
    ('RSRP', [-115, -105, -95, -85, -75], '-110,-100,-90,-80'),
    ('SINR', [-1, 0, 5, 13, 20], '0,5,13,20'),
])
def test_quality_maps_use_measurement_buckets_without_replacing_coordinates(field, values, edges) -> None:
    from src.modules.cdr_reporting import CatalogEntry, catalog_chart_payload

    entry = CatalogEntry(
        slide=1, slide_title='RF map', slide_subtitle='', layout='Title and 1 column + Comments',
        chart_title='Quality', cdr_source='CDR-Data', kpi=f'Latitude vs Longitude vs {field}',
        chart_type='Map', legend='Value Bucket', filters=f'Buckets = {edges}',
        grouping_rows='Value Bucket', grouping_columns='',
        legend_position='Right', exclude_null_empty=True,
    )
    frame = pd.DataFrame({
        'Latitude': [51.5 + index * .01 for index in range(7)],
        'Longitude': [-.1] * 7,
        field: [*values, None, 'invalid'],
    })
    model = catalog_chart_payload(frame, entry)
    assert model['type'] == 'map'
    assert sum(len(series['points']) for series in model['series']) == 5
    assert {series['colour'] for series in model['series']} == {
        '#D7263D', '#F08A24', '#F2C230', '#8BC34A', '#2E8B57',
    }
    points = [point for series in model['series'] for point in series['points']]
    assert sorted(point[1] for point in points) == pytest.approx([51.5, 51.51, 51.52, 51.53, 51.54])


def test_rf_histogram_counts_samples_in_ordered_quality_ranges() -> None:
    from src.modules.cdr_reporting import catalog_chart_payload

    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    entry = next(entry for entry in catalogue if entry.chart_type == 'Histogram Line' and entry.kpi == 'NR_PCell_RSRP_Avg')
    frame = pd.DataFrame({
        'Operator': ['EE'] * 9,
        'Campaign': ['2026-Q1'] * 9,
        'NR_PCell_RSRP_Avg': [-115, -105, -100, -95, -85, -75, None, 'invalid', 0],
    })
    model = catalog_chart_payload(frame, entry)
    assert model['histogram'] is True
    assert model['bin_size'] == 5
    series = model['series'][0]
    assert sum(series['counts']) == series['samples'] == 6
    assert series['counts'][series['bin_edges'].index(-100)] == 1
    assert series['counts'][series['bin_edges'].index(-95)] == 1
    assert sum(series['counts']) / series['samples'] == 1
    assert [band['colour'] for band in model['quality_bands']] == [
        '#D7263D', '#F08A24', '#F2C230', '#8BC34A', '#2E8B57',
    ]


def test_network_insights_page_and_analysis(client, tmp_path) -> None:
    _login(client)
    _add_ready_dataset(tmp_path, 'CDR_Data_NSA_2026-Q1.xlsx', 'data', _data_rows('2026-Q1'))
    _add_ready_dataset(tmp_path, 'CDR_Data_NSA_2026-Q2.xlsx', 'data', _data_rows('2026-Q2', rsrp_offset=5))

    page = client.get('/network-insights')
    assert page.status_code == 200
    assert 'id="ni-config"' in page.text
    assert 'module-tab-network-insights active' in page.text
    config = json.loads(page.text.split('<script id="ni-config" type="application/json">', 1)[1].split('</script>', 1)[0])
    ids = [row['id'] for row in config['datasets'] if row['kind'] == 'data']
    assert len(ids) == 2

    response = client.post('/api/network-insights/analysis', json={'datasets': {'data': ids}, 'technology': 'lte', 'group': 'campaign'})
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload['comparison'] is None

    operator_comparison = client.post('/api/network-insights/analysis', json={
        'datasets': {'data': ids}, 'technology': 'lte', 'group': 'operator',
    })
    assert operator_comparison.status_code == 200, operator_comparison.text
    assert operator_comparison.json()['comparison'] == {'previous': '2026-Q1', 'latest': '2026-Q2'}
    # Operator is always grouped unless Vendor replaces it.
    assert {row['operator'] for row in payload['overview']} == {'EE · 2026-Q1', 'EE · 2026-Q2', 'VF · 2026-Q1', 'VF · 2026-Q2'}
    operator_payload = operator_comparison.json()
    assert {row['operator'] for row in operator_payload['overview']} == {'EE', 'VF'}
    ee = next(row for row in operator_payload['overview'] if row['operator'] == 'EE')
    assert ee['deltas']['rsrp_median'] is not None
    assert payload['charts']['rsrp_cdf']['type'] == 'cdf'
    assert len(payload['charts']['rsrp_cdf']['series']) == 4
    assert payload['maps']['coverage']['type'] == 'map'
    assert payload['options']['cities'] == ['Leeds', 'York']
    assert payload['spectrum']['licensed'] == []

    filtered = client.post('/api/network-insights/analysis', json={
        'datasets': {'data': ids}, 'technology': 'lte', 'group': 'city', 'cities': ['York'], 'map_operator': 'VF · York',
    }).json()
    assert {row['operator'] for row in filtered['rf_rows']} == {'VF · York'}
    assert filtered['maps']['operator'] == 'VF · York'

    empty = client.post('/api/network-insights/analysis', json={'datasets': {'data': ids}, 'operators': ['Nobody']})
    assert empty.status_code == 400
    assert client.post('/api/network-insights/analysis', json={'datasets': {}}).status_code == 400


def test_network_deployment_counts_inventory_sites(client, tmp_path) -> None:
    _login(client)
    inventory = pd.DataFrame({
        'Site_ID': ['S1', 'S1', 'S2', 'S3', 'S4'],
        'CId___ECI': ['1', '2', '3', '4', '5'],
        'eMOCNScenario': ['NNS', 'NNS', 'S1', 'NNS', ''],
        'OP_Vendor': ['Ericsson', 'Ericsson', 'Nokia', 'Ericsson', 'Nokia'],
    })
    _add_ready_dataset(tmp_path, 'VF_inventory.xlsx', 'mapping_vodafone', inventory, nr_mode=None)

    payload = client.get('/api/network-insights/deployment?group=scenario').json()
    assert payload['group_label'] == 'eMOCN Scenario'
    vodafone = payload['inventories'][0]
    assert vodafone['operator'] == 'Vodafone'
    assert {row['group']: row['sites'] for row in vodafone['rows']} == {'NNS': 2, 'S1': 1, 'Not set': 1}
    assert vodafone['totals'] == {'sites': 4, 'cells': 5}
    assert 'vendor' in vodafone['available_groups'] and 'band' not in vodafone['available_groups']


def test_spectrum_holdings_api_workspace_config_and_archive(client) -> None:
    _login(client)
    saved = client.put('/api/network-insights/spectrum', json={'holdings': [{'operator': 'EE', 'band': 'n78', 'bandwidth_mhz': 80}]})
    assert saved.status_code == 200
    assert saved.json()['summary'][0]['High (TDD)'] == 80
    assert client.put('/api/network-insights/spectrum', json={'holdings': [{'operator': 'EE'}]}).status_code == 400

    form = client.post('/workspace-config/spectrum-holdings/save', data={
        'holdings_csv': 'Operator,Band,Bandwidth MHz\nO2,B20,20\nO2,B40,40',
    }, follow_redirects=False)
    assert form.status_code == 303
    assert 'spectrum_holdings_notice' in form.headers['location']
    config_page = client.get('/workspace-config')
    assert 'id="spectrum-holdings"' in config_page.text
    assert 'O2,B40,TDD,High (TDD),40,' in config_page.text
    invalid = client.post('/workspace-config/spectrum-holdings/save', data={'holdings_csv': 'Operator,Band\nO2,'}, follow_redirects=False)
    assert 'spectrum_holdings_error' in invalid.headers['location']

    workspace = app_module.active_workspace
    archive = json.loads(app_module._operator_mappings_archive_payload(workspace))
    assert archive['version'] == 3
    assert [row['band'] for row in archive['spectrum_holdings']] == ['B20', 'B40']

    ni.save_spectrum_holdings(app_module.repository, [])
    app_module._restore_workspace_operator_mappings(workspace, json.dumps(archive).encode())
    assert len(ni.load_spectrum_holdings(app_module.repository)) == 2

    legacy = {key: value for key, value in archive.items() if key != 'spectrum_holdings'} | {'version': 2}
    app_module._restore_workspace_operator_mappings(workspace, json.dumps(legacy).encode())
    assert len(ni.load_spectrum_holdings(app_module.repository)) == 2
    with pytest.raises(ValueError, match='Spectrum Holdings'):
        app_module._restore_workspace_operator_mappings(workspace, json.dumps(archive | {'spectrum_holdings': [{'operator': 'EE'}]}).encode())


def test_combined_rf_source_pools_samples_instead_of_type_means() -> None:
    from src.modules.rf_catalog_source import pool_rf_frames
    from src.modules.cdr_reporting import catalog_chart_payload

    frames = {
        'data': pd.DataFrame({'Operator': ['EE'] * 3, 'Campaign': ['2026-Q1'] * 3,
                              'LTE_PCell_RSRP_Avg': [-100, -90, -80], 'NR_PCell_RSRP_Avg': [-95, None, None]}),
        'voice': pd.DataFrame({'Operator': ['EE'], 'Campaign': ['2026-Q1'], '4G_RSRP_Avg_A': [-60]}),
        'speech': pd.DataFrame({'Operator': ['EE'], 'Campaign': ['2026-Q1'], 'Playing_RSRP_Avg': [-70]}),
    }
    combined = pool_rf_frames(frames)
    assert combined['LTE_RSRP'].mean() == -80
    assert combined['NR_RSRP'].count() == 1
    assert combined['CDR_Type'].tolist() == ['Data', 'Data', 'Data', 'Voice', 'Speech']
    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    aggregate = [entry for entry in catalogue if entry.source_kind == 'all']
    assert len({entry.slide for entry in aggregate}) == 12
    assert len(aggregate) == 28
    entry = next(entry for entry in aggregate if entry.chart_type == 'Histogram Line' and entry.kpi == 'LTE_RSRP')
    model = catalog_chart_payload(combined, entry)
    assert model['series'][0]['samples'] == 5
    assert sum(model['series'][0]['counts']) == 5


def test_histogram_legacy_png_renderer_accepts_combined_rf_samples(monkeypatch) -> None:
    from PIL import Image
    from io import BytesIO
    from src.modules.cdr_reporting import render_catalog_chart_preview
    from src.modules.rf_catalog_source import pool_rf_frames

    monkeypatch.setenv('DASHBOARD_ANALYTIC_REPORT_CHART_RENDERER', 'pil')
    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    entry = next(entry for entry in catalogue if entry.source_kind == 'all' and entry.chart_type == 'Histogram Line' and entry.kpi == 'LTE_RSRP')
    combined = pool_rf_frames({'data': pd.DataFrame({
        'Operator': ['EE'] * 4, 'Campaign': ['2026-Q1', '2026-Q1', '2026-Q2', '2026-Q2'],
        'LTE_PCell_RSRP_Avg': [-115, -95, -90, -85],
    })})
    image = Image.open(BytesIO(render_catalog_chart_preview(combined, entry)))
    assert image.width >= 1500
    assert image.height >= 900



@pytest.mark.parametrize('layout,columns', [
    ('2 rows + dynamic columns, comments down', True),
    ('2 rows + dynamic columns, comments right', True),
    ('2 columns + dynamic rows, comments down', False),
    ('2 columns + dynamic rows, comments right', False),
])
def test_dynamic_histogram_grids_keep_radio_positions_and_all_operators(layout, columns):
    from dataclasses import replace
    from src.modules.cdr_reporting import expand_dynamic_layouts, _layout_chart_frames, _named_slide_layout, catalogue_csv, parse_catalog_csv
    from pptx import Presentation
    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    base = [entry for entry in catalogue if entry.chart_type == 'Histogram Bars'][:2]
    base = [replace(
        entry, layout=layout, dynamic_rows_field='' if columns else 'Operator',
        dynamic_columns_field='Operator' if columns else '', dynamic_field='',
    ) for entry in base]
    expanded = expand_dynamic_layouts(base, {'Operator': ['A', 'B', 'C', 'D', 'E', 'F']})
    assert len(expanded) == 12
    assert len({entry.slide for entry in expanded}) == 1
    if columns:
        assert all(entry.chart_title.startswith('LTE ') for entry in expanded[:6])
        assert all(entry.chart_title.startswith('NR ') for entry in expanded[6:])
    else:
        assert all(entry.chart_title.startswith('LTE ') for entry in expanded[::2])
        assert all(entry.chart_title.startswith('NR ') for entry in expanded[1::2])
    deck = Presentation('assets/ppt-templates/Template_CDR_analysis.pptx')
    positions = _layout_chart_frames(_named_slide_layout(deck, expanded[0].layout))
    assert len(positions) == 12
    for count in (1, 2, 3, 4, 5, 6, 7):
        actual = expand_dynamic_layouts(base, {'Operator': [str(index) for index in range(count)]})
        assert len(actual) == count * 2
        assert len(_layout_chart_frames(_named_slide_layout(deck, actual[0].layout))) == count * 2
    restored = parse_catalog_csv(catalogue_csv(base), 'nsa')
    if columns:
        assert [entry.dynamic_columns_field for entry in restored] == ['Operator', 'Operator']
        assert [entry.dynamic_rows_field for entry in restored] == ['', '']
    else:
        assert [entry.dynamic_rows_field for entry in restored] == ['Operator', 'Operator']
        assert [entry.dynamic_columns_field for entry in restored] == ['', '']
    from src.modules.report_layouts import canonical_layout_name
    assert [entry.layout for entry in restored] == [canonical_layout_name(layout)] * 2


@pytest.mark.parametrize(
    'dynamic_layout,expected_layout',
    [
        ('Title + 2 rows + dynamic columns', 'Title + 2 rows + 3 columns'),
        ('Title + dynamic rows + 2 columns', 'Title + 3 rows + 2 columns'),
    ],
)
def test_dynamic_grids_without_comments_expand_to_matching_frames(dynamic_layout, expected_layout):
    from dataclasses import replace
    from pptx import Presentation
    from src.modules.cdr_reporting import (
        _layout_chart_frames,
        _named_slide_layout,
        expand_dynamic_layouts,
    )

    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    base = [
        replace(
            entry, layout=dynamic_layout,
            dynamic_rows_field='' if dynamic_layout.startswith('Title + 2 rows') else 'Operator',
            dynamic_columns_field='Operator' if dynamic_layout.startswith('Title + 2 rows') else '',
            dynamic_field='',
        )
        for entry in catalogue if entry.chart_type == 'Histogram Bars'
    ][:2]

    expanded = expand_dynamic_layouts(base, {'Operator': ['A', 'B', 'C']})

    assert len(expanded) == 6
    assert {entry.layout for entry in expanded} == {expected_layout}
    deck = Presentation('assets/ppt-templates/Template_CDR_analysis.pptx')
    layout = _named_slide_layout(deck, expected_layout)
    assert layout is not None
    assert 'comments' not in layout.name.casefold()
    assert len(_layout_chart_frames(layout)) == 6


def test_two_axis_dynamic_expansion_filters_vendor_only_and_paginates_grid_frames():
    from src.modules.cdr_reporting import (
        CatalogEntry,
        _layout_chart_frames,
        _named_slide_layout,
        _select_dynamic_chart_frame,
        expand_dynamic_layouts,
    )
    from pptx import Presentation

    definition = CatalogEntry(
        slide=1, slide_title='Vendor campaign quality', slide_subtitle='',
        layout='Title + dynamic rows + dynamic columns + comments down',
        chart_title='Signal', cdr_source='CDR-Data', kpi='LTE_RSRP',
        chart_type='Histogram Bars', legend='', filters='',
        grouping_rows='Operator_Vendor', grouping_columns='Campaign',
        dynamic_rows_field='Operator_Vendor', dynamic_columns_field='Campaign',
    )
    source = pd.DataFrame({
        'Vendor': ['Ericsson', 'Ericsson', 'Huawei', 'Huawei'],
        'Campaign': ['Q1', 'Q2', 'Q1', 'Q2'],
        'LTE_RSRP': [-100, -95, -90, -85],
    })
    pair = expand_dynamic_layouts([definition], {
        'Vendor': ['Ericsson', 'Huawei'], 'Campaign': ['Q1', 'Q2'],
    }, multivendor=True, vendor_comparison='vendor_only')
    selected = _select_dynamic_chart_frame(
        source, next(entry for entry in pair if entry.dynamic_row_value == 'Ericsson' and entry.dynamic_column_value == 'Q2'),
    )
    assert selected[['Vendor', 'Campaign', 'LTE_RSRP']].to_dict('records') == [
        {'Vendor': 'Ericsson', 'Campaign': 'Q2', 'LTE_RSRP': -95},
    ]

    vendors = [f'Vendor {index}' for index in range(7)]
    campaigns = [f'Q{index}' for index in range(7)]
    paged = expand_dynamic_layouts([definition], {
        'Vendor': vendors, 'Campaign': campaigns,
    }, multivendor=True, vendor_comparison='vendor_only')
    assert len(paged) == 49
    pages = {}
    for entry in paged:
        pages.setdefault(entry.slide, []).append(entry)
    assert len(pages) == 4
    deck = Presentation('assets/ppt-templates/Template_CDR_analysis.pptx')
    page_dimensions = []
    for entries in pages.values():
        row_count = len({entry.dynamic_row_value for entry in entries})
        column_count = len({entry.dynamic_column_value for entry in entries})
        page_dimensions.append((row_count, column_count))
        layout = _named_slide_layout(deck, entries[0].layout)
        assert layout is not None
        assert len(_layout_chart_frames(layout)) == row_count * column_count
    assert sorted(page_dimensions) == [(1, 1), (1, 6), (6, 1), (6, 6)]


def test_operator_histogram_campaign_bars_use_ordered_shades_and_exact_counts():
    from src.modules.cdr_reporting import catalog_chart_payload, expand_dynamic_layouts
    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    expanded = expand_dynamic_layouts(catalogue, {'Operator': ['EE', 'O2']})
    entry = next(entry for entry in expanded if entry.chart_type == 'Histogram Bars' and entry.kpi == 'LTE_RSRP')
    frame = pd.DataFrame({'Operator': ['EE'] * 6 + ['O2'],
        'Campaign': ['2026-Q3', '2026-Q3', '2026-Q1', '2026-Q1', '2026-Q2', '2026-Q2', '2026-Q3'],
        'LTE_RSRP': [-95, -90, -105, -100, -100, -95, -75]})
    model = catalog_chart_payload(frame, entry)
    assert model['histogram_bars'] is True
    assert [item['name'] for item in model['series']] == ['2026-Q1', '2026-Q2', '2026-Q3']
    assert all(item['samples'] == 2 and sum(item['counts']) == 2 for item in model['series'])
    assert all(sum(item['bin_ratios']) == 1 for item in model['series'])
    from PIL import ImageColor
    channels = [ImageColor.getrgb(item['colour']) for item in model['series']]
    assert all(channels[0][i] >= channels[1][i] >= channels[2][i] for i in range(3))
    contour = next(row for row in catalogue if row.chart_type == 'Histogram Line' and row.kpi == 'LTE_RSRP')
    newest = next(item for item in catalog_chart_payload(frame, contour)['series'] if item['key'] == ['EE', '2026-Q3'])
    assert model['series'][-1]['colour'].lower() == newest['colour'].lower()


def test_legacy_dynamic_template_field_is_optional_for_static_rows_and_required_for_dynamic_layouts():
    import csv, io
    from src.modules.cdr_reporting import SINGLE_DYNAMIC_CATALOG_HEADERS, catalogue_csv, parse_catalog_csv
    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    static = [entry for entry in catalogue if entry.chart_type != 'Histogram Bars']
    reader = csv.DictReader(io.StringIO(catalogue_csv(static).decode()))
    new_rows = list(reader)
    headers = SINGLE_DYNAMIC_CATALOG_HEADERS
    output = io.StringIO(); writer = csv.DictWriter(output, headers, lineterminator='\n')
    writer.writeheader()
    writer.writerows({**{field: row[field] for field in headers[:-1]}, 'Dynamic Field': ''} for row in new_rows)
    restored = parse_catalog_csv(output.getvalue(), 'nsa')
    assert len(restored) == len(static)
    assert all(entry.dynamic_field == '' for entry in restored)
    from dataclasses import replace
    pair = [entry for entry in catalogue if entry.chart_type == 'Histogram Bars'][:2]
    missing = [replace(entry, dynamic_field='', dynamic_rows_field='', dynamic_columns_field='') for entry in pair]
    with pytest.raises(ValueError, match='requires Dynamic .* Field'):
        parse_catalog_csv(catalogue_csv(missing), 'nsa')


def test_multivendor_dynamic_histograms_page_by_vendor_and_keep_template_indexes():
    from dataclasses import replace
    from src.modules.cdr_reporting import expand_dynamic_layouts

    catalogue = app_module.load_template_catalogue(TEMPLATE_PATH.read_bytes(), 'nsa')
    base = [
        replace(entry, layout='2 rows + dynamic columns, comments down')
        for entry in catalogue if entry.chart_type == 'Histogram Bars'
    ][:2]
    template_indexes = set(range(len(base)))
    values = [
        'EE_Ericsson', 'O2_Ericsson', '3_Huawei', 'VF_Huawei',
        'A_NSN', 'B_NSN', 'C_Samsung',
    ]

    expanded = expand_dynamic_layouts(
        base, {'Operator_Vendor': values}, multivendor=True, vendor_comparison='operator_vendor',
    )

    assert {entry.dynamic_field for entry in expanded} == {'Operator_Vendor'}
    assert {entry.template_index for entry in expanded} == template_indexes
    pages = {}
    for entry in expanded:
        pages.setdefault(entry.slide, []).append(entry)
    assert sorted(len({entry.dynamic_value for entry in rows}) for rows in pages.values()) == [1, 6]
    assert sorted(len(rows) for rows in pages.values()) == [2, 12]
    page_by_vendor = {}
    for slide, rows in pages.items():
        for value in {entry.dynamic_value for entry in rows}:
            page_by_vendor.setdefault(value.rsplit('_', 1)[1], set()).add(slide)
    assert all(len(slides) == 1 for slides in page_by_vendor.values())
    for rows in pages.values():
        for index in template_indexes:
            assert sum(entry.template_index == index for entry in rows) == len({entry.dynamic_value for entry in rows})

    pooled = expand_dynamic_layouts(
        base, {'Vendor': ['Ericsson', 'Huawei']},
        multivendor=True, vendor_comparison='vendor_only',
    )
    assert {entry.dynamic_field for entry in pooled} == {'Vendor'}
    assert {entry.dynamic_value for entry in pooled} == {'Ericsson', 'Huawei'}


def test_grouped_cdf_curves_differ_by_line_style_and_lte_nr_stays_separate(client, tmp_path) -> None:
    _login(client)
    frames = [_data_rows('2026-Q1'), _data_rows('2026-Q2', rsrp_offset=5)]
    for frame in frames:
        frame['NR_PCell_RSRP_Avg'] = frame['LTE_PCell_RSRP_Avg'] - 8
    ids = [_add_ready_dataset(tmp_path, f'CDR_Data_NSA_{campaign}.xlsx', 'data', frame)
           for campaign, frame in zip(('2026-Q1', '2026-Q2'), frames, strict=True)]

    grouped = client.post('/api/network-insights/analysis', json={
        'datasets': {'data': ids}, 'technology': 'lte', 'group': ['operator', 'campaign'],
    }).json()
    # Operator is always grouped unless Vendor replaces it.
    assert client.post('/api/network-insights/analysis', json={
        'datasets': {'data': ids}, 'technology': 'lte', 'group': ['campaign'],
    }).json()['group'] == ['operator', 'campaign']
    pooled = client.post('/api/network-insights/analysis', json={
        'datasets': {'data': ids}, 'technology': 'lte', 'group': ['vendor'],
    }).json()
    assert pooled['group'] == ['vendor']
    # Campaigns keep the Operator colour: the latest is the thickest line.
    series = grouped['charts']['rsrp_cdf']['series']
    assert len(series) == 4 and all(not item['dash'] for item in series)
    widths = {item['name']: item['width'] for item in series}
    assert widths == {'EE · 2026-Q1': 1, 'EE · 2026-Q2': 4, 'VF · 2026-Q1': 1, 'VF · 2026-Q2': 4}
    assert len({item['colour'] for item in series}) == 2
    styles = {(item['colour'], str(item['dash'])) for item in client.post('/api/network-insights/analysis', json={
        'datasets': {'data': ids}, 'technology': 'lte', 'group': ['operator', 'city'],
    }).json()['charts']['rsrp_cdf']['series']}
    # Other groups (EE in Leeds, VF in York) differ by line style.
    assert len({dash for _colour, dash in styles}) == len(styles) == 2

    # LTE and NR are never pooled: LTE+NR returns one section per technology
    # with its own thresholds, CDFs, overview and maps.
    combined = client.post('/api/network-insights/analysis', json={
        'datasets': {'data': ids}, 'technology': 'lte_nr', 'group': ['operator'],
        'coverage_threshold': -100, 'nr_coverage_threshold': -90,
    }).json()
    assert combined['group'] == ['operator']
    lte, nr = combined['sections']
    assert (lte['technology'], nr['technology']) == ('lte', 'nr')
    assert (lte['coverage_threshold'], nr['coverage_threshold']) == (-100, -90)
    assert nr['interference_threshold'] == ni.DEFAULT_NR_INTERFERENCE_THRESHOLD
    assert lte['charts']['rsrp_cdf']['title'] == 'LTE RSRP' and nr['charts']['rsrp_cdf']['title'] == 'NR RSRP'
    assert lte['maps']['coverage']['type'] == nr['maps']['coverage']['type'] == 'map'
    lte_ee = next(row for row in lte['overview'] if row['operator'] == 'EE')
    nr_ee = next(row for row in nr['overview'] if row['operator'] == 'EE')
    assert nr_ee['rsrp_median'] == lte_ee['rsrp_median'] - 8
    assert lte_ee['technology'] == 'LTE' and nr_ee['technology'] == 'NR'
    assert len(combined['overview']) == len(lte['overview']) + len(nr['overview'])
    assert sorted(series['name'] for series in nr['charts']['rsrp_cdf']['series']) == ['EE', 'VF']


@pytest.mark.parametrize('kind', ['mapping_vodafone', 'mapping_three'])
def test_full_inventory_pages_and_csv_preserve_all_records(client, tmp_path, kind) -> None:
    _login(client)
    rows = pd.DataFrame({
        'Site ID': ['S1'] * 102 + [None],
        'GCID': list(range(103)),
        'source_sheet': ['4G'] * 50 + ['5G'] * 53,
        'Additional attribute': ['Quoted "value", with comma\nand newline'] * 103,
        'Coordinates': [51.1234567890123] * 103,
    })
    dataset_id = _add_ready_dataset(tmp_path, f'{kind}.xlsx', kind, rows, nr_mode=None)
    payload = client.get('/api/network-insights/deployment?group=inventory').json()
    inventory = payload['inventories'][0]
    assert inventory['columns'] == list(rows.columns)
    assert inventory['total_rows'] == 103
    assert len(inventory['rows']) == 50
    last = client.get(f'/api/network-insights/inventory/{dataset_id}?page=2').json()
    assert last['page'] == 2 and len(last['rows']) == 3
    assert last['rows'][-1][0] is None
    assert client.get(f'/api/network-insights/inventory/{dataset_id}?page=99').json()['page'] == 2
    assert client.get(f'/api/network-insights/inventory/{dataset_id}?page=-1').status_code == 422
    response = client.get(f'/api/network-insights/inventory/{dataset_id}?download=true')
    assert response.status_code == 200 and 'text/csv' in response.headers['content-type']
    exported = list(csv.reader(io.StringIO(response.text)))
    assert exported[0] == list(rows.columns)
    assert len(exported) == 104
    assert exported[1][3] == rows.iloc[0]['Additional attribute']
    assert exported[-1][0] == '' and exported[-1][1] == '102'
    assert float(exported[1][4]) == rows.iloc[0]['Coordinates']
    page = client.get('/network-insights').text
    assert 'Full Sites/Cells Inventory' in page
    assert page.index('class="module-tabs-secondary"') < page.index('class="module-tabs-primary"')


@pytest.mark.parametrize('group', list(ni.DEPLOYMENT_GROUPINGS))
def test_deployment_csv_matches_grouped_table_and_totals_fallback(client, tmp_path, group) -> None:
    _login(client)
    if group == 'cluster':
        _add_ready_polygon(tmp_path, 'clusters', 'Polygon Cluster')
    rows = pd.DataFrame({
        'Site_ID': ['S1', 'S1', 'S2', 'S3'],
        'CId___ECI': ['1', '2', '3', '4'],
        'eMOCNScenario': ['NNS', 'NNS', 'S1', ''],
        'Vendor': ['Ericsson', 'Ericsson', 'Nokia', 'Nokia'],
        'Host Network': ['VF', 'VF', '3UK', '3UK'],
    })
    dataset_id = _add_ready_dataset(tmp_path, 'inventory.xlsx', 'mapping_three', rows, nr_mode=None)
    table = client.get(f'/api/network-insights/deployment?group={group}').json()['inventories'][0]
    response = client.get(f'/api/network-insights/deployment/{dataset_id}/export?group={group}')
    assert response.status_code == 200
    exported = list(csv.reader(io.StringIO(response.text)))
    assert exported[0] == [ni.DEPLOYMENT_GROUPINGS[group] if group in table['available_groups'] else 'Inventory', 'Sites', 'Cells']
    assert exported[1:] == [[row['group'], str(row['sites']), str(row['cells'])] for row in table['rows']]


def test_inventory_exports_reject_non_inventory_unready_and_missing_sources(client, tmp_path) -> None:
    _login(client)
    rows = pd.DataFrame({'Site_ID': ['S1']})
    cdr_id = _add_ready_dataset(tmp_path, 'cdr.xlsx', 'data', rows)
    pending_id = _add_ready_dataset(tmp_path, 'pending.xlsx', 'mapping_three', rows, nr_mode=None)
    app_module.repository.update_dataset_profile(pending_id, status='pending')
    for dataset_id in [cdr_id, pending_id, 999999]:
        assert client.get(f'/api/network-insights/inventory/{dataset_id}').status_code == 404
        assert client.get(f'/api/network-insights/inventory/{dataset_id}?download=true').status_code == 404
        assert client.get(f'/api/network-insights/deployment/{dataset_id}/export').status_code == 404
    assert client.get(f'/api/network-insights/deployment/{pending_id}/export?group=invalid').status_code == 400
    client.get('/logout')
    assert client.get(f'/api/network-insights/inventory/{pending_id}?download=true', follow_redirects=False).status_code != 200
    assert client.get(f'/api/network-insights/deployment/{pending_id}/export', follow_redirects=False).status_code != 200


@pytest.mark.parametrize('path', [
    '/api/network-insights/deployment?group=inventory',
    '/api/network-insights/inventory/1?download=true',
    '/api/network-insights/deployment/1/export',
])
def test_inventory_endpoints_require_workspace_access(client, monkeypatch, path) -> None:
    _login(client)
    monkeypatch.setattr(app_module.repository, 'user_has_workspace_access', lambda *_args: False)
    assert client.get(path).status_code == 403
    monkeypatch.setattr(app_module, 'active_workspace', None)
    assert client.get(path).status_code == 400


def test_module_tab_rows_remain_siblings_outside_configuration_popover(client) -> None:
    from html.parser import HTMLParser

    class NavigationParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.stack = []
            self.parents = {}

        def handle_starttag(self, tag, attrs):
            values = dict(attrs)
            classes = values.get('class', '').split()
            for row in ['module-tabs-secondary', 'module-tabs-primary']:
                if row in classes:
                    self.parents[row] = list(self.stack)
            if tag not in {'area', 'base', 'br', 'col', 'embed', 'hr', 'img', 'input', 'link', 'meta', 'param', 'source', 'track', 'wbr'}:
                self.stack.append((tag, values.get('id'), values.get('class')))

        def handle_endtag(self, tag):
            for index in range(len(self.stack) - 1, -1, -1):
                if self.stack[index][0] == tag:
                    del self.stack[index:]
                    break

    _login(client)
    parser = NavigationParser()
    parser.feed(client.get('/network-insights').text)
    assert list(parser.parents) == ['module-tabs-secondary', 'module-tabs-primary']
    assert parser.parents['module-tabs-primary'] == parser.parents['module-tabs-secondary']
    assert parser.parents['module-tabs-primary'][-1] == ('nav', None, 'module-tabs')


def _combined_inventory_sources(tmp_path):
    vf = pd.DataFrame({
        'Site_ID': ['VF1'] * 55 + ['VF5G'], 'GCID': list(range(100, 155)) + [555],
        'OP/ Vendor': ['VF_Ericsson'] * 56, 'Region': ['North'] * 56,
        'City': ['Leeds'] * 56, 'Cluster': ['Cluster A'] * 56,
        'Source_Sheet': ['4G'] * 55 + ['5G'], 'Extra VF': ['keep "all",\nfields'] * 56,
    })
    three = pd.DataFrame({
        'Site_ID': ['31', '32'], 'Cid__ECI': [100, 999],
        'Vendor': ['3_Nokia', '3_Nokia'], 'Vendor_Only': ['Nokia', 'Nokia'],
        'Region': ['South', 'South'], 'City': ['London', 'London'], 'Technology': ['4G', '4G'],
        'Extra Three': [1, 2],
    })
    _add_ready_dataset(tmp_path, 'vf-full.xlsx', 'mapping_vodafone', vf, nr_mode=None)
    _add_ready_dataset(tmp_path, 'three-full.csv', 'mapping_three', three, nr_mode=None)
    cdr = pd.DataFrame({
        'Operator': ['Vodafone UK'] * 55 + ['Three', 'Three', 'Vodafone UK', 'Vodafone UK'],
        'Cell_ID': list(range(100, 155)) + [100, 999, 555, 100],
        'Campaign': ['2026 Q1'] * 56 + ['2026 Q2', '2026 Q1', '2026 Q1'],
        'Technology': ['LTE'] * 57 + ['NR', 'LTE'],
    })
    data_id = _add_ready_dataset(tmp_path, 'selected.xlsx', 'data', cdr)
    sa_id = _add_ready_dataset(tmp_path, 'sa.xlsx', 'data', cdr.iloc[[57]], nr_mode='SA')
    return {'datasets': {'data': [data_id, sa_id]}, 'nr_mode': 'NSA', 'technology': 'lte_nr', 'campaigns': ['2026-Q1']}


def test_combined_inventory_preserves_sources_and_exports_filtered_pages(client, tmp_path):
    _login(client)
    selection = _combined_inventory_sources(tmp_path)
    response = client.post('/api/network-insights/inventory', json=selection)
    assert response.status_code == 200
    first = response.json()
    assert first['columns'][:7] == list(ni.INVENTORY_KEY_COLUMNS)
    assert first['key_columns'] == list(ni.INVENTORY_KEY_COLUMNS)
    assert first['total_rows'] == 58 and len(first['rows']) == 50
    assert first['rows'][0][:7] == ['VF', 'VF_Ericsson', 'Ericsson', 'North', 'Cluster A', 'Leeds', 'LTE']
    second = client.post('/api/network-insights/inventory?page=1', json=selection).json()
    assert len(second['rows']) == 8
    assert second['rows'][-1][:7] == ['3', '3_Nokia', 'Nokia', 'South', '', 'London', 'LTE']
    response = client.post('/api/network-insights/inventory/export', json=selection)
    exported = list(csv.reader(io.StringIO(response.text)))
    assert exported[0] == first['columns'] and len(exported) == 59
    assert exported[1][first['columns'].index('Extra VF')] == 'keep "all",\nfields'
    assert exported[-1][first['columns'].index('Source_Vendor')] == '3_Nokia'
    assert exported[-1][first['columns'].index('Source_Technology')] == '4G'
    assert exported[-1][first['columns'].index('Extra VF')] == ''
    assert {row[first['columns'].index('Source_Dataset_Name')] for row in exported[1:]} == {'vf-full.xlsx', 'three-full.csv'}
    assert '__Inventory_Cell_ID' not in first['columns']
    assert client.post('/api/network-insights/inventory?page=99', json=selection).json()['page'] == 1
    assert client.post('/api/network-insights/inventory?page=-1', json=selection).status_code == 422


@pytest.mark.parametrize('filters, expected', [
    ({'operators': ['Vodafone UK']}, 56),
    ({'operators': ['3']}, 2),
    ({'vendors': ['Ericsson']}, 56),
    ({'vendors': ['3_Nokia']}, 2),
    ({'regions': ['south']}, 2),
    ({'cities': ['London']}, 2),
    ({'technology': 'nr'}, 1),
    ({'technology': 'lte'}, 57),
    ({'campaigns': ['2026-Q2']}, 58),
    ({'nr_mode': 'SA', 'technology': 'nr'}, 1),
    ({'operators': ['EE']}, 0),
    ({'datasets': {}}, 58),
])
def test_complete_inventory_filters_mapping_attributes_independently_of_cdrs(client, tmp_path, filters, expected):
    _login(client)
    selection = {**_combined_inventory_sources(tmp_path), **filters}
    response = client.post('/api/network-insights/inventory', json=selection)
    assert response.status_code == 200
    payload = response.json()
    assert payload['total_rows'] == expected
    exported = list(csv.reader(io.StringIO(client.post('/api/network-insights/inventory/export', json=selection).text)))
    assert len(exported) == expected + 1
    if filters.get('technology') == 'nr' and payload['rows']:
        assert all(row[6] == 'NR' for row in payload['rows'])


def test_combined_inventory_endpoints_require_workspace_access(client, monkeypatch):
    _login(client)
    monkeypatch.setattr(app_module.repository, 'user_has_workspace_access', lambda *_args: False)
    for path in ['/api/network-insights/inventory', '/api/network-insights/inventory/export']:
        assert client.post(path, json={}).status_code == 403
    monkeypatch.setattr(app_module, 'active_workspace', None)
    for path in ['/api/network-insights/inventory', '/api/network-insights/inventory/export']:
        assert client.post(path, json={}).status_code == 400


def test_required_aggregation_choices_stay_selected_without_closing_menu():
    import shutil
    import subprocess

    node = shutil.which('node')
    if not node:
        pytest.skip('Node.js is unavailable.')
    source = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/network_insights.js').read_text()
    grouping_code = source[source.index('  const syncTechnologyGrouping ='):source.index('  const kinds =')]
    script = r'''
const assert = require('node:assert/strict');
class Choice {
  constructor(text, value, defaultSelected = false, selected = false) {
    this.text = text; this.value = value; this.selected = selected;
    this.disabled = false; this.dataset = {};
  }
  remove() { group.options.splice(group.options.indexOf(this), 1); }
}
global.Option = Choice;
const events = [];
const listeners = [];
const group = {
  options: [new Choice('Operator', 'operator', true, true), new Choice('Vendor', 'vendor'), new Choice('Campaign', 'campaign', true, true)],
  add(option, before) { this.options.splice(before ? this.options.indexOf(before) : this.options.length, 0, option); },
  addEventListener(name, callback) { if (name === 'change') listeners.push(callback); },
  dispatchEvent(event) { events.push(event.type); if (event.type === 'change') listeners.forEach(callback => callback()); },
};
const technology = {value: 'lte'};
const thresholds = [{dataset: {niRadio: 'lte'}, hidden: false}, {dataset: {niRadio: 'nr'}, hidden: false}];
global.document = {querySelectorAll: () => thresholds};
const $ = id => id === 'ni-group' ? group : technology;
'''
    assertions = r'''
const choice = value => group.options.find(option => option.value === value);
assert.equal(choice('operator').disabled, true);
assert.equal(choice('operator').selected, true);
events.length = 0;
choice('vendor').selected = true;
group.dispatchEvent(new Event('change'));
assert.equal(choice('operator').disabled, false);
assert.equal(events.includes('multiselect:options-updated'), false);
choice('operator').selected = false;
choice('vendor').selected = false;
group.dispatchEvent(new Event('change'));
assert.equal(choice('operator').disabled, true);
assert.equal(choice('operator').selected, true);
assert.equal(events.includes('multiselect:options-updated'), false);
group.options.push(new Choice('Technology', 'technology'));
technology.value = 'lte_nr'; syncTechnologyGrouping();
assert.equal(choice('technology'), undefined);
assert.deepEqual(thresholds.map(item => item.hidden), [false, false]);
group.options.filter(option => !option.disabled).forEach(option => {option.selected = false;});
group.dispatchEvent(new Event('change'));
assert.equal(choice('operator').selected, true);
technology.value = 'nr'; syncTechnologyGrouping();
assert.deepEqual(thresholds.map(item => item.hidden), [true, false]);
technology.value = 'lte'; syncTechnologyGrouping();
assert.deepEqual(thresholds.map(item => item.hidden), [false, true]);
'''
    subprocess.run([node, '-e', script + grouping_code + assertions], check=True, capture_output=True, text=True)


def test_full_inventory_tables_and_individual_exports_are_scoped_by_operator(client, tmp_path):
    _login(client)
    selection = _combined_inventory_sources(tmp_path)
    response = client.post('/api/network-insights/inventory/tables', json=selection)
    assert response.status_code == 200
    tables = response.json()['inventories']
    assert {table['operator'] for table in tables} == {'VF', '3'}
    assert {table['operator_label'] for table in tables} == {'Vodafone', 'Three'}
    assert {table['operator']: table['total_rows'] for table in tables} == {'VF': 56, '3': 2}
    for table in tables:
        assert all(row[0] == table['operator'] for row in table['rows'])
        response = client.post(f"/api/network-insights/inventory/export?operator={table['operator']}", json=selection)
        exported = list(csv.reader(io.StringIO(response.text)))
        assert len(exported) == table['total_rows'] + 1
        assert all(row[0] == table['operator'] for row in exported[1:])
    page = client.post('/api/network-insights/inventory?operator=VF&page=1', json=selection).json()
    assert page['total_rows'] == 56 and len(page['rows']) == 6 and page['page'] == 1
    filtered = client.post('/api/network-insights/inventory/tables', json={**selection, 'operators': ['3']}).json()
    assert [row['operator'] for row in filtered['inventories']] == ['3']
    # A table-specific export cannot widen the general Operator filter.
    response = client.post('/api/network-insights/inventory/export?operator=VF', json={**selection, 'operators': ['3']})
    assert len(list(csv.reader(io.StringIO(response.text)))) == 1
    page = client.get('/network-insights').text
    assert 'Export All to CSV' in page


@pytest.mark.parametrize('group', list(ni.DEPLOYMENT_GROUPINGS))
def test_export_all_deployment_csv_combines_every_displayed_inventory(client, tmp_path, group):
    _login(client)
    if group == 'cluster':
        _add_ready_polygon(tmp_path, 'clusters', 'Polygon Cluster')
    rows = pd.DataFrame({'Site_ID': ['S1', 'S1', 'S2'], 'CId___ECI': ['1', '2', '3'],
                         'eMOCNScenario': ['NNS', 'NNS', 'S1'], 'Vendor': ['Ericsson', 'Ericsson', 'Nokia']})
    _add_ready_dataset(tmp_path, 'vf-inventory.csv', 'mapping_vodafone', rows, nr_mode=None)
    _add_ready_dataset(tmp_path, 'three-inventory.csv', 'mapping_three', rows, nr_mode=None)
    displayed = client.get(f'/api/network-insights/deployment?group={group}').json()['inventories']
    response = client.get(f'/api/network-insights/deployment/export-all?group={group}')
    assert response.status_code == 200
    exported = list(csv.reader(io.StringIO(response.text)))
    assert exported[0] == ['Operator', 'Source_Dataset_ID', 'Source_Dataset_Name', ni.DEPLOYMENT_GROUPINGS[group], 'Sites', 'Cells']
    expected = [[table['operator'], str(table['id']), table['file_name'], row['group'], str(row['sites']), str(row['cells'])]
                for table in displayed for row in table['rows']]
    assert exported[1:] == expected
    assert {row[0] for row in exported[1:]} == {'Vodafone', 'Three'}
    assert client.get('/api/network-insights/deployment/export-all?group=invalid').status_code == 400


def test_full_inventory_uses_real_mapping_geography_headers_before_filtering(client, tmp_path):
    _login(client)
    vf = pd.DataFrame({'Site ID': ['V1'], 'GCID': [100], 'source_sheet': ['4G'],
                       'OP/ Vendor': ['Ericsson'], 'Region': ['North'], 'City': [''],
                       'Beacon2Town': ['Leeds'], 'Cluster': [''], 'Engineering Polygon': ['VF Cluster']})
    three = pd.DataFrame({'Site_ID': ['T1'], 'CId___ECI': [200], 'Vendor': ['Ericsson'],
                          'Region': [''], 'UK "Regional"': ['North'], 'City': [''], 'Town': ['Leeds'],
                          'Cluster': [''], 'Regional Optimisation Polygon': ['Three Cluster']})
    _add_ready_dataset(tmp_path, 'vf-real-headers.xlsx', 'mapping_vodafone', vf, nr_mode=None)
    _add_ready_dataset(tmp_path, 'three-real-headers.csv', 'mapping_three', three, nr_mode=None)
    cdr_id = _add_ready_dataset(tmp_path, 'cdr.xlsx', 'data', pd.DataFrame({
        'Operator': ['VF', '3'], 'Cell_ID': [100, 200], 'Campaign': ['2026-Q1', '2026-Q1'], 'Technology': ['LTE', 'LTE'],
    }))
    selection = {'datasets': {'data': [cdr_id]}, 'technology': 'lte', 'cities': ['Leeds'], 'regions': ['North']}
    response = client.post('/api/network-insights/inventory/tables', json=selection)
    assert response.status_code == 200
    tables = response.json()['inventories']
    assert {table['operator']: table['total_rows'] for table in tables} == {'VF': 1, '3': 1}
    assert {row[4] for table in tables for row in table['rows']} == {'VF Cluster', 'Three Cluster'}
    assert all(row[3] == 'North' and row[5:7] == ['Leeds', 'LTE'] for table in tables for row in table['rows'])
    exported = list(csv.reader(io.StringIO(client.post('/api/network-insights/inventory/export', json=selection).text)))
    assert len(exported) == 3
    grouped = client.get('/api/network-insights/deployment?group=region').json()['inventories']
    assert all({row['group'] for row in table['rows']} == {'North'} for table in grouped)


def test_cluster_inventory_counts_deduplicate_uploads_and_do_not_require_cdrs(client, tmp_path):
    _login(client)
    for name, kind, rows in [
        ('vf-one.xlsx', 'mapping_vodafone', {'Site_ID': ['S1', 'S1'], 'Cell ID': ['C1', 'C2']}),
        ('vf-two.xlsx', 'mapping_vodafone', {'Site_ID': ['S1', 'S2'], 'Cell ID': ['C2', 'C3']}),
        ('three.csv', 'mapping_three', {'MBNL_ID': ['T1'], 'Cell ID': ['T01']}),
    ]:
        _add_ready_dataset(tmp_path, name, kind, pd.DataFrame(rows), None)
    response = client.get('/api/network-insights/cluster-inventories')
    assert response.status_code == 200
    assert response.json()['rows'] == [{'operator': 'Three', 'sites': 1, 'cells': 1},
                                       {'operator': 'Vodafone', 'sites': 2, 'cells': 3}]
    html = client.get('/network-insights').text
    assert 'id="ni-sites-source"' in html
    assert html.index('id="ni-sites-pending"') < html.index('Cluster polygons required.')


def test_combined_inventory_export_prepares_all_rows_before_delivery(client, tmp_path):
    _login(client)
    selection = _combined_inventory_sources(tmp_path)
    _add_ready_dataset(tmp_path, 'three-large.csv', 'mapping_three', pd.DataFrame({
        'Site_ID': ['31'] * 10050, 'Cid__ECI': [100] * 10050,
        'Technology': ['LTE'] * 10050,
    }), nr_mode=None)
    response = client.post('/api/network-insights/inventory/export', json=selection)
    assert response.status_code == 200
    assert int(response.headers['content-length']) == len(response.content)
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert len(rows) == 10108
    assert sum(row['Operator'] == 'VF' for row in rows) == 56
    assert sum(row['Operator'] == '3' for row in rows) == 10052
    assert {row['Source_Dataset_Name'] for row in rows} == {
        'vf-full.xlsx', 'three-full.csv', 'three-large.csv'}


def test_cluster_and_region_prefer_polygons_and_fallback_to_inventory(client, tmp_path):
    _login(client)
    selection = _combined_inventory_sources(tmp_path)
    vf = next(row for row in app_module.repository.list_datasets() if row['dataset_kind'] == 'mapping_vodafone')
    app_module.repository.replace_dataset_rows(vf['id'], pd.DataFrame({
        'Site_ID': ['S1', 'S1', 'S2'], 'GCID': [100, 101, 102], 'Technology': ['LTE'] * 3,
        'Region': ['Source Region'] * 3, 'Cluster': ['Source Cluster'] * 3,
        'Longitude': [-1, -1, 10], 'Latitude': [54, 54, 10],
    }))
    # Without boundary inputs, mapping labels remain available, including Cluster.
    for group, name in [('region', 'Source Region'), ('cluster', 'Source Cluster')]:
        response = client.get(f'/api/network-insights/deployment?group={group}')
        table = next(row for row in response.json()['inventories'] if row['operator'] == 'Vodafone')
        assert table['rows'] == [{'group': name, 'sites': 2, 'cells': 3}]
    page = client.get('/network-insights').text
    assert 'value="cluster" disabled' not in page
    _add_ready_polygon(tmp_path, 'mapping_region', 'Polygon Region')
    _add_ready_polygon(tmp_path, 'clusters', 'Polygon Cluster')
    for group, name in [('region', 'Polygon Region'), ('cluster', 'Polygon Cluster')]:
        response = client.get(f'/api/network-insights/deployment?group={group}')
        table = next(row for row in response.json()['inventories'] if row['operator'] == 'Vodafone')
        assert {row['group']: (row['sites'], row['cells']) for row in table['rows']} == {
            name: (1, 2), 'Not set': (1, 1)}
        csv_response = client.get(f"/api/network-insights/deployment/{vf['id']}/export?group={group}")
        assert f'{name},1,2' in csv_response.text
    inventory = client.post('/api/network-insights/inventory/tables', json={**selection, 'regions': ['Polygon Region']}).json()
    table = next(row for row in inventory['inventories'] if row['operator'] == 'VF')
    assert table['total_rows'] == 2
    assert all(row[3] == 'Polygon Region' and row[4] == 'Polygon Cluster' for row in table['rows'])
    exported = client.post('/api/network-insights/inventory/export', json={**selection, 'regions': ['Polygon Region']})
    records = list(csv.DictReader(io.StringIO(exported.text)))
    assert len(records) == 2
    assert all(row['Region'] == 'Polygon Region' and row['Cluster'] == 'Polygon Cluster' for row in records)
    assert all(row['Source_Region'] == 'Source Region' and row['Source_Cluster'] == 'Source Cluster' for row in records)


def test_cluster_grouping_disabled_without_polygons_or_inventory_cluster_fields(client, tmp_path):
    _login(client)
    _add_ready_dataset(tmp_path, 'no-cluster.csv', 'mapping_three', pd.DataFrame({'Site_ID': ['S1']}), None)
    assert 'value="cluster" disabled' in client.get('/network-insights').text
    for path in ['/api/network-insights/deployment?group=cluster',
                 '/api/network-insights/deployment/export-all?group=cluster']:
        assert client.get(path).status_code == 400


def test_full_inventory_without_cdrs_includes_unobserved_cells_and_both_operator_exports(client, tmp_path):
    _login(client)
    _combined_inventory_sources(tmp_path)
    selection = {'datasets': {}, 'technology': 'lte_nr', 'campaigns': ['No matching campaign']}
    tables = client.post('/api/network-insights/inventory/tables', json=selection).json()['inventories']
    assert {table['operator']: table['total_rows'] for table in tables} == {'VF': 56, '3': 2}
    exported = client.post('/api/network-insights/inventory/export', json=selection)
    rows = list(csv.DictReader(io.StringIO(exported.text)))
    assert len(rows) == 58
    assert {row['Operator'] for row in rows} == {'VF', '3'}
    assert any(row['Operator'] == '3' and row['Cid__ECI'] == '999' for row in rows)


def test_pending_inventory_names_follow_actual_filtered_operators_not_group_labels(client, tmp_path):
    _login(client)
    cdr = _add_ready_dataset(tmp_path, 'analysis.xlsx', 'data', _data_rows('2026-Q1'))
    _add_ready_dataset(tmp_path, 'vf-inventory.csv', 'mapping_vodafone', pd.DataFrame({
        'Site_ID': ['V1'], 'CId___ECI': [30001922],
    }), None)
    selection = {'datasets': {'data': [cdr]}, 'technology': 'lte', 'group': ['operator', 'city', 'campaign']}
    response = client.post('/api/network-insights/analysis', json=selection)
    assert response.status_code == 200
    assert response.json()['missing_inventory_operators'] == ['EE']
    response = client.post('/api/network-insights/analysis', json={**selection, 'operators': ['VF']})
    assert response.status_code == 200
    assert response.json()['missing_inventory_operators'] == []
    response = client.post('/api/network-insights/analysis', json={**selection, 'operators': ['EE']})
    assert response.json()['missing_inventory_operators'] == ['EE']


def test_site_tables_filter_columns_excel_style_and_export_with_operator_suffix(client, tmp_path):
    _login(client)
    selection = _combined_inventory_sources(tmp_path)
    tables = client.post('/api/network-insights/sites/tables', json={**selection, 'source': 'inventory'}).json()['tables']
    assert {table['operator']: table['total_rows'] for table in tables} == {'VF': 56, '3': 2}
    request = {**selection, 'source': 'inventory', 'operator': 'VF', 'column': 'Site_ID'}
    values = client.post('/api/network-insights/sites/values', json=request).json()['values']
    assert {row['value']: row['count'] for row in values} == {'VF1': 55, 'VF5G': 1}
    page = client.post('/api/network-insights/sites/page', json={**request, 'filters': {'Site_ID': ['VF5G']}}).json()
    assert page['total_rows'] == 1 and page['filters'] == {'Site_ID': ['VF5G']}
    response = client.post('/api/network-insights/sites/export', json={**request, 'filters': {'Site_ID': ['VF5G']}})
    assert 'site-cell-inventory_VF.csv' in response.headers['content-disposition']
    assert len(list(csv.reader(io.StringIO(response.text)))) == 2
    response = client.post('/api/network-insights/sites/export', json={**selection, 'source': 'inventory', 'operator': '3'})
    assert 'site-cell-inventory_3.csv' in response.headers['content-disposition']
    response = client.post('/api/network-insights/sites/export', json={**selection, 'source': 'inventory'})
    assert 'site-cell-inventory_VF_3.csv' in response.headers['content-disposition']
    assert client.post('/api/network-insights/sites/values', json={**request, 'column': 'Missing'}).status_code == 400


def test_observed_site_tables_aggregate_cdr_cells_per_operator(client, tmp_path):
    _login(client)
    cdr = _add_ready_dataset(tmp_path, 'analysis.xlsx', 'data', _data_rows('2026-Q1'))
    selection = {'datasets': {'data': [cdr]}, 'technology': 'lte', 'source': 'observed'}
    response = client.post('/api/network-insights/sites/tables', json=selection)
    assert response.status_code == 200
    tables = response.json()['tables']
    assert tables and all(table['source'] == 'observed' for table in tables)
    assert tables[0]['columns'] == list(ni.OBSERVED_COLUMNS)
    assert sum(row[tables[0]['columns'].index('Samples')] for table in tables for row in table['rows']) > 0
    vf = next(table for table in tables if table['operator'] == 'VF')
    exported = client.post('/api/network-insights/sites/export', json={**selection, 'operator': 'VF'})
    assert 'cdr-observed-sites-cells_VF.csv' in exported.headers['content-disposition']
    assert len(list(csv.reader(io.StringIO(exported.text)))) == vf['total_rows'] + 1


def test_repeated_analysis_reuses_cached_result_until_cdrs_change(client, tmp_path, monkeypatch):
    _login(client)
    cdr = _add_ready_dataset(tmp_path, 'analysis.xlsx', 'data', _data_rows('2026-Q1'))
    selection = {'datasets': {'data': [cdr]}, 'technology': 'lte'}
    first = client.post('/api/network-insights/analysis', json=selection)
    assert first.status_code == 200
    calls = []
    original = ni.normalise_samples
    monkeypatch.setattr(ni, 'normalise_samples', lambda *args, **kwargs: calls.append(1) or original(*args, **kwargs))
    second = client.post('/api/network-insights/analysis', json=selection)
    assert second.json() == first.json() and calls == []
    app_module.repository.replace_dataset_rows(cdr, _data_rows('2026-Q1', rsrp_offset=-20))
    app_module.repository.update_dataset_profile(cdr, processed_at=local_now_iso())
    third = client.post('/api/network-insights/analysis', json=selection)
    assert third.status_code == 200 and calls


def test_lte_nr_summary_export_has_separate_technology_sections(client, tmp_path):
    from docx import Document

    _login(client)
    rows = _data_rows('2026-Q1')
    rows['NR_PCell_RSRP_Avg'] = rows['LTE_PCell_RSRP_Avg'] - 8
    cdr = _add_ready_dataset(tmp_path, 'analysis.xlsx', 'data', rows)
    response = client.post('/api/network-insights/export/word', json={
        'datasets': {'data': [cdr]}, 'technology': 'lte_nr', 'nr_coverage_threshold': -120})
    assert response.status_code == 200
    document = Document(io.BytesIO(response.content))
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith('Heading')]
    assert 'Overview · LTE' in headings and 'Overview · NR' in headings
    # One maps section per technology and group.
    assert any(heading.startswith('Coverage & Interference Maps · NR · ') for heading in headings)
    assert 'RF Quality · LTE' in headings and 'Overview · NR' in headings
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert 'NR thresholds: low coverage below -120.0 dBm RSRP' in text
