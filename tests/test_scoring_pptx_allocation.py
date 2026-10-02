from io import BytesIO
import math
from types import SimpleNamespace

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches
import pytest

from src.modules.scoring_pptx_allocation import (
    _add_percentage_labels,
    add_maximum_allocation_donut,
    category_maximum_allocations,
    maximum_allocations_from_configuration,
)


def _configuration():
    return {
        'scope': {'environments': {
            'DriveCity': {'total_points': 65},
            'DriveRoad': {'total_points': 35},
            'Unused': {'total_points': 0},
        }},
        'metrics': [
            {'source_kind': 'voice', 'contexts': {
                'DriveCity': {'max_points': 10}, 'DriveRoad': {'max_points': 4},
            }},
            {'calculation': {'source_kind': 'speech'}, 'contexts': {
                'DriveCity': {'max_points': 2}, 'DriveRoad': {'max_points': 1},
            }},
            {'source_kind': 'data', 'contexts': {
                'DriveCity': {'max_points': 20}, 'DriveRoad': {'max_points': 9},
            }},
        ],
    }


def _assert_percentage_glyph_fits(label, chart_shape, hole, share, start_angle):
    font_size = label.text_frame.paragraphs[0].runs[0].font.size.pt
    text_width = sum(.556 if character.isdigit() else
                     .278 if character == '.' else .889
                     for character in label.text_frame.text) * font_size / 72
    text_height = font_size * 1.2 / 72
    center_x = chart_shape.left.inches + chart_shape.width.inches / 2
    center_y = chart_shape.top.inches + chart_shape.height.inches / 2
    label_x = label.left.inches + label.width.inches / 2
    label_y = label.top.inches + label.height.inches / 2
    outer_radius = chart_shape.width.inches * .45
    inner_radius = outer_radius * hole / 100
    middle_radius = (outer_radius + inner_radius) / 2
    midpoint = start_angle + share * math.pi
    expected_x = center_x + middle_radius * math.sin(midpoint)
    expected_y = center_y - middle_radius * math.cos(midpoint)
    assert label_x == pytest.approx(expected_x, abs=.02)
    assert label_y == pytest.approx(expected_y, abs=.02)
    padding = 1 / 72
    for offset_x in (-text_width / 2, text_width / 2):
        for offset_y in (-text_height / 2, text_height / 2):
            x = label_x + offset_x - center_x
            y = -(label_y + offset_y - center_y)
            radius = math.hypot(x, y)
            angle = math.atan2(x, y)
            angle_error = abs(math.atan2(math.sin(angle - midpoint), math.cos(angle - midpoint)))
            assert inner_radius + padding <= radius <= outer_radius - padding
            assert angle_error <= share * math.pi - padding / middle_radius


@pytest.mark.parametrize('size, values, expected_labels', [
    (1.909, [35, 65], ['35.0%', '65.0%']),
    (1.733, [20.3, 13.3, 1.4, 35.1, 13, 13, 3.9],
     ['20.3%', '35.1%', '13.0%', '13.0%']),
])
def test_inner_percentage_labels_fit_native_ring_after_save(size, values, expected_labels):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_percentage_labels(slide, values, ['4472C4'] * len(values), 1, 1, size, 62)
    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    labels = list(Presentation(output).slides[0].shapes)
    assert [label.text for label in labels] == expected_labels
    ring = SimpleNamespace(left=Inches(1), top=Inches(1),
                           width=Inches(size), height=Inches(size))
    total = sum(values)
    label_index = 0
    for index, value in enumerate(values):
        if label_index >= len(labels):
            break
        if labels[label_index].text == f'{value / total * 100:.1f}%':
            _assert_percentage_glyph_fits(
                labels[label_index], ring, 62, value / total,
                sum(values[:index]) / total * 2 * math.pi,
            )
            label_index += 1
    assert label_index == len(labels)


def test_configuration_allocations_sum_metrics_once_per_positive_environment():
    assert maximum_allocations_from_configuration(_configuration()) == {
        'DriveCity': {'voice': 12.0, 'data': 20.0, 'max': 32.0},
        'DriveRoad': {'voice': 5.0, 'data': 9.0, 'max': 14.0},
    }
    assert maximum_allocations_from_configuration(None) == {}


def test_combined_allocation_has_environment_outer_ring_and_one_family_ring():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    allocations = maximum_allocations_from_configuration(_configuration())

    chart = add_maximum_allocation_donut(
        slide, {'context': {'environment': 'Combined'}}, allocations,
        left=1, top=1, width=4, height=3,
    )

    assert chart.chart_type == XL_CHART_TYPE.DOUGHNUT
    assert len(chart.series) == 1
    assert [category.label for category in chart.plots[0].categories] == ['DriveCity', 'DriveRoad']
    assert list(chart.series[0].values) == [32.0, 14.0]
    assert [str(point.format.fill.fore_color.rgb) for point in chart.series[0].points] == [
        '176E77', 'E6A81D',
    ]
    charts = [shape.chart for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 2
    assert [category.label for category in charts[1].plots[0].categories] == ['Voice', 'Data']
    assert list(charts[1].series[0].values) == [17.0, 29.0]
    assert [str(point.format.fill.fore_color.rgb) for point in charts[1].series[0].points] == [
        '4472C4', '7030A0',
    ]
    ring_shapes = [shape for shape in slide.shapes if shape.has_chart]
    assert ring_shapes[1].width < ring_shapes[0].width
    for environment_chart in charts:
        chart_space_children = list(environment_chart._chartSpace)
        chart_index = next(i for i, child in enumerate(chart_space_children) if child.tag == qn('c:chart'))
        shape_index = next(i for i, child in enumerate(chart_space_children) if child.tag == qn('c:spPr'))
        assert shape_index > chart_index
        text_properties = [i for i, child in enumerate(chart_space_children) if child.tag == qn('c:txPr')]
        assert not text_properties or shape_index < text_properties[0]
    text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame).replace('\x0b', '\n')
    assert '46.00\nmax points' in text
    assert 'Drive - City: 32 pts (69.6%)' in text
    assert 'Voice: 17 pts (37.0%)' in text and 'Data: 29 pts (63.0%)' in text
    assert 'Drive - Connecting Roads: 14 pts (30.4%)' in text
    assert 'Points per Environment:' in text and 'Points per Service:' in text
    labels = [shape for shape in slide.shapes if shape.has_text_frame
              and shape.text.startswith(('Drive - City', 'Drive - Connecting Roads', 'Voice: ', 'Data: '))]
    assert [shape.text.splitlines()[0] for shape in labels] == [
        'Drive - City: 32 pts (69.6%)', 'Drive - Connecting Roads: 14 pts (30.4%)',
        'Voice: 17 pts (37.0%)', 'Data: 29 pts (63.0%)',
    ]
    assert all(first.top < second.top for first, second in zip(labels, labels[1:]))
    detail_run = labels[0].text_frame.paragraphs[0].runs[-1]
    assert detail_run.text == '32 pts (69.6%)'
    assert str(detail_run.font.color.rgb) == '8A3D0A'
    chart_shape = next(shape for shape in slide.shapes if shape.has_chart)
    assert labels[-1].top >= chart_shape.top + chart_shape.height


def test_single_environment_allocation_keeps_outer_environment_ring_and_inner_service_ring():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    chart = add_maximum_allocation_donut(
        slide, {'context': {'environment': 'DriveRoad'}},
        maximum_allocations_from_configuration(_configuration()),
        left=1, top=1, width=4, height=3,
    )

    assert chart.chart_type == XL_CHART_TYPE.DOUGHNUT
    assert [category.label for category in chart.plots[0].categories] == ['DriveRoad']
    assert list(chart.series[0].values) == [14.0]
    assert [str(point.format.fill.fore_color.rgb) for point in chart.series[0].points] == [
        'E6A81D',
    ]
    charts = [shape.chart for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 2
    assert [category.label for category in charts[1].plots[0].categories] == ['Voice', 'Data']
    assert list(charts[1].series[0].values) == [5.0, 9.0]
    assert [str(point.format.fill.fore_color.rgb) for point in charts[1].series[0].points] == [
        '4472C4', '7030A0',
    ]
    drive_road_swatch = next(
        shape for shape in slide.shapes
        if shape.name.startswith('Maximum Allocation Swatch Drive - Connecting Roads')
    )
    assert str(drive_road_swatch.fill.fore_color.rgb) == 'E6A81D'
    assert '14.00\nmax points' in '\n'.join(
        shape.text for shape in slide.shapes if shape.has_text_frame
    ).replace('\x0b', '\n')


def test_many_environment_rings_keep_hole_size_within_powerpoint_limits():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    allocations = {
        f'Environment{index}': {'voice': 10 + index, 'data': 20 + index}
        for index in range(6)
    }

    add_maximum_allocation_donut(
        slide, {'context': {'environment': 'Combined'}}, allocations,
        left=1, top=1, width=4, height=3,
    )

    charts = [shape.chart for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 2
    hole_sizes = [int(chart.plots[0]._element.find(qn('c:holeSize')).get('val')) for chart in charts]
    assert all(10 <= hole <= 90 for hole in hole_sizes)


def test_category_allocations_count_each_matrix_row_once_and_render_category_ring():
    matrix = {
        'context': {'environment': 'DriveRoad'},
        'rows': [
            {'category': 'Voice', 'max_points': 12, 'values': {'operator-a': {'points': 4}, 'operator-b': {'points': 6}}},
            {'category': 'Data', 'max_points': 20, 'values': {'operator-a': {'points': 10}, 'operator-b': {'points': 9}}},
            {'category': 'Voice', 'max_points': 3, 'values': {'operator-a': {'points': 1}, 'operator-b': {'points': 2}}},
        ],
    }
    categories = category_maximum_allocations(matrix)
    assert categories == {'Voice': 15.0, 'Data': 20.0}

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart = add_maximum_allocation_donut(
        slide, matrix, {'DriveRoad': {'voice': 15, 'data': 20, 'max': 35}},
        left=1, top=1, width=3, height=3.2, category_allocations=categories,
    )

    assert chart.chart_type == XL_CHART_TYPE.DOUGHNUT
    assert [category.label for category in chart.plots[0].categories] == ['DriveRoad']
    assert list(chart.series[0].values) == [35]
    assert [str(point.format.fill.fore_color.rgb) for point in chart.series[0].points] == [
        '176E77',
    ]
    charts = [shape for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 2
    assert [category.label for category in charts[0].chart.plots[0].categories] == ['DriveRoad']
    assert list(charts[0].chart.series[0].values) == [35.0]
    assert [category.label for category in charts[1].chart.plots[0].categories] == ['Voice', 'Data']
    assert list(charts[1].chart.series[0].values) == [15.0, 20.0]
    assert [str(point.format.fill.fore_color.rgb) for point in charts[1].chart.series[0].points] == [
        '4472C4', '7030A0',
    ]
    assert all(chart.chart.series[0]._element.find(qn('c:dLbls')) is None for chart in charts)
    percentage_labels = [shape for shape in slide.shapes
                         if shape.name == 'Maximum Allocation Percentage Label']
    assert [shape.text_frame.text for shape in percentage_labels] == ['100.0%']
    assert all(7.5 <= shape.text_frame.paragraphs[0].runs[0].font.size.pt <= 8.5
               for shape in percentage_labels)
    chart_shapes = [shape for shape in slide.shapes if shape.has_chart]
    _assert_percentage_glyph_fits(percentage_labels[0], chart_shapes[0], 72, 1, 0)
    chart_shape = charts[0]
    swatches = [shape for shape in slide.shapes if shape.name.startswith('Maximum Allocation Swatch')]
    assert len(swatches) == 3
    texts = [shape for shape in slide.shapes if shape.has_text_frame]
    legend_texts = [shape for shape in texts if ' pts' in shape.text]
    assert all(shape.top >= chart_shape.top + chart_shape.height
               for shape in legend_texts)


def test_percentage_labels_hide_category_slices_that_are_too_small():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_maximum_allocation_donut(
        slide, {'context': {'environment': 'DriveCity'}},
        {'DriveCity': {'voice': 1, 'data': 99, 'max': 100}},
        left=1, top=1, width=3, height=3.2,
        category_allocations={'Multi RAB': 3.9, 'Video Stream': 96.1},
    )

    charts = [shape.chart for shape in slide.shapes if shape.has_chart]
    assert all(chart.series[0]._element.find(qn('c:dLbls')) is None for chart in charts)
    percentage_labels = [shape.text_frame.text for shape in slide.shapes
                         if shape.name == 'Maximum Allocation Percentage Label']
    assert percentage_labels == ['100.0%', '96.1%']


@pytest.mark.parametrize('environment', ['Combined', 'DriveCity'])
@pytest.mark.parametrize('category_mode', [False, True])
def test_saved_inner_ring_retains_fitting_35_65_percentages(environment, category_mode):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    allocations = {'DriveCity': {'voice': 227.5, 'data': 422.5},
                   'DriveRoad': {'voice': 122.5, 'data': 227.5}}
    total = 1000 if environment == 'Combined' else 650
    add_maximum_allocation_donut(
        slide, {'context': {'environment': environment}}, allocations,
        left=1, top=1, width=3, height=5,
        category_allocations={'Classic Calls': total * .35, 'Data': total * .65}
        if category_mode else None,
    )
    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    saved_slide = Presentation(output).slides[0]
    charts = [shape for shape in saved_slide.shapes if shape.has_chart]
    labels = [shape for shape in saved_slide.shapes
              if shape.name == 'Maximum Allocation Percentage Label']

    # The outer ring can contain the same percentages; require the separate
    # inner labels as well, after a complete native PowerPoint round trip.
    assert [label.text for label in labels] == (
        ['65.0%', '35.0%', '35.0%', '65.0%'] if environment == 'Combined'
        else ['100.0%', '35.0%', '65.0%']
    )
    _assert_percentage_glyph_fits(labels[-2], charts[1], 62, .35, 0)
    _assert_percentage_glyph_fits(labels[-1], charts[1], 62, .65, .35 * 2 * math.pi)


def test_saved_seven_category_widget_keeps_outer_and_major_inner_percentages():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    categories = {'Classic Calls': 203, 'WhatsApp Calls': 133, 'Multi RAB': 14,
                  'File Transfer': 351, 'Browsing': 130, 'Video': 130,
                  'Interactivity': 39}
    add_maximum_allocation_donut(
        slide, {'context': {'environment': 'Combined'}},
        {'DriveCity': {'voice': 227.5, 'data': 422.5},
         'DriveRoad': {'voice': 122.5, 'data': 227.5}},
        left=1, top=1.75, width=3, height=5, category_allocations=categories,
    )
    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    saved_slide = Presentation(output).slides[0]
    charts = [shape for shape in saved_slide.shapes if shape.has_chart]
    labels = [shape for shape in saved_slide.shapes
              if shape.name == 'Maximum Allocation Percentage Label']
    assert charts[0].width.inches >= 2.49 - .001
    assert [label.text for label in labels] == [
        '65.0%', '35.0%', '20.3%', '13.3%', '35.1%', '13.0%', '13.0%',
    ]
    for label, share, start in zip(labels[:2], [.65, .35], [0, .65 * 2 * math.pi]):
        _assert_percentage_glyph_fits(label, charts[0], 72, share, start)
    inner_hole = int(charts[1].chart.plots[0]._element.find(qn('c:holeSize')).get('val'))
    for label, category_index in zip(labels[2:], [0, 1, 3, 4, 5]):
        values = list(categories.values())
        _assert_percentage_glyph_fits(
            label, charts[1], inner_hole, values[category_index] / 1000,
            sum(values[:category_index]) / 1000 * 2 * math.pi,
        )
    legends = [shape for shape in saved_slide.shapes
               if shape.has_text_frame and ' pts (' in shape.text]
    assert all(shape.top >= charts[0].top + charts[0].height for shape in legends)
    assert all(shape.top.inches + shape.height.inches <= 6.75 for shape in legends)


@pytest.mark.parametrize('environment', ['Combined', 'DriveCity'])
@pytest.mark.parametrize('category_mode', [False, True])
def test_saved_allocation_matches_web_legend_and_center(environment, category_mode):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    allocations = {'DriveCity': {'voice': 200, 'data': 300},
                   'DriveRoad': {'voice': 100, 'data': 400}}
    total = 1000 if environment == 'Combined' else 500
    categories = {'Classic Calls': total * .2, 'WhatsApp Calls': total * .1,
                  'Multi RAB': total * .1, 'File Transfer': total * .1,
                  'Browsing': total * .1, 'Video': total * .1,
                  'Interactivity': total * .1, 'Other': total * .2} if category_mode else None
    add_maximum_allocation_donut(
        slide, {'context': {'environment': environment}}, allocations,
        left=1, top=1, width=3, height=5, category_allocations=categories,
    )
    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    saved_slide = Presentation(output).slides[0]
    chart_shapes = [shape for shape in saved_slide.shapes if shape.has_chart]
    charts = [shape.chart for shape in chart_shapes]
    assert len(charts) == 2
    assert sum(charts[0].series[0].values) == total
    assert sum(charts[1].series[0].values) == total
    percentage_labels = [shape for shape in saved_slide.shapes
                         if shape.name == 'Maximum Allocation Percentage Label']
    outer_shares = [.5, .5] if environment == 'Combined' else [1.0]
    inner_shares = ([.2, .1, .1, .1, .1, .1, .1, .2] if category_mode else
                    ([.3, .7] if environment == 'Combined' else [.4, .6]))
    for label in percentage_labels:
        label_x = label.left.inches + label.width.inches / 2
        label_y = label.top.inches + label.height.inches / 2
        candidate_rings = []
        for chart_shape, hole, shares in zip(chart_shapes, (72, 62), (outer_shares, inner_shares)):
            center_x = chart_shape.left.inches + chart_shape.width.inches / 2
            center_y = chart_shape.top.inches + chart_shape.height.inches / 2
            outer_radius = chart_shape.width.inches * .45
            middle_radius = (outer_radius + outer_radius * hole / 100) / 2
            distance = math.hypot(label_x - center_x, label_y - center_y)
            candidate_rings.append((abs(distance - middle_radius), chart_shape, hole, shares))
        _, chart_shape, hole, shares = min(candidate_rings, key=lambda item: item[0])
        label_share = float(label.text_frame.text.rstrip('%')) / 100
        matching = [share for share in shares if round(share * 100, 1) == round(label_share * 100, 1)]
        assert matching
        center_x = chart_shape.left.inches + chart_shape.width.inches / 2
        center_y = chart_shape.top.inches + chart_shape.height.inches / 2
        actual_angle = math.atan2(label_x - center_x, center_y - label_y)
        starts = [sum(shares[:index]) * 2 * math.pi for index in range(len(shares))]
        share, start_angle = min(
            ((share, start) for share, start in zip(shares, starts) if share in matching),
            key=lambda item: abs(math.atan2(math.sin(actual_angle - (item[1] + item[0] * math.pi)),
                                            math.cos(actual_angle - (item[1] + item[0] * math.pi)))),
        )
        _assert_percentage_glyph_fits(label, chart_shape, hole, share, start_angle)
    text_shapes = [shape for shape in saved_slide.shapes if shape.has_text_frame and shape.text]
    legends = [shape for shape in text_shapes if ' pts (' in shape.text]
    texts = [shape.text for shape in text_shapes]
    assert ('Total Points: 1000 pts (100.0%)' in texts) == (environment == 'Combined')
    assert ('Points per KPI Category:' if category_mode else 'Points per Service:') in texts
    city_share = '50.0%' if environment == 'Combined' else '100.0%'
    assert f'Drive - City: 500 pts ({city_share})' in texts
    for shape in legends:
        runs = shape.text_frame.paragraphs[0].runs
        assert runs[0].text.endswith(': ')
        assert str(runs[0].font.color.rgb) == '465565'
        assert str(runs[1].font.color.rgb) == '8A3D0A'
        assert runs[0].font.bold == shape.text.startswith('Total Points:')
        assert shape.top + shape.height <= presentation.slide_height
    if environment == 'Combined':
        assert legends[0].left < legends[1].left
        assert legends[0].top < legends[1].top
    center = next(shape for shape in saved_slide.shapes if shape.name == 'Maximum Score Allocation Total')
    number, unit = center.text_frame.paragraphs
    assert number.text == f'{total:.2f}'
    assert unit.text == 'max points'
    assert number.font.bold and unit.font.bold
    assert number.font.name == unit.font.name == 'Arial'
    assert number.font.size.pt > unit.font.size.pt
    assert str(number.font.color.rgb) == '1C3745'
    assert str(unit.font.color.rgb) == '5C707A'
    icons = [shape for shape in saved_slide.shapes if shape.name.startswith('Maximum Allocation Icon')]
    assert len(icons) == len(legends)
    segment_icons = [shape for shape in saved_slide.shapes
                     if shape.name.startswith('Maximum Allocation Environment Segment Icon')]
    assert len(segment_icons) == len(charts[0].plots[0].categories)
    assert all(icon.width.inches == pytest.approx(.45) and icon.height.inches == pytest.approx(.45)
               for icon in segment_icons)
    slide_width = presentation.slide_width / 914400
    slide_height = presentation.slide_height / 914400
    outer_radius = chart_shapes[0].width.inches * .45
    for icon in segment_icons:
        left, top = icon.left.inches, icon.top.inches
        right, bottom = left + icon.width.inches, top + icon.height.inches
        assert 0 <= left < right <= slide_width
        assert 0 <= top < bottom <= slide_height
        center_x = chart_shapes[0].left.inches + chart_shapes[0].width.inches / 2
        center_y = chart_shapes[0].top.inches + chart_shapes[0].height.inches / 2
        closest_x = max(abs(left + icon.width.inches / 2 - center_x) - icon.width.inches / 2, 0)
        closest_y = max(abs(top + icon.height.inches / 2 - center_y) - icon.height.inches / 2, 0)
        assert math.hypot(closest_x, closest_y) - outer_radius == pytest.approx(.08, abs=.015)
        assert all(right <= legend.left / 914400 or left >= (legend.left + legend.width) / 914400
                   or bottom <= legend.top / 914400 or top >= (legend.top + legend.height) / 914400
                   for legend in legends)
    for icon in icons:
        path = icon._element.spPr.find(qn('a:custGeom')).find(qn('a:pathLst')).find(qn('a:path'))
        assert path.get('fill') == 'none'
        assert path.get('w') == path.get('h') == '24000'
        assert icon._element.spPr.find(qn('a:effectLst')) is not None
        assert not list(icon._element.spPr.find(qn('a:effectLst')))
        assert icon._element.spPr.find(qn('a:ln')).get('cap') == 'rnd'
        assert all(reference.get('idx') == '0' for reference in icon._element.xpath('./p:style/a:effectRef'))
    city = next(icon for icon in icons if icon.name.endswith('city'))
    first_point = city._element.spPr.find(qn('a:custGeom')).find(qn('a:pathLst'))[0][0][0]
    assert (first_point.get('x'), first_point.get('y')) == ('3000', '21000')
    if category_mode:
        multirab = next(icon for icon in icons if icon.name.endswith('multirab'))
        assert len(multirab._element.xpath('.//a:arcTo')) == 2
        voice = next(icon for icon in icons if icon.name.endswith('voice'))
        assert len(voice._element.xpath('.//a:cubicBezTo')) == 3
    swatches = [shape for shape in saved_slide.shapes if shape.name.startswith('Maximum Allocation Swatch')]
    assert all(shape._element.spPr.find(qn('a:effectLst')) is not None for shape in swatches)


def test_legend_preserves_fractional_points_without_unnecessary_zeroes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_maximum_allocation_donut(
        slide, {'context': {'environment': 'DriveCity'}},
        {'DriveCity': {'voice': 1.125, 'data': 2.5}}, left=1, top=1, width=3, height=5,
    )
    text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert 'Voice: 1.125 pts (31.0%)' in text
    assert 'Data: 2.5 pts (69.0%)' in text


def test_environment_segment_icons_are_larger_than_legend_icons():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    allocations = {'DriveCity': {'voice': 227.5, 'data': 422.5},
                   'DriveRoad': {'voice': 122.5, 'data': 227.5}}
    add_maximum_allocation_donut(
        slide, {'context': {'environment': 'Combined'}}, allocations,
        left=1, top=1, width=3, height=5,
    )
    icons = [shape for shape in slide.shapes
             if shape.name.startswith('Maximum Allocation Environment Segment Icon')]
    assert len(icons) == 2
    assert all(max(icon.width.inches, icon.height.inches) > .18 for icon in icons)
