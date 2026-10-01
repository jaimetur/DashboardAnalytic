from collections import Counter
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

from src.modules.scoring_exports import export_scoring_powerpoint
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_exports import TEMPLATE, _mapping_groups, _result


NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}


def _presentation_bytes() -> bytes:
    job = {
        'aggregation_levels': ['Operator'],
        'nr_mode': 'NSA',
        'baseline_operator': 'EE',
        'campaigns': ['UK_Q2_2026'],
        'configuration': scoring_configuration(),
    }
    return export_scoring_powerpoint(
        job,
        _result(operators=('EE', 'O2 UK')),
        TEMPLATE,
        operator_mapping_groups=_mapping_groups(order=('EE', 'O2 UK')),
    )


def _normalize_part_path(path: str) -> str:
    parts = []
    for part in path.split('/'):
        if part == '..':
            parts.pop()
        elif part and part != '.':
            parts.append(part)
    return '/'.join(parts)


def test_generated_scoring_pptx_serializes_valid_table_and_combo_chart_structure():
    pptx_bytes = _presentation_bytes()
    presentation = Presentation(BytesIO(pptx_bytes))
    assert any(shape.has_table for slide in presentation.slides for shape in slide.shapes)
    best_network_slide = next(
        slide for slide in presentation.slides
        if slide.shapes.title.text.startswith('Best Network Scoring')
    )
    assert any(shape.has_chart for shape in best_network_slide.shapes)

    with ZipFile(BytesIO(pptx_bytes)) as package:
        names = set(package.namelist())
        part_roots = {}
        for name in names:
            if name.endswith(('.xml', '.rels')):
                part_roots[name] = etree.fromstring(package.read(name))

        # Table-cell border elements must occur once, in schema order, before fill/effects.
        edge_names = ('lnL', 'lnR', 'lnT', 'lnB')
        edge_tags = {f'{{{NS["a"]}}}{name}' for name in edge_names}
        all_table_cell_properties = [
            properties
            for name, root in part_roots.items()
            if name.startswith('ppt/slides/slide') and name.endswith('.xml')
            for properties in root.xpath('.//a:tbl//a:tcPr', namespaces=NS)
        ]
        table_cell_properties = [
            properties for properties in all_table_cell_properties
            if any(child.tag in edge_tags for child in properties)
        ]
        assert all_table_cell_properties and table_cell_properties
        for properties in table_cell_properties:
            children = list(properties)
            present_edges = [child for child in children if child.tag in edge_tags]
            assert [etree.QName(child).localname for child in present_edges] == list(edge_names)
            assert all(sum(child.tag == tag for child in children) == 1 for tag in edge_tags)
            assert max(children.index(child) for child in present_edges) < min(
                (index for index, child in enumerate(children)
                 if etree.QName(child).localname in {'noFill', 'solidFill', 'gradFill', 'blipFill', 'pattFill', 'effectLst'}),
                default=len(children),
            )

        chart_roots = [root for name, root in part_roots.items()
                       if name.startswith('ppt/charts/chart') and name.endswith('.xml')]
        assert chart_roots
        found_combo_chart = False
        for root in chart_roots:
            assert all(
                0 <= int(axis.get('val')) <= 0xFFFFFFFF
                for axis in root.xpath('.//c:axId | .//c:crossAx', namespaces=NS)
            )
            bar_plots = root.xpath('.//c:barChart', namespaces=NS)
            line_plots = root.xpath('.//c:lineChart', namespaces=NS)
            if bar_plots and line_plots:
                found_combo_chart = True
            if not bar_plots and not line_plots:
                continue
            axis_ids = {
                axis.get('val')
                for axis in root.xpath(
                    './/c:plotArea/c:valAx/c:axId | .//c:plotArea/c:catAx/c:axId',
                    namespaces=NS,
                )
            }
            assert axis_ids
            assert all(
                cross_axis.get('val') in axis_ids
                for cross_axis in root.xpath('.//c:plotArea//c:crossAx', namespaces=NS)
            )
            for plot in root.xpath('.//c:barChart | .//c:lineChart', namespaces=NS):
                series_indices = [
                    index.get('val') for index in plot.xpath('./c:ser/c:idx', namespaces=NS)
                ]
                assert len(series_indices) == len(set(series_indices))
            for series in root.xpath('.//c:barChart/c:ser | .//c:lineChart/c:ser', namespaces=NS):
                children = list(series)
                child_names = [etree.QName(child).localname for child in children]
                assert child_names.count('spPr') <= 1
                assert child_names.count('marker') <= 1
                if 'spPr' in child_names and 'marker' in child_names:
                    assert child_names.index('spPr') < child_names.index('marker')
        assert found_combo_chart

        # Relationship IDs and targets, along with per-slide shape IDs, must be unambiguous.
        for name, root in part_roots.items():
            if not name.endswith('.rels'):
                continue
            rel_ids = [relationship.get('Id') for relationship in root]
            assert len(rel_ids) == len(set(rel_ids)), name
            source_part = str(PurePosixPath(name).parent.parent / PurePosixPath(name).name.removesuffix('.rels'))
            for relationship in root:
                if relationship.get('TargetMode') == 'External':
                    continue
                target = _normalize_part_path(
                    str(PurePosixPath(source_part).parent / relationship.get('Target')),
                )
                assert target in names, (name, relationship.get('Id'), target)

        for name, root in part_roots.items():
            if not name.startswith('ppt/slides/slide') or not name.endswith('.xml'):
                continue
            shape_ids = [shape.get('id') for shape in root.xpath('.//p:cNvPr', namespaces=NS)]
            assert len(shape_ids) == len(set(shape_ids)), name

        slide_ids = [slide_id.get('id') for slide_id in part_roots['ppt/presentation.xml'].xpath(
            './/p:sldId', namespaces=NS,
        )]
        assert len(slide_ids) == len(set(slide_ids))
