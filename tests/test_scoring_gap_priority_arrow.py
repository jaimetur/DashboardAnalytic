import pytest
from pptx.enum.text import PP_ALIGN

from tests.test_scoring_exports import _export, _result


@pytest.mark.parametrize('levels', [('Operator',), ('Operator', 'Region', 'Campaign')])
@pytest.mark.parametrize('operators', [('EE', 'O2 UK'), ('EE', 'O2 UK', 'Three UK')])
def test_priority_arrow_is_only_on_individual_gap_slides(levels, operators):
    result = _result(operators=operators)
    if len(levels) > 1:
        result.update({'aggregation_contract_version': 2, 'aggregation_levels': list(levels)})
    presentation = _export(result, levels=levels)
    combined = individual = 0
    for slide in presentation.slides:
        title = slide.shapes.title.text
        assert all('Environment:' not in paragraph.text
                   for paragraph in slide.shapes.title.text_frame.paragraphs[1:])
        if not title.startswith('GAP Analysis'):
            continue
        all_operators = title.startswith('GAP Analysis — All vs ')
        combined += all_operators
        individual += not all_operators
        for name in ('GAP KPI Priority Arrow', 'GAP KPI Priority Label'):
            assert sum(shape.name == name for shape in slide.shapes) == (0 if all_operators else 1)
        assert any(shape.name.startswith('GAP Color Scale Segment ') for shape in slide.shapes)
        table_shape = next(shape for shape in slide.shapes if shape.has_table)
        table = table_shape.table
        header_rows = len(levels) + 1 if len(levels) > 1 else 1
        # KPI rows are followed by the Total KPI GAP row.
        assert len(table.rows) == header_rows + len(result['configuration']['metrics']) + 1
        assert table.cell(len(table.rows) - 1, 0).text == 'Total KPI GAP'
        kpi_names = [row.cells[1].text for row in list(table.rows)[header_rows:-1]]
        assert not any(name.lower().endswith(' total') for name in kpi_names)
        assert any('Average KPI GAP:' in shape.text
                   for shape in slide.shapes if shape.has_text_frame) == (not all_operators)
        visible_category_cells = [
            row.cells[0] for row in list(table.rows)[1:-1]
            if row.cells[0].text and row.cells[0].text != 'Category'
        ]
        category_fills = [str(cell.fill.fore_color.rgb).upper() for cell in visible_category_cells]
        assert category_fills
        assert category_fills == [
            'E6F0F7' if index % 2 == 0 else 'D7E5EE'
            for index in range(len(category_fills))
        ]
        scale_heading = next(shape for shape in slide.shapes
                             if shape.has_text_frame and shape.text == 'GAP color scale')
        assert table_shape.top + table_shape.height < scale_heading.top
        if not all_operators:
            label = next(shape for shape in slide.shapes if shape.name == 'GAP KPI Priority Label')
            arrow = next(shape for shape in slide.shapes if shape.name == 'GAP KPI Priority Arrow')
            assert arrow.top > label.top + label.height
            assert arrow.top + arrow.height == table_shape.top + table_shape.height
        for cell in visible_category_cells:
            for paragraph in cell.text_frame.paragraphs:
                assert paragraph.alignment == PP_ALIGN.LEFT
                assert paragraph.runs
                assert all(
                    run.font.bold if run.font.bold is not None else paragraph.font.bold
                    for run in paragraph.runs
                )
    assert combined == 1
    assert individual == len(operators) - 1
