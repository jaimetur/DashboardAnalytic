"""Most Reliable Network scoring: a subset of KPIs with its own maximum points."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation

from src.modules.scoring import calculate_scoring, interpolate_score, most_reliable_result
from src.modules.scoring_config import (
    configuration_hash, has_most_reliable_scoring, most_reliable_configuration, scoring_profiles_document,
    unwrap_scoring_profiles_payload, validate_scoring_configuration,
)
from src.modules.scoring_exports import export_scoring_csv, export_scoring_powerpoint
from src.modules.scoring_views import build_scoring_views
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_api import scoring_api  # noqa: F401  (fixture)
from tests.test_scoring_exports import TEMPLATE

# NetCheck 2026 Most Reliable allocation: each KPI's global points split 65% City / 35% Road.
RELIABLE_POINTS = {
    'K1': (92.1375, 49.6125), 'K4': (44.3625, 23.8875), 'K7': (61.425, 33.075), 'K8': (29.575, 15.925),
    'K12': (63.375, 34.125), 'K14': (38.025, 20.475), 'K16': (101.4, 54.6), 'K21': (50.7, 27.3),
    'K27': (84.5, 45.5), 'K28': (84.5, 45.5),
}
# Netcheck_Score_Mapping_2026Q2_Mostreliable_Drive_City_Road.xlsx rounds them to two decimals.
WORKBOOK_POINTS = {
    'K1': (92.14, 49.61), 'K4': (44.36, 23.89), 'K7': (61.43, 33.08), 'K8': (29.58, 15.93),
    'K12': (63.38, 34.13), 'K14': (38.03, 20.48), 'K16': (101.4, 54.6), 'K21': (50.7, 27.3),
    'K27': (84.5, 45.5), 'K28': (84.5, 45.5),
}
# Workbook KPI values per operator (O2, Vodafone, EE, Three) for City and Road.
WORKBOOK_VALUES = {
    'DriveCity': {
        'K1': (99.279661017, 99.280270957, 99.576809141, 98.215044624),
        'K4': (0.685724735, 0.794689888, 0.433938217, 0.797157423),
        'K7': (99.741100324, 99.675535367, 99.935979513, 99.739752765),
        'K8': (1.912365392, 1.089030577, 0.439761796, 3.865004901),
        'K12': (99.220639613, 99.188481675, 99.771457593, 99.012016021),
        'K14': (96.12590799, 99.108313664, 98.296465802, 98.797434527),
        'K16': (99.126290707, 98.969868658, 99.799599198, 98.61147498),
        'K21': (97.58507135, 99.52241974, 99.588900308, 99.316379546),
        'K27': (99.564038492, 99.580223881, 99.773550725, 99.494124466),
        'K28': (99.414582225, 99.635796046, 99.597281651, 99.471319059),
    },
    'DriveConnectionroad': {
        'K1': (96.918767507, 97.881355932, 99.298737728, 97.331460674),
        'K4': (2.712215321, 2.17772732, 0.909274601, 2.283293729),
        'K7': (99.315068493, 98.627002288, 99.092970522, 98.594847775),
        'K8': (3.367833087, 1.830474934, 0.666558283, 4.33901739),
        'K12': (96.272134203, 96.794871795, 99.743808711, 97.623400366),
        'K14': (94.642857143, 97.150735294, 96.318493151, 97.058823529),
        'K16': (95.551436515, 96.415770609, 99.49452401, 98.109810981),
        'K21': (96.62487946, 97.641509434, 99.04097646, 97.238095238),
        'K27': (98.575656678, 98.899314327, 99.544841537, 99.025270758),
        'K28': (98.235840297, 98.914027149, 99.408783784, 99.002719855),
    },
}
OPERATORS = ('O2 UK', 'Vodafone UK', 'EE', 'Three UK')
# Workbook totals: City voice + data, Road voice + data.
WORKBOOK_TOTALS = {
    'DriveCity': (561.752462776642, 587.9043264550877, 612.9768804671703, 559.0465470730871),
    'DriveConnectionroad': (217.99831293993665, 247.0313104871609, 315.10300726054186, 262.5496255047821),
}


def _reliable_configuration(points=RELIABLE_POINTS):
    configuration = scoring_configuration()
    for metric in configuration['metrics']:
        if metric['code'] in points:
            city, road = points[metric['code']]
            metric['contexts']['DriveCity']['most_reliable_points'] = city
            metric['contexts']['DriveConnectionroad']['most_reliable_points'] = road
    return validate_scoring_configuration(configuration)


def _workbook_result(configuration):
    """A saved result whose KPI values are the Most Reliable workbook values."""
    metrics = {metric['code']: metric for metric in configuration['metrics']}
    rows = []
    for environment, values in WORKBOOK_VALUES.items():
        for code, measurements in values.items():
            metric = metrics[code]
            context = metric['contexts'][environment]
            for operator, value in zip(OPERATORS, measurements):
                score = interpolate_score(value, context['thresholds'], None, context['score_mapping'])
                rows.append({
                    'campaign': None, 'operator': operator, 'environment': environment, 'kpi_code': code,
                    'kpi': metric['kpi'], 'category': metric['category'],
                    'dataset_type': metric['source_kind'].title(), 'kpi_type': metric['kpi_type'],
                    'value': value, 'sample_count': 100, 'score': score,
                    'max_points': context['max_points'], 'weighted_points': score * context['max_points'],
                })
    return {'scoring': rows, 'global_kpis': [], 'warnings': [], 'aggregation_levels': ['Operator'],
            'configuration': configuration, 'baseline_aliases': ['EE']}


def _overall(result, environment):
    return {row['operator']: row for row in result['totals']
            if row['environment'] == environment and row['category'] == 'Overall'}


def test_most_reliable_points_are_validated_and_only_change_identity_when_set():
    plain = validate_scoring_configuration(scoring_configuration())
    assert not has_most_reliable_scoring(plain)
    assert most_reliable_configuration(plain) is None

    zero = scoring_configuration()
    zero['metrics'][0]['contexts']['DriveCity']['most_reliable_points'] = 0
    assert 'most_reliable_points' not in validate_scoring_configuration(zero)['metrics'][0]['contexts']['DriveCity']
    assert configuration_hash(validate_scoring_configuration(zero)) == configuration_hash(plain)

    negative = scoring_configuration()
    negative['metrics'][0]['contexts']['DriveCity']['most_reliable_points'] = -1
    with pytest.raises(ValueError, match='most_reliable_points'):
        validate_scoring_configuration(negative)

    reliable = _reliable_configuration()
    assert has_most_reliable_scoring(reliable)
    assert configuration_hash(reliable) != configuration_hash(plain)
    assert reliable['metrics'][0]['contexts']['DriveCity']['most_reliable_points'] == 92.1375
    # Best Network keeps its own maximum points.
    assert reliable['scope']['total_max_points'] == pytest.approx(1000)


def test_most_reliable_configuration_keeps_the_subset_with_its_points_and_thresholds():
    reliable = _reliable_configuration()
    derived = most_reliable_configuration(reliable)
    assert [metric['code'] for metric in derived['metrics']] == list(RELIABLE_POINTS)
    assert derived['gap_priority'] == [code for code in reliable['gap_priority'] if code in RELIABLE_POINTS]
    k1 = derived['metrics'][0]
    assert k1['contexts']['DriveCity']['max_points'] == 92.1375
    assert k1['contexts']['DriveConnectionroad']['max_points'] == 49.6125
    assert k1['contexts']['Walk']['max_points'] == 0
    assert k1['contexts']['DriveCity']['thresholds'] == reliable['metrics'][0]['contexts']['DriveCity']['thresholds']
    environments = derived['scope']['environments']
    assert environments['DriveCity']['total_points'] == pytest.approx(650)
    assert environments['DriveConnectionroad']['total_points'] == pytest.approx(350)
    voice = sum(context['max_points'] for metric in derived['metrics'] if metric['source_kind'] != 'data'
                for context in metric['contexts'].values())
    assert voice == pytest.approx(350)
    assert derived['scoring'] == 'most_reliable'


def test_most_reliable_scoring_matches_the_netcheck_workbook():
    result = most_reliable_result(_workbook_result(_reliable_configuration()), 'EE')
    # The workbook rounds the points to two decimals: totals differ by less than 0.05 points.
    for environment, expected in WORKBOOK_TOTALS.items():
        totals = _overall(result, environment)
        for operator, points in zip(OPERATORS, expected):
            assert totals[operator]['weighted_points'] == pytest.approx(points, abs=.05)
            assert totals[operator]['complete_coverage'] is True
    combined = _overall(result, 'Combined')
    assert combined['EE']['max_points'] == pytest.approx(1000)
    for index, operator in enumerate(OPERATORS):
        assert combined[operator]['weighted_points'] == pytest.approx(
            WORKBOOK_TOTALS['DriveCity'][index] + WORKBOOK_TOTALS['DriveConnectionroad'][index], abs=.1)
    # The GAP subtracts the reference's Most Reliable points.
    o2_gap = next(row for row in result['gap'] if row['operator'] == 'O2 UK' and row['kpi_code'] == 'K1'
                  and row['environment'] == 'DriveCity')
    assert o2_gap['gap_points'] == pytest.approx(85.50279661063796 - 88.24071942517398, abs=.01)
    assert {row['kpi_code'] for row in result['scoring']} == set(RELIABLE_POINTS)


def test_rounded_points_are_scaled_to_the_environment_totals():
    rounded = _reliable_configuration(WORKBOOK_POINTS)
    derived = most_reliable_configuration(rounded)
    environments = derived['scope']['environments']
    assert environments['DriveCity']['total_points'] == pytest.approx(650, abs=1e-6)
    assert environments['DriveConnectionroad']['total_points'] == pytest.approx(350, abs=1e-6)
    k1 = rounded['metrics'][0]['contexts']['DriveCity']['most_reliable_points']
    assert k1 == pytest.approx(RELIABLE_POINTS['K1'][0], abs=.001)
    # Validation is idempotent, so the methodology keeps its identity.
    assert configuration_hash(validate_scoring_configuration(rounded)) == configuration_hash(rounded)
    # Totals far from the Best Network totals are kept as they are.
    custom = scoring_configuration()
    custom['metrics'][0]['contexts']['DriveCity']['most_reliable_points'] = 92.14
    assert validate_scoring_configuration(custom)['metrics'][0]['contexts']['DriveCity']['most_reliable_points'] == 92.14


def test_only_reliable_kpis_take_part_in_most_reliable():
    configuration = scoring_configuration()
    k1, k2 = configuration['metrics'][0], configuration['metrics'][1]
    assert (k1['kpi_type'], k2['kpi_type']) == ('Reliable', 'Diff')
    k1['contexts']['DriveCity']['most_reliable_points'] = 100
    k2['contexts']['DriveCity']['most_reliable_points'] = 50
    validated = validate_scoring_configuration(configuration)
    assert 'most_reliable_points' not in validated['metrics'][1]['contexts']['DriveCity']
    derived = most_reliable_configuration(validated)
    assert 'K2' not in {metric['code'] for metric in derived['metrics']}
    # Every Reliable KPI takes part; those without points have no maximum.
    reliable = {metric['code'] for metric in configuration['metrics'] if metric['kpi_type'] == 'Reliable'}
    assert {metric['code'] for metric in derived['metrics']} == reliable
    # A Diff KPI with points gives no Most Reliable scoring.
    only_diff = scoring_configuration()
    only_diff['metrics'][1]['contexts']['DriveCity']['most_reliable_points'] = 50
    assert most_reliable_configuration(only_diff) is None


def test_most_reliable_result_reuses_saved_scores_and_is_absent_without_points():
    configuration = _reliable_configuration()
    rows = [{'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': level_2,
             'Session_Type': 'CALL', 'Call_Status': status}
            for level_2 in ('City', 'Connectionroad')
            for operator, status in (('EE', 'Completed'), ('O2 UK', 'Completed'), ('O2 UK', 'Failed'))]
    result = calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator'], configuration=configuration)
    reliable = most_reliable_result(result, 'EE')
    best_k1 = next(row for row in result['scoring'] if row['kpi_code'] == 'K1' and row['operator'] == 'EE'
                   and row['environment'] == 'DriveCity')
    reliable_k1 = next(row for row in reliable['scoring'] if row['kpi_code'] == 'K1' and row['operator'] == 'EE'
                       and row['environment'] == 'DriveCity')
    assert reliable_k1['score'] == best_k1['score'] == 1
    assert best_k1['weighted_points'] == pytest.approx(73.4825)
    assert reliable_k1['weighted_points'] == pytest.approx(92.1375)
    # Warnings about KPIs outside Most Reliable are left out.
    assert any('K2 requires missing column' in warning for warning in result['warnings'])
    assert not any('K2 requires missing column' in warning for warning in reliable['warnings'])
    assert most_reliable_result({**result, 'configuration': scoring_configuration()}, 'EE') is None
    # The saved result is not changed.
    assert best_k1['max_points'] == pytest.approx(73.4825)


def test_most_reliable_views_and_csv_use_the_subset():
    configuration = _reliable_configuration()
    result = _workbook_result(configuration)
    job = {'levels': ['Operator'], 'baseline_operator': 'EE', 'configuration': configuration}
    reliable = most_reliable_result(result, 'EE')
    views = build_scoring_views(job, reliable)
    city = next(table for table in views['score_tables'] if table['context']['environment'] == 'DriveCity')
    assert {row['kpi_code'] for row in city['expanded_rows'] if row.get('kpi_code')} == set(RELIABLE_POINTS)
    csv_text = export_scoring_csv(job, reliable, 'scoring', 'expanded')
    assert 'FDTT DL THROUGHPUT > 5Mbit/s' not in csv_text
    assert 'FDTT DL THROUGHPUT > 2Mbit/s' in csv_text


def _slide_titles(content):
    presentation = Presentation(BytesIO(content))
    return [slide.shapes.title.text_frame.text.replace('\n', ' | ') if slide.shapes.title is not None else ''
            for slide in presentation.slides]


def test_powerpoint_carries_both_scorings_in_one_document():
    configuration = _reliable_configuration()
    job = {'levels': ['Operator'], 'baseline_operator': 'EE', 'configuration': configuration}
    titles = _slide_titles(export_scoring_powerpoint(job, _workbook_result(configuration), TEMPLATE))
    assert titles[0].startswith('Scoring & GAP Analysis')
    best = titles.index(next(title for title in titles if title.startswith('Best Network Scoring')))
    reliable = titles.index(next(title for title in titles if title.startswith('Most Reliable Network Scoring')))
    assert 0 < best < reliable
    assert 'Best Network Scoring per Service | Best Network' in titles[best:reliable]
    assert any(title.startswith('Best Network Scoring per Service | Best Network — ') for title in titles[best:reliable])
    assert any(title.startswith('Most Reliable Network Scoring per Service | Most Reliable Network — ')
               for title in titles[reliable:])
    assert any(title.startswith('GAP Analysis — O2 UK vs EE | Most Reliable Network') for title in titles[reliable:])


def test_powerpoint_without_most_reliable_points_is_unchanged():
    configuration = validate_scoring_configuration(scoring_configuration())
    job = {'levels': ['Operator'], 'baseline_operator': 'EE', 'configuration': configuration}
    titles = _slide_titles(export_scoring_powerpoint(job, _workbook_result(configuration), TEMPLATE))
    assert not any('Most Reliable' in title for title in titles)
    assert any(title.startswith('Best Network Scoring per Service') for title in titles)


def test_most_reliable_points_survive_methodology_export_and_import():
    configuration = _reliable_configuration()
    document = scoring_profiles_document({
        'active_profile_id': 'netcheck-2026',
        'profiles': [{'id': 'netcheck-2026', 'name': 'NetCheck 2026', 'configuration': deepcopy(configuration)}],
    })
    restored = unwrap_scoring_profiles_payload(document)['profiles'][0]['configuration']
    assert restored['metrics'][0]['contexts']['DriveCity']['most_reliable_points'] == 92.1375
    assert restored['metrics'][0]['contexts']['DriveCity']['max_points'] == pytest.approx(73.4825)
    assert configuration_hash(most_reliable_configuration(restored)) == configuration_hash(
        most_reliable_configuration(configuration))


def test_scoring_job_api_returns_and_exports_the_most_reliable_scoring(scoring_api, monkeypatch):
    from types import ModuleType

    from src.modules import scoring as scoring_engine
    from src.modules import scoring_jobs

    client = scoring_api['client']
    repository = scoring_api['repository']
    repository.replace_scoring_configuration(_reliable_configuration())
    rows = [{'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': level_2,
             'Session_Type': 'CALL', 'Call_Status': status}
            for level_2 in ('City', 'Connectionroad')
            for operator, status in (('EE', 'Completed'), ('O2 UK', 'Completed'), ('O2 UK', 'Failed'))]
    engine = ModuleType('src.modules.scoring')
    engine.METHOD_VERSION = 'most-reliable-api-test'
    engine.required_input_columns = lambda kind, levels: [*levels, 'score']
    engine.calculate_scoring = lambda frames, levels, *, baseline_operator, configuration=None: (
        scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, levels, baseline_operator=baseline_operator,
                                         configuration=configuration))
    monkeypatch.setattr(scoring_jobs, '_scoring_engine', lambda: engine)
    created = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator'],
        'nr_mode': 'NSA', 'baseline_operator': 'EE',
    })
    assert created.status_code == 200, created.text
    job_id = created.json()['job']['id']
    assert scoring_jobs.run_scoring_job(repository, job_id)['status'] == 'completed'

    payload = client.get(f'/api/scoring/jobs/{job_id}').json()
    reliable = payload['scorings']['most_reliable']
    assert reliable['label'] == 'Most Reliable Network'
    assert [metric['code'] for metric in reliable['configuration']['metrics']] == list(RELIABLE_POINTS)
    city = next(table for table in reliable['views']['score_tables'] if table['context']['environment'] == 'DriveCity')
    assert {row['kpi_code'] for row in city['expanded_rows'] if row.get('kpi_code')} == set(RELIABLE_POINTS)
    assert any(row['category'] == 'Overall' for row in reliable['totals'])

    best_csv = client.get(f'/scoring/jobs/{job_id}/export/scoring?table_mode=expanded').content.decode('utf-8-sig')
    reliable_csv = client.get(
        f'/scoring/jobs/{job_id}/export/scoring?table_mode=expanded&scoring=most_reliable').content.decode('utf-8-sig')
    assert 'CALL SETUP TIME' in best_csv
    assert 'CALL SETUP TIME' not in reliable_csv and 'CALL SUCCESS RATIO' in reliable_csv
    assert client.get(f'/scoring/jobs/{job_id}/export/scoring?scoring=unknown').status_code == 400

    ppt = client.get(f'/scoring/jobs/{job_id}/export/ppt')
    assert ppt.status_code == 200, ppt.text
    titles = _slide_titles(ppt.content)
    assert any(title.startswith('Best Network Scoring') for title in titles)
    assert any(title.startswith('Most Reliable Network Scoring') for title in titles)
