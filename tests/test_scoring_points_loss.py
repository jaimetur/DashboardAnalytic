"""Where the scoring points are lost: points-lost maps per City, ITL3 area and custom Clusters."""
from __future__ import annotations

import json
import random
from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation

from src.modules import scoring_points_loss as points_loss
from src.modules.scoring import calculate_scoring
from src.modules.scoring_exports import export_scoring_powerpoint, export_scoring_report
from src.modules.scoring_reports import default_scenario
from src.modules.scoring_views import build_scoring_views
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_api import scoring_api  # noqa: F401  (fixture)
from tests.test_scoring_exports import TEMPLATE

PLACES = {'Leeds': (53.80, -1.55), 'Manchester': (53.48, -2.24), 'London': (51.51, -0.13)}


def _result():
    """Three operators; Vodafone and Three fail calls mostly in Leeds and on a road towards it."""
    random.seed(7)
    rows = []
    for operator in ('EE', 'Vodafone UK', 'Three UK'):
        for city, (latitude, longitude) in PLACES.items():
            for index in range(40):
                failed = operator != 'EE' and ((city == 'Leeds' and index < 8) or (city == 'Manchester' and index < 2))
                rows.append({'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': 'City',
                             'City': city, 'Session_Type': 'CALL', 'Call_Status': 'Failed' if failed else 'Completed',
                             'Call_Start_Latitude_A': latitude + random.uniform(-.02, .02),
                             'Call_Start_Longitude_A': longitude + random.uniform(-.02, .02)})
        for index in range(30):
            rows.append({'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': 'Connectionroad',
                         'City': 'Sheffield to Leeds', 'Session_Type': 'CALL',
                         'Call_Status': 'Failed' if operator != 'EE' and index % 6 == 0 else 'Completed',
                         'Call_Start_Latitude_A': 53.38 + index * .014, 'Call_Start_Longitude_A': -1.47 - index * .003})
    result = calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator'], configuration=scoring_configuration())
    job = {'id': 1, 'levels': ['Operator'], 'baseline_operator': 'EE', 'configuration': result['configuration'],
           'campaigns': ['2026-Q2']}
    return job, result


def test_points_lost_are_spread_over_the_cities_routes_and_itl3_areas():
    job, result = _result()
    maps = build_scoring_views(job, result)['insights']['points_loss_maps']
    vodafone = {item['field']: item for item in maps if item['operator'] == 'Vodafone UK' and item['environment'] == 'Combined'}
    assert {'City', 'Area'} <= set(vodafone) and vodafone['Area']['area_label'] == 'ITL3 area'
    cities = {area['name']: area for area in vodafone['City']['areas']}
    # Leeds loses most of the city points; the road is a route between cities.
    assert next(iter(cities)) in {'Leeds', 'Sheffield to Leeds'}
    assert cities['Leeds']['points'] > cities['Manchester']['points'] and 'London' not in cities
    assert cities['Sheffield to Leeds']['kind'] == 'route' and cities['Sheffield to Leeds']['route']
    # The map areas (ITL3 in the UK) share the same total: the tests are placed in the polygon that contains them.
    areas = {area['name']: area['points'] for area in vodafone['Area']['areas']}
    assert vodafone['Area']['total'] == pytest.approx(vodafone['City']['total'])
    # Leeds also takes the end of the road that enters it.
    assert cities['Leeds']['points'] < areas['Leeds'] < cities['Leeds']['points'] + cities['Sheffield to Leeds']['points']
    # EE completes every call: it loses no points and gets no map.
    assert not any(item['operator'] == 'EE' for item in maps)
    # Each environment has its own maps; All Environments names, beside each area, where it loses points.
    environments = {item['environment'] for item in maps} - {'Combined'}
    assert environments and all(cities[name]['environments'] for name in cities)
    for environment in environments:
        city = next(item for item in maps if item['operator'] == 'Vodafone UK' and item['environment'] == environment
                    and item['field'] == 'City')
        assert all(not area['environments'] for area in city['areas'])
    # The map areas of each City's tests (of every operator), to show it alone on the map.
    assert cities['Leeds']['area_tests'] == {'Leeds': 120}
    assert set(cities['Sheffield to Leeds']['area_tests']) >= {'Sheffield', 'Leeds'}
    # The points each City loses in each area add up to its points, and a route loses them along its way.
    assert cities['Leeds']['area_losses'] == {'Leeds': pytest.approx(cities['Leeds']['points'])}
    road = cities['Sheffield to Leeds']['area_losses']
    assert len(road) > 1 and sum(road.values()) == pytest.approx(cities['Sheffield to Leeds']['points'])
    assert areas['Leeds'] == pytest.approx(cities['Leeds']['points'] + road.get('Leeds', 0))


def test_map_areas_take_the_points_of_the_city_or_route_that_loses_most_in_them():
    areas = [
        {'name': 'London', 'points': 12.5, 'kind': 'place',
         'area_tests': {'Westminster and City of London': 60, 'Camden': 39, 'West Surrey': 1},
         'area_losses': {'Westminster and City of London': 8.0, 'Camden': 4.4, 'West Surrey': .1}},
        {'name': 'Coventry', 'points': 5.5, 'kind': 'place', 'area_tests': {'Coventry': 50}, 'area_losses': {'Coventry': 5.5}},
        {'name': 'Coventry to Corby', 'points': 1.7, 'kind': 'route', 'area_tests': {'Coventry': 10, 'Leicestershire': 30},
         'area_losses': {'Coventry': 1.2, 'Leicestershire': .5}},
    ]
    # The points of the city or route that loses most among those measured in each area (2% of its tests or more).
    assert points_loss.row_area_values(areas) == {
        'Westminster and City of London': 12.5, 'Camden': 12.5, 'Coventry': 5.5, 'Leicestershire': 1.7}



def test_bundled_itl3_boundaries_cover_the_uk():
    boundaries = points_loss.bundled_boundaries('ITL3')
    assert len(boundaries) == 182 and 'Leeds' in boundaries
    names = points_loss.boundary_names('ITL3', pd.Series([53.80, 51.51, 0.0]), pd.Series([-1.55, -0.13, 0.0]))
    assert names.iloc[0] == 'Leeds' and names.iloc[1] == 'Westminster and City of London' and pd.isna(names.iloc[2])


def test_custom_cluster_polygons_are_read_from_a_geojson(tmp_path):
    path = tmp_path / 'clusters.geojson'
    path.write_text(json.dumps({'type': 'FeatureCollection', 'crs': {'type': 'name', 'properties': {'name': 'EPSG:4326'}},
                                'features': [{'type': 'Feature', 'properties': {'Cluster': 'North'},
                                              'geometry': {'type': 'Polygon', 'coordinates': [[[-2.5, 53], [-1, 53], [-1, 54.5], [-2.5, 54.5], [-2.5, 53]]]}}]}))
    boundaries = points_loss.mapping_boundaries(path, 'Cluster')
    assert list(boundaries) == ['North'] and boundaries['North'][0][0] == [-2.5, 53.0]
    assert points_loss.map_boundaries({'points_loss': {'boundaries': {'Cluster': boundaries}}}, 'Cluster') == boundaries


def _titles(content):
    return [slide.shapes.title.text_frame.text.split('\n')[0] for slide in Presentation(BytesIO(content)).slides
            if slide.shapes.title is not None]


def test_powerpoint_maps_the_points_lost_of_the_compared_operators():
    job, result = _result()
    content = export_scoring_powerpoint(job, result, TEMPLATE)
    titles = _titles(content)
    assert 'Points Lost per City — Vodafone UK' in titles and 'Points Lost per City — Three UK' in titles
    assert not any(title.startswith('Points Lost per City — EE') for title in titles)
    deck = Presentation(BytesIO(content))
    slide = next(slide for slide in deck.slides if slide.shapes.title is not None
                 and slide.shapes.title.text_frame.text.startswith('Points Lost per City — Vodafone UK'))
    names = {shape.name for shape in slide.shapes}
    assert {'Scoring Points Lost Map', 'Scoring Points Lost Bars', 'Scoring Points Lost Scale'} <= names
    # The operator, large, above the map and the ranking.
    heading = next(shape for shape in slide.shapes if shape.name == 'Scoring Points Lost Operator')
    assert heading.text_frame.text == 'Vodafone UK' and heading.text_frame.paragraphs[0].runs[0].font.size.pt == 24
    scenario = default_scenario(operators=['Vodafone UK', 'Three UK'])
    for options in scenario['scorings'].values():
        options['gap']['points_loss_map'] = False
    without = export_scoring_report([{'job': job, 'result': result, 'scenario': scenario}], TEMPLATE)
    assert not any(title.startswith('Points Lost per') for title in _titles(without))


def test_job_api_serves_the_map_layers_without_the_raw_shares(scoring_api, monkeypatch):
    from types import ModuleType

    from src.modules import scoring_jobs

    client = scoring_api['client']
    repository = scoring_api['repository']
    repository.replace_scoring_configuration(scoring_configuration())
    _job, computed = _result()
    engine = ModuleType('src.modules.scoring')
    engine.METHOD_VERSION = 'points-loss-api-test'
    engine.required_input_columns = lambda kind, levels: [*levels, 'score']
    engine.calculate_scoring = lambda frames, levels, *, baseline_operator, configuration=None: computed
    monkeypatch.setattr(scoring_jobs, '_scoring_engine', lambda: engine)
    created = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator'],
        'nr_mode': 'NSA', 'baseline_operator': 'EE',
    })
    assert created.status_code == 200, created.text
    job_id = created.json()['job']['id']
    assert scoring_jobs.run_scoring_job(repository, job_id)['status'] == 'completed'

    payload = client.get(f'/api/scoring/jobs/{job_id}').json()
    assert 'points_loss' not in payload
    assert any(item['field'] == 'Area' for item in payload['views']['insights']['points_loss_maps'])
    # UK tests use the bundled ITL3 areas; no country outline is needed under them.
    layer = client.get(f'/api/scoring/jobs/{job_id}/boundaries/Area').json()
    assert layer['field'] == 'Area' and len(layer['boundaries']) == 182 and layer['background'] == []
    assert client.get(f'/api/scoring/jobs/{job_id}/boundaries/ITL3').status_code == 404
    # Without a Clusters dataset applied to the CDRs the Cluster layer has no polygons.
    assert client.get(f'/api/scoring/jobs/{job_id}/boundaries/Cluster').json()['boundaries'] == {}
    assert client.get(f'/api/scoring/jobs/{job_id}/boundaries/Unknown').status_code == 404


def test_the_web_map_links_the_ranking_and_the_areas():
    from pathlib import Path

    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    # Each column is headed by its operator; rows and areas choose each other, and chosen rows centre the map.
    assert "title.className = 'scoring-loss-map-operator';" in script
    assert 'function lossMapSelection(card, svg, rows, initialLinks)' in script and 'selection.zoom.fit({x, y, width, height});' in script
    # Every area keeps its colour: choosing rows or areas only dims the others.
    assert "marks.forEach(mark => mark.classList.toggle('is-dimmed', !lit.has(mark.dataset.area)));" in script
    # One choice above the maps, the city or route that loses most by default; changing it keeps the chosen rows.
    assert "for (const [value, text] of [['top', 'Points of the city or route that loses most in each area']," in script
    assert "cards.forEach(card => card.setColourMode?.(modeSelect.value));" in script
    assert "mode = real && insightSelections.get('points-loss-colour') === 'losses' ? 'losses' : 'top';" in script
    assert 'if (event.shiftKey && onSelect) onSelect(area, event); else apply(area);' in script
    # All Environments names the environments of each area.
    assert "...(withEnvironments ? ['Environment'] : [])" in script
