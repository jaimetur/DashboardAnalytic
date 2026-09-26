import json
import re
import time
import zipfile
from io import BytesIO
from pathlib import Path
from threading import Event, Thread

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

import src.DashboardAnalytic as core
from src.modules.e2e_dashboards import DashboardDefinition, filter_frame


def definition(**changes):
    return DashboardDefinition(name='Comparison', template='Dashboard test', datasets={'data': [1]}, **changes)


def test_dashboard_uses_combined_tables_without_projection_or_warmup_queue():
    source = (Path(__file__).parents[1] / 'src/modules/e2e_dashboards.py').read_text(encoding='utf-8')
    assert 'dashboard-analytics.sqlite3' not in source
    assert 'def ensure_projection' not in source
    assert 'def enqueue_prefetch' not in source
    assert 'def prefetch_workspace_dashboards' not in source
    assert 'def reporting_source(snapshot, kind, task_repository):' in source


def test_dashboard_filter_panel_state_is_scoped_to_the_authenticated_session():
    first_user = core.SessionUser(username='super', role='super-admin')
    second_user = core.SessionUser(username='super', role='super-admin')
    assert first_user.session_marker
    assert first_user.session_marker != second_user.session_marker


def test_dashboard_library_can_change_nr_mode_and_template_with_confirmation_controls():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/e2e_dashboards.js').read_text(encoding='utf-8')
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')

    assert "technologySelect.className = 'ds-dashboard-definition-select ds-dashboard-nr-mode-select'" in script
    assert "templateSelect.className = 'ds-dashboard-definition-select ds-dashboard-template-select'" in script
    assert 'will invalidate every report and chart created inside this Dashboard' in script
    assert 'Its saved Dataset Universe and filters will be preserved.' in script
    assert "api(`/${encodeURIComponent(id)}/template`, 'PATCH'" in script
    assert '.ds-dashboard-definition-select{' in stylesheet

    app_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')
    assert "activeCell.textContent = operation ? `${operation}(${selected[0]})` : selected[0];" in app_script
    assert "if (editedCell?.dataset.catalogueField === 'CDR source' && !assistanceFocused) normaliseCatalogueRows();" in app_script
    assert "const sessionMarker = document.body.dataset.authenticatedSession || 'anonymous';" in app_script
    assert "${sessionScoped ? `:${sessionMarker}` : ''}" in app_script

    dashboard_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/e2e_dashboards.js').read_text(encoding='utf-8')
    assert "const authenticatedSession = document.body.dataset.authenticatedSession || 'anonymous';" in dashboard_script
    assert ':open:${authenticatedSession}`' in dashboard_script
    assert ':filters-open:${authenticatedSession}`' in dashboard_script
    assert "sessionStorage.setItem(filtersOpenStorageKey, isOpen ? 'open' : 'closed')" in dashboard_script
    assert 'await openDashboard(last, {showFilters: rememberedFiltersOpen()});' in dashboard_script


def test_template_visual_controls_are_available_for_every_chart_type():
    root = Path(__file__).parents[1]
    app_script = (root / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')
    reporting_template = (root / 'src/web_interface/templates/reporting.html').read_text(encoding='utf-8')
    chart_script = (root / 'src/web_interface/static/js/dashboard_charts.js').read_text(encoding='utf-8')

    assert 'const syncConditionalVisualCells = (row, {clear = true} = {}) =>' in app_script
    assert "'Axis X Range': true" in app_script
    assert "'Axis Y Range': true" in app_script
    assert "'Label Position': true" in app_script
    assert "'Label Format': true" in app_script
    assert "'Legend Format': true" in app_script
    assert "cell.contentEditable = enabled ? 'true' : 'false';" in app_script
    assert "if (!enabled && clear && cell.textContent.trim()) cell.textContent = '';" in app_script
    assert 'const syncConditionalVisualControls = () =>' in app_script
    assert "axis_x_range: true" in app_script
    assert "axis_y_range: true" in app_script
    assert "label_position: true" in app_script
    assert "if (control.name === 'chart_type') syncConditionalVisualControls();" in app_script
    assert "['axis_x_range', 'Axis X Range']" in reporting_template
    assert "['label_position', 'Label Position']" in reporting_template
    assert "['label_format', 'Label Format']" in reporting_template
    assert "['legend_format', 'Legend Format']" in reporting_template
    assert "['exclude_null_empty', 'Exclude Null/Empty']" in app_script
    assert "['exclude_zero', 'Exclude Zero']" in app_script
    assert 'function drawConfiguredPointLabel(' in chart_script
    assert "drawConfiguredPointLabel(context, series.name" in chart_script


def test_saving_the_embedded_template_rebuilds_the_dashboard_immediately():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/e2e_dashboards.js').read_text(encoding='utf-8')

    assert 'const rebuildDashboardAfterTemplateSave = async () =>' in script
    assert "const expandedChartIndex = !$('ds-chart-expanded-overlay').hidden" in script
    assert "await openExpandedChart(refreshedChart, null, 'dashboard', {preserveFocus: true});" in script
    assert 'function expandedChartOverlay(show, {preserveFocus = false} = {})' in script
    assert "if (event.data?.type === 'dashboard-analytic:template-saved')" in script
    assert 'await rebuildDashboardAfterTemplateSave();' in script
    assert 'templateEditorSaved = false;' in script


def test_compact_landscape_dashboard_comments_are_docked_to_the_bottom():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')

    assert '@media (orientation:landscape) and (max-height:600px) {' in stylesheet
    assert '#ds-viewer.ds-overlay {' in stylesheet
    assert 'grid-template-columns:minmax(10rem,1fr) auto;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-content {' in stylesheet
    assert 'flex-direction:column;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-main {' in stylesheet
    assert 'grid-template-rows:auto minmax(0,1fr);' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-positioned .ds-chart {' in stylesheet
    assert 'left:var(--ds-chart-left)!important;' in stylesheet
    assert 'height:var(--ds-chart-height)!important;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-panel {' in stylesheet
    assert 'padding:.5rem .65rem 3.75rem!important;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-content>.ds-slide-comments {' in stylesheet
    assert 'position:fixed;' in stylesheet
    assert 'bottom:max(.35rem,env(safe-area-inset-bottom,0px));' in stylesheet
    assert 'left:max(.35rem,env(safe-area-inset-left,0px));' in stylesheet
    assert 'right:max(.35rem,env(safe-area-inset-right,0px));' in stylesheet
    assert 'max-height:min(70dvh,16rem);' in stylesheet


def test_dashboard_job_panels_use_the_stack_spacing_without_an_empty_filter_row():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')

    assert '#ds-filter-home:not(:has(> #ds-filter-panel:not([hidden]))) {' in stylesheet
    assert '.e2e-dashboards :is(.ds-ppt-jobs-panel,.ds-ppt-charts-panel) {' in stylesheet
    assert 'margin-top:0;' in stylesheet


def test_dashboard_discard_restores_saved_filters_and_universe_before_navigation():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/e2e_dashboards.js').read_text(encoding='utf-8')

    assert "const discardPart = (part) => {" in script
    assert "copyDefinitionFields(definition, savedDashboardDefinition(), fields);" in script
    assert "if (part === 'universe') rememberUniverse();" in script
    assert "else if (choice === 'secondary') discardPart('filters');" in script
    assert "else if (choice === 'secondary') discardPart('universe');" in script


def test_dashboard_prepare_retries_one_transient_proxy_501_response():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/e2e_dashboards.js').read_text(encoding='utf-8')

    assert "path.startsWith('/prepare') && response.status === 501" in script
    assert "await new Promise(resolve => window.setTimeout(resolve, 250));" in script


def test_auto_calculated_field_editor_uses_wide_content_aware_dialog_geometry():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/app.css').read_text(encoding='utf-8')
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')

    dialog_rule = stylesheet.split('.confirm-panel.calculated-dimensions-dialog {', 1)[1].split('}', 1)[0]
    assert 'width: 90vw;' in dialog_rule
    assert 'height: fit-content !important;' in dialog_rule
    assert 'max-width: none;' in dialog_rule
    assert 'max-height: calc(100dvh - 2rem) !important;' in dialog_rule
    assert '.confirm-overlay.calculated-dimensions-overlay { display: flex; align-items: center;' in stylesheet
    assert script.count("overlay.className = 'confirm-overlay calculated-dimensions-overlay';") == 2
    assert '.calculated-dimensions-list { display: grid; flex: 0 1 auto;' in stylesheet
    assert '.calculated-dimension-editor { display: grid; flex: 0 1 auto;' in stylesheet
    assert '.calculated-dimension-editor[hidden] { display: none; }' in stylesheet
    assert 'grid-template-rows: auto auto minmax(12rem, 1fr) auto;' in stylesheet
    assert '.calculated-dimensions-manager-actions { display: flex; flex: 0 0 auto; align-items: center; justify-content: flex-end;' in stylesheet
    assert '.calculated-dimensions-manager-actions[hidden] { display: none; }' in stylesheet
    assert '.calculated-dimension-editor > .confirm-actions.calculated-dimension-rules > button { flex: 0 0 auto;' in stylesheet
    assert script.count("save.textContent = 'Save'; save.disabled = true;") == 2
    assert script.count("saveAndMaterialize.textContent = 'Save & Materialize'; saveAndMaterialize.disabled = true;") == 2
    assert script.count("discard.textContent = 'Discard';") == 2
    assert script.count("const apply = document.createElement('button'); apply.type = 'button'; apply.textContent = 'Apply';") == 2
    assert script.count("orderActions.className = 'calculated-dimension-order-actions';") == 2
    assert "const formatCalculatedDimensionAliases = (value) => String(value || '')" in script
    assert script.count('default_from: formatCalculatedDimensionAliases(') == 2
    assert script.count('when: formatCalculatedDimensionRuleAliases(line.slice(0, separator))') == 2
    assert script.count('Default source fields (optional, use [Field A] OR [Field B])') == 2
    assert 'function configureCalculatedDimensionFieldAutocomplete(input, getColumns)' in script
    assert "input.closest('.calculated-dimensions-overlay') || ownerDocument.body" in script
    assert script.count('configureCalculatedDimensionFieldAutocomplete(field, selectedColumns)') == 2
    assert script.count("...(availableDimensionColumns[checkbox.value] || [])") == 2
    assert ".map((alias) => `[${alias}]`)" in script
    assert '.calculated-dimension-field-suggestions { position: fixed;' in stylesheet
    assert script.count('updateSaveActions();\n        });\n        moveDown.addEventListener') == 2
    assert ".auto-calculated-field-job:not([data-materialization-job-key])" in script
    assert "if (status === 'ready') return 100;" in script
    assert '.calculated-dimension-order-actions { display: grid;' in stylesheet
    assert 'body: JSON.stringify({dimensions: next, renames, materialize}),' in script
    assert 'body: JSON.stringify({dimensions: calculatedDimensions, renames, materialize}),' in script
    assert script.count('updateSaveActions();\n        finish();') == 2


def test_dashboard_warmup_retries_contention_and_compact_panel_headers_stay_aligned():
    dashboard_module = (Path(__file__).parents[1] / 'src/modules/e2e_dashboards.py').read_text(encoding='utf-8')
    app_stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/app.css').read_text(encoding='utf-8')
    dashboard_stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')

    assert 'dashboard_warmup_pending: dict[tuple[str, str], tuple[dict, str]] = {}' in dashboard_module
    assert 'def retry_later() -> None:' in dashboard_module
    assert 'core.submit_background_task(delayed_retry)' in dashboard_module
    assert '.collapsible-summary .collapse-chip {' in app_stylesheet
    assert 'position: absolute; top: .7rem; right: .75rem;' in app_stylesheet
    assert '.report-charts-panel .report-charts-chart-count .pill {' in app_stylesheet
    assert '#ds-viewer.ds-presentation-active .ds-viewer-primary-controls>.ds-viewer-tool-actions {' in dashboard_stylesheet
    assert 'visibility:visible!important;' in dashboard_stylesheet


def test_compact_landscape_presentation_settings_are_vertically_scrollable():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')

    assert '#ds-presentation-overlay.ds-overlay {' in stylesheet
    assert '#ds-presentation-overlay .ds-presentation-dialog {' in stylesheet
    assert 'max-height:100%;' in stylesheet
    assert 'overflow-y:auto;' in stylesheet
    assert 'overscroll-behavior:contain;' in stylesheet
    assert '-webkit-overflow-scrolling:touch;' in stylesheet
    assert '#ds-presentation-overlay .ds-presentation-dialog>.section-head {' in stylesheet
    assert 'position:sticky;' in stylesheet
    assert '#ds-presentation-overlay .ds-presentation-options {' in stylesheet
    assert 'grid-template-columns:repeat(2,minmax(0,1fr));' in stylesheet


def test_compact_portrait_dashboard_fills_the_available_viewport():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')

    assert '@media (orientation:portrait) and (max-width:640px) {' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-content {' in stylesheet
    assert 'flex:1 1 auto;' in stylesheet
    assert 'overflow:hidden;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-main {' in stylesheet
    assert 'grid-template-columns:minmax(0,3fr) minmax(0,4fr);' in stylesheet
    assert 'grid-template-rows:auto minmax(0,1fr);' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-chart-stage.ds-positioned,' in stylesheet
    assert 'height:100%;' in stylesheet
    assert 'aspect-ratio:auto;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-positioned .ds-chart {' in stylesheet
    assert 'left:var(--ds-chart-left)!important;' in stylesheet
    assert 'width:var(--ds-chart-width)!important;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-primary-controls {' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-navigation-footer {' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-right-actions {' in stylesheet
    assert 'grid-template-columns:repeat(3,minmax(0,1fr));' in stylesheet


def test_dashboard_end_navigation_controls_use_line_svgs():
    template = (Path(__file__).parents[1] / 'src/web_interface/templates/e2e_dashboards.html').read_text(encoding='utf-8')

    assert 'id="ds-first" class="ds-slide-nav-button" title="First slide" aria-label="First slide"><svg class="ds-slide-arrow-icon"' in template
    assert 'id="ds-last" class="ds-slide-nav-button" title="Last slide" aria-label="Last slide"><svg class="ds-slide-arrow-icon"' in template
    assert 'id="ds-chart-expanded-first" title="First chart" aria-label="First chart"><svg class="ds-chart-expanded-arrow-icon"' in template
    assert 'id="ds-chart-expanded-last" title="Last chart" aria-label="Last chart"><svg class="ds-chart-expanded-arrow-icon"' in template
    assert 'M5 5v14M18 6l-6 6 6 6M12 6l-6 6 6 6' in template
    assert 'M19 5v14M6 6l6 6-6 6M12 6l6 6-6 6' in template
    assert '⏮' not in template
    assert '⏭' not in template


def setup_dashboard(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('sample.csv', BytesIO(b'Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\nA,London,10,HTTP DL,2026-09-01\nB,Leeds,20,HTTP DL,2026-09-02\nA,London,30,HTTP DL,2026-09-03\n'), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('nsa', 'Dashboard test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Aggregation,Column Aggregation,Legend,Legend Position\n'
        '1,Comparison,,Title and 2 columns + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
        '1,Comparison,,Title and 2 columns + Comments,Rate again,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
        '2,Next,,Title and 1 column + Comments,Rate next,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    return definition().model_dump(mode='json')


def test_filter_empty_missing_and_inclusive_dates():
    frame = pd.DataFrame({'City': ['London', 'Leeds'], 'Test_Start_Time': ['2026-09-01 23:59', '2026-09-02']})
    assert filter_frame(frame, definition(filters={'City': []})).empty
    assert filter_frame(frame, definition(filters={'Missing': ['Yes']})).empty
    assert len(filter_frame(frame, definition(date_from='2026-09-01', date_to='2026-09-01'))) == 1
    assert len(filter_frame(frame, definition(filters={'city': ['London']}))) == 1
    radio = pd.DataFrame({'RAT_A': ['LTE', 'NR'], 'technology_primary': ['4G', '5G']})
    assert len(filter_frame(radio, definition(filters={'RAT': ['NR']}))) == 1
    assert filter_frame(radio, definition(filters={'Technology': ['5G']})).empty
    geography = pd.DataFrame({
        'city': ['', None], 'G_Level_4': ['London', 'Leeds'],
        'region': ['', None], 'G_Level_2': ['England', 'England'],
    })
    assert len(filter_frame(geography, definition(filters={'City': ['London'], 'Region': ['England']}))) == 1


def test_filter_frame_ignores_date_range_but_keeps_timestamp_field_filters(monkeypatch):
    monkeypatch.setenv('IGNORE_EVENT_TIME_FILTERING', 'true')
    frame = pd.DataFrame({
        'Event_Start_Time': ['2026-09-01 10:00:00', None],
        'Event_End_Time': ['2026-09-01 11:00:00', None],
        'Value': [1, 2],
    })

    filtered = filter_frame(frame, definition(
        date_from='2026-09-01', date_to='2026-09-01',
        filters={'Event_Start_Time': ['2026-09-01 10:00:00']},
    ))

    assert len(filtered) == 1


def test_dashboard_page_exposes_disabled_event_time_filtering(client, monkeypatch):
    monkeypatch.setenv('IGNORE_EVENT_TIME_FILTERING', 'true')
    setup_dashboard(client)

    response = client.get('/e2e-dashboards')

    assert response.status_code == 200
    assert '"ignore_event_time_filtering": true' in response.text
    assert 'data-ignore-event-time-filtering="true"' in response.text


def test_dashboard_export_uses_the_admin_import_archive_format(client):
    payload = setup_dashboard(client)
    dashboard_id = 'portable-dashboard'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200

    exported = client.get(f'/api/e2e-dashboards/{dashboard_id}/export')

    assert exported.status_code == 200
    with zipfile.ZipFile(BytesIO(exported.content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        assert manifest['format'] == 'dashboard-analytic-export'
        assert manifest['kind'] == 'dashboards'
        assert manifest['workspace_components'] == ['dashboards']
        document = json.loads(archive.read(manifest['archive_path']))
    assert document['format'] == 'dashboard-analytic-dashboards'
    assert document['dashboards'] == {dashboard_id: payload}


def test_dashboard_persists_dataset_universe_separately_from_filters(client):
    payload = setup_dashboard(client)
    payload['date_from'], payload['date_to'] = 'Oldest', 'Newest'

    saved = client.put('/api/e2e-dashboards/symbolic-dates', json=payload)

    assert saved.status_code == 200, saved.text
    assert {key: saved.json()['definition'][key] for key in ('scope', 'datasets', 'date_from', 'date_to')} == {
        key: payload[key] for key in ('scope', 'datasets', 'date_from', 'date_to')
    }
    stored = client.get('/api/e2e-dashboards').json()['symbolic-dates']
    assert {key: stored[key] for key in ('scope', 'datasets', 'date_from', 'date_to')} == {
        key: payload[key] for key in ('scope', 'datasets', 'date_from', 'date_to')
    }
    prepared = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['date_bounds'] == {'min': '2026-09-01', 'max': '2026-09-03'}
    assert prepared.json()['rows']['data'] == 3


def test_dashboard_all_values_expands_when_saved_universe_gains_a_new_value(client):
    payload = setup_dashboard(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    payload['custom_fields'] = ['Test_Name']
    payload['filters'] = {}
    dashboard_id = 'expand-all-values'
    saved = client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload)
    assert saved.status_code == 200, saved.text
    assert 'Test_Name' not in saved.json()['definition']['filters']

    initial = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert initial.status_code == 200, initial.text
    assert initial.json()['options']['Test_Name'] == ['HTTP DL']
    assert initial.json()['rows']['data'] == 3

    uploaded = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('new-test.csv', BytesIO(
            b'Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\n'
            b'C,Bristol,40,HTTP UL,2026-09-04\n'
        ), 'text/csv'),
    })
    assert uploaded.status_code == 200, uploaded.text
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(2)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(2)['status'] == 'ready'

    expanded = client.get('/api/e2e-dashboards').json()[dashboard_id]
    expanded['datasets']['data'] = [1, 2]
    updated = client.put(f'/api/e2e-dashboards/{dashboard_id}', json=expanded)
    assert updated.status_code == 200, updated.text
    assert 'Test_Name' not in updated.json()['definition']['filters']
    prepared = client.post('/api/e2e-dashboards/prepare', json=updated.json()['definition'])
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['options']['Test_Name'] == ['HTTP DL', 'HTTP UL']
    assert prepared.json()['rows']['data'] == 4


def test_dashboard_background_warmup_prepares_the_four_automatic_date_universes(client):
    payload = setup_dashboard(client)
    for index in (2, 3):
        response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
            'dataset_files': (
                f'sample-{index}.csv',
                BytesIO(f'Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\nA,London,{index},HTTP DL,2026-09-0{index}\n'.encode()),
                'text/csv',
            ),
        })
        assert response.status_code == 200
    dashboard_id = 'four-universes'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200

    deadline = time.monotonic() + 15
    dashboard_status = {}
    while time.monotonic() < deadline:
        dashboard_status = client.get('/api/e2e-dashboards/statuses').json()[dashboard_id]
        if dashboard_status['state'] == 'ready':
            break
        assert dashboard_status['state'] == 'loading-data'
        assert dashboard_status['label'].startswith('Preparing ')
        time.sleep(0.05)

    assert dashboard_status == {'state': 'ready', 'label': 'Ready'}
    manifests = [
        json.loads(path.read_text(encoding='utf-8'))
        for path in (Path(core.repository.db_path).parent / '.dashboard-data-cache' / 'dashboard-previews').glob('*.json')
    ]
    definitions = [manifest['definition'] for manifest in manifests if manifest.get('dashboard_id') == dashboard_id]
    assert {tuple(definition['datasets']['data']) for definition in definitions} == {
        (3, 2, 1), (3,), (3, 2), (3, 1),
    }
    assert {
        tuple(definition['datasets']['data']): (definition['date_from'], definition['date_to'])
        for definition in definitions
    } == {
        (3, 2, 1): ('2026-09-01', '2026-09-03'),
        (3,): ('2026-09-03', '2026-09-03'),
        (3, 2): ('2026-09-02', '2026-09-03'),
        (3, 1): ('2026-09-01', '2026-09-03'),
    }


def test_dashboard_template_change_preserves_saved_universe_and_filters(client):
    payload = setup_dashboard(client)
    core.repository.add_report_template('sa', 'Dashboard SA test', core.repository.report_template_content(
        'nsa', 'Dashboard test',
    ), is_default=False)
    payload.update({
        'scope': 'multivendor',
        'datasets': {'data': [1], 'voice': [], 'speech': []},
        'filters': {'City': ['London']},
        'custom_fields': ['Mean_Data_Rate'],
        'hidden_filters': ['Vendor'],
        'date_from': '2026-09-01',
        'date_to': '2026-09-03',
        'slide_comments': {'1': ['Keep this note']},
    })
    assert client.put('/api/e2e-dashboards/change-template', json=payload).status_code == 200
    saved_before_change = client.get('/api/e2e-dashboards').json()['change-template']

    changed = client.patch('/api/e2e-dashboards/change-template/template', json={
        'technology': 'sa', 'template': 'Dashboard SA test',
    })

    assert changed.status_code == 200, changed.text
    definition_payload = changed.json()['definition']
    assert changed.json()['invalidated'] is True
    assert definition_payload['technology'] == definition_payload['template_technology'] == 'sa'
    assert definition_payload['template'] == 'Dashboard SA test'
    for field in ('scope', 'datasets', 'filters', 'custom_fields', 'hidden_filters', 'date_from', 'date_to', 'slide_comments'):
        assert definition_payload[field] == saved_before_change[field]


def test_dashboard_library_ppt_scope_builds_its_automatic_dataset_universe(client):
    payload = setup_dashboard(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'], 'City': ['London', 'Leeds', 'London'],
        'Mean_Data_Rate': [10, 20, 30], 'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
        'Event_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    dashboard_id = 'library-ppt-scope'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200

    scope_only = {**payload, 'scope': 'single', 'datasets': {}, 'date_from': None, 'date_to': None}
    prepared = client.post(
        f'/api/e2e-dashboards/prepare?dashboard_id={dashboard_id}&use_scope_universe=1', json=scope_only,
    )

    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['rows']['data'] == 3
    assert prepared.json()['date_bounds'] == {'min': '2026-09-01', 'max': '2026-09-03'}


def test_dashboard_library_geography_options_are_loaded_in_one_request(client):
    payload = setup_dashboard(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'

    response = client.post('/api/e2e-dashboards/geography-options', json=payload)

    assert response.status_code == 200, response.text
    assert response.json() == {
        'operators': ['A', 'B'], 'vendors': [], 'regions': [], 'cities': ['Leeds', 'London'],
    }


def test_dashboard_filter_catalogue_includes_values_beyond_legacy_profile_limit(client, monkeypatch):
    setup_dashboard(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    cities = [f'City {index:03d}' for index in range(225)]
    rows = '\n'.join(
        f'A,{city},{index + 1},HTTP DL,2026-09-01'
        for index, city in enumerate(cities)
    )
    csv_content = ('Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\n' + rows + '\n').encode()
    uploaded = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('many-cities.csv', BytesIO(csv_content), 'text/csv'),
    })
    assert uploaded.status_code == 200
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(2)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(2)['status'] == 'ready'
    processed_options = json.loads(core.repository.get_dataset(2)['filter_options_json'])
    assert len(processed_options['city']) == 225
    assert core.repository.list_distinct_dataset_row_values(2, 'City', limit=None) == cities
    core.repository.replace_reporting_rows(2, 'data', pd.read_csv(BytesIO(csv_content)))
    with core.repository.connection() as connection:
        connection.execute(
            'UPDATE dataset_profiles SET filter_options_json = ? WHERE dataset_id = 2',
            (json.dumps({'city': cities[:50]}),),
        )
    monkeypatch.setattr('src.modules.e2e_dashboards.DASHBOARD_PROFILE_SELECTION_THRESHOLD', 1)
    payload = definition().model_dump(mode='json')
    payload['datasets'] = {'data': [2]}

    prepared = client.post('/api/e2e-dashboards/prepare', json=payload)

    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['options']['City'] == cities


def test_dashboard_ppt_filename_summarizes_complete_geography_selections(client, monkeypatch):
    payload = setup_dashboard(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        if core.repository.get_dataset(1)['status'] == 'ready':
            break
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    dashboard_id = 'geography-ppt-name'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=['North', 'South'], cities=['Leeds', 'London'],
    )
    monkeypatch.setattr(core, 'submit_background_task', lambda *args, **kwargs: None)

    def queued_name(regions, cities, operators=None):
        filters = {'Region': regions, 'City': cities}
        if operators is not None:
            filters['Operator'] = operators
        export_definition = {**payload, 'filters': filters}
        response = client.post(f'/api/e2e-dashboards/{dashboard_id}/export-ppt', json={
            'definition': export_definition,
            'selected_regions': regions,
            'selected_cities': cities,
        })
        assert response.status_code == 202, response.text
        with core.repository.connection() as connection:
            row = connection.execute(
                'SELECT output_file FROM dashboard_ppt_jobs WHERE id = ?',
                (response.json()['job_id'],),
            ).fetchone()
        return row['output_file']

    assert re.fullmatch(
        r'\d{8}_\d{6} - Comparison - Operator Comparison - All Regions\.pptx',
        queued_name(['South', 'North'], ['London', 'Leeds']),
    )
    first_north_export = queued_name(['North'], ['London'])
    assert re.fullmatch(
        r'\d{8}_\d{6} - Comparison - Operator Comparison - North\.pptx',
        first_north_export,
    )
    second_north_export = queued_name(['North'], ['London'], ['A'])
    assert re.fullmatch(
        r'\d{8}_\d{6} - Comparison - Operator Comparison - North\.pptx',
        second_north_export,
    )
    assert second_north_export != first_north_export


def test_dashboard_ppt_filename_summarizes_partial_selections_and_fits_filesystem(client, monkeypatch):
    payload = setup_dashboard(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    payload['name'] = 'Validation Dashboard ' + 'Long Name ' * 9
    dashboard_id = 'compact-ppt-name'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200
    core.repository.replace_cdr_catalogue(
        1,
        vendors=['Vendor A', 'Vendor B', 'Vendor C'],
        regions=['North', 'South', 'East'],
        cities=['Belfast', 'Bristol', 'Cardiff', 'Edinburgh', 'Leeds', 'London', 'Sheffield', 'York'],
    )
    monkeypatch.setattr(core, 'submit_background_task', lambda *args, **kwargs: None)
    selections = {
        'selected_operators': ['A'],
        'selected_vendors': ['Vendor A', 'Vendor B'],
        'selected_regions': ['North', 'South'],
        'selected_cities': ['Belfast', 'Bristol', 'Cardiff', 'Edinburgh', 'Leeds', 'London', 'Sheffield'],
    }
    queued = client.post(f'/api/e2e-dashboards/{dashboard_id}/export-ppt', json={
        'definition': payload, **selections,
    })
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    with core.repository.connection() as connection:
        row = connection.execute(
            'SELECT output_file, output_path, filters_json FROM dashboard_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()
    output_path = Path(row['output_path'])
    assert row['output_file'].endswith(
        ' - Operator Comparison - North + South.pptx'
    )
    assert len(output_path.parent.name.encode('utf-8')) <= 240
    assert len(output_path.name.encode('utf-8')) <= 255
    output_path.parent.mkdir(parents=True)
    assert 'Vendor: Vendor A, Vendor B' in json.loads(row['filters_json'])
    assert 'Region: North, South' in json.loads(row['filters_json'])
    assert 'City: Belfast, Bristol, Cardiff, Edinburgh, Leeds, London, Sheffield' in json.loads(row['filters_json'])

    with core.repository.connection() as connection:
        connection.execute('UPDATE dashboard_ppt_jobs SET status = ? WHERE id = ?', ('failed', job_id))
    relaunched = client.post(f'/api/e2e-dashboards/ppt-jobs/{job_id}/retry')
    assert relaunched.status_code == 202, relaunched.text
    with core.repository.connection() as connection:
        retried = connection.execute(
            'SELECT output_file, filters_json FROM dashboard_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()
    assert retried['output_file'].endswith(
        ' - Operator Comparison - North + South.pptx'
    )
    assert 'City: Belfast, Bristol, Cardiff, Edinburgh, Leeds, London, Sheffield' in json.loads(retried['filters_json'])

    unicode_definition = {**payload, 'name': 'É' * 120}
    unicode_job = client.post(f'/api/e2e-dashboards/{dashboard_id}/export-ppt', json={
        'definition': unicode_definition, **selections,
    })
    assert unicode_job.status_code == 202, unicode_job.text
    with core.repository.connection() as connection:
        unicode_row = connection.execute(
            'SELECT output_path FROM dashboard_ppt_jobs WHERE id = ?', (unicode_job.json()['job_id'],),
        ).fetchone()
    unicode_path = Path(unicode_row['output_path'])
    assert len(unicode_path.parent.name.encode('utf-8')) <= 240
    unicode_path.parent.mkdir(parents=True)


def test_dashboard_ppt_dialog_selections_override_saved_dashboard_filters(client, monkeypatch):
    payload = setup_dashboard(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    dashboard_id = 'dialog-ppt-filters'
    payload['filters'] = {
        'Operator': ['A'], 'Vendor': ['Vendor A'],
        'Region': ['North'], 'City': ['London'],
    }
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=['North', 'South'], cities=['Leeds', 'London'],
    )
    submitted = []
    monkeypatch.setattr(core, 'submit_background_task', lambda *args, **kwargs: submitted.append(args))

    queued = client.post(f'/api/e2e-dashboards/{dashboard_id}/export-ppt', json={
        'definition': payload,
        'selected_operators': ['B'],
        'selected_vendors': ['Vendor B'],
        'selected_regions': ['South'],
        'selected_cities': ['Leeds'],
    })

    assert queued.status_code == 202, queued.text
    job = next(item for item in client.get('/api/e2e-dashboards/ppt-jobs').json()['jobs'] if item['id'] == queued.json()['job_id'])
    assert job['filters'][-4:] == [
        'Operator: B', 'Vendor: Vendor B', 'Region: South', 'City: Leeds',
    ]
    assert submitted[0][7].name.endswith(' - Operator Comparison - South.pptx')
    assert submitted[0][5].filters == {
        'Operator': ['B'], 'Vendor': ['Vendor B'], 'Region': ['South'], 'City': ['Leeds'],
    }
    assert client.get('/api/e2e-dashboards').json()[dashboard_id]['filters'] == payload['filters']


def test_dashboard_ppt_cover_uses_scope_and_catalogue_geography(client):
    payload = setup_dashboard(client)
    core.repository.add_report_template('nsa', 'Structural dashboard', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Aggregation,Column Aggregation,Legend,Legend Position\n'
        '1,Quarterly review,Template subtitle,Title Page,,,,Title Slide,,,,,Top\n'
    ).encode(), is_default=False)
    payload['template'] = 'Structural dashboard'
    dashboard_id = 'structural-ppt-cover'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=['North', 'South'], cities=['Leeds', 'London'],
    )

    queued = client.post(f'/api/e2e-dashboards/{dashboard_id}/export-ppt', json={'definition': payload})
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    while time.monotonic() < deadline:
        job = next(item for item in client.get('/api/e2e-dashboards/ppt-jobs').json()['jobs'] if item['id'] == job_id)
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.05)
    assert job['status'] == 'ready', job
    assert 'Operator: All Operators' in job['filters']
    assert 'Vendor: All Vendors' in job['filters']
    assert 'Region: All Regions' in job['filters']
    assert 'City: All Cities' in job['filters']
    with core.repository.connection() as connection:
        row = connection.execute('SELECT output_path FROM dashboard_ppt_jobs WHERE id = ?', (job_id,)).fetchone()
    slide = Presentation(row['output_path']).slides[0]
    placeholders = {shape.placeholder_format.type: shape for shape in slide.placeholders}
    title = placeholders.get(1) or placeholders[3]
    assert title.text == 'Quarterly review'
    assert placeholders[4].text == 'Operator Comparison'
    details = {shape.name: shape for shape in slide.shapes if shape.name.startswith('dashboard-ppt-')}
    assert details['dashboard-ppt-region'].text == 'All Regions'
    assert details['dashboard-ppt-city'].text == 'All Cities'
    assert details['dashboard-ppt-region'].left == title.left
    assert details['dashboard-ppt-city'].left == title.left
    assert details['dashboard-ppt-region'].text_frame.paragraphs[0].font.color.rgb not in {
        RGBColor(255, 255, 255), RGBColor(255, 255, 0),
    }


def test_expanded_dashboard_chart_apply_builds_a_new_temporary_model(client):
    payload = setup_dashboard(client)
    prepared = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    token = prepared.json()['token']

    original = client.get(f'/api/e2e-dashboards/chart/{token}/0')
    assert original.status_code == 200, original.text
    assert original.headers['cache-control'] == 'no-store'
    assert len(original.json()['series']) == 2
    context = client.get(f'/api/e2e-dashboards/chart/{token}/0/filter-context')
    assert context.status_code == 200, context.text
    assert context.json()['kpi'] == 'Mean_Data_Rate'
    assert context.json()['exclude_null_empty'] is False
    assert context.json()['exclude_zero'] is False

    preview = client.post(f'/api/e2e-dashboards/chart/{token}/0/filter-preview', json={
        'chart_title': 'Filtered operator preview',
        'filters': 'Operator = A',
        'grouping_rows': 'Test_Name × City × Operator',
        'grouping_columns': '',
        'legend': 'Operator',
        'legend_position': 'Right',
    })
    assert preview.status_code == 200, preview.text
    assert preview.json()['title'] == 'Filtered operator preview'
    assert len(preview.json()['series']) == 1
    assert preview.json()['series'][0]['name'] == 'HTTP DL · London · A'
    assert preview.json()['legend']['position'] == 'right'

    retained_kpi = client.post(f'/api/e2e-dashboards/chart/{token}/0/filter-preview', json={
        'kpi': '', 'grouping_rows': 'Operator', 'grouping_columns': '',
    })
    assert retained_kpi.status_code == 200, retained_kpi.text
    assert retained_kpi.json()['metric'] == 'Mean Data Rate'

    updated = client.post(f'/api/e2e-dashboards/chart/{token}/0/update-template', json={
        'chart_title': 'Updated template chart', 'grouping_rows': 'Test_Name × Operator',
        'grouping_columns': 'City', 'legend_position': 'Right',
        'exclude_null_empty': 'Yes', 'exclude_zero': 'Yes',
    })
    assert updated.status_code == 200, updated.text
    entry = core.load_template_catalogue(core.repository.report_template_content('nsa', 'Dashboard test'), 'nsa')[0]
    assert (entry.chart_title, entry.grouping_rows, entry.grouping_columns, entry.legend_position) == (
        'Updated template chart', 'Test_Name × Operator', 'City', 'right',
    )
    assert entry.exclude_null_empty is True
    assert entry.exclude_zero is True
    refreshed_context = client.get(f'/api/e2e-dashboards/chart/{token}/0/filter-context')
    assert refreshed_context.status_code == 200, refreshed_context.text
    assert refreshed_context.json()['chart_title'] == 'Updated template chart'
    assert refreshed_context.json()['grouping_rows'] == 'Test_Name × Operator'
    assert refreshed_context.json()['exclude_null_empty'] is True
    assert refreshed_context.json()['exclude_zero'] is True
    assert updated.json()['updated_at']


def test_expanded_dashboard_chart_apply_can_render_outside_the_proxy_request(client):
    payload = setup_dashboard(client)
    prepared = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    token = prepared.json()['token']

    queued = client.post(
        f'/api/e2e-dashboards/chart/{token}/0/filter-preview?background=true',
        json={'filters': 'Operator = A', 'grouping_rows': 'Operator'},
    )

    assert queued.status_code == 202, queued.text
    assert queued.json()['state'] == 'processing'
    completed = client.get(
        f"/api/e2e-dashboards/chart-preview-jobs/{queued.json()['job_id']}"
    )
    assert completed.status_code == 200, completed.text
    assert completed.json()['series'][0]['name'] == 'A'


def test_update_template_synchronizes_chart_definition_across_live_snapshots(client):
    payload = setup_dashboard(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'], 'City': ['London', 'Leeds', 'London'],
        'Mean_Data_Rate': [10, 20, 30], 'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
        'Event_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    first = client.post('/api/e2e-dashboards/prepare', json=payload)
    second = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert first.status_code == second.status_code == 200
    first_token = first.json()['token']
    second_token = second.json()['token']

    filtered = client.post(
        f'/api/e2e-dashboards/chart/{first_token}/0/update-template',
        json={'filters': 'Operator = A'},
    )
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()['synced_snapshots'] == 2
    second_context = client.get(
        f'/api/e2e-dashboards/chart/{second_token}/0/filter-context'
    )
    assert second_context.status_code == 200, second_context.text
    assert second_context.json()['filters'] == 'Operator = A'

    cleared = client.post(
        f'/api/e2e-dashboards/chart/{first_token}/0/update-template',
        json={'filters': ''},
    )
    assert cleared.status_code == 200, cleared.text
    reopened_context = client.get(
        f'/api/e2e-dashboards/chart/{second_token}/0/filter-context'
    )
    assert reopened_context.status_code == 200, reopened_context.text
    assert reopened_context.json()['filters'] == ''


def test_dashboard_refresh_rebuilds_all_models_and_chart_refresh_rebuilds_only_one(client, monkeypatch):
    payload = setup_dashboard(client)
    preview = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    token = preview['token']
    indexes = [
        chart['index']
        for slide in preview['slides']
        for chart in slide['charts']
        if chart['available']
    ]
    for index in indexes:
        assert client.get(f'/api/e2e-dashboards/chart/{token}/{index}').status_code == 200

    import src.modules.e2e_dashboards as dashboards_module
    original = dashboards_module.catalog_chart_payload
    calls = []

    def tracked(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(dashboards_module, 'catalog_chart_payload', tracked)

    current = indexes[1]
    core.repository.replace_operator_mapping_group(None, 'Alpha', ['A'], '#123456')
    core.repository.replace_operator_mapping_group(None, 'Beta', ['B'], '#654321')
    core.repository.move_chart_mapping_group('operator', 'Beta', 'up')
    refreshed_chart = client.post(f'/api/e2e-dashboards/chart/{token}/{current}/refresh')
    assert refreshed_chart.status_code == 200, refreshed_chart.text
    assert len(calls) == 1
    assert [series['name'] for series in refreshed_chart.json()['series']] == ['Beta', 'Alpha']
    assert {
        series['name']: series['colour'] for series in refreshed_chart.json()['series']
    } == {'Alpha': '#123456', 'Beta': '#654321'}
    assert client.get(f'/api/e2e-dashboards/chart/{token}/{indexes[0]}').status_code == 200
    assert len(calls) == 1

    calls.clear()
    refreshed_dashboard = client.post(f'/api/e2e-dashboards/charts/{token}/refresh')
    assert refreshed_dashboard.status_code == 200, refreshed_dashboard.text
    assert refreshed_dashboard.json() == {'refreshed': len(indexes)}
    assert len(calls) == len(indexes)


def test_dashboard_editing_controls_match_workspace_editor_roles(client):
    controls = (
        'id="ds-edit"',
        'data-workspace-manage-calculated-dimensions>Auto-Calculated Fields',
        'id="ds-chart-expanded-edit"',
        'id="ds-chart-expanded-auto-fields"',
    )
    for role in ('user-viewer', 'user-editor', 'admin', 'super-admin'):
        session_id = f'dashboard-controls-{role}'
        if role != 'super-admin':
            core.repository.create_user(session_id, 'test-password', role)
            account = next(row for row in core.repository.list_users() if row['username'] == session_id)
            core.repository.set_user_workspace_access(int(account['id']), [core.active_workspace.id])
        core.SESSIONS[session_id] = core.SessionUser(username=session_id, role=role)
        client.cookies.set(core.SESSION_COOKIE, session_id)
        page = client.get('/e2e-dashboards')
        assert page.status_code == 200
        for control in controls:
            assert (control in page.text) == (role != 'user-viewer')
        assert ('id="ds-ppt-jobs-delete-all"' in page.text) == (role != 'user-viewer')
        assert ('"can_manage": true' in page.text) == (role != 'user-viewer')
        expected_status = 403 if role == 'user-viewer' else 404
        assert client.post('/api/e2e-dashboards/ppt-jobs/999999/delete').status_code == expected_status
        expected_chart_status = 403 if role == 'user-viewer' else 410
        assert client.post('/api/e2e-dashboards/chart/missing/0/update-template', json={}).status_code == expected_chart_status
        if role != 'user-viewer':
            assert client.post('/api/e2e-dashboards/ppt-jobs/delete-all').status_code == 200


def test_dashboards_lifecycle_and_layout(client):
    payload = setup_dashboard(client)
    page = client.get('/e2e-dashboards')
    assert page.status_code == 200
    assert re.search(r'data-authenticated-session="[^"]+"', page.text)
    assert 'id="page-panel-navigator"' in page.text
    assert 'data-page-panel-navigator-list' in page.text
    assert page.text.index('>Datasets Analysis<') < page.text.index('>E2E Dashboards<') < page.text.index('>E2E Reporting<')
    assert client.get('/e2e-reporting').status_code == 200
    legacy_reporting = client.get('/reporting', follow_redirects=False)
    assert legacy_reporting.status_code == 307
    assert legacy_reporting.headers['location'] == '/e2e-reporting'
    assert 'id="ds-nr-mode"' in page.text
    assert 'id="ds-dashboards-body"' in page.text
    assert '>PowerPoint Generation Jobs<' in page.text
    assert 'id="ds-ppt-jobs-body"' in page.text
    assert '>Charts Panel<' in page.text
    assert 'id="ds-ppt-chart-job"' in page.text
    assert 'id="ds-ppt-charts-body"' in page.text
    assert 'id="ds-ppt-chart-viewer"' not in page.text
    assert page.text.count('id="ds-chart-expanded-overlay"') == 1
    assert page.text.index('<th>Status</th>') < page.text.index('<th>Actions</th>')
    assert 'colspan="5" class="form-note">Loading Dashboards' in page.text
    assert 'id="ds-library"' not in page.text
    assert 'id="ds-save"' in page.text
    assert '>Save Filters<' in page.text
    assert page.text.index('id="ds-unapplied-filters-badge"') < page.text.index('id="ds-unsaved-filters-badge"')
    assert 'id="confirm-secondary"' in page.text
    assert 'id="confirm-tertiary"' in page.text
    assert '"filter_aliases"' in page.text
    assert '"Region": ["Region", "G_Level_2", "G Level 2"]' in page.text
    assert '>Apply Filters<' in page.text
    assert '>Apply Universe</span>' in page.text
    assert '>Save Universe</span>' in page.text
    assert page.text.index('id="ds-apply-universe"') < page.text.index('id="ds-apply-filters"')
    assert page.text.index('>Apply Filters<') < page.text.index('>Save Filters<')
    assert 'id="ds-apply-filters" title="Apply the current Dashboard filters without saving them" disabled' in page.text
    assert 'id="ds-refresh"' not in page.text
    assert '>Import Dashboard<' in page.text
    assert 'Total Dashboards: 0' in page.text
    assert '>Dashboard Datasets & Filters<' in page.text
    assert 'id="ds-active-dashboard-heading">Dashboard Filters<' in page.text
    filter_panel = page.text.split('id="ds-filter-panel"', 1)[1].split('>', 1)[0]
    assert ' open' in filter_panel
    assert ' hidden' in filter_panel
    assert 'data-panel-state-storage="session"' in filter_panel
    assert 'id="ds-ppt-dataset-overlay"' in page.text
    assert 'id="ds-ppt-dataset-choices"' in page.text
    assert 'id="ds-ppt-dataset-scope"' in page.text
    assert '>Select Dashboard Datasets Universe<' in page.text
    assert 'id="ds-ppt-date-from"' in page.text
    assert 'id="ds-ppt-date-to"' in page.text
    assert '>Dataset Universe<' in page.text
    assert '>Select Dataset Universe<' in page.text
    assert '>Select Dataset Filters<' in page.text
    assert page.text.index('>Select Dataset Universe<') < page.text.index('>Dashboard Scope<')
    assert page.text.index('>Select Dataset Filters<') < page.text.index('>Default Filters<')
    assert '>Dashboard Scope<' in page.text
    assert '>Select Comparison Scope<' in page.text
    assert page.text.index('class="ds-scope-control"') < page.text.index('id="ds-sources"')
    assert '>Default Filters<' in page.text
    assert '>Additional Filters<' in page.text
    assert '>Clear Filters<' in page.text
    assert '>Reload Saved Filters<' in page.text
    assert 'id="ds-dashboard-name"' not in page.text
    assert 'title="Save and apply the current Dashboard filters"' in page.text
    assert 'title="Remove all filter restrictions"' in page.text
    assert 'title="Restore filters and additional fields from the last saved Dashboard"' in page.text
    assert 'id="ds-preparing-rows"' in page.text
    assert 'class="form-note ds-filter-help"' in page.text
    assert 'id="ds-default-facets"' in page.text
    assert 'id="ds-additional-facets"' in page.text
    assert 'Select field to add new filter' in page.text
    assert 'id="ds-custom-field" multiple size="1" data-multiselect-single="true"' in page.text
    assert 'hidden_filters' in DashboardDefinition.model_fields
    assert 'slide_comments' in DashboardDefinition.model_fields
    assert 'id="ds-view" class="ds-view-dashboard-action" title="Open the Dashboard viewer" disabled' in page.text
    assert 'id="ds-generate-ppt" class="ds-generate-ppt-action" title="Generate a PowerPoint presentation for this Dashboard" disabled' in page.text
    assert page.text.index('id="ds-generate-ppt"') < page.text.index('id="ds-view"')
    assert 'id="ds-preparing"' in page.text
    assert 'id="ds-preparing-title"' in page.text
    assert '>Preparing Dashboard dataset<' in page.text
    assert 'id="ds-preparing-progress"' in page.text
    assert 'id="ds-preparing-progress-bar"' in page.text
    assert 'id="ds-preparing-progress-label">0%</span>' in page.text
    assert 'id="ds-preparing-refresh"' in page.text
    assert 'id="ds-preparing-out-of-sync"' in page.text
    assert 'id="ds-viewer-preparing"' in page.text
    assert 'id="ds-viewer-preparing-progress"' in page.text
    assert 'id="ds-viewer-preparing-detail"' in page.text
    assert '>Auto-Calculated Fields<' in page.text
    assert 'class="ds-viewer-icon-action ds-viewer-refresh-action"' in page.text
    assert 'class="ds-viewer-tool-actions" role="group" aria-label="Dashboard actions"' in page.text
    assert 'id="ds-viewer-close" class="report-chart-viewer-close" title="Close"' in page.text
    assert page.text.index('id="ds-floating-filters"') < page.text.index('id="ds-viewer-export-ppt"')
    assert page.text.index('id="ds-viewer-export-ppt"') < page.text.index('id="ds-viewer-refresh"')
    assert page.text.index('id="ds-viewer-refresh"') < page.text.index('id="ds-presentation"')
    assert 'id="ds-viewer-export-ppt" class="ds-viewer-top-action ds-viewer-top-ppt"' in page.text
    assert '>Generate PPT</button>' in page.text
    assert 'id="ds-viewer-refresh" class="ds-viewer-top-action ds-viewer-top-refresh"' in page.text
    assert '>Refresh Dashboard</button>' in page.text
    assert 'id="ds-floating-filters" class="ds-viewer-top-action ds-viewer-top-filters"' in page.text
    assert 'class="ds-viewer-top-action ds-viewer-top-auto-fields" data-workspace-manage-calculated-dimensions' in page.text
    assert 'id="ds-edit" class="ds-viewer-top-action ds-viewer-top-edit"' in page.text
    assert page.text.index('id="ds-presentation"') < page.text.index('id="ds-first"')
    assert '>Dashboard Filters</button>' in page.text
    assert 'id="ds-prev" class="ds-slide-nav-button" title="Previous slide" aria-label="Previous slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-next" class="ds-slide-nav-button" title="Next slide" aria-label="Next slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-first" class="ds-slide-nav-button" title="First slide" aria-label="First slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-last" class="ds-slide-nav-button" title="Last slide" aria-label="Last slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-overlay"' in page.text
    assert 'class="ds-chart-expanded-meta" id="ds-chart-expanded-meta"' in page.text
    assert 'id="ds-chart-expanded-close" class="report-chart-viewer-close" title="Close"' in page.text
    assert 'id="ds-chart-expanded-canvas"' in page.text
    assert 'id="ds-chart-expanded-data"' in page.text
    assert 'id="ds-chart-expanded-zoom"' in page.text
    assert 'id="ds-chart-expanded-filters"' in page.text
    assert 'id="ds-chart-expanded-edit"' in page.text
    assert 'id="ds-chart-expanded-first"' in page.text
    assert 'id="ds-chart-expanded-last"' in page.text
    assert 'id="ds-chart-expanded-pan-left"' in page.text
    assert 'id="ds-chart-expanded-pan-right"' in page.text
    assert 'id="ds-chart-expanded-first" title="First chart" aria-label="First chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-prev" title="Previous chart" aria-label="Previous chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-next" title="Next chart" aria-label="Next chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-last" title="Last chart" aria-label="Last chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert '⏮' not in page.text
    assert '⏭' not in page.text
    assert 'id="ds-chart-expanded-canvas-shell"' in page.text
    assert 'id="ds-chart-expanded-position"' in page.text
    assert 'id="ds-data-table"' in page.text
    assert '<span>Export CSV</span>' in page.text
    assert 'class="ds-chart-expanded-footer"' in page.text
    assert 'id="ds-chart-expanded-controls"' in page.text
    assert 'id="ds-chart-expanded-controls"><button type="button" id="ds-chart-expanded-data"' in page.text
    assert 'data-dashboard-ppt-chart-filter="nr_mode"' in page.text
    assert 'data-dashboard-ppt-chart-filter="dashboard"' in page.text
    assert 'data-dashboard-ppt-chart-filter="template"' in page.text
    assert 'data-dashboard-ppt-chart-filter="scope"' in page.text
    assert 'id="ds-ppt-charts-filters"' in page.text
    assert '>View Filters</button>' in page.text
    assert 'id="ds-ppt-chart-job-picker"' in page.text
    assert 'id="ds-multivendor-overlay"' not in page.text
    assert '>Keep Current Selection</button>' not in page.text
    assert '>Use Selected Datasets</button>' not in page.text
    assert '>Use Latest Datasets</button>' not in page.text
    assert 'ds-viewer-refresh-action' in page.text
    dashboard_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/e2e_dashboards.js').read_text(encoding='utf-8')
    assert "controls.append(data, expand, zoom)" in dashboard_script
    assert "savedDefinition = definitionFingerprint(savedRuntimeDashboardDefinition(dashboards[id])); dirty = false;" in dashboard_script
    assert "await warmDashboardModels();" not in dashboard_script
    assert "Rendering Dashboard Charts" in dashboard_script
    assert "overlay('ds-viewer', true);\n    if (needsPreparation) await prepare();" in dashboard_script
    assert "window.addEventListener('beforeunload', rememberScroll);" in dashboard_script
    assert "document.addEventListener('visibilitychange'" in dashboard_script
    assert "document.documentElement.scrollHeight - window.innerHeight" in dashboard_script
    assert "const restorePageState = navigationEntry?.type === 'reload';" in dashboard_script
    assert "if (!restorePageState) resetScroll();" in dashboard_script
    assert "if (restorePageState) restoreScroll(); else resetScroll();" in dashboard_script
    assert "sessionStorage.removeItem(scrollStorageKey);" in dashboard_script
    assert "last = sessionStorage.getItem(openStorageKey) || '';" in dashboard_script
    assert 'const restoreOpenDashboard = restorePageState;' not in dashboard_script
    assert "if (dashboards[last]) await openDashboard(last, {showFilters: rememberedFiltersOpen()});" in dashboard_script
    assert dashboard_script.index("action('View Dashboard', 'View Dashboard'") < dashboard_script.index("filtersAreOpen ? 'Close Filters' : 'Open Filters'")
    assert "preview_snapshot = replace(" in (Path(__file__).parents[1] / 'src/modules/e2e_dashboards.py').read_text(encoding='utf-8')
    assert "The template owns these required chart attributes." in (Path(__file__).parents[1] / 'src/modules/e2e_dashboards.py').read_text(encoding='utf-8')
    app_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')
    assert 'function setupPagePanelNavigator()' in app_script
    assert "mainMenuLabel.textContent = 'Main Menu';" in app_script
    assert "mainMenuPath.setAttribute('d', 'M3 11.5 12 4l9 7.5M5.5 10v10h13V10M9.5 20v-6h5v6');" in app_script
    assert "window.scrollTo({top: 0, left: 0, behavior: 'smooth'});" in app_script
    assert "'module-tab-e2e-dashboards': 'dashboards'" in app_script
    assert "panel.scrollIntoView({behavior: 'smooth', block: 'start'});" in app_script
    assert "const topLevelPanels = Array.from(main.querySelectorAll('article.panel, details.panel, section.panel'))" in app_script
    assert 'const visiblePanels = topLevelPanels.filter((panel) => (' in app_script
    assert "index === 0 ? eyebrow?.textContent || heading?.textContent" in app_script
    assert 'const explicitLabel = panel.dataset.pagePanelLabel;' in app_script
    assert 'panel.getClientRects().length > 0' in app_script
    assert "attributeFilter: ['hidden', 'class', 'style']" in app_script
    assert 'childList: true' in app_script
    assert "window.addEventListener('resize', scheduleRebuild, {passive: true});" in app_script
    assert "if (opening) rebuild();" in app_script
    assert "document.addEventListener('page-panel-navigation:update', scheduleRebuild);" in app_script
    assert "panel.dataset.pagePanelNavigation !== 'exclude'" in app_script
    assert 'setupPagePanelNavigator();' in app_script
    reporting_template = (Path(__file__).parents[1] / 'src/web_interface/templates/reporting.html').read_text(encoding='utf-8')
    assert 'id="report-charts-panel"' in reporting_template
    assert "document.dispatchEvent(new CustomEvent('page-panel-navigation:update'));" in reporting_template
    dashboard_template = (Path(__file__).parents[1] / 'src/web_interface/templates/e2e_dashboards.html').read_text(encoding='utf-8')
    assert dashboard_template.index('id="ds-apply-filters"') < dashboard_template.index('id="ds-generate-ppt"')
    assert 'id="ds-ppt-charts-panel" open data-panel-state-key="e2e-dashboards:ppt-charts" hidden' in dashboard_template
    assert "$('ds-ppt-charts-panel').hidden = false;" in dashboard_script
    assert "$('ds-ppt-charts-panel').hidden = true;" in dashboard_script
    assert "dashboardFiltersOpen = true;" in dashboard_script
    assert "$('ds-filter-panel').hidden = false;" in dashboard_script
    assert "dashboardFiltersOpen = false;" in dashboard_script
    assert "$('ds-filter-panel').hidden = true;" in dashboard_script
    chart_builder_template = (Path(__file__).parents[1] / 'src/web_interface/templates/chart_builder.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">Chart Builder</p><h2>Ad-Hoc Analysis</h2>' in chart_builder_template
    assert 'data-page-panel-label="Interactive Preview"' in chart_builder_template
    assert '<p class="eyebrow">Chart Builder</p>\n        <h2>Interactive Preview</h2>' in chart_builder_template
    assert '<p class="eyebrow">Chart Definition</p>' in chart_builder_template
    workspace_template = (Path(__file__).parents[1] / 'src/web_interface/templates/workspace.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">Workspaces Management</p>\n        <h2>Select Workspace</h2>' in workspace_template
    datasets_analysis_template = (Path(__file__).parents[1] / 'src/web_interface/templates/datasets_analysis.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">Dataset Analysis</p>\n            <h2>{{ selected_dataset.file_name if selected_dataset else \'Select Dataset\' }}</h2>' in datasets_analysis_template
    app_logs_template = (Path(__file__).parents[1] / 'src/web_interface/templates/app_logs.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">Application activity</p>\n        <h2>App Logs</h2>' in app_logs_template
    documentation_template = (Path(__file__).parents[1] / 'src/web_interface/templates/doc_view.html').read_text(encoding='utf-8')
    assert 'class="panel doc-panel" data-page-panel-label="{{ doc_name }}"' in documentation_template
    assert "const selectAllOrNone = () => {\n      cancelAutoClose();" in app_script
    assert "dispatchNativeChange();\n      menu.hidden = true;\n      syncTrigger();\n      trigger.focus();" not in app_script
    assert 'window.createUnifiedDatasetViewer' in app_script
    assert 'window.createUnifiedDatasetViewer' in dashboard_script
    assert "exportControl: $('ds-data-download')" in dashboard_script
    assert "window.showLoadingOverlay('Loading Filtered Dataset'" in dashboard_script
    assert "await renderData();\n      overlay('ds-data-overlay', true);" in dashboard_script
    app_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/app.css').read_text(encoding='utf-8')
    assert '.page-panel-navigator{position:fixed;' in app_styles
    assert '.page-panel-navigator-tab{' in app_styles
    assert '.page-panel-navigator[data-theme="datasets"]' in app_styles
    assert '.page-panel-navigator[data-theme="dashboards"]' in app_styles
    assert '.page-panel-navigator[data-theme="reporting"]' in app_styles
    assert '.page-panel-navigator .page-panel-main-menu{' in app_styles
    assert '.preview-column-filter-menu { position: fixed; z-index: 10010;' in app_styles
    assert "\\s*×\\s*|\\s+\\bx\\b\\s+" in app_script
    assert "const configuredValues = multiFields.has(key)" in app_script
    assert "columns_by_source" in dashboard_script
    assert "control.dataset.previewDisplay = parsed.value.trim();" in app_script
    assert "control.dispatchEvent(new Event('preview-selection-change'));" in app_script
    assert "would erase text the user has just entered but has not completed." in app_script
    reporting_source = (Path(__file__).parents[1] / 'src/modules/cdr_reporting.py').read_text(encoding='utf-8')
    assert 'This chart type has no interactive renderer' not in reporting_source
    app_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/app.css').read_text(encoding='utf-8')
    dashboard_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')
    assert '.ds-chart-expanded-canvas.ds-hover .ds-chart-filter-panel' in dashboard_styles
    assert '#ds-data-overlay .ds-data-dialog { position: absolute; inset: 5%;' in dashboard_styles
    assert '.ds-chart-filter-panel.is-open {' in dashboard_styles
    assert '.ds-chart-expanded-overlay.ds-overlay { position: fixed; z-index: 9000;' in dashboard_styles
    assert '#ds-viewer.ds-overlay{z-index:8800}' in dashboard_styles
    assert '#ds-data-overlay.ds-overlay, #ds-filter-overlay.ds-overlay, #ds-presentation-overlay.ds-overlay { z-index: 9100 !important; }' in dashboard_styles
    assert '#ds-editor-overlay.ds-overlay { z-index: 9200 !important; }' in dashboard_styles
    assert '.e2e-dashboards .ds-chart-expanded-navigation { gap: 0.65rem; }' in dashboard_styles
    assert '.e2e-dashboards .ds-chart-expanded-arrow-icon {' in dashboard_styles
    assert '.e2e-dashboards .ds-slide-navigation { gap: 0.65rem; }' in dashboard_styles
    assert '.e2e-dashboards .ds-slide-arrow-icon {' in dashboard_styles
    assert '.ds-viewer-tool-actions { display: flex;' in dashboard_styles
    assert '.e2e-dashboards .ds-viewer-tool-actions .ds-viewer-icon-action {' in dashboard_styles
    assert '#ds-viewer.e2e-dashboards .ds-viewer-top-action::before {' in dashboard_styles
    assert '#ds-viewer.e2e-dashboards .ds-viewer-top-filters::before {' in dashboard_styles
    assert '#ds-viewer.e2e-dashboards .ds-viewer-top-auto-fields::before {' in dashboard_styles
    assert '#ds-viewer.e2e-dashboards .ds-viewer-top-edit::before {' in dashboard_styles
    assert 'class="ds-presentation-navigation-separator"' in page.text
    assert '#ds-viewer.e2e-dashboards:not(.ds-presentation-active) .ds-presentation-navigation-separator {' in dashboard_styles
    assert 'border-right:1px solid #b8a8ca;' in dashboard_styles
    assert '.e2e-dashboards .ds-chart-expanded-pan-up {' in dashboard_styles
    assert '.e2e-dashboards .ds-chart-expanded-pan-down {' in dashboard_styles
    assert '.e2e-dashboards .ds-chart-pan-button {' in dashboard_styles
    assert '.ds-chart-expanded-overlay .ds-chart-expanded-dialog { position: relative;' in dashboard_styles
    assert '.ds-chart-expanded-overlay .ds-chart-expanded-canvas { border:' in dashboard_styles
    assert '#ds-chart-expanded-close.report-chart-viewer-close, .e2e-dashboards .ds-viewer-header #ds-viewer-close.report-chart-viewer-close {' in dashboard_styles
    assert '#ds-subtitle{min-height:1.4em}.ds-viewer-panel>.ds-actions{display:grid;' in dashboard_styles
    assert '#ds-subtitle:empty{display:none}' not in dashboard_styles
    assert 'text-transform: uppercase;' in dashboard_styles
    assert 'height: 14rem;' in dashboard_styles
    assert '.multiselect-option[hidden], .multiselect-group-label[hidden] { display: none !important; }' in app_styles
    task_panel_start = app_script.index("const root = document.getElementById('background-task-panels');")
    task_listener = app_script.index("window.addEventListener('dashboard-analytic:background-task'")
    assert task_listener > task_panel_start
    assert "This Auto-calculated Field has unsaved changes. Close without saving them?" in app_script
    assert "const confirmTertiary = document.getElementById('confirm-tertiary');" in app_script
    assert "const handleTertiary = () => close('tertiary');" in app_script
    assert "overlay.addEventListener('click', (event) => { if (event.target === overlay) void requestFinish(); });" in app_script
    assert "window.parent.postMessage({type: 'dashboard-analytic:template-saved'}, window.location.origin);" in app_script
    assert "openTemplateEditor(expandedChart?.focus_row, sourceDefinition)" in dashboard_script
    assert 'const hasAppliedUnsavedFilterChanges = () => Boolean(' in dashboard_script
    assert "confirmLabel: 'Use Current Filters'" in dashboard_script
    assert "secondaryLabel: 'Use Saved Filters'" in dashboard_script
    assert 'wideActions: true' in dashboard_script
    assert 'preparation_token: preparationToken' in dashboard_script
    assert "const preview = node('div', undefined, 'ds-ppt-chart-thumbnail');" in dashboard_script
    assert 'globalThis.renderDashboardChart(canvas, model);' in dashboard_script
    assert 'preview.replaceChildren(cachedImage(chart))' in dashboard_script
    assert 'const rememberRenderedChartPayload = (chart, payload) =>' in dashboard_script
    assert 'rememberRenderedChartPayload(chart, payload);' in dashboard_script
    assert 'if (!live && expandedChartSlide(chart) === prepared?.slides[slideIndex]) renderSlide();' in dashboard_script
    assert "card.ondblclick = safe(async event =>" in dashboard_script
    assert "const syncExpandedChartNavigation" in dashboard_script
    assert 'const setExpandedChartHeader = (chart, title = \'\') =>' in dashboard_script
    assert "$('ds-chart-expanded-meta').textContent" in dashboard_script
    assert "navigateExpandedChart(expandedCharts().length - 1)" in dashboard_script
    assert "const chartPanDirections = ['left', 'right', 'up', 'down'];" in dashboard_script
    assert "expandedPanUp.id = 'ds-chart-expanded-pan-up';" in dashboard_script
    assert "expandedPanDown.id = 'ds-chart-expanded-pan-down';" in dashboard_script
    assert 'bindChartPanControls(expandedCanvas, {' in dashboard_script
    assert 'bindChartPanControls(canvas, panButtons);' in dashboard_script
    assert 'globalThis.panDashboardChart?.(canvas, direction);' in dashboard_script
    assert "visible === 'ds-chart-expanded-overlay' && !editing && event.key === 'ArrowLeft'" in dashboard_script
    assert "visible === 'ds-chart-expanded-overlay' && !editing && event.key === 'ArrowRight'" in dashboard_script
    assert "ds-chart-filter-panel" in page.text
    assert "Chart Definition" in page.text
    assert "/filter-preview" in dashboard_script
    assert "editableGroupingInputs: true" in dashboard_script
    assert "previewDefinition = expandedChartFilterControls.definition()" in dashboard_script
    assert "operation && value ? `${operation}(${value})` : value" in app_script
    assert "operation.dataset.previewKpiAggregation = ''" in app_script
    assert "['COUNTD', 'COUNTD']" in app_script
    assert "id = 'ds-chart-filter-update'" in dashboard_script
    assert "/update-template" in dashboard_script
    assert 'refreshEmbeddedTemplateEditor(updated, expandedChart?.focus_row);' in dashboard_script
    assert "parameters.set('template_revision', String(updatedTemplate.updated_at));" in dashboard_script
    assert 'const renderExpandedChartDefinition = async ({closePanel = true} = {}) =>' in dashboard_script
    assert 'const rendered = await renderExpandedChartDefinition({closePanel: false});' in dashboard_script
    assert dashboard_script.index('const rendered = await renderExpandedChartDefinition({closePanel: false});') < dashboard_script.index("/update-template`, 'POST', rendered.previewDefinition")
    assert "const preparedContext = await api(`${expandedChartFilterContextPath}?prepare=true`);" in dashboard_script
    assert "if (open && !expandedChartFilterControls) await loadExpandedChartFilters();" in dashboard_script
    assert 'const discardExpandedChartDefinitionDraft = () =>' in dashboard_script
    assert "$('ds-chart-filter-close').onclick = discardExpandedChartDefinitionDraft;" in dashboard_script
    assert 'request !== expandedChartFilterLoadRequest' in dashboard_script
    assert "expandedCanvasShell.classList.add('ds-hover')" in dashboard_script
    assert 'const scheduleExpandedChartFiltersClose = () =>' in dashboard_script
    assert '}, 10000);' in dashboard_script
    assert "$('ds-chart-filter-panel').addEventListener('pointerleave', scheduleExpandedChartFiltersClose);" in dashboard_script
    assert "addEventListener('focusout', event =>" not in dashboard_script
    assert "document.addEventListener('pointerdown', event => {" in dashboard_script
    assert "event.target.closest('.report-chart-preview-select-menu')" in dashboard_script
    assert "`Chart ${index + 1} / ${charts.length}`" in dashboard_script
    assert "if (event.target === event.currentTarget) expandedChartOverlay(false);" in dashboard_script
    assert 'const closeOnOutsidePointer = (id, close) =>' in dashboard_script
    assert "closeOnOutsidePointer('ds-filter-overlay', closeFilters);" in dashboard_script
    assert "closeOnOutsidePointer('ds-editor-overlay', closeTemplateEditor);" in dashboard_script
    assert "$('ds-unapplied-filters-badge').hidden = !hasUnappliedFilterChanges();" in dashboard_script
    assert "$('ds-unsaved-filters-badge').hidden = !hasUnsavedFilterChanges();" in dashboard_script
    assert "$('ds-unapplied-universe-badge').hidden = !hasUnappliedUniverseChanges();" in dashboard_script
    assert "$('ds-unsaved-universe-badge').hidden = !hasUnsavedUniverseChanges();" in dashboard_script
    assert "if (!$('ds-filter-overlay').hidden) await closeFilters();" in dashboard_script
    assert "This Dashboard has unsaved changes. Close Adaptative Filters without saving them?" not in dashboard_script
    assert "This Report Template has unsaved changes. Close the editor without saving them?" in dashboard_script
    assert 'const templateChanged = templateEditorSaved;' in dashboard_script
    assert "if (templateChanged && expandedChartMode !== 'ppt') {" in dashboard_script
    assert dashboard_script.count('forgetPrepared();') >= 2
    assert "'Reload the current Report Template, rebuild its slides and render every chart in this Dashboard again?'" in dashboard_script
    assert "event.data?.type === 'dashboard-analytic:template-saved'" in dashboard_script
    chart_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/dashboard_charts.js').read_text(encoding='utf-8')
    assert "const percent = (value, digits = 1) => `${(Number(value) * 100).toFixed(digits)}%`;" in chart_script
    assert 'function drawInsideHorizontalBarLabel(' in chart_script
    assert "const tooltipPercent = value => `${(Number(value) * 100).toFixed(2)}%`;" in chart_script
    assert 'value: tooltipPercent(ratio)' in chart_script
    assert 'lines.push(tooltipPercent(point.cumulative))' in chart_script
    assert "const dynamic = Boolean(payload.dynamic)" in chart_script
    assert "const repeatedHierarchyValue = dynamic" in chart_script
    assert "const boundaryStyle = (level, levels) =>" in chart_script
    assert "left + changedLevel * columnWidth" in chart_script
    assert "const columnKeys = Array.isArray(payload.column_keys)" in chart_script
    assert 'function sameImmediateHierarchyParent(left, right)' in chart_script
    assert 'cell.row >= drag.parent.start && cell.row < drag.parent.end' in chart_script
    assert 'sameImmediateHierarchyParent(drag.sourceKey, targetKey)' in chart_script
    assert 'function chartPanState(canvas)' in chart_script
    assert 'function panChart(canvas, direction)' in chart_script
    assert "['left', 'right', 'up', 'down'].includes(direction)" in chart_script
    assert "if (direction === 'up') camera.panY += LOGICAL_HEIGHT * 0.16;" in chart_script
    assert "if (direction === 'down') camera.panY -= LOGICAL_HEIGHT * 0.16;" in chart_script
    assert 'globalThis.panDashboardChart = panChart;' in chart_script
    assert 'function legendLayout(legend, fontSize = LEGEND_TEXT_SIZE)' in chart_script
    assert "position = items.length ? String(legend?.position || 'none').toLowerCase() : 'none';" in chart_script
    assert "left: position === 'left' ? sideLegendWidth + 20 : 70" in chart_script
    assert "right: position === 'right' ? LOGICAL_WIDTH - sideLegendWidth - 20 : 1540" in chart_script
    assert "top: position === 'top' ? 80 + rows * rowHeight + 18 : 82" in chart_script
    assert "bottom: position === 'bottom' ? 900 - rows * rowHeight - 22 : 860" in chart_script
    assert "function labelAngle(context, value, availableWidth, size)" in chart_script
    assert "function bottomAxisReserve(context, keys, width, size = 22)" in chart_script
    assert "function hierarchyRowLabelSize()" in chart_script
    assert "function drawFullHierarchyLabel(context, value, x, y, width, size" in chart_script
    assert "const aggregationFont = (context, format = {}, level = 0) =>" in chart_script
    assert "bar.colour, 24, payload.label_position" in chart_script
    assert "series.colour, 21, payload.label_position" in chart_script
    assert "bucket.colour, 17, payload.label_position" in chart_script
    assert "drawFullHierarchyLabel(context, value, x + 4" in chart_script
    assert "const rowLabelGap = rowLevels > 1 ? 20 : 0;" in chart_script
    assert "const rowLabelArea = rowLevels ? Math.min(440, Math.max(230" in chart_script
    assert "drawFullHierarchyLabel(context, value, labelLeft + 3" in chart_script
    assert "const usableBottom = layout.position === 'bottom' ? layout.bottom : 884;" in chart_script
    assert 'Math.min(250, columnWidth * .72)' in chart_script
    assert "const colour = payload.cell_colours?.[rowIndex]?.[columnIndex] || '#4E79A7';" in chart_script
    assert "context.textAlign = 'center'; aggregationFont(context, payload.legend_format, level); context.fillText(fittedText(context, value" in chart_script
    assert "aggregationFont(context, payload.legend_format); context.fillText(fittedText(context, String(columnKey.at(-1)" in chart_script
    assert "line(context, left, layout.top + (level + 1) * 28 - 3, right" in chart_script
    assert "const lineTop = layout.top + Math.min(changed, upperLevels) * 28;" in chart_script
    assert "else dashedVertical(context, cellLeft, lineTop, chartTop + chartHeight + 22);" in chart_script
    assert 'function drawOutsideBarLabel(context, value, x, y, colour, size = 12, format = {})' in chart_script
    assert "context.fillStyle = 'rgba(255, 255, 255, 0.94)'" in chart_script
    assert "context.fillText(`${(cumulative * 100).toFixed(0)}%`, plot.left - 14, y - 10)" in chart_script
    assert 'function drawConfiguredBarLabel(' in chart_script
    assert 'labelWidth + 8 > width && size + 8 <= width && labelWidth + 8 <= height' in chart_script
    assert "verticalLabel(context, value, x + width / 2, centreY, format.color || '#FFFFFF', size, false, format);" in chart_script
    assert 'function drawAdjacentStackLabel(' in chart_script
    assert 'sideSpace, colour, occupied, size = 10' in chart_script
    assert "if (ratio > 0) drawAdjacentStackLabel(" in chart_script
    assert 'series.colour, sideLabels' in chart_script
    assert 'Math.floor((layout.right - left) / Math.max(headers.length, 1))' in chart_script
    assert 'function selectionStartAllowed(canvas, event)' in chart_script
    assert 'return logicalY >= 90;' in chart_script
    assert "canvas.classList.toggle('ds-chart-selection-blocked', !selectionStartAllowed(canvas, event));" in chart_script
    assert 'function applyDateBounds(bounds)' in dashboard_script
    assert "button.disabled = Boolean((minimum && value < minimum) || (maximum && value > maximum));" in dashboard_script
    assert 'applyDateBounds(payload.date_bounds)' in dashboard_script
    assert 'function datePicker(key, label)' in dashboard_script
    assert "previous.addEventListener('click', () => { month.setMonth(month.getMonth() - 1); render(); });" in dashboard_script
    assert "button.addEventListener('click', () => { input.value = iso; definition[key] = iso; automatic.setAttribute('aria-pressed', 'false'); updateFilterControlState(wrapper, dateState(key), true); menu.hidden = true; filterChanged(); });" in dashboard_script
    assert 'const current = (hasStoredSelection ? selected : available).filter(Boolean);' in dashboard_script
    assert "const filterControlState = (current, applied, saved, equal = sameFilterValues) =>" in dashboard_script
    assert "if (!equal(current, applied)) return 'unapplied';" in dashboard_script
    assert "return equal(applied, saved) ? '' : 'applied-unsaved';" in dashboard_script
    assert "const sourceState = kind => filterControlState(" in dashboard_script
    assert "const scopeState = () => filterControlState(" in dashboard_script
    assert "const dateState = key => filterControlState(" in dashboard_script
    assert "const hasUnappliedFilterChanges = () =>" in dashboard_script
    assert "filterStateFingerprint(definition) !== filterStateFingerprint(appliedDefinition())" in dashboard_script
    assert "const hasUnappliedUniverseChanges = () =>" in dashboard_script
    assert "universeStateFingerprint(definition) !== universeStateFingerprint(appliedDefinition())" in dashboard_script
    assert 'const resolveUnappliedFilterChanges = async () =>' in dashboard_script
    assert "title: 'Unapplied Dashboard filters'" in dashboard_script
    assert "confirmLabel: 'Apply Filters and Continue'" in dashboard_script
    assert "secondaryLabel: 'Discard and Continue'" in dashboard_script
    assert "tertiaryLabel: 'Save and Continue'" in dashboard_script
    assert "await prepare();\n      return 'applied';" in dashboard_script
    assert "if (hasUnsavedFilterChanges()) await save();" in dashboard_script
    assert "return 'saved';" in dashboard_script
    assert 'restoreAppliedFilterSelection();' in dashboard_script
    assert 'const openActiveDashboardViewer = async () =>' in dashboard_script
    assert "if (!$('ds-filter-overlay').hidden) await closeFilters();" in dashboard_script
    assert 'const needsPreparation = !prepared || preparationStateFingerprint(definition) !== appliedFilterState;' in dashboard_script
    assert "setPreparationState('preparing', needsDataPreparation ? 'data' : 'rendering');" in dashboard_script
    assert "bind('ds-view', openActiveDashboardViewer);" in dashboard_script
    assert "async function queueDashboardPptExport(id, item, {chooseScope = false} = {})" in dashboard_script
    assert "const universeChoice = await chooseDashboardPptUniverse(item, id);" in dashboard_script
    assert "title: 'Choose PowerPoint Scope'" not in dashboard_script
    assert "const host = node('label', label, 'ds-source-filter')" in dashboard_script
    assert "updateFilterControlState(facet, filterState(field));" in dashboard_script
    assert "updateFilterControlState(wrapper, dateState(key), true);" in dashboard_script
    assert 'savedDefinition = definitionFingerprint(result.definition); rememberUniverse(); updateDirtyState(); sources(); facets(); library();' in dashboard_script
    assert 'await prepare();' in dashboard_script
    assert 'if (next.length === current.length && next.every(value => current.includes(value))) return;' in dashboard_script
    assert 'function filterChanged() {' in dashboard_script
    assert "? 'Filter changes are ready to apply.'" in dashboard_script
    assert ": 'All Dashboard filters and Dataset Universe controls are applied.'" in dashboard_script
    assert 'if (preparing && cached && currentFilterState !== preparingFilterState)' not in dashboard_script
    assert "status('Restored the previously prepared filters.');" in dashboard_script
    assert "const payload = await api('/filter-options', 'POST', {definition, field});" in dashboard_script
    assert "facetOptionRequests.has(field) ? 'Loading values…'" in dashboard_script
    assert "const filterAliases = config.filter_aliases || {};" in dashboard_script
    assert "facet.dataset.aliasTooltip = aliasTooltip;" in dashboard_script
    dashboard_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')
    assert '.ds-facet[data-alias-tooltip]::after' in dashboard_styles
    assert 'const preparedPayloadKey = (id, fingerprint) =>' in dashboard_script
    assert 'const preparedStateFingerprint = value =>' in dashboard_script
    assert 'const canonicalSetValues = values =>' in dashboard_script
    assert 'const canonicalSelectionDefinition = value =>' in dashboard_script
    assert 'selection.filters = Object.fromEntries' in dashboard_script
    assert "if (!definitionValue.date_from) definitionValue.date_from = 'Oldest';" in dashboard_script
    assert "if (!definitionValue.date_to) definitionValue.date_to = 'Newest';" in dashboard_script
    assert "const automaticLabel = key === 'date_from' ? 'Use oldest' : 'Use newest';" in dashboard_script
    assert "const automaticValue = key === 'date_from' ? 'Oldest' : 'Newest';" in dashboard_script
    assert 'const eventTimeFilteringDisabled = Boolean(config.ignore_event_time_filtering);' in dashboard_script
    assert "wrapper.classList.add('is-event-time-filtering-disabled');" in dashboard_script
    assert "eventTimeFilteringDisabled && ['eventstarttime', 'eventendtime']" not in dashboard_script
    assert "if (value === 'Oldest') return dateBounds?.min ? `Oldest (${dateBounds.min})` : 'Oldest';" in dashboard_script
    assert "if (value === 'Newest') return dateBounds?.max ? `Newest (${dateBounds.max})` : 'Newest';" in dashboard_script
    assert "input.value = dateInputDisplayValue(key, definition[key]);" in dashboard_script
    assert 'const persistedDashboardDefinition = value => canonicalDashboardDefinition(value);' in dashboard_script
    assert "delete persisted.scope;" not in dashboard_script
    assert "delete persisted.datasets;" not in dashboard_script
    assert 'delete preparedDefinition.name;' in dashboard_script
    assert 'delete preparedDefinition.slide_comments;' in dashboard_script
    assert 'entries.push({fingerprint, token: payload.token});' in dashboard_script
    assert 'const selectionStateFingerprint = value =>' in dashboard_script
    assert 'appliedSelectionState = selectionStateFingerprint(definition);' in dashboard_script
    assert 'const needsDataPreparation = selectionStateFingerprint(definition) !== appliedSelectionState;' in dashboard_script
    assert "params.set('rendering_only', '1');" in dashboard_script
    assert "params.set('preparation_id', preparationToken);" in dashboard_script
    assert "dashboard_name: definition?.name || 'Dashboard'," in dashboard_script
    assert "let preparationToken = '';" in dashboard_script
    assert "window.dispatchEvent(new Event('dashboard-analytic:refresh-background-tasks'));" in dashboard_script
    assert 'let dashboardPptJobsLoaded = false, dashboardPptJobsRefreshing = false;' in dashboard_script
    assert "previous.status !== 'ready'" in dashboard_script
    assert "renderDashboardPptJobs(jobs, newlyReady[0]?.id || '');" in dashboard_script
    assert "const previous = String(preferredJobId || select.value);" in dashboard_script
    dashboard_module = (Path(__file__).parents[1] / 'src/modules/e2e_dashboards.py').read_text(encoding='utf-8')
    assert 'for saved_id, saved_definition in read_dashboards(task_repository).items()' not in dashboard_module
    assert "'dashboard_name': task['name']," in dashboard_module
    assert "'Waiting to prepare Dashboard dataset'" in dashboard_module
    assert "'progress': task.get('progress', 0)," in dashboard_module
    assert "@app.get('/api/e2e-dashboards/preparation-progress/{preparation_id}')" in dashboard_module
    assert "'Counting reduced and filtered Universe rows in the combined CDR tables'" in dashboard_module
    assert "f'Counting Filtered Universe CDR-{kind.title()} rows in the combined table'" in dashboard_module
    assert "progress(78, 'Restoring cached row counts and filter options')" in dashboard_module
    assert "update_preparation_progress(94, 'Saving the reusable Dashboard preparation cache')" in dashboard_module
    assert 'with lock, task_repository.connection() as connection:' not in dashboard_module
    assert "job['progress'] = round(82 + job['completed'] * 18 / max(job['total'], 1))" not in dashboard_module
    assert 'direct_preparation_tasks: dict[str, dict] = {}' in dashboard_module
    assert "preparation_id: str | None = None," in dashboard_module
    assert "'label': 'Rendering Dashboard Charts' if job['total'] else 'Preparing Dashboard dataset'," not in dashboard_module
    assert "phase === 'rendering' ? 'Rendering Dashboard Charts' : 'Preparing Dashboard dataset'" in dashboard_script
    assert "if (preparingFilterState === requestedFilterState) return preparing;" in dashboard_script
    assert "if (backgroundPreparationToken === preparationToken) dismissPreparationStatus();" in dashboard_script
    assert "bind('ds-apply-filters', async () => {" in dashboard_script
    assert "if (!hasUnappliedFilterChanges()) return;" in dashboard_script
    assert "bind('ds-apply-universe', async () => {" in dashboard_script
    assert "if (!hasUnappliedUniverseChanges()) return;" in dashboard_script
    assert "bind('ds-reload-universe', () => {" in dashboard_script
    assert "copyDefinitionFields(definition, savedDashboardDefinition(), universeDefinitionFields);" in dashboard_script
    assert "`/prepare${params.size ? `?${params}` : ''}`,'POST',definition,controller.signal," in dashboard_script
    assert "window.dispatchEvent(new Event('dashboard-analytic:refresh-background-tasks'));" in dashboard_script
    assert "title: 'Unsaved Dashboard filters'" in dashboard_script
    assert "confirmLabel: 'Save Filters'" in dashboard_script
    assert "title: 'Unsaved Dataset Universe'" in dashboard_script
    assert "confirmLabel: 'Save Universe'" in dashboard_script
    assert "if (link.classList.contains('topnav-link-logout'))" in dashboard_script
    assert 'for (const key of [openStorageKey, filtersOpenStorageKey, scrollStorageKey, preparedStorageKey, universeStorageKey])' in dashboard_script
    assert "secondaryLabel: 'Discard'" in dashboard_script
    assert "window.location.assign(target.href);" in dashboard_script
    assert "$('ds-apply-filters').disabled = !hasUnappliedFilterChanges() || filterActionBusy;" in dashboard_script
    assert "$('ds-apply-universe').disabled = !hasUnappliedUniverseChanges() || filterActionBusy;" in dashboard_script
    assert "$('ds-reload-universe').disabled = !hasUnsavedUniverseChanges() || filterActionBusy;" in dashboard_script
    assert "const dashboardStatuses = new Map();" in dashboard_script
    assert 'const dashboardPreparationTokens = new Map();' in dashboard_script
    assert "if (!dashboardPreparationTokens.has(id)) dashboardStatuses.set(id, value);" in dashboard_script
    assert "bind('ds-generate-ppt', async () => {" in dashboard_script
    assert "heading.replaceChildren(document.createTextNode(name ? 'Dashboard Filters: ' : 'Dashboard Filters'));" in dashboard_script
    assert "heading.append(node('span', name.toUpperCase(), 'ds-active-dashboard-name'))" in dashboard_script
    assert 'setActiveDashboardHeading(definition.name);' in dashboard_script
    assert "setActiveDashboardHeading('');" in dashboard_script
    assert "if (activePrepared) dashboardStatuses.set(id, {state: 'ready', label: 'Ready'});" in dashboard_script
    assert "setDashboardStatus(activeId, 'ready', 'Ready');" in dashboard_script
    assert "view.dataset.dashboardViewId = id;" in dashboard_script
    assert 'button.disabled = false;' in dashboard_script
    assert "const setViewEnabled = enabled => { $('ds-view').disabled = !definition; syncDashboardViewActions(); syncDashboardPptActions(); };" in dashboard_script
    assert "const payload = await api('/statuses', 'POST', statusDefinitions);" in dashboard_script
    assert "window.setInterval(refreshDashboardStatuses, 2000);" in dashboard_script
    assert "const dashboardName = String(task.dashboard_name || '');" in app_script
    assert "? `Dashboard “${dashboardName}”: ${String(task.label || 'Background task')}`" in app_script
    assert 'completedByWorkspace.forEach((entries) =>' in app_script
    assert 'entry.group.tasks.splice(Math.min(entry.position + offset, entry.group.tasks.length), 0, entry.task);' in app_script
    selection_key_source = dashboard_module[dashboard_module.index('def persistent_selection_key'):dashboard_module.index('def selected_date_bounds')]
    assert "'scope': definition.scope," not in selection_key_source
    assert "'schema': DASHBOARD_SELECTION_CACHE_VERSION," in selection_key_source
    assert 'DASHBOARD_SELECTION_CACHE_VERSION = 12' in dashboard_module
    assert "kind: sorted([" in selection_key_source
    assert "kind: sorted(set(dataset_ids))" in selection_key_source
    assert "field: sorted(set(values))" in selection_key_source
    assert '{entry_key}:{operator_mapping_key}' in dashboard_module
    assert "'rendering_only': rendering_only," in dashboard_module
    assert 'def materialize_selection(' in dashboard_module
    assert 'use_profile_options=False, known_full_row_counts=None, progress=None,' in dashboard_module
    assert 'use_profile_options=use_profile_options,' in dashboard_module
    assert 'DASHBOARD_CHART_RENDER_WORKERS = max(1, min(2, (os.cpu_count() or 2) - 1))' in dashboard_module
    assert 'DASHBOARD_PREVIEW_MANIFEST_VERSION = 8' in dashboard_module
    assert "thread_name_prefix='e2e-dashboard-chart'," not in dashboard_module
    assert 'def schedule_next_prefetch() -> None:' not in dashboard_module
    assert 'def enqueue_prefetch(' not in dashboard_module
    assert "void api(`/prefetched/${encodeURIComponent(activeId)}/priority`" not in dashboard_script
    assert 'const proximityIndexes = (length, origin) =>' in dashboard_script
    assert 'if (origin + distance < length) indexes.push(origin + distance);' in dashboard_script
    assert 'if (origin - distance >= 0) indexes.push(origin - distance);' in dashboard_script
    assert "charts.map(chart => loadChartPayload(chart, 'low'))" in dashboard_script
    assert "await loadChartPayload(charts[target], 'low').catch(() => undefined);" in dashboard_script
    assert 'scheduleNearbySlidePreload(preloadToken, preloadOrigin, visibleLoads);' in dashboard_script
    assert 'scheduleNearbyExpandedChartPreload(contextKey, chart);' in dashboard_script
    assert 'template_where, template_parameters, template_filters_applied = chart_filter_sql(' in dashboard_module
    assert 'template_filters_applied=template_filters_applied,' in dashboard_module
    assert 'database_path, table_name, source_columns = reporting_source(' in dashboard_module
    assert 'ATTACH DATABASE' not in dashboard_module
    assert 'sqlite3.connect(database_path, timeout=120.0)' in dashboard_module
    assert 'aggregation_columns = chart_aggregation_columns(' in dashboard_module
    assert "thread_name_prefix='e2e-dashboard-data'," not in dashboard_module
    dashboard_css = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')
    assert '.e2e-dashboards :is(.ds-unsaved-filters-badge,.ds-unsaved-universe-badge)' in dashboard_css
    assert '.e2e-dashboards :is(.ds-unapplied-filters-badge,.ds-unapplied-universe-badge)' in dashboard_css
    assert '.ds-scope-control.ds-filter-applied-unsaved select' in dashboard_css
    assert '.ds-date-picker.ds-date-picker-applied-unsaved .ds-date-picker-control>input' in dashboard_css
    assert '.e2e-dashboards .ds-dashboard-close::after' in dashboard_css
    assert '.e2e-dashboards .ds-active-dashboard-name{color:#12664f;font-family:inherit;font-size:1.12em;font-weight:900}' in dashboard_css
    assert '.e2e-dashboards .ds-dashboard-close{background:linear-gradient(135deg,#a9273a,#dc5361)' in dashboard_css
    assert '.e2e-dashboards .ds-dashboard-open{background:linear-gradient(145deg,#e87912,#ffad2f)' in dashboard_css
    assert '.e2e-dashboards .ds-dashboard-view{background:linear-gradient(145deg,#167957,#29ae7d)' in dashboard_css
    assert '.ds-dashboard-status-loading-data{border-color:#d3aa45;background:#fff1c9;color:#77570a}' in dashboard_css
    assert '.ds-dashboard-status-data-queued,.ds-dashboard-status-charts-queued{border-color:#aaa3b2;background:#f0edf2;color:#655e6c}' in dashboard_css
    assert '.ds-filter-groups{display:grid;grid-template-rows:max-content max-content minmax(0,1fr);gap:14px}' in dashboard_css
    assert '.ds-default-filter-panel{border-color:#afd1e8;background:#f1f8fd}' in dashboard_css
    assert "d='M12 2v10'" in dashboard_css
    assert "title: 'Refresh Dashboard?'" in dashboard_script
    assert "api(`/charts/${encodeURIComponent(token)}/refresh`, 'POST')" in dashboard_script
    assert "api(`/chart/${encodeURIComponent(token)}/${chart.index}/refresh`, 'POST')" in dashboard_script
    assert "function resetAutomaticDatesForDatasetChange()" in dashboard_script
    assert "definition[key] = automaticValue;" in dashboard_script
    assert "input.value = dateInputDisplayValue(key, automaticValue);" in dashboard_script
    assert "const latestDatasetsForScope = scope =>" in dashboard_script
    assert ".slice(0, scope === 'multivendor' ? 1 : 2)" in dashboard_script
    assert "datasetRecency(right) - datasetRecency(left)" in dashboard_script
    assert "if (selectedScope !== definition.scope)" in dashboard_script
    assert "definition.datasets = latestDatasetsForScope(selectedScope);" in dashboard_script
    assert "chooseScopeDatasets" not in dashboard_script
    assert 'async function restorePrepared(id) {' in dashboard_script
    assert 'const preparedPayloads = new Map();' in dashboard_script
    assert 'const restoreRememberedPrepared = async id =>' in dashboard_script
    assert 'const inMemory = preparedPayloads.get(preparedPayloadKey(id, fingerprint));' in dashboard_script
    assert 'if (await restorePrepared(activeId))' in dashboard_script
    assert "const payload = await api(`/prepared/${encodeURIComponent(inMemory.payload.token)}`);" in dashboard_script
    assert 'forgetPreparedToken(inMemory.payload.token);' in dashboard_script
    assert 'forgetPrepared();' in dashboard_script
    assert "await prepare();\n      token = prepared?.token || '';" in dashboard_script
    assert "api(`/prefetched/${encodeURIComponent(id)}`, 'POST', definition)" in dashboard_script
    assert 'if (error.status === 409) return false;' in dashboard_script
    assert 'No CDR ${chart.source[0].toUpperCase()}${chart.source.slice(1)} dataset has been selected for this chart.' in dashboard_script
    assert "api(`/prepared/${encodeURIComponent(cachedEntry.token)}`)" in dashboard_script
    assert 'const monitorPreparationProgress = token =>' in dashboard_script
    assert 'const dashboardNeedsRefresh = () => Boolean(' in dashboard_script
    assert "$('ds-preparing-out-of-sync').hidden = !stale;" in dashboard_script
    assert "bind('ds-preparing-refresh', async () => {" in dashboard_script
    assert 'label.textContent = known ? `${detail} · ${percent}%` : detail;' in dashboard_script
    assert "const known = progress !== null && progress !== undefined && progress !== '' && Number.isFinite(numeric);" in dashboard_script
    assert "$('ds-preparing-detail').textContent = detail;" in dashboard_script
    assert "emitPreparationStatus(payload.status || 'processing', payload.detail, token, queued ? null : payload.progress);" in dashboard_script
    assert 'api(`/preparation-progress/${encodeURIComponent(token)}`)' in dashboard_script
    assert 'monitorPreparationProgress(preparationToken);' in dashboard_script
    assert dashboard_script.index('const preparationRequest = api(') < dashboard_script.index('monitorPreparationProgress(preparationToken);')
    assert "dashboard_work_gate.acquire(timeout=0.2)" in dashboard_module
    assert "status(''); await prepare();" in dashboard_script
    assert "setPreparationState('preparing', needsDataPreparation ? 'data' : 'rendering');" in dashboard_script
    assert 'void refreshDashboardStatuses();' in dashboard_script
    assert 'void refreshDashboardPptJobs();' in dashboard_script
    assert 'await Promise.all([refreshDashboardStatuses(), refreshDashboardPptJobs()]);' not in dashboard_script
    assert 'savedDefinition = definitionFingerprint(definition); updateDirtyState();' not in dashboard_script
    assert 'the controls so their state agrees with the Apply and Save buttons.\n    sources(); facets();' in dashboard_script
    assert "setPreparationState('preparing');" in dashboard_script
    assert "bind('ds-refresh',prepare);" not in dashboard_script
    assert 'const setPreparationRows = payload =>' in dashboard_script
    assert 'setPreparationRows(payload);' in dashboard_script
    assert 'applyPreparedPayload(cached.payload);' not in dashboard_script
    assert "['Dataset Universe', universe, 'ds-preparing-universe-label']" in dashboard_script
    assert "['Filtered Universe', filtered, 'ds-preparing-filtered-label']" in dashboard_script
    assert 'rows.hidden = !rows.textContent;' in dashboard_script
    assert 'const zoomResetIcon = () =>' in dashboard_script
    assert "if (request.filter_column) parameters.set('filter_column', request.filter_column);" in dashboard_script
    assert "Clear ${activeFilters} Filter${activeFilters === 1 ? '' : 's'}" in app_script
    assert "window.createUnifiedDatasetViewer" in dashboard_script
    assert "const reset = node('button', undefined, 'ds-chart-zoom-button ds-chart-zoom-reset');" in dashboard_script
    assert "bind('ds-clear-filters', () => { definition.filters = {}; facets(); filterChanged(); });" in dashboard_script
    assert "bind('ds-last-saved-filters', () => {" in dashboard_script
    assert 'id="ds-reload-universe"' in page.text
    assert "definition.datasets = structuredClone(saved.datasets || {});" not in dashboard_script
    assert "definition.scope = saved.scope || 'single';" not in dashboard_script
    assert "definition.filters = structuredClone(saved.filters || {});" in dashboard_script
    assert "definition.custom_fields = structuredClone(saved.custom_fields || []);" in dashboard_script
    assert "definition.hidden_filters = structuredClone(saved.hidden_filters || []);" in dashboard_script
    dashboard_css = (Path(__file__).parents[1] / 'src/web_interface/static/css/e2e_dashboards.css').read_text(encoding='utf-8')
    assert '.ds-dashboard-status-ready{' in dashboard_css
    assert '.ds-dashboard-status-rendering{' in dashboard_css
    assert '.ds-chart-controls button:not(:disabled){cursor:pointer!important}' in dashboard_css
    assert '.ds-preparing .ds-preparing-universe-label{color:#60408d}' in dashboard_css
    assert '.ds-preparing .ds-preparing-filtered-label{color:#d7a921}' in dashboard_css
    assert '.e2e-dashboards .ds-chart-zoom-reset svg' in dashboard_css
    assert '.e2e-dashboards .ds-chart-data{width:2rem;min-width:2rem;height:2rem;min-height:2rem}' in dashboard_css
    assert '.ds-source-filter.ds-filter-unsaved select,.ds-source-filter.ds-filter-unsaved .multiselect-trigger' in dashboard_css
    assert '.ds-date-picker.ds-date-picker-unsaved .ds-date-picker-control>input,.ds-facet.ds-filter-unsaved .multiselect-trigger' in dashboard_css
    assert '.e2e-dashboards .ds-date-auto-action{position:absolute;right:.32rem;top:50%' in dashboard_css
    assert '.e2e-dashboards .ds-generate-ppt-action{' in dashboard_css
    assert ':is(.ds-view-dashboard-action,.ds-generate-ppt-action)::before{width:1.35rem' in dashboard_css
    assert '#ds-apply-filters::before' in dashboard_css
    assert '#ds-save::before' in dashboard_css
    assert '#ds-clear-filters::before' in dashboard_css
    assert '#ds-last-saved-filters::before' in dashboard_css
    assert '.ds-slide-content.ds-comments-right .ds-viewer-navigation-footer {' in dashboard_css
    assert 'grid-column:1 / -1;' in dashboard_css
    assert '.ds-slide-content.ds-comments-right>.ds-slide-comments:not([open]) #ds-comments-title {' in dashboard_css
    assert 'transform:translate(-50%,-50%) rotate(90deg);' in dashboard_css
    assert 'clip-path:polygon(100% 0,32% 50%,100% 100%,76% 100%,8% 50%,76% 0);' in dashboard_css
    assert '.ds-slide-comments,\n.ds-slide-comments>summary,\n.ds-slide-comments-body {\n  background:#c4b0dc;' in dashboard_css
    assert "panel?.classList.toggle('ds-has-comments', comments.length > 0);" in dashboard_script
    assert 'if (panel && presentation.active) panel.open = presentation.showComments && comments.length > 0;' in dashboard_script
    assert 'commentsPanel.open = presentationCommentsWasOpen;' in dashboard_script
    assert '#ds-viewer.ds-presentation-active.ds-presentation-comments-enabled .ds-slide-comments.ds-has-comments {' in dashboard_css
    assert '#ds-viewer.ds-presentation-active .ds-slide-comments { display:none!important; }' in dashboard_css
    assert 'id="ds-presentation-comments"' in page.text
    assert 'id="ds-presentation-delay" type="number" min="1" max="300"' in page.text
    assert 'id="ds-presentation-transition-duration" type="number" min="0.1" max="10" step="0.1" value="1"' in page.text
    assert '<option value="slide-left">Slide from left</option>' in page.text
    assert '<option value="zoom">Zoom</option>' in page.text
    assert '<option value="rise">Rise</option>' in page.text
    assert '<option value="blur">Blur</option>' in page.text
    assert '<option value="rotate">Soft rotate</option>' in page.text
    assert '<option value="flip" selected>Flip</option>' in page.text
    assert '<option value="bounce">Bounce</option>' in page.text
    assert '<option value="wipe">Wipe</option>' in page.text
    assert '<option value="mosaic">Mosaic</option>' in page.text
    assert '<option value="curtain">Curtain</option>' in page.text
    assert '<option value="blinds">Blinds</option>' in page.text
    assert '<option value="random">Random</option>' in page.text
    assert '<option value="none">None</option><option value="random">Random</option><optgroup label="Effects">' in page.text
    assert '<option value="flip" selected>Flip</option>' in page.text
    assert 'id="ds-presentation-stop-viewer"' in page.text
    assert 'id="ds-presentation-toggle-viewer"' in page.text
    assert "presentation.showComments = $('ds-presentation-comments').value === 'yes';" in dashboard_script
    assert "$('ds-viewer').style.setProperty('--ds-presentation-transition-duration', `${presentation.transitionDuration}s`);" in dashboard_script
    assert 'animation:ds-slide-fade var(--ds-presentation-transition-duration,1s) ease both' in dashboard_css
    assert "const randomPresentationEffects = ['fade', 'slide', 'slide-left', 'zoom', 'rise', 'blur', 'rotate', 'flip', 'bounce', 'wipe', 'mosaic', 'curtain', 'blinds'];" in dashboard_script
    assert "if (presentation.effect !== 'random') return presentation.effect;" in dashboard_script
    assert "lastRandomPresentationEffect = choices[Math.floor(Math.random() * choices.length)] || 'fade';" in dashboard_script
    assert "localStorage.getItem(presentationEffectStorageKey)" in dashboard_script
    assert "localStorage.setItem(presentationEffectStorageKey, event.target.value)" in dashboard_script
    assert "lastRandomPresentationEffect = '';\n    slideIndex = 0;" in dashboard_script
    assert "bind('ds-presentation-stop-viewer', stopPresentation);" in dashboard_script
    assert "bind('ds-presentation-toggle-viewer', togglePresentation);" in dashboard_script
    assert "if (presentation.running) pausePresentation(); else resumePresentation();" in dashboard_script
    assert "event.code === 'Space' || event.key === ' '" in dashboard_script
    assert "if (visible === 'ds-viewer' && !editing && event.key === 'F8')" in dashboard_script
    assert 'if (!presentation.active) startPresentation();' in dashboard_script
    assert "presentation.active && event.key === 'F7'" in dashboard_script
    assert "presentation.active && event.key === 'F9'" in dashboard_script
    assert 'if (presentation.active) navigatePresentationSlide(-1);' in dashboard_script
    assert 'if (presentation.active) navigatePresentationSlide(1);' in dashboard_script
    assert '[data-presentation-effect="slide-left"]' in dashboard_css
    assert '[data-presentation-effect="zoom"]' in dashboard_css
    assert '[data-presentation-effect="rise"]' in dashboard_css
    assert '[data-presentation-effect="blur"]' in dashboard_css
    assert ':is(#ds-presentation,#ds-presentation-toggle-viewer,#ds-presentation-stop-viewer)' in dashboard_css
    assert '.ds-presentation-stop-action::after' in dashboard_css
    assert ':is(#ds-presentation-toggle-viewer,#ds-presentation-stop-viewer)[hidden] {' in dashboard_css
    assert "function dashboardViewerBrand(className = 'ds-structural-brand')" in dashboard_script
    assert "const brand = dashboardViewerBrand('ds-chart-brand'); card.append(brand);" in dashboard_script
    assert '.ds-chart-brand{' in dashboard_css
    assert 'position:absolute;z-index:2;top:2.5rem;right:12px;' in dashboard_css
    assert '.ds-chart>img,.ds-chart>canvas{z-index:1}' in dashboard_css
    assert '.ds-chart>.ds-chart-controls{top:calc(2.9rem + 9px)}' in dashboard_css
    assert '.ds-chart>.ds-chart-brand>img{' in dashboard_css
    assert 'position:static!important;inset:auto!important;width:2.65rem!important;height:2.65rem!important;object-fit:contain!important' in dashboard_css
    assert "canvas.addEventListener('dashboardchartlayout', event => {" in dashboard_script
    assert "brand.style.top = `${top + height / 2}px`;" in dashboard_script
    assert 'filter:none' in dashboard_css
    assert 'id="ds-filter-close-action"' in page.text
    assert '#ds-filter-close-action[hidden]{display:none!important}' in dashboard_css
    assert "bind('ds-filter-close-action', closeFilters);" in dashboard_script
    assert "$('ds-generate-ppt').hidden = true; $('ds-view').hidden = true; $('ds-filter-close-action').hidden = false; overlay('ds-filter-overlay', true);" in dashboard_script
    assert "panel.hidden = !dashboardFiltersOpen;" in dashboard_script
    assert '#ds-add-filter::before' in dashboard_css
    assert '#ds-ppt-jobs-delete-all::before' in dashboard_css
    assert '.e2e-dashboards .ds-ppt-charts-filters::before{' in dashboard_css
    assert '#ds-filter-float .ds-filter-help{margin:26px 0 5px}' in dashboard_css
    assert '#ds-filter-float .ds-data-panel>summary{pointer-events:none;cursor:default}' in dashboard_css
    assert "panel.querySelector('summary').tabIndex = -1;" in dashboard_script
    assert '.ds-chart-expanded-canvas.ds-hover .ds-chart-controls,.ds-chart-expanded-canvas:focus-within .ds-chart-controls{opacity:1;visibility:visible;transform:translateY(0);transition-delay:0s;pointer-events:auto}' in dashboard_css
    saved = client.put('/api/e2e-dashboards/test', json=payload)
    assert saved.status_code == 200
    assert saved.json()['definition']['hidden_filters'] == []
    assert saved.json()['definition']['slide_comments'] == {}
    listed = client.get('/api/e2e-dashboards')
    assert listed.json()['test']['name'] == 'Comparison'
    assert listed.headers['cache-control'] == 'no-store, max-age=0, must-revalidate'
    renamed = client.patch('/api/e2e-dashboards/test/name', json={'name': 'Renamed comparison'})
    assert renamed.status_code == 200
    assert renamed.json()['name'] == 'Renamed comparison'
    assert client.get('/api/e2e-dashboards').json()['test']['name'] == 'Renamed comparison'
    comments = client.patch('/api/e2e-dashboards/test/comments', json={'slide_comments': {'1': ['Review city outliers']}})
    assert comments.status_code == 200
    assert comments.json()['slide_comments'] == {'1': ['Review city outliers']}
    assert client.get('/api/e2e-dashboards').json()['test']['slide_comments'] == {'1': ['Review city outliers']}
    result = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert result.status_code == 200, result.text
    preview = result.json()
    assert preview['filter_fields'] == ['Market', 'Region', 'City', 'Campaign', 'Operator', 'Vendor', 'RAT', 'Session Type', 'Call Status']
    assert 'Technology' not in preview['options']
    assert preview['options']['Operator'] == ['A', 'B']
    assert preview['options']['City'] == ['Leeds', 'London']
    assert preview['date_bounds'] == {'min': '2026-09-01', 'max': '2026-09-03'}
    assert len(preview['slides']) == 2
    assert len(preview['slides'][0]['charts']) == 2
    assert preview['slides'][0]['charts'][0]['focus_row'] == 0
    assert preview['slides'][0]['charts'][0]['cdr_source']
    assert preview['slides'][0]['charts'][0]['chart_type']
    assert preview['slides'][0]['charts'][1]['focus_row'] == 1
    assert preview['slides'][0]['charts'][0]['position'][0] < preview['slides'][0]['charts'][1]['position'][0]
    token = preview['token']
    restored = client.get(f'/api/e2e-dashboards/prepared/{token}')
    assert restored.status_code == 200
    assert restored.json() == preview
    image = client.get(f'/api/e2e-dashboards/preview/{token}/0.png')
    assert image.status_code == 200, image.text if image.status_code != 200 else ''
    assert image.content.startswith(b'\x89PNG')
    interactive = client.get(f'/api/e2e-dashboards/chart/{token}/0')
    assert interactive.status_code == 200, interactive.text
    assert interactive.json()['type'] == 'cdf'
    assert interactive.json()['renderer'] == 'catalog-v2'
    assert interactive.json()['series']
    assert list((Path(core.repository.db_path).parent / '.dashboard-data-cache' / 'charts-pil').glob('*.png'))
    entry = core.load_template_catalogue(next(row['content'] for row in core.repository.list_report_templates('nsa') if row['name'] == 'Dashboard test'), 'nsa')[0]
    assert not core.is_empty_catalog_chart(image.content, entry)
    data = client.get(f'/api/e2e-dashboards/data/{token}/0').json()
    assert data['total'] == 3
    assert data['page_size'] == 100
    assert data['unfiltered_total'] == 3
    assert set(data['column_metadata']) == set(data['columns'])
    assert all({'label', 'kind', 'rule', 'pinned', 'class_name'} <= set(item) for item in data['column_metadata'].values())
    city_values = client.get(f'/api/e2e-dashboards/data/{token}/0', params={
        'column_filters': json.dumps({'Operator': ['A']}), 'filter_column': 'City',
    }).json()
    assert city_values['filter_values'] == ['London']
    assert client.get(f'/api/e2e-dashboards/data/{token}/0?download=true').headers['content-type'].startswith('text/csv')
    assert client.delete('/api/e2e-dashboards/test').status_code == 200
    assert client.get('/api/e2e-dashboards').json() == {}


def test_e2e_reporting_is_restricted_to_super_admins_and_ejaitur(client):
    client.post('/login', data={'username': 'admin', 'password': 'admin123'})
    workspace_page = client.get('/workspace')
    assert workspace_page.status_code == 200
    assert 'href="/e2e-reporting"' not in workspace_page.text
    assert 'E2E Reporting' not in workspace_page.text
    assert client.get('/e2e-reporting').status_code == 403
    assert client.get('/api/e2e-reporting/jobs').status_code == 403

    super_admin_token = 'e2e-reporting-super-admin'
    core.SESSIONS[super_admin_token] = core.SessionUser(username='someone', role='super-admin')
    client.cookies.set(core.SESSION_COOKIE, super_admin_token)
    super_admin_page = client.get('/e2e-reporting')
    assert super_admin_page.status_code == 200
    assert 'href="/e2e-reporting"' in super_admin_page.text

    token = 'e2e-reporting-allowed-user'
    core.SESSIONS[token] = core.SessionUser(username='EJAITUR', role='user')
    client.cookies.set(core.SESSION_COOKIE, token)
    allowed_page = client.get('/e2e-reporting')
    assert allowed_page.status_code == 200
    assert 'href="/e2e-reporting"' in allowed_page.text
    assert client.get('/api/e2e-reporting/jobs').status_code == 200


def test_ready_dashboard_exports_ppt_and_persistent_chart_files(client, monkeypatch):
    payload = setup_dashboard(client)
    payload['slide_comments'] = {'1': ['Review city outliers', 'Validate campaign coverage']}
    payload['filters'] = {'City': ['London']}
    dashboard_id = 'ppt-dashboard'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200

    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    initial_status = client.get('/api/e2e-dashboards/statuses').json()[dashboard_id]
    assert initial_status['state'] in {'loading-data', 'ready'}
    assert initial_status['label'] == 'Ready' or initial_status['label'].startswith('Preparing ')

    with core.repository.connection() as connection:
        connection.execute('UPDATE dataset_profiles SET vendor_mapping_applied = 1')
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=[], cities=['Leeds', 'London'],
    )
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'],
        'Vendor': ['Vendor A', 'Vendor B', 'Vendor A'],
        'City': ['London', 'Leeds', 'London'],
        'Mean_Data_Rate': [10, 20, 30],
        'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
        # Ingestion always normalizes this field; a reused snapshot filters
        # its resolved calendar dates against it.
        'event_start_time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    export_payload = json.loads(json.dumps(payload))
    export_payload['scope'] = 'multivendor'
    prepared = client.post(
        f'/api/e2e-dashboards/prepare?dashboard_id={dashboard_id}', json=export_payload,
    )
    assert prepared.status_code == 200, prepared.text
    prepared_payload = prepared.json()
    for slide in prepared_payload['slides']:
        for chart in slide['charts']:
            if chart['available']:
                rendered = client.get(f"/api/e2e-dashboards/chart/{prepared_payload['token']}/{chart['index']}")
                assert rendered.status_code == 200, rendered.text

    queued = client.post(f'/api/e2e-dashboards/{dashboard_id}/export-ppt', json={
        'definition': export_payload,
        'preparation_token': prepared_payload['token'],
    })
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    job = None
    while time.monotonic() < deadline:
        jobs = client.get('/api/e2e-dashboards/ppt-jobs').json()['jobs']
        job = next(item for item in jobs if item['id'] == job_id)
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.05)
    assert job is not None and job['status'] == 'ready', job
    assert job['slides'] == 2
    assert job['charts'] == 3
    assert job['nr_mode'] == 'NSA'
    assert job['scope'] == 'Multivendor Comparison'
    assert job['template'] == 'Dashboard test'
    assert re.fullmatch(r'\d{8}_\d{6}', job['timestamp'])
    assert job['filters'] == [
        'CDR Data: sample.csv', 'Date: Oldest to Newest',
        'Operator: All Operators', 'Vendor: All Vendors', 'City: London',
    ]
    assert job['duration_seconds'] is not None
    background_groups = client.get('/api/background-tasks').json()['groups']
    background_task = next(
        task for group in background_groups for task in group['tasks']
        if task['id'] == f'dashboard-ppt:{core.active_workspace.id}:{job_id}'
    )
    assert background_task['dashboard_name'] == 'Comparison'
    assert background_task['label'] == 'Generating Dashboard PPT'
    assert background_task['detail'] == 'Completed'
    assert background_task['duration_seconds'] is not None

    with core.repository.connection() as connection:
        row = connection.execute('SELECT output_path FROM dashboard_ppt_jobs WHERE id = ?', (job_id,)).fetchone()
    output_path = Path(row['output_path'])
    charts_dir = output_path.parent / 'dashboard-charts'
    assert output_path.is_file()
    assert re.fullmatch(
        r'\d{8}_\d{6} - Comparison - Multivendor Comparison\.pptx',
        output_path.name,
    )
    assert output_path.parent.name == output_path.stem
    assert output_path.parent.parent == Path(core.repository.db_path).parent / 'output' / 'dashboards'
    assert len(list(charts_dir.glob('*.png'))) == 3
    assert len(list(charts_dir.glob('*.hover.json'))) == 3
    assert len(list(charts_dir.glob('*.model.json'))) == 3
    manifest = json.loads((charts_dir / 'manifest.json').read_text(encoding='utf-8'))
    assert manifest['generate_tooltips'] is True
    assert manifest['dashboard_id'] == dashboard_id
    assert re.fullmatch(r'[0-9a-f]{64}', manifest['preview_fingerprint'])
    assert manifest['definition']['filters'] == {'City': ['London']}
    assert manifest['definition']['scope'] == 'multivendor'
    assert len(manifest['charts']) == 3
    assert manifest['charts'][0]['entry_index'] == 0
    assert manifest['charts'][0]['focus_row'] == 0
    presentation = Presentation(output_path)
    assert not [
        shape for slide in presentation.slides for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        and any((shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom))
    ]
    for shape in (
        shape for slide in presentation.slides for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ):
        with Image.open(BytesIO(shape.image.blob)) as chart_image:
            assert abs(chart_image.width / chart_image.height - shape.width / shape.height) < 0.01
    commentary = next(
        shape for shape in presentation.slides[0].shapes
        if shape.is_placeholder and shape.placeholder_format.idx == 10
    )
    assert [paragraph.text for paragraph in commentary.text_frame.paragraphs] == [
        'Review city outliers', 'Validate campaign coverage',
    ]

    ppt = client.get(job['download_url'])
    assert ppt.status_code == 200
    assert ppt.content.startswith(b'PK')
    assert client.get(job['charts_url']).status_code == 200
    charts_manifest = client.get(job['charts_api_url'])
    assert charts_manifest.status_code == 200
    charts_payload = charts_manifest.json()
    assert charts_payload['job']['id'] == job_id
    assert len(charts_payload['charts']) == 3
    assert client.get(charts_payload['charts'][0]['image_url']).status_code == 200
    assert charts_payload['charts'][0]['payload_url']
    assert charts_payload['charts'][0]['data_url']
    assert charts_payload['charts'][0]['focus_row'] == 0
    chart_model = client.get(charts_payload['charts'][0]['payload_url'])
    assert chart_model.status_code == 200
    import src.modules.e2e_dashboards as dashboards_module
    monkeypatch.setattr(
        dashboards_module, 'prepare_catalog_chart_preview_frame',
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError('The combined table should serve dataset pages.')),
    )
    chart_data = client.get(f"/api/e2e-dashboards{charts_payload['charts'][0]['data_url']}")
    assert chart_data.status_code == 200
    assert chart_data.json()['total'] == 2
    filtered_data = client.get(
        f"/api/e2e-dashboards{charts_payload['charts'][0]['data_url']}",
        params={
            'column_filters': json.dumps({'Operator': ['A']}),
            'include_filter_values': 'true',
        },
    )
    assert filtered_data.status_code == 200
    assert filtered_data.json()['chart_total'] == 2
    assert filtered_data.json()['total'] == 2
    assert filtered_data.json()['filter_values']['Operator'] == ['A']
    empty_data = client.get(
        f"/api/e2e-dashboards{charts_payload['charts'][0]['data_url']}",
        params={'column_filters': json.dumps({'Operator': ['B']})},
    )
    assert empty_data.status_code == 200
    assert empty_data.json()['chart_total'] == 2
    assert empty_data.json()['total'] == 0
    assert empty_data.json()['rows'] == []
    assert chart_model.json()['title'] == charts_payload['charts'][0]['title']
    archive = client.get(job['charts_download_url'])
    assert archive.status_code == 200
    with zipfile.ZipFile(BytesIO(archive.content)) as bundle:
        assert len([name for name in bundle.namelist() if name.endswith('.png')]) == 3

    deleted = client.post('/api/e2e-dashboards/ppt-jobs/delete-all')
    assert deleted.status_code == 200
    assert deleted.json()['deleted'] == 1
    assert client.get('/api/e2e-dashboards/ppt-jobs').json()['jobs'] == []
    assert not output_path.parent.exists()


def test_relaunching_dashboard_ppt_uses_the_modified_report_template(client):
    payload = setup_dashboard(client)
    dashboard_id = 'updated-template-ppt'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200

    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    with core.repository.connection() as connection:
        connection.execute('UPDATE dataset_profiles SET vendor_mapping_applied = 1')
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=[], cities=['Leeds', 'London'],
    )
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'], 'Vendor': ['Vendor A', 'Vendor B', 'Vendor A'],
        'City': ['London', 'Leeds', 'London'], 'Mean_Data_Rate': [10, 20, 30],
        'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    export_payload = {**payload, 'scope': 'multivendor'}
    prepared = client.post(
        f'/api/e2e-dashboards/prepare?dashboard_id={dashboard_id}', json=export_payload,
    )
    assert prepared.status_code == 200, prepared.text
    prepared_payload = prepared.json()
    for slide in prepared_payload['slides']:
        for chart in slide['charts']:
            if chart['available']:
                assert client.get(
                    f"/api/e2e-dashboards/chart/{prepared_payload['token']}/{chart['index']}"
                ).status_code == 200
    queued = client.post(f'/api/e2e-dashboards/{dashboard_id}/export-ppt', json={
        'definition': export_payload,
        'preparation_token': prepared_payload['token'],
    })
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']

    def wait_for_job():
        while time.monotonic() < deadline:
            job = next(
                item for item in client.get('/api/e2e-dashboards/ppt-jobs').json()['jobs']
                if item['id'] == job_id
            )
            if job['status'] in {'ready', 'failed'}:
                return job
            time.sleep(0.05)
        return job

    original_job = wait_for_job()
    assert original_job['status'] == 'ready'
    assert original_job['scope'] == 'Multivendor Comparison'
    original_charts = client.get(
        f'/api/e2e-dashboards/ppt-jobs/{job_id}/charts.json'
    ).json()['charts']
    assert original_charts[0]['title'] == 'Rate'

    template = core.repository.report_template_content('nsa', 'Dashboard test')
    core.repository.set_report_template_content(
        'nsa', 'Dashboard test', template.replace(b',Rate,CDR-Data,', b',Updated rate,CDR-Data,', 1),
    )

    relaunched = client.post(f'/api/e2e-dashboards/ppt-jobs/{job_id}/retry')
    assert relaunched.status_code == 202, relaunched.text
    deadline = time.monotonic() + 15
    job = wait_for_job()
    assert job['status'] == 'ready', job
    assert job['scope'] == 'Multivendor Comparison'
    refreshed_charts = client.get(job['charts_api_url']).json()['charts']
    assert refreshed_charts[0]['title'] == 'Updated rate'
    refreshed_model = client.get(refreshed_charts[0]['payload_url'])
    assert refreshed_model.status_code == 200, refreshed_model.text
    assert refreshed_model.json()['title'] == 'Updated rate'
    with core.repository.connection() as connection:
        output_path = Path(connection.execute(
            'SELECT output_path FROM dashboard_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()['output_path'])
    manifest = json.loads((output_path.parent / 'dashboard-charts' / 'manifest.json').read_text(encoding='utf-8'))
    assert manifest['definition']['scope'] == 'multivendor'
    assert client.get(job['download_url']).content.startswith(b'PK')


def test_dashboard_is_prepared_only_when_opened_and_then_reuses_manifest(client):
    payload = setup_dashboard(client)
    dashboard_id = 'on-demand-dashboard'
    saved = client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload)
    assert saved.status_code == 200

    response = client.get(f'/api/e2e-dashboards/prefetched/{dashboard_id}')
    assert response.status_code == 409
    assert not any(
        task.get('dashboard_name') == payload['name']
        for group in client.get('/api/background-tasks').json()['groups']
        for task in group['tasks']
    )

    prepared = client.post(f'/api/e2e-dashboards/prepare?dashboard_id={dashboard_id}', json=payload)
    assert prepared.status_code == 200, prepared.text
    response = client.get(f'/api/e2e-dashboards/prefetched/{dashboard_id}')
    assert response.status_code == 200, response.text
    assert response.json()['slides']
    manifests = list(
        (Path(core.repository.db_path).parent / '.dashboard-data-cache' / 'dashboard-previews').glob('*.json')
    )
    assert manifests
    token = prepared.json()['token']
    for _ in range(130):
        assert client.get('/api/e2e-dashboards/statuses').status_code == 200
    assert client.get(f'/api/e2e-dashboards/prepared/{token}').status_code == 200


def test_closed_dashboard_status_uses_its_remembered_session_universe(client):
    payload = setup_dashboard(client)
    dashboard_id = 'session-universe-dashboard'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200

    session_definition = json.loads(json.dumps(payload))
    session_definition['date_from'] = '2026-09-02'
    session_definition['date_to'] = '2026-09-03'
    prepared = client.post(
        f'/api/e2e-dashboards/prepare?dashboard_id={dashboard_id}', json=session_definition,
    )
    assert prepared.status_code == 200, prepared.text

    background_status = client.get('/api/e2e-dashboards/statuses').json()[dashboard_id]
    assert background_status['state'] in {'loading-data', 'ready'}
    session_status = client.post('/api/e2e-dashboards/statuses', json={
        dashboard_id: session_definition,
    }).json()[dashboard_id]
    assert session_status['state'] in {'loading-data', 'ready'}


def test_saving_dashboard_prepares_data_in_background_without_rendering_charts(client, monkeypatch):
    payload = setup_dashboard(client)
    import src.modules.e2e_dashboards as dashboards_module
    monkeypatch.setattr(
        dashboards_module,
        'catalog_chart_payload',
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError('Saving must not render charts.')),
    )
    dashboard_id = 'saved-dashboard'
    assert client.put(f'/api/e2e-dashboards/{dashboard_id}', json=payload).status_code == 200
    assert client.get(f'/api/e2e-dashboards/prefetched/{dashboard_id}').status_code == 409
    saved_status = client.get('/api/e2e-dashboards/statuses').json()[dashboard_id]
    assert saved_status['state'] in {'loading-data', 'ready'}
    assert saved_status['label'] == 'Ready' or saved_status['label'].startswith('Preparing ')


def test_direct_dashboard_preparation_separates_queue_and_execution_timestamps(client, monkeypatch):
    payload = setup_dashboard(client)
    entered = Event()
    release = Event()
    original = core.Repository.list_dataset_row_columns

    def delayed(repository, *args, **kwargs):
        entered.set()
        assert release.wait(5)
        return original(repository, *args, **kwargs)

    monkeypatch.setattr(core.Repository, 'list_dataset_row_columns', delayed)
    responses = {}
    running = Thread(target=lambda: responses.setdefault(
        'running', client.post(
            '/api/e2e-dashboards/prepare?preparation_id=direct-running', json=payload,
        ),
    ))
    queued = Thread(target=lambda: responses.setdefault(
        'queued', client.post(
            '/api/e2e-dashboards/prepare?preparation_id=direct-queued', json=payload,
        ),
    ))
    try:
        running.start()
        assert entered.wait(5)
        running_task = next(
            task for group in client.get('/api/background-tasks').json()['groups']
            for task in group['tasks'] if task['id'] == 'direct-running'
        )
        assert running_task['status'] == 'processing'
        assert running_task['started_at'] >= running_task['queued_at']

        queued.start()
        deadline = time.monotonic() + 5
        queued_task = None
        while time.monotonic() < deadline:
            queued_task = next((
                task for group in client.get('/api/background-tasks').json()['groups']
                for task in group['tasks'] if task['id'] == 'direct-queued'
            ), None)
            if queued_task is not None:
                break
            time.sleep(0.02)
        assert queued_task is not None
        assert queued_task['status'] == 'queued'
        assert queued_task['queued_at'] is not None
        assert queued_task['started_at'] is None
    finally:
        release.set()
        running.join(5)
        if queued.ident is not None:
            queued.join(5)

    assert not running.is_alive()
    assert not queued.is_alive()


def test_applying_filters_prepares_only_data_and_renders_charts_on_demand(client, monkeypatch):
    payload = setup_dashboard(client)
    import src.modules.e2e_dashboards as dashboards_module

    core.repository.set_workspace_state('e2e_dashboards_v2', json.dumps({'filtered-dashboard': payload}))
    uncached = client.get('/api/e2e-dashboards/statuses')
    assert uncached.status_code == 200
    assert uncached.json()['filtered-dashboard']['state'] in {'loading-data', 'ready'}

    calls = []
    original = dashboards_module.catalog_chart_payload

    def tracked(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(dashboards_module, 'catalog_chart_payload', tracked)
    response = client.post('/api/e2e-dashboards/prepare?dashboard_id=filtered-dashboard', json=payload)
    assert response.status_code == 200, response.text

    cache_dir = Path(core.repository.db_path).parent / '.dashboard-data-cache' / 'charts-canvas'
    assert not list(cache_dir.glob('*.json'))
    assert calls == []
    assert client.get('/api/e2e-dashboards/statuses').json()['filtered-dashboard'] == {
        'state': 'ready', 'label': 'Ready',
    }

    token = response.json()['token']
    for index in range(3):
        rendered = client.get(f'/api/e2e-dashboards/chart/{token}/{index}')
        assert rendered.status_code == 200, rendered.text
    first_models = {path.name for path in cache_dir.glob('*.json')}
    assert len(first_models) == 3
    assert len(calls) == 3

    payload['filters'] = {'City': ['London']}
    filtered = client.post('/api/e2e-dashboards/prepare?dashboard_id=filtered-dashboard', json=payload)
    assert filtered.status_code == 200, filtered.text
    assert {path.name for path in cache_dir.glob('*.json')} == first_models
    assert len(calls) == 3

    payload['filters'] = {}
    calls_before_restore = len(calls)
    restored = client.post('/api/e2e-dashboards/prefetched/filtered-dashboard', json=payload)
    assert restored.status_code == 200, restored.text
    assert len(calls) == calls_before_restore
    assert {path.name for path in cache_dir.glob('*.json')} == first_models

    core.repository.update_dataset_profile(1, progress=100)
    refreshed_status = client.get('/api/e2e-dashboards/statuses').json()['filtered-dashboard']
    assert refreshed_status['state'] in {'loading-data', 'ready'}


def test_adding_a_dataset_builds_a_new_chart_model_with_every_campaign(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    for file_name, campaign, date_value in (
        ('campaign-q1.csv', '2026-Q1', '2026-02-01'),
        ('campaign-q2.csv', '2026-Q2', '2026-05-01'),
    ):
        response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
            'dataset_files': (
                file_name,
                BytesIO(
                    f'Operator,Campaign,Mean_Data_Rate,Test_Start_Time\nA,{campaign},10,{date_value}\n'.encode()
                ),
                'text/csv',
            ),
        })
        assert response.status_code == 200
    core.repository.add_report_template('nsa', 'Campaign cache test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Aggregation,Column Aggregation,Legend,Legend Position\n'
        '1,Campaigns,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,Average Vertical Bars,,Operator,Campaign,,Right\n'
    ).encode(), is_default=False)
    datasets = sorted(core.repository.list_datasets(), key=lambda row: int(row['id']))
    first_id, second_id = (int(row['id']) for row in datasets[-2:])
    payload = DashboardDefinition(
        name='Campaign cache test', template='Campaign cache test', datasets={'data': [first_id]},
    ).model_dump(mode='json')

    first = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert first.status_code == 200, first.text
    first_model = client.get(f"/api/e2e-dashboards/chart/{first.json()['token']}/0")
    assert first_model.status_code == 200, first_model.text
    assert {bar['key'][-1] for bar in first_model.json()['bars']} == {'2026-Q1'}

    payload['datasets']['data'] = [first_id, second_id]
    payload['date_from'], payload['date_to'] = 'Oldest', 'Newest'
    combined = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert combined.status_code == 200, combined.text
    combined_model = client.get(f"/api/e2e-dashboards/chart/{combined.json()['token']}/0")
    assert combined_model.status_code == 200, combined_model.text
    assert {bar['key'][-1] for bar in combined_model.json()['bars']} == {'2026-Q1', '2026-Q2'}

    cache_dir = Path(core.repository.db_path).parent / '.dashboard-data-cache' / 'charts-canvas'
    assert len(list(cache_dir.glob('*.json'))) >= 2


def test_dashboard_api_session_expires_on_application_process_restart(client):
    setup_dashboard(client)
    import src.DashboardAnalytic as app_module

    app_module.SESSIONS.clear()

    response = client.get('/api/e2e-dashboards', follow_redirects=False)
    assert response.status_code == 401
    assert response.json()['detail'] == 'Your session has expired. Please sign in again.'


def test_dashboard_preview_identifies_title_and_transition_slides(client):
    payload = setup_dashboard(client)
    core.repository.add_report_template('nsa', 'Structural dashboard', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Aggregation,Column Aggregation,Legend,Legend Position\n'
        '1,Quarterly review,Network results,Title Page,,,,Title Slide,,,,,Top\n'
        '2,Voice performance,,Title Only,,,,Transition Slide,,,,,Top\n'
    ).encode(), is_default=False)
    payload['template'] = 'Structural dashboard'

    preview = client.post('/api/e2e-dashboards/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    slides = preview.json()['slides']
    assert [(slide['title'], slide['structural_type'], slide['charts']) for slide in slides] == [
        ('Quarterly review', 'title slide', []),
        ('Voice performance', 'transition slide', []),
    ]


def test_dashboard_state_migrates_from_legacy_storage(client):
    payload = setup_dashboard(client)
    core.repository.set_workspace_state('e2e_dashboard_sets_v1', '{"legacy": ' + json.dumps(payload) + '}')
    result = client.get('/api/e2e-dashboards')
    assert result.status_code == 200
    assert result.json()['legacy']['name'] == 'Comparison'
    assert json.loads(core.repository.get_workspace_state('e2e_dashboards_v2')) == result.json()


def test_dashboard_migrates_all_saved_definitions_to_current_default_filters(client):
    payload = setup_dashboard(client)
    payload['filters'] = {'Technology': ['5G'], 'Zone': ['North'], 'Operator': ['A']}
    payload['custom_fields'] = ['Technology', 'Zone', 'Region', 'Mean_Data_Rate']
    payload['hidden_filters'] = ['Technology', 'Region', 'Zone', 'City', 'Campaign', 'RAT', 'Call Status', 'Former custom filter']
    with core.repository.connection() as connection:
        connection.execute("DELETE FROM workspace_state WHERE key = 'e2e_dashboard_default_filters_v6'")
    core.repository.set_workspace_state('e2e_dashboards_v2', json.dumps({'legacy': payload}))

    result = client.get('/api/e2e-dashboards')

    assert result.status_code == 200
    migrated = result.json()['legacy']
    assert migrated['filters'] == {'Operator': ['A'], 'Region': ['North']}
    assert migrated['custom_fields'] == ['Mean_Data_Rate']
    assert migrated['hidden_filters'] == ['Former custom filter']
    assert {key: migrated[key] for key in ('scope', 'datasets', 'date_from', 'date_to')} == {
        key: payload[key] for key in ('scope', 'datasets', 'date_from', 'date_to')
    }
    assert core.repository.get_workspace_state('e2e_dashboard_default_filters_v6') == '1'
    assert json.loads(core.repository.get_workspace_state('e2e_dashboards_v2'))['legacy'] == migrated


def test_dashboard_custom_fields_and_snapshot_filters(client):
    payload = setup_dashboard(client)
    core.write_workspace_calculated_dimensions([{'name': '7-cities', 'sources': ['cdr-data'], 'rules': [{'when': 'City IN (London)', 'value': 'Yes'}], 'default': 'No'}])
    payload['custom_fields'] = ['7-cities']
    preview = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['7-cities'] == ['No', 'Yes']
    original = preview.json()['token']
    payload['filters'] = {'7-cities': ['Yes']}
    preview = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    assert preview['rows']['data'] == 2
    assert preview['universe_rows']['data'] == 3
    assert preview['options']['7-cities'] == ['No', 'Yes']
    for index in (0, 1, 2):
        assert client.get(f"/api/e2e-dashboards/data/{preview['token']}/{index}").json()['total'] == 2
    assert client.get(f'/api/e2e-dashboards/data/{original}/0').json()['total'] == 3
    payload['filters'] = {'7-cities': []}
    assert client.post('/api/e2e-dashboards/prepare', json=payload).json()['rows']['data'] == 0
    payload['filters'] = {'Test_Name': ['HTTP DL']}
    payload['custom_fields'] = ['Test_Name']
    preview = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    assert 'Test_Name' in preview['available_fields']
    assert preview['options']['Test_Name'] == ['HTTP DL']
    payload['hidden_filters'] = ['Market']
    preview = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    assert 'Market' not in preview['options']


def test_dashboard_source_specific_custom_fields_do_not_repeat_combined_preparation(client, monkeypatch):
    payload = setup_dashboard(client)
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'voice'}, files={
        'dataset_files': ('voice.csv', BytesIO(
            b'Operator,City,Mean_Call_Setup_Time,Session_Type\nA,London,1.2,VoLTE\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.write_workspace_calculated_dimensions([
        {
            'name': 'Test Family', 'sources': ['cdr-data'],
            'rules': [{'when': 'Test_Name CONTAINS HTTP', 'value': 'HTTP'}], 'default': '',
        },
        {
            'name': 'Call Family', 'sources': ['cdr-voice'],
            'rules': [{'when': 'Session_Type CONTAINS VoLTE', 'value': 'VoLTE'}], 'default': 'Call',
        },
    ])
    payload['datasets']['voice'] = [2]
    payload['custom_fields'] = ['Test Family', 'Call Family']
    first = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert first.status_code == 200, first.text

    from src.modules.repository import Repository

    def unexpected_combined_rebuild(*_args, **_kwargs):
        raise AssertionError('Fields from another CDR type must not prepare source columns.')

    monkeypatch.setattr(Repository, 'copy_dataset_rows_to_reporting', unexpected_combined_rebuild)
    payload['filters'] = {'City': ['London']}
    repeated = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert repeated.status_code == 200, repeated.text


def test_dashboard_loads_new_filter_values_without_preparing_a_snapshot(client):
    payload = setup_dashboard(client)
    payload['custom_fields'] = ['mean data rate']
    payload['filters'] = {'CITY': ['lOnDoN']}

    response = client.post('/api/e2e-dashboards/filter-options', json={
        'definition': payload,
        'field': 'MEAN-DATA-RATE',
    })

    assert response.status_code == 200, response.text
    assert response.json() == {'field': 'MEAN-DATA-RATE', 'values': ['10', '30']}


def test_dashboard_rat_filter_uses_dataset_preview_column_precedence(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('radio.csv', BytesIO(
            b'Operator,City,Campaign,RAT_A,RAT,Sample_RAT_A,Call_Status,Mean_Data_Rate\n'
            b'A,London,Spring,ENDC,LTE,NR,Completed,10\n'
            b'B,Leeds,Summer,NR,LTE,ENDC,Failed,20\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('nsa', 'Radio Dashboard', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Aggregation,Column Aggregation,Legend,Legend Position\n'
        '1,Radio,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    payload = definition().model_dump(mode='json')
    payload['template'] = 'Radio Dashboard'

    preview = client.post('/api/e2e-dashboards/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['Campaign'] == ['Spring', 'Summer']
    assert preview.json()['options']['RAT'] == ['ENDC', 'NR']
    assert preview.json()['options']['Call Status'] == ['Completed', 'Failed']


def test_dashboard_falls_back_to_combined_rows_for_values_missing_from_profiles(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('profile-gap.csv', BytesIO(
            b'Operator,G_Level_2,G_Level_4,Campaign,RAT,Call_Status,Mean_Data_Rate,Test_Start_Time\n'
            b'A,England,London,Spring,ENDC,Completed,10,2026-09-01\n'
            b'B,England,Leeds,Summer,NR,Failed,20,2026-09-02\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    with core.repository.connection() as connection:
        connection.execute("UPDATE dataset_profiles SET filter_options_json = '{}' WHERE dataset_id = 1")
    core.repository.add_report_template('nsa', 'Profile fallback Dashboard', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Aggregation,Column Aggregation,Legend,Legend Position\n'
        '1,Profile fallback,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    payload = definition().model_dump(mode='json')
    payload['template'] = 'Profile fallback Dashboard'

    preview = client.post('/api/e2e-dashboards/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['Region'] == ['England']
    assert preview.json()['options']['City'] == ['Leeds', 'London']
    assert preview.json()['options']['Campaign'] == ['Spring', 'Summer']
    assert preview.json()['options']['RAT'] == ['ENDC', 'NR']
    assert preview.json()['options']['Call Status'] == ['Completed', 'Failed']


def test_dashboard_validates_template_dates_and_sources(client):
    payload = setup_dashboard(client)
    payload['date_from'], payload['date_to'] = '2026-09-03', '2026-09-01'
    assert client.post('/api/e2e-dashboards/prepare', json=payload).status_code == 400
    payload['date_from'] = payload['date_to'] = None
    payload['datasets'] = {}
    assert client.post('/api/e2e-dashboards/prepare', json=payload).status_code == 400
    payload['template'] = 'Unknown'
    assert client.put('/api/e2e-dashboards/test', json=payload).status_code == 400
    assert client.get('/api/e2e-dashboards/data/missing/0').status_code == 410


def test_dashboard_sql_selection_leaves_nr_mode_to_explicit_user_filters(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'voice'}, files={
        'dataset_files': ('voice.csv', BytesIO(
            b'Operator,Mean_Call_Setup_Time,RAT_A,Session_Type,L1_Call_Mode_A\n'
            b'A,1.0,EN-DC,WhatsApp Voice,\n'
            b'A,2.0,NR,WhatsApp Voice,\n'
            b'A,3.0,LTE,Native Voice,VoLTE\n'
            b'A,4.0,NR,Native Voice,VoNR\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('nsa', 'Voice Dashboard', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Aggregation,Column Aggregation,Legend,Legend Position\n'
        '1,Voice,,Title and 1 column + Comments,Calls,CDR-Voice,Mean_Call_Setup_Time,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    payload = DashboardDefinition(
        name='Voice', template='Voice Dashboard', datasets={'voice': [1]}, technology='nsa',
    ).model_dump(mode='json')
    nsa = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    assert nsa['rows']['voice'] == 4
    assert nsa['universe_rows']['voice'] == 4
    payload['technology'] = 'sa'
    sa = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    assert sa['rows']['voice'] == 4
    assert sa['universe_rows']['voice'] == 4
    dashboard_module = (Path(__file__).parents[1] / 'src/modules/e2e_dashboards.py').read_text(encoding='utf-8')
    assert 'def nr_mode_sql(' not in dashboard_module


def test_dashboard_snapshot_access_and_legacy_redirect(client):
    payload = setup_dashboard(client)
    response = client.get('/dashboard?dataset_id=1', follow_redirects=False)
    assert response.status_code == 307
    assert response.headers['location'] == '/datasets-analysis?dataset_id=1'
    assert client.get(response.headers['location']).status_code == 200
    preview = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    client.get('/logout')
    client.post('/login', data={'username': 'demo', 'password': 'demo123'})
    assert client.get(f"/api/e2e-dashboards/data/{preview['token']}/0").status_code == 410
    assert client.get(f"/api/e2e-dashboards/preview/{preview['token']}/0.png").status_code == 410
    assert client.get(f"/api/e2e-dashboards/chart/{preview['token']}/0").status_code == 410
    assert client.get(f"/api/e2e-dashboards/prepared/{preview['token']}").status_code == 410


def test_dashboard_reuses_persistent_sql_selection_and_invalidates_dataset_versions(client, monkeypatch):
    payload = setup_dashboard(client)
    first = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert first.status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 1
    repeated = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert repeated.status_code == 200
    assert repeated.json()['token'] != first.json()['token']
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 1
    payload['name'] = 'Renamed comparison'
    assert client.post('/api/e2e-dashboards/prepare', json=payload).status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 1

    from src.modules.repository import Repository

    def unexpected_combined_rebuild(*_args, **_kwargs):
        raise AssertionError('Default Filter value changes must not prepare source columns.')

    monkeypatch.setattr(Repository, 'copy_dataset_rows_to_reporting', unexpected_combined_rebuild)
    payload['filters'] = {'City': ['London']}
    assert client.post('/api/e2e-dashboards/prepare', json=payload).json()['rows']['data'] == 2
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 2
    payload['filters'] = {'City': ['Leeds']}
    assert client.post('/api/e2e-dashboards/prepare', json=payload).json()['rows']['data'] == 1
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 3
    payload['filters'] = {'City': ['London']}
    assert client.post('/api/e2e-dashboards/prepare', json=payload).json()['rows']['data'] == 2
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 3
    core.repository.set_workspace_state('dashboard_cache_test', 'updated')
    assert client.post('/api/e2e-dashboards/prepare', json=payload).status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 3
    core.repository.update_dataset_profile(1, progress=100)
    assert client.post('/api/e2e-dashboards/prepare', json=payload).status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM dashboard_filter_selections').fetchone()['count'] == 4


def test_dashboard_selection_cache_ignores_filter_value_order(client):
    payload = setup_dashboard(client)
    payload['filters'] = {'City': ['London', 'Leeds']}

    first = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert first.status_code == 200
    payload['filters']['City'].reverse()
    repeated = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert repeated.status_code == 200

    with core.repository.connection() as connection:
        assert connection.execute(
            'SELECT COUNT(*) AS count FROM dashboard_filter_selections'
        ).fetchone()['count'] == 1


def test_dashboard_selection_uses_direct_sql_predicate(client, monkeypatch):
    payload = setup_dashboard(client)
    import src.modules.e2e_dashboards as dashboards_module

    monkeypatch.setattr(dashboards_module, 'DASHBOARD_PROFILE_SELECTION_THRESHOLD', 1)
    payload['filters'] = {'City': ['London']}
    preview = client.post('/api/e2e-dashboards/prepare', json=payload)
    assert preview.status_code == 200
    assert preview.json()['rows']['data'] == 2
    assert preview.json()['rows_exact'] is True
    with core.repository.connection() as connection:
        selection = connection.execute(
            'SELECT id FROM dashboard_filter_selections ORDER BY id DESC LIMIT 1'
        ).fetchone()
        assert selection is not None
        assert connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'dashboard_filter_selection_rows'"
        ).fetchone() is None
        assert 'materialized' not in {
            row['name'] for row in connection.execute('PRAGMA table_info(dashboard_filter_selections)')
        }
    data = client.get(f"/api/e2e-dashboards/data/{preview.json()['token']}/0")
    assert data.status_code == 200
    assert data.json()['total'] == 2


def test_dashboard_profile_facets_show_values_outside_the_saved_filter(client, monkeypatch):
    payload = setup_dashboard(client)
    import src.modules.e2e_dashboards as dashboards_module

    monkeypatch.setattr(dashboards_module, 'DASHBOARD_PROFILE_SELECTION_THRESHOLD', 1)
    with core.repository.connection() as connection:
        connection.execute("UPDATE dataset_profiles SET filter_options_json = '{}' WHERE dataset_id = 1")
    payload['filters'] = {'Operator': ['A']}

    preview = client.post('/api/e2e-dashboards/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    assert preview.json()['rows']['data'] == 2
    assert preview.json()['options']['Operator'] == ['A', 'B']


def test_dashboard_operator_facets_keep_the_values_stored_in_the_combined_table(client):
    payload = setup_dashboard(client)
    core.repository.replace_operator_mapping_group(None, 'Alpha', ['A'])

    preview = client.post('/api/e2e-dashboards/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['Operator'] == ['A', 'B']
    options = client.post('/api/e2e-dashboards/filter-options', json={
        'definition': payload, 'field': 'Operator',
    })
    assert options.status_code == 200, options.text
    assert options.json()['values'] == ['A', 'B']


def test_dashboard_reuses_normalized_snapshot_for_every_chart(client, monkeypatch):
    payload = setup_dashboard(client)
    import src.modules.e2e_dashboards as dashboards_module

    calls = []
    original = dashboards_module.normalise_operator_aliases

    def tracked(frame):
        calls.append(len(frame))
        return original(frame)

    monkeypatch.setattr(dashboards_module, 'normalise_operator_aliases', tracked)
    preview = client.post('/api/e2e-dashboards/prepare', json=payload).json()
    for index in (0, 1, 2):
        assert client.get(f"/api/e2e-dashboards/preview/{preview['token']}/{index}.png").status_code == 200
    assert calls == [3]
    assert client.get(f"/api/e2e-dashboards/preview/{preview['token']}/0.png").status_code == 200
    assert calls == [3]


def test_dashboard_is_available_to_workspace_users(client):
    setup_dashboard(client)
    client.get('/logout')
    client.post('/login', data={'username': 'admin', 'password': 'admin123'})
    assert client.get('/e2e-dashboards').status_code == 200
    assert client.get('/api/e2e-dashboards').status_code == 200
    assert client.put('/api/e2e-dashboards/test', json=definition().model_dump(mode='json')).status_code == 200
    assert 'href="/e2e-dashboards"' in client.get('/datasets-analysis').text
    client.get('/logout')
    client.post('/login', data={'username': 'demo', 'password': 'demo123'})
    assert client.get('/e2e-dashboards').status_code == 200
    assert 'href="/e2e-dashboards"' in client.get('/datasets-analysis').text
