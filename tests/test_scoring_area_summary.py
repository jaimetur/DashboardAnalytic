"""National & Area Summary: the National scoring next to chosen cities and each Region, Cluster or City."""
from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation

from src.modules import scoring as scoring_engine
from src.modules.scoring import most_reliable_result
from src.modules.scoring_area_summary import (
    area_summary_levels, build_area_summary, calculate_area_summary, normalize_area_summary_request,
)
from src.modules.scoring_exports import export_scoring_report
from src.modules.scoring_reports import default_scenario, has_geographic_filters, normalize_report_configuration
from tests.scoring_synthetic import synthetic_result
from tests.test_scoring_exports import TEMPLATE
from tests.test_scoring_most_reliable import _reliable_configuration

SUMMARY_TITLE = 'Best Network — National & Areas'


def _calls(city, region, statuses):
    return [{'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': 'City',
             'Session_Type': 'CALL', 'Call_Status': status, 'City': city, 'Region': region}
            for operator, values in statuses.items() for status in values]


ROWS = [
    *_calls('London', 'Vendor A', {'EE': ['Completed'] * 4, 'Three UK': ['Completed', 'Completed', 'Completed', 'Failed']}),
    *_calls('Leeds', 'VMO2 North', {'EE': ['Completed', 'Failed'], 'Three UK': ['Completed', 'Completed']}),
]


def _calculation():
    configuration = _reliable_configuration()
    kwargs = {'baseline_operator': 'EE', 'configuration': configuration}
    frames = {'voice': pd.DataFrame(ROWS)}
    result = scoring_engine.calculate_scoring(frames, ['Operator'], **kwargs)
    job = {'levels': ['Operator'], 'aggregation_levels': ['Operator'], 'aggregation_contract_version': 2,
           'baseline_operator': 'EE', 'configuration': configuration}
    return job, result, frames, kwargs


def _kpi_values(result):
    return {(row.get('city') or row.get('region'), row['operator'], row['environment'], row['kpi_code']): row['value']
            for row in result['scoring'] if row['value'] is not None}


def test_area_summary_request_needs_the_whole_country():
    assert normalize_area_summary_request(None, {}) is None
    assert normalize_area_summary_request({'breakdown': 'None', 'cities': []}, {}) is None
    assert normalize_area_summary_request({'breakdown': 'region', 'cities': ['London', 'london', ' ']}, {}) == {
        'breakdown': 'Region', 'cities': ['London'], 'time_split': 'Campaign'}
    with pytest.raises(ValueError, match='whole country'):
        normalize_area_summary_request({'breakdown': 'Region'}, {'City': ['London']})
    # Operator or Campaign filters keep the whole country.
    assert normalize_area_summary_request({'cities': ['London']}, {'Operator': ['EE']}) == {
        'breakdown': None, 'cities': ['London'], 'time_split': 'Campaign'}
    assert area_summary_levels(['Operator', 'Campaign'], {'breakdown': 'Region', 'cities': ['London']}) == ['Region', 'City']
    assert area_summary_levels(['Operator', 'Region'], {'breakdown': 'Region', 'cities': []}) == []


def test_each_area_is_scored_with_its_own_measurements():
    job, result, frames, kwargs = _calculation()
    summary = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                     {'breakdown': 'Region', 'cities': ['london', 'Bristol'], 'time_split': 'All'}, kwargs)
    assert [item['kind'] for item in summary['passes']] == ['National', 'City', 'Region']
    national, cities, regions = summary['passes']
    # Every selected CDR in one value, as the job: the National calculation is the job's own, not calculated again.
    assert national['reuses_job'] and 'result' not in national
    assert {row['city'] for row in cities['result']['scoring']} == {'London'}
    assert {row['region'] for row in regions['result']['scoring']} == {'Vendor A', 'VMO2 North'}
    assert 'points_loss' not in cities['result'] and 'configuration' not in cities['result']
    # London scores with its own calls only, the same as a calculation filtered to London.
    london = scoring_engine.calculate_scoring({'voice': pd.DataFrame([row for row in ROWS if row['City'] == 'London'])},
                                              ['Operator', 'City'], **kwargs)
    assert _kpi_values(cities['result']) == _kpi_values(london)
    # A chosen city without measurements is left out with a warning.
    assert summary['warnings'] == []
    missing = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                     {'breakdown': None, 'cities': ['Bristol']}, kwargs)
    assert [item['kind'] for item in missing['passes']] == ['National']
    assert missing['warnings'] == ['National & Area Summary: the selected CDRs have no Bristol measurements.']


def test_the_summary_lists_national_the_chosen_cities_and_the_breakdown():
    job, result, frames, kwargs = _calculation()
    result['area_summary'] = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                                    {'breakdown': 'Region', 'cities': ['London']}, kwargs)
    summaries = build_area_summary(job, result)
    combined = next(item for item in summaries if item['environment'] == 'Combined')
    assert [area['label'] for area in combined['areas']] == ['National', 'London', 'Vendor A', 'VMO2 North']
    national = {row['operator']: row for row in combined['areas'][0]['rows']}
    # The reference first; the GAP is each operator's points minus the reference's.
    assert combined['areas'][0]['rows'][0]['operator'] == 'EE' and national['EE']['gap'] is None
    assert national['Three UK']['gap'] == pytest.approx(national['Three UK']['points'] - national['EE']['points'])
    # One campaign: no change from a previous campaign.
    assert combined['previous_campaign'] is None and national['EE']['delta'] is None
    # The Most Reliable scoring rates the area calculations with its own points.
    reliable = most_reliable_result(result, 'EE')
    reliable_passes = [item for item in reliable['area_summary']['passes'] if 'result' in item]
    assert reliable_passes and all(item['result']['scoring_kind'] == 'most_reliable' for item in reliable_passes)


def _synthetic_report(campaigns=('2026-Q1', '2026-Q2'), with_summary=True):
    levels = ('Operator', 'Campaign')
    job, result = synthetic_result(levels=levels, campaigns=campaigns)
    _job, london = synthetic_result(levels=('Operator', 'City', 'Campaign'), cities=('London',), campaigns=campaigns)
    _job, areas = synthetic_result(levels=('Operator', 'City', 'Campaign'), cities=('Vendor A', 'VMO2 North'),
                                   campaigns=campaigns, seed=5)
    rename = lambda row: {('region' if key == 'city' else key): value for key, value in row.items()}
    regions = {**areas, 'scoring': [rename(row) for row in areas['scoring']],
               'totals': [rename(row) for row in areas['totals']]}
    summary = {'breakdown': 'Region', 'cities': ['London'], 'passes': [
        {'kind': 'National', 'field': None, 'levels': list(levels), 'reuses_job': True},
        {'kind': 'City', 'field': 'city', 'levels': [*levels, 'City'], 'cities': ['London'], 'reuses_job': False,
         'result': {key: value for key, value in london.items() if key != 'configuration'}},
        {'kind': 'Region', 'field': 'region', 'levels': [*levels, 'Region'], 'reuses_job': False,
         'result': {key: value for key, value in regions.items() if key != 'configuration'}},
    ]}
    # The job of a scenario without the National & area summary has none.
    if with_summary:
        result['area_summary'] = summary
    scenario = default_scenario(name='National', operators=['Vodafone UK', 'Three UK'])
    scenario['scorings']['most_reliable']['enabled'] = False
    scenario = normalize_report_configuration({'scenarios': [scenario]})['scenarios'][0]
    return Presentation(BytesIO(export_scoring_report([{'scenario': scenario, 'job': job, 'result': result}], TEMPLATE)))


def test_report_adds_the_national_and_areas_slide():
    presentation = _synthetic_report()
    slides = [slide for slide in presentation.slides
              if slide.shapes.title is not None and slide.shapes.title.text_frame.text.startswith(SUMMARY_TITLE)]
    assert len(slides) == 1
    table = next(shape.table for shape in slides[0].shapes if shape.has_table)
    assert [table.cell(0, column).text for column in range(len(table.columns))] == [
        'Area', 'Operator', 'Score\n2026-Q2', 'Δ vs\n2026-Q1', 'Δ vs\nEE']
    areas = [table.cell(row, 0).text for row in range(1, len(table.rows)) if table.cell(row, 0).text]
    assert areas == ['National', 'London', 'Vendor A', 'VMO2 North']
    # The reference and the operators to compare (Vodafone and Three), not O2.
    operators = {table.cell(row, 1).text for row in range(1, len(table.rows))}
    assert operators == {'EE', 'Vodafone UK', 'Three UK'}
    charts = [shape.name for shape in slides[0].shapes if shape.has_chart]
    assert charts == [f'Scoring Area Summary Chart {area}' for area in areas]


def test_jobs_without_area_summary_have_no_slide():
    presentation = _synthetic_report(with_summary=False)
    assert not any(slide.shapes.title is not None and slide.shapes.title.text_frame.text.startswith(SUMMARY_TITLE)
                   for slide in presentation.slides)


def test_scenarios_of_the_whole_country_calculate_their_area_summary(monkeypatch):
    import src.DriveTestAnalyzer as core

    class Repository:
        def list_main_cities(self):
            return ['Leeds', 'Bristol']

    scenario = normalize_report_configuration({'scenarios': [default_scenario(name='National')]})['scenarios'][0]
    # By default a scenario of the whole country has it, with the four regions and London.
    assert scenario['area_summary'] == {'enabled': True, 'time_split': 'Campaign', 'breakdown': 'Region',
                                        'cities': ['London'], 'main_cities': False}
    assert core._scenario_area_summary(Repository(), scenario) == {'breakdown': 'Region', 'cities': ['London'],
                                                                   'time_split': 'Campaign'}
    scenario['area_summary'] = {'enabled': True, 'breakdown': 'None', 'cities': ['London'], 'main_cities': True}
    assert core._scenario_area_summary(Repository(), scenario) == {
        'breakdown': None, 'cities': ['London', 'Leeds', 'Bristol'], 'time_split': 'Campaign'}
    # Without the whole country there is no National scoring.
    for filtered in ({'context_filters': {'Region': ['Vendor A']}}, {'main_cities': True}):
        assert has_geographic_filters({**scenario, **filtered})
        assert core._scenario_area_summary(Repository(), {**scenario, **filtered}) is None
    london = normalize_report_configuration({'scenarios': [default_scenario(name='London', context_filters={
        'City': ['London']})]})['scenarios'][0]
    assert london['area_summary']['enabled'] is False
    # Nor when it is not checked.
    scenario['area_summary']['enabled'] = False
    assert core._scenario_area_summary(Repository(), scenario) is None

    calculated = []

    def create(_repository, _selected, levels, *_args, area_summary=None, **_kwargs):
        calculated.append(area_summary)
        return {'id': 1, 'status': 'completed', 'levels': list(levels)}, True

    monkeypatch.setattr(core, 'validate_complete_scoring_cdr_selection', lambda *_args, **_kwargs: [1])
    monkeypatch.setattr(core, 'create_scoring_job', create)
    monkeypatch.setattr(core, 'get_scoring_job', lambda _repository, job_id, include_result=False: {
        'id': job_id, 'status': 'completed', 'result': {}})
    core._scenario_scoring_job(Repository(), {'dataset_ids': [1], 'nr_mode': 'NSA'},
                               normalize_report_configuration({'scenarios': [default_scenario()]})['scenarios'][0], 'super')
    # Also when the scenario is calculated again without a Campaign level of a single value.
    assert calculated and all(item == {'breakdown': 'Region', 'cities': ['London'], 'time_split': 'Campaign'}
                              for item in calculated)


def _timed_calls():
    rows = []
    for day, statuses in (('2026-02-10 09:00:00', {'EE': ['Completed', 'Completed'], 'Three UK': ['Failed', 'Completed']}),
                          ('2026-02-17 09:00:00', {'EE': ['Completed', 'Failed'], 'Three UK': ['Completed', 'Completed']})):
        for row in _calls('London', 'Vendor A', statuses):
            rows.append({**row, 'Campaign': 'UK_Q1_2026', 'Event_Start_Time': day})
    return rows


def test_the_time_split_divides_the_campaign_level_by_period():
    from src.modules.scoring import normalize_time_split

    configuration = _reliable_configuration()
    rows = _timed_calls()
    kwargs = {'baseline_operator': 'EE', 'configuration': configuration}
    weekly = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator', 'Campaign'],
                                              time_split='Weekly', **kwargs)
    assert {row['campaign'] for row in weekly['scoring']} == {'2026-W07', '2026-W08'}
    # The campaigns of the CDRs stay listed as they are.
    assert weekly['campaigns'] == ['UK_Q1_2026']
    # Each week is scored with its own tests.
    first_week = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows[:4])}, ['Operator'], **kwargs)
    week_values = {(row['operator'], row['kpi_code']): row['value'] for row in weekly['scoring']
                   if row['campaign'] == '2026-W07' and row['environment'] == 'DriveCity' and row['value'] is not None}
    assert week_values == {(row['operator'], row['kpi_code']): row['value'] for row in first_week['scoring']
                           if row['environment'] == 'DriveCity' and row['value'] is not None}
    monthly = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator', 'Campaign'],
                                               time_split='monthly', **kwargs)
    assert {row['campaign'] for row in monthly['scoring']} == {'2026-02'}
    # Without the Campaign level every selected CDR is aggregated into one value.
    pooled = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator'], time_split='Weekly', **kwargs)
    assert {row['campaign'] for row in pooled['scoring']} == {None}
    # Rows without a start time cannot be placed in a period.
    untimed = scoring_engine.calculate_scoring({'voice': pd.DataFrame(ROWS)}, ['Operator', 'Campaign'],
                                               time_split='Weekly', **kwargs)
    assert 'Voice: no test start time; its rows cannot be split by weekly period.' in untimed['warnings']
    assert normalize_time_split('Campaign') is None and normalize_time_split('YEARLY') == 'Yearly'


def test_the_area_summary_has_its_own_time_split(monkeypatch):
    import src.DriveTestAnalyzer as core

    configuration = _reliable_configuration()
    rows = _timed_calls()
    kwargs = {'baseline_operator': 'EE', 'configuration': configuration}
    frames = {'voice': pd.DataFrame(rows)}
    # The job itself splits by campaign; only the area summary is split by week.
    request = normalize_area_summary_request({'cities': ['London'], 'time_split': 'weekly'}, {})
    assert request == {'breakdown': None, 'cities': ['London'], 'time_split': 'Weekly'}
    assert set(area_summary_levels(['Operator', 'Campaign'], request)) == {'City', 'Event_Start_Time', 'Call_Start_Time',
                                                                            'Test_Start_Time'}
    summary = calculate_area_summary(scoring_engine.calculate_scoring, frames, ['Operator', 'Campaign'], request, kwargs)
    national, london = summary['passes']
    assert not national['reuses_job'] and summary['time_split'] == 'Weekly'
    assert {row['campaign'] for row in national['result']['scoring']} == {'2026-W07', '2026-W08'}
    assert {row['campaign'] for row in london['result']['scoring']} == {'2026-W07', '2026-W08'}
    # All selected CDRs: one value; the job split by campaign keeps its own result for itself.
    pooled = calculate_area_summary(scoring_engine.calculate_scoring, frames, ['Operator'],
                                    {'breakdown': None, 'cities': ['London'], 'time_split': 'All'}, kwargs)
    assert pooled['passes'][0]['reuses_job'] and pooled['passes'][0]['levels'] == ['Operator']
    assert normalize_area_summary_request({'cities': ['London']}, {})['time_split'] == 'Campaign'

    scenario = normalize_report_configuration({'scenarios': [default_scenario()]})['scenarios'][0]
    assert 'time_split' not in scenario and scenario['area_summary']['time_split'] == 'Campaign'
    scenario['area_summary']['time_split'] = 'Monthly'
    assert core._scenario_area_summary(None, scenario) == {'breakdown': 'Region', 'cities': ['London'],
                                                           'time_split': 'Monthly'}
