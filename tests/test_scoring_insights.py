"""NetCheck-style scoring insights: location cards, KPI GAP profiles, trends and campaign comparisons."""
from __future__ import annotations

from io import BytesIO

import pytest
from pptx import Presentation

from src.modules.scoring_exports import export_scoring_powerpoint
from src.modules.scoring_views import build_scoring_views
from tests.scoring_synthetic import CITIES, most_reliable_result, synthetic_result
from tests.test_scoring_exports import TEMPLATE


def _insights(levels=('Operator', 'City', 'Campaign'), **kwargs):
    job, result = synthetic_result(levels=levels, **kwargs)
    return job, result, build_scoring_views(job, result)['insights']


def _combined(items):
    return [item for item in items if item['environment'] == 'Combined']


def test_location_cards_scale_each_city_to_the_full_maximum():
    _job, result, insights = _insights(levels=('Operator', 'City'))
    groups = _combined(insights['location_cards'])
    assert len(groups) == 1
    group = groups[0]
    assert group['level'] == 'City'
    assert [card['label'] for card in group['cards']] == sorted(CITIES)
    assert group['maximum'] == pytest.approx(1000)
    assert group['max_voice'] + group['max_data'] == pytest.approx(1000)
    london = next(card for card in group['cards'] if card['label'] == 'London')
    ee = next(item for item in london['operators'] if item['operator'] == 'EE')
    expected = next(row for row in result['totals'] if row['city'] == 'London' and row['operator'] == 'EE'
                    and row['environment'] == 'Combined' and row['category'] == 'Overall')
    assert ee['total'] == pytest.approx(expected['weighted_points'])
    assert ee['voice'] + ee['data'] == pytest.approx(ee['total'])


def test_location_cards_need_two_to_twelve_locations():
    _job, _result, insights = _insights(levels=('Operator', 'City'), cities=('London',))
    assert insights['location_cards'] == []
    many = tuple(f'City {index}' for index in range(13))
    _job, _result, insights = _insights(levels=('Operator', 'City'), cities=many)
    assert insights['location_cards'] == []


def test_location_cards_scale_a_city_measured_in_one_environment():
    job, result = synthetic_result(levels=('Operator', 'City'), cities=('London', 'Leeds'))
    result['scoring'] = [row for row in result['scoring']
                         if not (row['city'] == 'London' and row['environment'] == 'DriveConnectionroad')]
    insights = build_scoring_views(job, result)['insights']
    group = _combined(insights['location_cards'])[0]
    london = next(card for card in group['cards'] if card['label'] == 'London')
    ee = next(item for item in london['operators'] if item['operator'] == 'EE')
    city_points = sum(row['weighted_points'] for row in result['scoring']
                      if row['city'] == 'London' and row['operator'] == 'EE' and row['environment'] == 'DriveCity')
    assert group['scaled'] is True
    assert ee['total'] == pytest.approx(city_points * 1000 / 650)


def test_campaign_trend_needs_more_than_four_campaigns():
    _job, _result, insights = _insights(campaigns=('2026-Q1', '2026-Q2', '2025-Q4', '2025-Q3'))
    assert insights['campaign_trends'] == []
    _job, _result, insights = _insights()
    trends = _combined(insights['campaign_trends'])
    assert len(trends) == len(CITIES)
    assert trends[0]['campaigns'] == ['2025-Q2', '2025-Q3', '2025-Q4', '2026-Q1', '2026-Q2']
    assert {series['operator'] for series in trends[0]['series']} == {'EE', 'O2 UK', 'Three UK', 'Vodafone UK'}


def test_campaign_comparison_subtracts_the_previous_campaign():
    _job, result, insights = _insights(campaigns=('2026-Q1', '2026-Q2'))
    comparison = next(item for item in _combined(insights['campaign_comparisons'])
                      if item['operator'] == 'EE' and 'London' in item['title'])
    assert (comparison['previous_campaign'], comparison['latest_campaign']) == ('2026-Q1', '2026-Q2')
    deltas = [row['delta'] for row in comparison['rows']]
    assert deltas == sorted(deltas)
    assert comparison['total_delta'] == pytest.approx(sum(deltas))

    def total(campaign):
        return next(row['weighted_points'] for row in result['totals']
                    if row['city'] == 'London' and row['operator'] == 'EE' and row['campaign'] == campaign
                    and row['environment'] == 'Combined' and row['category'] == 'Overall')
    assert comparison['total_delta'] == pytest.approx(total('2026-Q2') - total('2026-Q1'))


def test_campaign_comparison_keeps_every_campaign_to_compare_any_two():
    _job, result, insights = _insights()
    comparison = next(item for item in _combined(insights['campaign_comparisons'])
                      if item['operator'] == 'EE' and 'London' in item['title'])
    assert comparison['campaigns'] == ['2025-Q2', '2025-Q3', '2025-Q4', '2026-Q1', '2026-Q2']
    # The points of each KPI in every campaign add up to the scoring of that campaign.

    def total(campaign):
        return next(row['weighted_points'] for row in result['totals']
                    if row['city'] == 'London' and row['operator'] == 'EE' and row['campaign'] == campaign
                    and row['environment'] == 'Combined' and row['category'] == 'Overall')
    for index, campaign in enumerate(comparison['campaigns']):
        points = sum(kpi['values'][index]['points'] for kpi in comparison['kpis'] if kpi['values'][index])
        assert points == pytest.approx(total(campaign))
    # The latest and previous campaigns give the rows of the default comparison.
    latest = {row['kpi_code']: row['delta'] for row in comparison['rows']}
    for kpi in comparison['kpis']:
        before, after = kpi['values'][-2], kpi['values'][-1]
        if before and after:
            assert latest[kpi['kpi_code']] == pytest.approx(after['points'] - before['points'])


def test_kpi_gap_profiles_rank_points_lost_and_reference_gaps():
    _job, _result, insights = _insights(levels=('Operator',))
    profile = next(item for item in _combined(insights['kpi_gap_profiles']) if item['operator'] == 'Three UK')
    assert profile['reference'] == 'EE'
    assert profile['columns'] == ['DriveCity', 'DriveConnectionroad']
    lost = [row['total_gap_to_maximum'] for row in profile['to_maximum']]
    assert lost == sorted(lost, reverse=True) and min(lost) >= 0
    gaps = [row['total_gap_to_reference'] for row in profile['to_reference']]
    assert gaps == sorted(gaps)
    row = profile['to_maximum'][0]
    assert row['total_gap_to_maximum'] == pytest.approx(sum(row['gap_to_maximum'].values()))
    # Best Network marks the KPIs of the Most Reliable scoring.
    assert profile['underline_most_reliable'] is True
    assert {item['kpi_code'] for item in profile['to_maximum'] if item['most_reliable']} == {
        'K1', 'K4', 'K7', 'K8', 'K12', 'K14', 'K16', 'K21', 'K27', 'K28'}


def test_most_reliable_insights_use_its_points():
    job, result = synthetic_result(levels=('Operator', 'City'))
    reliable = most_reliable_result(result, 'EE')
    insights = build_scoring_views(job, reliable)['insights']
    group = _combined(insights['location_cards'])[0]
    assert group['maximum'] == pytest.approx(1000)
    assert (group['max_voice'], group['max_data']) == (pytest.approx(350), pytest.approx(650))
    profile = _combined(insights['kpi_gap_profiles'])[0]
    assert len(profile['to_maximum']) == 10
    assert profile['underline_most_reliable'] is False


def test_powerpoint_adds_the_insight_slides_to_all_environments():
    job, result = synthetic_result()
    presentation = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE)))
    titles = [slide.shapes.title.text_frame.text if slide.shapes.title is not None else '' for slide in presentation.slides]
    first = [title.split('\n')[0] for title in titles]
    for heading in ('Best Network Scoring per City', 'Best Network Scoring Trend', 'Best Network Campaign Comparison',
                    'Most Reliable Network Scoring per City', 'Most Reliable Network Scoring Trend'):
        assert heading in first
    assert any(title.startswith('KPI GAP Profile — Three UK vs EE') for title in first)
    # Only All Environments gets them, and the cards only for the latest campaign.
    insight_titles = [title for title in titles if 'Scoring per City' in title or 'Profile' in title]
    # All Environments subtitles name only the scoring (and the context), never an environment.
    assert all('Drive - ' not in title for title in insight_titles)
    assert sum(1 for title in first if title == 'Best Network Scoring per City') == 1
    cards = next(slide for slide, title in zip(presentation.slides, first) if title == 'Best Network Scoring per City')
    charts = [shape for shape in cards.shapes if shape.name.startswith('Scoring Location Chart')]
    assert len(charts) == len(CITIES)
    # The cards show the operator labels as mapped, never shortened.
    assert set(charts[0].chart.plots[0].categories) == {'EE', 'O2 UK', 'Three UK', 'Vodafone UK'}
    assert next(shape for shape in cards.shapes if shape.name == 'Scoring Location Cards Panel').text_frame.text.count(
        'Total: 1,000 points') == 1


def test_profile_bars_line_up_with_their_rows():
    job, result = synthetic_result(levels=('Operator',))
    presentation = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE)))
    slide = next(slide for slide in presentation.slides
                 if slide.shapes.title is not None and slide.shapes.title.text_frame.text.startswith('KPI GAP Profile'))
    table_shape = next(shape for shape in slide.shapes if shape.name == 'Scoring KPI GAP Maximum Table')
    table = table_shape.table
    first_row_top = table_shape.top + table.rows[0].height + table.rows[1].height
    bars = sorted((shape for shape in slide.shapes if shape.name == 'Scoring GAP Bar' and shape.left < table_shape.left
                   + table_shape.width), key=lambda shape: shape.top)
    assert bars and bars[0].top > first_row_top
    assert bars[0].top < first_row_top + table.rows[2].height
    note = next(shape for shape in slide.shapes if shape.name == 'Scoring KPI GAP Profile Note')
    assert note.top > table_shape.top + sum(row.height for row in table.rows)
    bold = [run.text for run in note.text_frame.paragraphs[0].runs if run.font.bold]
    operator = slide.shapes.title.text_frame.text.split('\n')[0].removeprefix('KPI GAP Profile — ').split(' vs ')[0]
    assert bold[0] == operator and any('Best Network points' in text for text in bold)


def test_kpi_gap_profile_slides_cover_every_selected_city():
    from src.modules.scoring_insight_slides import MAX_PROFILE_SCOPES, _block_profiles

    cities = ['Belfast', 'Bristol', 'Cardiff', 'Edinburgh', 'Leeds', 'London', 'Sheffield']
    profiles = [{'operator': operator, 'reference': 'EE', 'context': {'city': city}}
                for city in cities for operator in ('VF', 'EE', '3')]
    kept = _block_profiles(profiles)
    # One profile per city and compared operator; the reference has none.
    assert [(item['context']['city'], item['operator']) for item in kept] == [
        (city, operator) for city in cities for operator in ('VF', '3')]
    assert MAX_PROFILE_SCOPES == 12
