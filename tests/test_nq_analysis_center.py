"""NQ Analysis Center: field catalog by section, automatic NQ Call Status, root cause catalog, RCA script results,
recommendation from the three root cause sources, lifecycle and RCA insights, the Excel layout and the catalog workbook."""

import io
import json
import sqlite3

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import src.DriveTestAnalyzer as core
from src.modules import non_qualified_calls as nq
from src.modules import nq_catalog, nq_rca
from tests.test_non_qualified_calls import add_cdr, enable_module, login, query


@pytest.fixture(autouse=True)
def indexing_runs_at_once(monkeypatch):
    monkeypatch.setattr(core, 'submit_background_task', lambda callback, *args: callback(*args))


def calls_frame() -> pd.DataFrame:
    """Three Non-Qualified Voice calls with a JOIN_ID and NetCheck RCA, and one completed call."""
    return pd.DataFrame({
        'JOIN_ID': ['0xA1', '0xB2', '0xC3', '0xD4'],
        'Session_ID_A': [1, 2, 3, 4],
        'Operator': ['Vodafone UK', 'Vodafone UK', 'EE', 'EE'],
        'Campaign': ['UK_Q3_2026'] * 4,
        'Session_Type': ['CALL', 'WhatsApp CALL', 'MultiRAB CALL', 'CALL'],
        'Call_Status': ['Failed', 'Dropped', 'Failed', 'Completed'],
        'status': ['Failed', 'Dropped', 'Failed', 'Completed'],
        'Call_Start_Time': ['2026-09-01 10:00:00', '2026-09-02 11:00:00', '2026-09-03 12:00:00', '2026-09-03 13:00:00'],
        'Technology': ['LTE', 'LTE', 'NR', 'LTE'],
        'Failure_Classification': ['RF_Problems', 'RF Problems', 'VoLTE Problems', ''],
        'Failure_Category': ['Coverage problem', 'Interference problem', 'VoLTE Core', ''],
        'Failure_Subcategory': ['DL coverage problems', 'DL interference problems', 'No QCI1 bearer activation', ''],
        'Cell_ID_A': ['123456', '123457', '9876543210', '1'],
        'Vendor': ['Ericsson', 'Huawei', 'Nokia', 'Nokia'],
        'Call_Start_Latitude_A': [51.5, 51.6, 51.7, 51.8],
        'Call_Start_Longitude_A': [-0.1, -0.2, -0.3, -0.4],
    })


def setup_calls(client, tmp_path) -> dict[str, dict]:
    enable_module()
    add_cdr(tmp_path, 'UK_Voice_Q3.xlsx', 'voice', calls_frame())
    login(client)
    return {call['join_id']: call for call in query(client)['calls']}


def patch(client, call, changes):
    response = client.patch(f"/api/non-qualified-calls/calls/{call['call_key']}", json={'changes': changes})
    assert response.status_code == 200, response.text
    return response.json()['call']


def test_new_workspaces_get_the_catalog_of_the_meeting_workbook(client, tmp_path):
    setup_calls(client, tmp_path)
    state = client.get('/api/non-qualified-calls/state').json()
    fields = state['fields']
    # The 37 fields of the meeting workbook, with the Last Cell ID and Failure Comment of the CDR in Failure Details.
    assert len(fields) == 39
    assert [field['key'] for field in fields if field['section'] == 'failure_event'][-2:] == ['vendor_when_failure', 'failure_comment']
    assert next(field for field in fields if field['key'] == 'last_cell_id')['source_ref'] == 'last_cell_id'
    assert [field['key'] for field in fields if field['section'] == 'analysis'][:3] == ['analysis_status', 'tunnel_failure', 'needed_technology']
    assert [section['label'] for section in state['sections']][2] == 'Failure Details'
    assert [field['label'] for field in fields[:3]] == ['Auto_RCA_Category_A', 'Auto_RCA_Subcategory_A', 'Auto_RCA_Category_B']
    by_key = {field['key']: field for field in fields}
    assert by_key['rca_suggested_cause']['source'] == 'rca' and by_key['rca_suggested_cause']['in_table']
    assert by_key['netcheck_suggested_cause']['source_ref'] == 'failure_subcategory'
    assert by_key['selected_root_category']['source'] == 'tracking' and by_key['nq_call_status']['section'] == 'general'
    assert by_key['analysis_status']['section'] == 'analysis'
    assert [option['name'] for option in by_key['planned_status']['options']] == ['Not Planned', 'Planned']
    assert [section['key'] for section in state['sections']] == list(nq_catalog.SECTION_KEYS)
    assert state['sections'][1]['color'] == '#FFC000'
    assert [item['name'] for item in state['options']['statuses']][:2] == ['Not Attended', 'Open']
    assert 'RAN Planning Coverage' in [item['name'] for item in state['options']['teams']]
    root = state['root_causes']
    assert 'Poor Coverage LTE' in [item['name'] for item in root['categories']]
    causes = {item['name']: item for item in root['causes']}
    # The causes of the meeting and those of the first catalog, each in a domain.
    assert causes['NLOS due to clutter or terrain profile']['domain'] == 'RF' and causes['DNS']['domain'] == 'Protocol'
    assert state['status_rules'][0]['status'] == 'Closed'


def test_status_follows_the_phase_fields_until_it_is_set_by_hand(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    call = calls['0xA1']
    assert call['status'] == 'Not Attended' and call['status_mode'] == 'auto'
    # A comment attends the call.
    assert client.post(f"/api/non-qualified-calls/calls/{call['call_key']}/comments", json={'body': 'Looking at it'}).status_code == 200
    assert {item['join_id']: item for item in query(client)['calls']}['0xA1']['status'] == 'Open'
    updated = patch(client, call, {'fields': {'analysis_status': 'Ongoing'}})
    assert updated['status'] == 'Under Analysis'
    updated = patch(client, call, {'fields': {'implementation_status': 'Under evaluation'}})
    assert updated['status'] == 'Under Implementation'
    updated = patch(client, call, {'fields': {'implementation_status': 'Implemented'}})
    assert updated['status'] == 'Closed'
    # Chosen by hand, the status no longer follows the fields...
    updated = patch(client, call, {'status': 'Under Planning'})
    assert updated['status'] == 'Under Planning' and updated['status_mode'] == 'manual'
    updated = patch(client, call, {'fields': {'analysis_status': 'Finished'}})
    assert updated['status'] == 'Under Planning'
    # ...until it goes back to automatic.
    updated = patch(client, call, {'status_mode': 'auto'})
    assert updated['status'] == 'Closed' and updated['status_mode'] == 'auto'
    detail = client.get(f"/api/non-qualified-calls/calls/{call['call_key']}").json()
    assert {'status_auto', 'status', 'status_mode'} <= {entry['field'] for entry in detail['history']}
    # A status set by the rules follows, in the history, the change that moved it.
    fields = [entry['field'] for entry in detail['history']]
    assert fields[fields.index('field:analysis_status') + 1] == 'status_auto' and fields[-2:] == ['status_mode', 'status_auto']
    # The rules adapt to renamed list values and the calls follow new rules at once.
    fields = client.get('/api/non-qualified-calls/state').json()['fields']
    for field in fields:
        if field['key'] == 'implementation_status':
            field['options'] = [{**option, 'previous': option['name']} for option in field['options']]
            field['options'][3]['name'] = 'Done'
    assert client.put('/api/non-qualified-calls/fields', json={'fields': fields}).status_code == 200
    rules = client.get('/api/non-qualified-calls/state').json()['status_rules']
    assert ['Done'] in [condition['values'] for rule in rules for condition in rule['conditions']]
    other = calls['0xB2']
    patch(client, other, {'fields': {'planned_status': 'Planned'}})
    rules = [{'status': 'Closed', 'conditions': [{'field': 'planned_status', 'op': 'in', 'values': ['Planned']}]}]
    assert client.put('/api/non-qualified-calls/status-rules', json={'rules': rules}).status_code == 200
    assert {item['join_id']: item for item in query(client)['calls']}['0xB2']['status'] == 'Closed'
    bad = client.put('/api/non-qualified-calls/status-rules', json={'rules': [{'status': 'Nope', 'conditions': []}]})
    assert bad.status_code == 400


def test_default_statuses_replace_statuses_the_rules_gave_but_not_those_set_by_hand(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    state = client.get('/api/non-qualified-calls/state').json()
    teams = [{**item, 'previous': item['name'], 'members': []} for item in state['options']['teams']]
    renamed = [{**item, 'previous': item['name'], 'name': 'Analysis Started' if item['name'] == 'Under Analysis' else item['name']}
               for item in state['options']['statuses']]
    assert client.put('/api/non-qualified-calls/options', json={'statuses': renamed, 'teams': teams}).status_code == 200
    by_rules = patch(client, calls['0xA1'], {'fields': {'analysis_status': 'Ongoing'}})
    by_hand = patch(client, calls['0xB2'], {'status': 'Analysis Started'})
    assert (by_rules['status'], by_rules['status_mode'], by_hand['status_mode']) == ('Analysis Started', 'auto', 'manual')
    current = {item['name'] for item in renamed}
    defaults = [{**item, 'previous': item['name'] if item['name'] in current else ''} for item in state['default_options']['statuses']]
    response = client.put('/api/non-qualified-calls/options', json={'statuses': defaults, 'teams': teams})
    assert response.status_code == 400 and 'set by hand on 1 call' in response.text
    patch(client, calls['0xB2'], {'status_mode': 'auto'})
    assert client.put('/api/non-qualified-calls/options', json={'statuses': defaults, 'teams': teams}).status_code == 200
    statuses = [item['name'] for item in client.get('/api/non-qualified-calls/state').json()['options']['statuses']]
    assert statuses == [item['name'] for item in state['default_options']['statuses']]
    # The call the rules had moved follows them again.
    assert {item['join_id']: item for item in query(client)['calls']}['0xA1']['status'] in statuses


def test_root_category_and_cause_are_independent_and_set_the_domain(client, tmp_path):
    call = setup_calls(client, tmp_path)['0xA1']
    updated = patch(client, call, {'root_category': 'Poor Coverage LTE', 'root_cause': 'Tunnel Issue'})
    assert (updated['root_category'], updated['root_cause'], updated['root_domain']) == ('Poor Coverage LTE', 'Tunnel Issue', 'RF')
    assert client.patch(f"/api/non-qualified-calls/calls/{call['call_key']}", json={'changes': {'root_cause': 'Unknown cause'}}).status_code == 400
    # Renamed catalog values follow on the calls; values in use cannot be removed.
    root = client.get('/api/non-qualified-calls/root-causes').json()
    for item in root['categories']:
        item['previous'] = item['name']
        if item['name'] == 'Poor Coverage LTE':
            item['name'], item['domain'] = 'LTE Coverage', 'RAN'
    response = client.put('/api/non-qualified-calls/root-causes', json=root)
    assert response.status_code == 200, response.text
    updated = client.get(f"/api/non-qualified-calls/calls/{call['call_key']}").json()['call']
    assert (updated['root_category'], updated['root_domain']) == ('LTE Coverage', 'RAN')
    root = client.get('/api/non-qualified-calls/root-causes').json()
    root['causes'] = [item for item in root['causes'] if item['name'] != 'Tunnel Issue']
    assert client.put('/api/non-qualified-calls/root-causes', json=root).status_code == 400


def test_rca_script_results_are_imported_by_join_id_and_recommend_the_root_cause(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    csv = ('JOIN_ID;Auto_RCA_Category_A;Auto_RCA_Subcategory_A;RCA_Suggested_Category;RCA_Suggested_Cause;Confidence\n'
           '0xa1;RF;Coverage;Poor Coverage LTE;NLOS due to clutter or terrain profile;0.9\n'
           '0xUNKNOWN;RF;Coverage;Poor Coverage LTE;LOS but weak signal;0.4\n').encode()
    preview = client.post('/api/non-qualified-calls/rca/import', data={'preview': '1'},
                          files={'file': ('rca.csv', io.BytesIO(csv), 'text/csv')}).json()
    assert (preview['rows'], preview['matched'], preview['unmatched']) == (2, 1, 1)
    assert preview['extra'] == ['Confidence'] and client.get('/api/non-qualified-calls/rca').json()['total'] == 0
    imported = client.post('/api/non-qualified-calls/rca/import', files={'file': ('rca.csv', io.BytesIO(csv), 'text/csv')})
    assert imported.status_code == 200 and client.get('/api/non-qualified-calls/rca').json()['matched'] == 1
    call = {item['join_id']: item for item in query(client)['calls']}['0xA1']
    assert call['values']['rca_suggested_category'] == 'Poor Coverage LTE'
    assert call['values']['netcheck_suggested_cause'] == 'DL coverage problems'
    assert call['values']['session_type'] == 'CALL'
    recommendation = call['rca_recommendation']
    sources = {source['source']: source for source in recommendation['sources']}
    # The RCA script and NetCheck (mapped by keywords) agree on the category.
    assert sources['script']['category'] == 'Poor Coverage LTE' and sources['netcheck']['category'] == 'Poor Coverage LTE'
    assert recommendation['recommendation']['category'] == 'Poor Coverage LTE'
    assert recommendation['recommendation']['cause'] == 'NLOS due to clutter or terrain profile'
    assert recommendation['recommendation']['confidence'] == 'high'
    # Accepting the recommendations of several calls at once.
    keys = [calls['0xA1']['call_key'], calls['0xB2']['call_key']]
    assert client.post('/api/non-qualified-calls/calls/bulk', json={'call_keys': keys, 'changes': {'apply_recommendation': True}}).json() == {'changed': 2}
    decided = {item['join_id']: item for item in query(client)['calls']}
    assert decided['0xA1']['root_category'] == 'Poor Coverage LTE'
    assert decided['0xB2']['root_category'] == 'DL Interference LTE' and decided['0xB2']['root_cause'] == 'DL interference'
    # Previous decisions teach the recommendation of calls with the same NetCheck values.
    learned = nq.learned_decisions(sqlite3.connect(core.repository.db_path, factory=_RowConnection))
    assert learned[nq_rca.learned_key('netcheck', 'Coverage problem', 'DL coverage problems')] == {
        ('Poor Coverage LTE', 'NLOS due to clutter or terrain profile'): {calls['0xA1']['call_key']}}
    # A decision never teaches the call that made it.
    assert 'learned' not in {source['source'] for source in decided['0xA1']['rca_recommendation']['sources']}
    # Filtering and sorting by catalog fields of any source.
    filtered = query(client, filters={'fields': {'rca_suggested_category': ['Poor Coverage LTE']}})
    assert [item['join_id'] for item in filtered['calls']] == ['0xA1']
    derived = query(client, filters={'fields': {'session_type': ['MULTI-RAB']}})
    assert [item['join_id'] for item in derived['calls']] == ['0xC3']
    ordered = query(client, sort='field:session_type', direction='asc')['calls']
    assert [item['values']['session_type'] for item in ordered] == ['CALL', 'MULTI-RAB', 'WhatsApp']
    assert client.delete('/api/non-qualified-calls/rca').json()['deleted'] == 2


class _RowConnection(sqlite3.Connection):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.row_factory = sqlite3.Row


def test_lifecycle_and_rca_insights_compare_the_three_sources(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    csv = 'JOIN_ID,RCA_Suggested_Category,RCA_Suggested_Cause\n0xA1,Poor Coverage LTE,LOS but weak signal\n0xB2,E2E,Potential E2E Issue\n'
    client.post('/api/non-qualified-calls/rca/import', files={'file': ('rca.csv', io.BytesIO(csv.encode()), 'text/csv')})
    patch(client, calls['0xA1'], {'root_category': 'Poor Coverage LTE', 'root_cause': 'LOS but weak signal'})
    patch(client, calls['0xB2'], {'root_category': 'DL Interference LTE', 'fields': {'analysis_status': 'Ongoing'}})
    insights = client.post('/api/non-qualified-calls/rca/insights', json={'filters': {}}).json()
    assert insights['coverage'] == {'total': 3, 'script': 2, 'netcheck': 3, 'selected': 2}
    assert insights['agreement']['selected_script'] == {'agree': 1, 'compared': 2, 'rate': 50.0}
    assert insights['agreement']['selected_netcheck']['agree'] == 2
    assert insights['confusion']['script']['rows'] and insights['domains'][0]['domain'] == 'RF'
    assert insights['hierarchies']['netcheck'][0]['causes']
    lifecycle = client.post('/api/non-qualified-calls/lifecycle', json={'filters': {}}).json()
    pipeline = {item['status']: item['count'] for item in lifecycle['pipeline']}
    assert pipeline['Under Analysis'] == 1 and pipeline['Open'] == 1 and pipeline['Not Attended'] == 1
    phases = {phase['key']: phase for phase in lifecycle['phases']}
    assert next(item['count'] for item in phases['analysis_status']['items'] if item['value'] == 'Ongoing') == 1
    assert lifecycle['modes'] == {'auto': 3, 'manual': 0}
    assert any(move['to'] == 'Under Analysis' for move in lifecycle['moves'])
    # Both panels are slides of the PowerPoint export.
    document = client.post('/api/non-qualified-calls/export/powerpoint', json={'filters': {}})
    assert document.status_code == 200


def test_excel_export_groups_the_catalog_by_section_with_its_colours(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    call = calls['0xA1']
    client.post(f"/api/non-qualified-calls/calls/{call['call_key']}/comments", json={'body': 'First look'})
    client.post(f"/api/non-qualified-calls/calls/{call['call_key']}/comments", json={'body': 'Coverage hole'})
    patch(client, call, {'fields': {'findings': 'Hill behind the site', 'tunnel_failure': 'No'}})
    sheet = load_workbook(io.BytesIO(client.post('/api/non-qualified-calls/export', json={'filters': {}}).content))['NQ Calls']
    sections = [cell.value for cell in sheet[1] if cell.value]
    assert sections[:3] == ['CALL', 'RCA', 'GENERAL'] and sections[-4:-1] == ['ANALYSIS', 'IMPLEMENTATION', 'PLANNING']
    headers = [cell.value for cell in sheet[2]]
    # The catalog in its order, each field after its section title, with the follow-up companions.
    assert headers.index('Auto_RCA_Category_A') < headers.index('Selected Root Cause') < headers.index('NQ Call Status') \
        < headers.index('Session Type') < headers.index('Findings') < headers.index('Planned Date')
    # JOIN_ID and the call key first; the root cause decision ends the RCA block; the last update last.
    assert headers[:3] == ['JOIN_ID', 'Call Key', 'CDR'] and headers[-2:] == ['Updated By', 'Updated At']
    selected = headers.index('Selected Root Category')
    assert headers[selected - 4:selected] == ['Root Domain', 'Recommended Root Category', 'Recommended Root Cause',
                                              'Recommendation Sources'] and headers[selected + 1] == 'Selected Root Cause'
    assert headers[headers.index('NQ Call Status') + 1] == 'Status Set'
    # The status of each block stands out, darker than its section.
    fill = lambda label: sheet.cell(2, headers.index(label) + 1).fill.fgColor.rgb[-6:]
    assert fill('NQ Call Status') == nq._tint('#0070C0', nq.KEY_FIELD_TINT) != fill('Status Set')
    assert fill('Analysis Status') == nq._tint('#7030A0', nq.KEY_FIELD_TINT) != fill('Findings')
    assert fill('Implementation Status') == nq._tint('#00B050', nq.KEY_FIELD_TINT)
    assert fill('Planned Status') == nq._tint('#00843D', nq.KEY_FIELD_TINT)
    rca_fill = sheet.cell(1, headers.index('Auto_RCA_Category_A') + 1).fill.fgColor.rgb
    assert rca_fill.endswith('FFC000')
    rows = [dict(zip(headers, (cell.value for cell in row))) for row in sheet.iter_rows(min_row=3)]
    row = next(item for item in rows if item['JOIN_ID'] == '0xA1')
    assert row['Findings'] == 'Hill behind the site' and row['Tunnel Failure'] == 'No'
    comments = row['Comments'].split('\n')
    assert len(comments) == 2 and comments[0].endswith('super: First look') and comments[1].endswith('super: Coverage hole')
    assert row['Netcheck Suggested Cause'] == 'DL coverage problems' and row['Session Type'] == 'CALL'


def test_catalog_workbook_round_trip_and_the_meeting_layout(client, tmp_path):
    setup_calls(client, tmp_path)
    exported = client.get('/api/non-qualified-calls/catalog/export')
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content))
    assert {'About', 'Sections', 'Fields', 'Lists', 'Root Catalog', 'Statuses and Teams', 'Status Rules'} <= set(workbook.sheetnames)
    fields = workbook['Fields']
    fields.append([38, 'Analysis', 'Drive Route', '', 'list', 'user', '', 'Yes', 'Yes', 'Yes', 'No', '', 'Route of the retest'])
    workbook['Lists'].append(['Drive Route', 'M25', '#123456', ''])
    workbook['Lists'].append(['Drive Route', 'A1', '', ''])
    workbook['Root Catalog'].append(['Root Cause', 'Fibre cut', 'Operational', '', 'fibre'])
    workbook['Statuses and Teams'].append(['Team', 'Field Team', '#0f6f7d', '', ''])
    output = io.BytesIO()
    workbook.save(output)
    preview = client.post('/api/non-qualified-calls/catalog/import', data={'preview': '1'},
                          files={'workbook': ('catalog.xlsx', io.BytesIO(output.getvalue()))}).json()
    assert preview['added'] == ['Drive Route'] and preview['root_added']['causes'] == ['Fibre cut']
    assert preview['teams_added'] == ['Field Team']
    imported = client.post('/api/non-qualified-calls/catalog/import', files={'workbook': ('catalog.xlsx', io.BytesIO(output.getvalue()))})
    assert imported.status_code == 200, imported.text
    route = next(field for field in imported.json()['fields'] if field['label'] == 'Drive Route')
    assert [option['name'] for option in route['options']] == ['M25', 'A1'] and route['section'] == 'analysis' and route['in_table']
    assert len(imported.json()['fields']) == 40
    # The workbook of the meeting: sections in the first row, fields in the second, values below.
    meeting = Workbook()
    sheet = meeting.active
    sheet.append(['FAILURE EVENT', None, 'PLANNING'])
    sheet.append(['Band when Failure', 'Weather', 'Planned Date'])
    sheet.append(['TO BE ADDED'] * 3)
    sheet.append(['Lista customizable', 'Lista customizable', 'Fecha'])
    sheet.append(['', '', ''])
    sheet.append(['LTE2100', 'Rain', 'User Define'])
    sheet.append(['NR3500', 'Snow', None])
    output = io.BytesIO()
    meeting.save(output)
    catalog = nq_catalog.read_catalog_workbook(output.getvalue())
    by_label = {field['label']: field for field in catalog['fields']}
    assert by_label['Band when Failure']['key'] == 'band_when_failure'
    assert 'NR3500' in [option['name'] for option in by_label['Band when Failure']['options']]
    assert by_label['Weather']['section'] == 'failure_event' and [option['name'] for option in by_label['Weather']['options']] == ['Rain', 'Snow']
    assert by_label['Planned Date']['type'] == 'date' and by_label['Planned Date']['section'] == 'planning'


def test_tracking_document_carries_the_analysis_center(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    client.post('/api/non-qualified-calls/rca/import',
                files={'file': ('rca.csv', io.BytesIO(b'JOIN_ID,RCA_Suggested_Category\n0xA1,E2E\n'), 'text/csv')})
    patch(client, calls['0xA1'], {'root_category': 'E2E', 'status': 'Under Planning'})
    sections = client.get('/api/non-qualified-calls/state').json()['sections']
    sections[0]['label'] = 'Overview'
    assert client.put('/api/non-qualified-calls/sections', json={'sections': sections}).status_code == 200
    document = json.loads(nq.export_tracking_document(core.repository))
    assert document['version'] == 3 and document['sections'][0]['label'] == 'Overview'
    assert document['rca_results'][0]['join_id'] == '0xA1' and document['status_rules']
    assert document['tracking'][0]['status_mode'] == 'manual' and document['tracking'][0]['root_category'] == 'E2E'
    with core.repository.connection() as connection:
        connection.execute(f'DELETE FROM {nq.NQ_RCA_RESULTS_TABLE}')
        connection.execute(f'DELETE FROM {nq.NQ_CALL_TRACKING_TABLE}')
    core.repository.set_workspace_state(nq.SECTIONS_STATE_KEY, '[]')
    assert nq.import_tracking_document(core.repository, json.dumps(document)) >= 1
    call = {item['join_id']: item for item in query(client)['calls']}['0xA1']
    assert (call['status'], call['status_mode'], call['root_category']) == ('Under Planning', 'manual', 'E2E')
    assert call['values']['rca_suggested_category'] == 'E2E'
    assert nq.list_sections(core.repository)[0]['label'] == 'Overview'


def test_first_version_workspaces_move_to_the_analysis_center(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'UK_Voice_Q3.xlsx', 'voice', calls_frame())
    login(client)
    repository = core.repository
    nq.ensure_nq_tables(repository)
    # The tables and defaults of the first version: its root cause taxonomy, statuses and analysis fields.
    with repository.connection() as connection:
        connection.execute(f'DROP TABLE IF EXISTS {nq.NQ_ROOT_CATALOG_TABLE}')
        connection.execute("CREATE TABLE nq_root_causes (domain TEXT, cause TEXT, color TEXT, keywords TEXT, position INTEGER)")
        connection.executemany('INSERT INTO nq_root_causes VALUES (?, ?, ?, ?, ?)', [
            ('RF', '', '#123456', '["rf"]', 0), ('RF', 'Coverage', '', '["coverage"]', 0),
            ('Transport', '', '#654321', '["backhaul"]', 1), ('Transport', 'Backhaul', '', '["backhaul"]', 0)])
        for table in (nq.NQ_FIELDS_TABLE, nq.NQ_CALL_OPTIONS_TABLE):
            connection.execute(f'DELETE FROM {table}')
        connection.executemany(f"INSERT INTO {nq.NQ_CALL_OPTIONS_TABLE} (kind, name, color, position, is_closed) VALUES (?, ?, ?, ?, ?)",
                               [('status', name, '#000000', index, int(name == 'Resolved'))
                                for index, name in enumerate(nq._FIRST_STATUSES)] +
                               [('team', name, '#000000', index, 0) for index, name in enumerate(nq._FIRST_TEAMS)])
        connection.execute(f"INSERT INTO {nq.NQ_FIELDS_TABLE} (field_key, label, field_type, options_json, position) "
                           "VALUES ('findings', 'Findings', 'long_text', '[]', 0)")
        connection.execute(f"INSERT INTO {nq.NQ_FIELD_VALUES_TABLE} (call_key, field_key, value) VALUES ('k1', 'findings', 'Old note')")
        connection.execute(f"INSERT INTO {nq.NQ_CALL_TRACKING_TABLE} (call_key, status, team, root_domain, root_cause, version) "
                           "VALUES ('k1', 'Resolved', 'Transport', 'Transport', 'Backhaul', 1)")
    repository.set_workspace_state(nq.ANALYSIS_CENTER_STATE_KEY, '')
    nq.ensure_nq_tables(repository)
    root = nq.list_root_causes(repository)
    assert {'RF', 'Transport'} <= {item['name'] for item in root['domains']}
    assert next(item for item in root['domains'] if item['name'] == 'RF')['color'] == '#123456'
    assert next(item for item in root['causes'] if item['name'] == 'Backhaul')['domain'] == 'Transport'
    with repository.connection() as connection:
        assert not connection.execute("SELECT 1 FROM sqlite_master WHERE name = 'nq_root_causes'").fetchone()
        tracking = connection.execute(f"SELECT status, status_mode, root_cause FROM {nq.NQ_CALL_TRACKING_TABLE} WHERE call_key = 'k1'").fetchone()
    assert tuple(tracking) == ('Closed', 'manual', 'Backhaul')
    assert [item['name'] for item in nq.list_options(repository)['statuses']][0] == 'Not Attended'
    # The workspace's field keeps its key and values and takes its place in the catalog.
    findings = next(field for field in nq.list_fields(repository) if field['label'] == 'Findings')
    assert findings['key'] == 'findings' and findings['section'] == 'analysis' and len(nq.list_fields(repository)) == 39


def test_recommendation_engine_maps_each_source_onto_the_catalog():
    root = {'domains': [{'name': name, 'color': color, 'keywords': list(keywords)} for name, color, keywords in nq_catalog.DEFAULT_DOMAINS],
            'categories': [{'name': name, 'domain': domain, 'keywords': nq_catalog.keywords(list(keywords))}
                           for name, domain, keywords in nq_catalog.DEFAULT_CATEGORIES],
            'causes': [{'name': name, 'domain': domain, 'keywords': nq_catalog.keywords(list(keywords))}
                       for name, domain, keywords in nq_catalog.DEFAULT_CAUSES]}
    rule = nq.DEFAULT_ROOT_CAUSE_RULE
    call = {'failure_classification': 'RF_Problems', 'failure_category': 'Interference problem',
            'failure_subcategory': 'DL interference problems', 'technology': 'NR', 'failure_technology': '5G', 'failure_comment': ''}
    suggestion = nq_rca.suggest_from_text(call, root, rule)
    assert suggestion == {'domain': 'RF', 'category': 'DL Interference NR', 'cause': 'DL interference', 'source': 'netcheck'}
    # "a+b" keywords need every word, as whole words.
    assert nq_catalog.keyword_found(['coverage+lte'], 'dl coverage problems | lte')
    assert not nq_catalog.keyword_found(['coverage+lte'], 'dl coverage problems | nr')
    result = nq_rca.recommend(call, {'rca_category': 'UL Interference NR', 'rca_cause': 'External UL interference'}, root, rule)
    # The script (weight 3) wins over NetCheck (weight 2) when they disagree, with medium confidence.
    assert result['recommendation']['category'] == 'UL Interference NR' and result['recommendation']['confidence'] == 'medium'
    # Two other calls with the same NetCheck values chose a category: it is proposed once per call, never for the call itself.
    choice = ('DL Interference NR', 'DL interference')
    learned = {nq_rca.learned_key('netcheck', 'Interference problem', 'DL interference problems'): {choice: {'k1', 'k2', 'own'}},
               nq_rca.learned_key('script', 'UL Interference NR', 'External UL interference'): {choice: {'k1'}}}
    taught = nq_rca.recommend({**call, 'call_key': 'own'}, {'rca_category': 'UL Interference NR', 'rca_cause': 'External UL interference'},
                              root, rule, learned=learned)
    sources = {source['source']: source for source in taught['sources']}
    assert sources['learned']['count'] == 2 and sources['learned']['category'] == 'DL Interference NR'
    assert nq_catalog.evaluate_status(nq_catalog.default_status_rules(), {'analysis_status': 'Finished'}, 'Not Attended') == 'Closed'
    assert nq_catalog.evaluate_status(nq_catalog.default_status_rules(), {'@attended': 'Yes'}, 'Not Attended') == 'Open'


def test_calls_are_read_from_their_latest_versions_kept_up_to_date(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    repository = core.repository
    # The latest version of every call is kept in a table, rebuilt when the indexed CDRs or the maps change.
    with repository.connection() as connection:
        latest = {row['call_key']: row for row in connection.execute(f'SELECT * FROM {nq.NQ_CALL_LATEST_TABLE}')}
    assert {call['call_key'] for call in calls.values()} == set(latest)
    assert latest[calls['0xA1']['call_key']]['last_cell_id'] == '123456'
    assert calls['0xA1']['values']['last_cell_id'] == '123456' and calls['0xA1']['values']['failure_comment'] == ''
    assert not nq.refresh_latest_calls(repository)
    repository.replace_operator_mapping_group('VF', 'VF Lab', ['Vodafone UK'])
    operators = {item['operator'] for item in query(client)['calls']}
    assert 'VF Lab' in operators and 'Vodafone UK' not in operators
    # The follow-up is read as it is now.
    updated = patch(client, calls['0xC3'], {'root_category': 'E2E'})
    assert {item['join_id']: item for item in query(client)['calls']}['0xC3']['root_category'] == updated['root_category'] == 'E2E'
    # A choice of CDRs reads their own latest versions.
    dataset = str(query(client)['calls'][0]['dataset_id'])
    assert query(client, filters={'datasets': [dataset]})['total'] == 3
    # Without the Summary on screen, its breakdowns are left out.
    assert query(client, breakdowns=False)['breakdowns'] == [] and query(client)['breakdowns']
    # Once in place, the schema is checked without writing.
    assert nq._schema_ready(repository)


def test_an_assignee_brings_the_call_to_their_team(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    core.repository.create_user('analyst', 'analyst123', 'user-editor')
    user_id = next(int(row['id']) for row in core.repository.list_users() if row['username'] == 'analyst')
    core.repository.set_user_workspace_access(user_id, [core.active_workspace.id])
    state = client.get('/api/non-qualified-calls/state').json()
    statuses = [{**status, 'previous': status['name']} for status in state['options']['statuses']]
    members = {'Netcheck': ['analyst'], 'Operations': ['super']}
    teams = [{**team, 'previous': team['name'], 'members': members.get(team['name'], [])} for team in state['options']['teams']]
    assert client.put('/api/non-qualified-calls/options', json={'statuses': statuses, 'teams': teams}).status_code == 200
    assert patch(client, calls['0xA1'], {'assignee': 'analyst'})['team'] == 'Netcheck'
    # Out of a team with other members, the assignee moves the call to their own team.
    other = patch(client, calls['0xB2'], {'team': 'Operations'})
    assert other['team'] == 'Operations' and patch(client, calls['0xB2'], {'assignee': 'analyst'})['team'] == 'Netcheck'


def test_first_layout_workspaces_take_the_second_one(client, tmp_path):
    setup_calls(client, tmp_path)
    repository = core.repository
    # A workspace of the first layout: its colours, names, Analysis order, no CDR fields in Failure Event.
    sections = [{**section, 'color': nq_catalog.FIRST_SECTION_COLORS.get(section['key'], section['color']),
                 'label': nq_catalog.FIRST_SECTION_LABELS.get(section['key'], section['label'])} for section in nq.list_sections(repository)]
    repository.set_workspace_state(nq.SECTIONS_STATE_KEY, json.dumps(sections))
    fields = [field for field in nq.list_fields(repository) if field['key'] not in {'last_cell_id', 'failure_comment'}]
    by_key = {field['key']: field for field in fields}
    keys = [field['key'] for field in fields]
    positions = sorted(keys.index(key) for key in nq_catalog.FIRST_ANALYSIS_ORDER)
    for position, key in zip(positions, nq_catalog.FIRST_ANALYSIS_ORDER):
        fields[position] = by_key[key]
    by_key['serving_cell_when_failure']['suggest_from'] = 'cell_id'
    nq._store_fields(repository, fields)
    repository.set_workspace_state(nq.ANALYSIS_CENTER_V2_STATE_KEY, '')
    nq.ensure_nq_tables(repository)
    upgraded = {section['key']: section for section in nq.list_sections(repository)}
    assert upgraded['failure_event'] == {'key': 'failure_event', 'label': 'Failure Details', 'color': '#FF0000'}
    assert upgraded['implementation']['color'] == '#00B050' and upgraded['planning']['color'] == '#00843D'
    fields = nq.list_fields(repository)
    failure = [field['key'] for field in fields if field['section'] == 'failure_event']
    assert failure.index('last_cell_id') == failure.index('serving_cell_when_failure') - 1 and failure[-1] == 'failure_comment'
    assert [field['key'] for field in fields if field['section'] == 'analysis'][-2:] == ['proposed_measure', 'findings']
    assert next(field for field in fields if field['key'] == 'serving_cell_when_failure')['suggest_from'] == 'last_cell_id'


def test_lifecycle_counts_the_rca_identification(client, tmp_path):
    calls = setup_calls(client, tmp_path)
    patch(client, calls['0xA1'], {'root_category': 'Poor Coverage LTE'})
    life = client.post('/api/non-qualified-calls/lifecycle', json={'filters': {}}).json()
    assert life['rca'] == {'identified': 1, 'not_identified': 2}
    assert [item['join_id'] for item in query(client, filters={'state': ['identified']})['calls']] == ['0xA1']
    assert query(client, filters={'state': ['not_identified']})['total'] == 2
