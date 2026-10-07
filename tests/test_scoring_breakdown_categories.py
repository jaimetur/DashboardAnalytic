import pytest
from pptx import Presentation

from src.modules.scoring_exports import _hierarchy_score_tables, _score_tables, _table_for_mode
from src.modules.scoring_views import build_scoring_views
from tests.test_scoring_table_exports import _job, _result


@pytest.mark.parametrize('hierarchy', [False, True])
def test_breakdown_category_spans_its_kpis_and_subtotal(hierarchy):
    job = _job()
    result = _result()
    if hierarchy:
        job['aggregation_levels'] = ['Operator', 'Region', 'Campaign']
        result.update({'aggregation_contract_version': 2, 'aggregation_levels': job['aggregation_levels']})
    views = build_scoring_views(job, result)
    matrices = views['hierarchy_score_tables' if hierarchy else 'score_tables']
    matrix = _table_for_mode(next(item for item in matrices
                                  if item['context']['environment'] == 'DriveCity'), 'expanded')
    presentation = Presentation()
    renderer = _hierarchy_score_tables if hierarchy else _score_tables
    renderer(presentation, [matrix], views.get('threshold_legend', []), show_gap_values=False)
    table = next(shape.table for shape in presentation.slides[0].shapes if shape.has_table)
    assert table.cell(0, 0).text == 'CATEGORY'
    subtotal_rows = [index for index in range(1, len(table.rows) - 1)
                     if table.cell(index, 1).text.endswith(' total')]
    assert subtotal_rows
    for end in subtotal_rows:
        category_cell = table.cell(end, 0)
        assert category_cell.is_spanned
        assert category_cell.text == ''
        origin = end - 1
        while not table.cell(origin, 0).is_merge_origin:
            origin -= 1
        cell = table.cell(origin, 0)
        assert origin + cell.span_height - 1 == end
        assert cell.text and not cell.text.endswith(' total')
        assert str(cell.fill.fore_color.rgb) == str(table.cell(end, 1).fill.fore_color.rgb)


def test_breakdown_slides_keep_up_to_fifty_bars_together():
    from src.modules.scoring_exports import BREAKDOWN_BARS_PER_SLIDE, _breakdown_pages

    categories = [f'C{index}' for index in range(1, 11)]
    assert BREAKDOWN_BARS_PER_SLIDE == 50
    # 10 categories × 5 operators = 50 bars: one slide.
    assert _breakdown_pages(categories, 5) == [categories]
    # 12 × 5 = 60 bars: two balanced slides of 30 bars.
    twelve = [f'C{index}' for index in range(1, 13)]
    assert _breakdown_pages(twelve, 5) == [twelve[:6], twelve[6:]]
    # Not splitting keeps every category together; a category wider than the limit stays alone.
    assert _breakdown_pages(twelve, 5, split=False) == [twelve]
    assert _breakdown_pages(['A', 'B'], 60) == [['A'], ['B']]
