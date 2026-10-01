"""Whole-value intro truncation and coherent download filenames."""
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from src.modules.cdr_report_filenames import build_scoring_report_filename, scoring_display_selections
from src.modules.scoring_exports import _scoring_filter_subtitle, _add_scoring_intro_slides, prepare_scoring_display_selections


def test_complete_unicode_values_and_filename_budget():
    values = [f'東京都 Ciudad {index:02d} with a complete name' for index in range(30)]
    filters = {field: values for field in ('operator', 'vendor', 'region', 'city')}
    plan = scoring_display_selections(filters)
    slide_values = {field: list(selected) for field, selected in plan['values'].items()}
    filename = build_scoring_report_filename(datetime(2026, 10, 2), 'NSA', filters, display_selections=plan)
    assert plan['values'] == slide_values
    assert len(filename.encode('utf-8')) <= 245
    assert '...' not in filename and '…' not in filename
    assert all(selected == values[:len(selected)] for field, selected in plan['values'].items()
               if field in filters)
    assert all(plan['omitted'][field] for field in filters)
    parts = filename.removesuffix('.pptx').split(' - ')[3:]
    for field, part in zip(('region', 'city', 'operator', 'vendor'), parts):
        retained = part.split(' + ') if part else []
        assert retained == plan['values'][field][:len(retained)]
        assert len(', '.join(plan['values'][field])) <= 110
    assert any(len(part.split(' + ')) < len(plan['values'][field])
               for field, part in zip(('region', 'city', 'operator', 'vendor'), parts))
    assert filters['city'] == values


def test_template_lines_never_wrap_and_share_retained_values():
    template = Path('assets/ppt-templates/Template_CDR_analysis.pptx')
    job = {'context_filters': {'city': [f'City {index} long complete value' for index in range(20)]},
           'aggregation_levels': ['Operator', 'Vendor', 'Region', 'City']}
    result = {}
    job['_scoring_display_selections'] = prepare_scoring_display_selections(job, result, template)
    subtitle = _scoring_filter_subtitle(job)
    selected = job['_scoring_display_selections']['values']['city']
    assert f'City: {", ".join(selected)}, ...' in subtitle
    presentation = Presentation(template)
    _add_scoring_intro_slides(presentation, job, result)
    shape = next(shape for shape in presentation.slides[-1].shapes
                 if shape.name == 'Scoring Aggregations and Filters')
    assert shape.text_frame.word_wrap is False
    assert len(shape.text_frame.paragraphs) == 6
    assert shape.text_frame.paragraphs[0].font.size.pt == 16
    assert all(paragraph.font.size.pt <= 16 for paragraph in shape.text_frame.paragraphs)
    assert 'Operator: All Operators' in shape.text


def test_character_limit_counts_values_and_separators_only():
    exact = scoring_display_selections({'city': ['i' * 110]})
    assert exact['values']['city'] == ['i' * 110]
    assert not exact['omitted']['city']
    overflow = scoring_display_selections({'city': ['i' * 111]})
    assert overflow['values']['city'] == []
    partial = scoring_display_selections({'city': ['x' * 104, 'four', 'extra value']})
    assert partial['values']['city'] == ['x' * 104, 'four']
    assert partial['omitted']['city']
    assert len(', '.join(partial['values']['city'])) == 110
    vendor = scoring_display_selections({'vendor': ['x' * 110]})
    assert vendor['values']['vendor'] == ['x' * 110]
