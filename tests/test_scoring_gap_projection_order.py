import pytest
from pptx import Presentation
from pptx.util import Inches
from src.modules.scoring_exports import _merge_category_cells

from src.modules.scoring_exports import _hierarchy_gap_projection


def _column(column_id, operator, *, is_reference=False):
    return {
        'id': column_id,
        'operator': operator,
        'is_reference': is_reference,
        'context': {},
    }


def _kpi(code, category, peer_gap):
    return {
        'row_type': 'kpi',
        'kpi_code': code,
        'category': category,
        'gaps': {'EE': None, 'O2 UK': peer_gap, 'Three UK': peer_gap},
        'gap_colors': {},
    }


def _subtotal(category, peer_gap):
    return {
        'row_type': 'category',
        'kpi_code': '',
        'category': category,
        'gaps': {'EE': None, 'O2 UK': peer_gap, 'Three UK': peer_gap},
        'gap_colors': {},
    }


def _matrix(rows, *, table_mode='expanded'):
    return {
        'operators': ['EE', 'O2 UK', 'Three UK'],
        'operator_styles': {'EE': {}, 'O2 UK': {}, 'Three UK': {}},
        'rows': rows,
        'total': {'gaps': {'EE': None, 'O2 UK': 2.0}},
        'table_mode': table_mode,
    }


def test_combined_hierarchy_projection_preserves_definition_order():
    source_rows = [
        _kpi('KPI-A', 'Category A', 1.0),
        _kpi('KPI-B', 'Category B', 5.0),
        _kpi('KPI-C', 'Category A', 4.0),
    ]
    projected = _hierarchy_gap_projection(
        _matrix(source_rows),
        [
            _column('EE', 'EE', is_reference=True),
            _column('O2 UK', 'O2 UK'),
            _column('Three UK', 'Three UK'),
        ],
    )

    assert [row['kpi_code'] for row in projected['rows']] == ['KPI-A', 'KPI-B', 'KPI-C']


def test_individual_expanded_projection_reranks_kpis_and_omits_subtotals():
    source_rows = [
        _kpi('KPI-A1', 'Category A', 1.0),
        _kpi('KPI-B1', 'Category B', 5.0),
        _kpi('KPI-A2', 'Category A', 4.0),
        _subtotal('Category A', 2.5),
        _subtotal('Category B', 5.0),
    ]
    projected = _hierarchy_gap_projection(
        _matrix(source_rows),
        [_column('EE', 'EE', is_reference=True), _column('O2 UK', 'O2 UK')],
    )

    assert [(row['row_type'], row['kpi_code'] or row['category']) for row in projected['rows']] == [
        ('kpi', 'KPI-A1'),
        ('kpi', 'KPI-A2'),
        ('kpi', 'KPI-B1'),
    ]


def test_individual_summary_projection_sorts_by_numeric_signed_gap_ascending():
    source_rows = [
        _kpi('KPI-NEGATIVE', 'Category A', -2.0),
        _kpi('KPI-TEN', 'Category B', 10.0),
        _kpi('KPI-ONE', 'Category C', 1.0),
    ]
    projected = _hierarchy_gap_projection(
        _matrix(source_rows, table_mode='summary'),
        [_column('EE', 'EE', is_reference=True), _column('O2 UK', 'O2 UK')],
    )

    assert [row['kpi_code'] for row in projected['rows']] == [
        'KPI-NEGATIVE', 'KPI-ONE', 'KPI-TEN',
    ]




@pytest.mark.parametrize('include_subtotals', [False, True])
def test_category_merge_includes_only_adjacent_subtotal_when_enabled(include_subtotals):
    rows = [
        _kpi('K1', 'Category A', -3),
        _kpi('K2', 'Category A', -2),
        _subtotal('Category A', -2.5),
        _kpi('K3', 'Category B', -1),
        _kpi('K4', 'Category A', 0),
    ]
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(6, 2, Inches(1), Inches(1), Inches(5), Inches(3)).table
    _merge_category_cells(table, rows, 1, include_subtotals=include_subtotals)
    assert table.cell(1, 0).is_merge_origin
    assert table.cell(2, 0).is_spanned
    assert table.cell(3, 0).is_spanned is include_subtotals
    assert not table.cell(4, 0).is_spanned
    assert not table.cell(5, 0).is_spanned
    assert table.cell(1, 0).text == 'Category A'
