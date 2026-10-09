import io
import json
import time
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import src.DriveTestAnalyzer as core
from src.modules import non_qualified_calls as nq
from src.modules.repository import local_now_iso


def login(client, username='super', password='super123'):
    client.cookies.clear()
    response = client.post('/login', data={'username': username, 'password': password}, follow_redirects=False)
    assert response.status_code == 303


def enable_module():
    settings = core.feature_activation_settings()
    settings['non-qualified-calls'] = {'default': 'all'}
    core.save_feature_activation_settings(settings)


def workspace_user(username, role):
    core.repository.create_user(username, f'{username}123', role)
    user_id = next(int(row['id']) for row in core.repository.list_users() if row['username'] == username)
    core.repository.set_user_workspace_access(user_id, [core.active_workspace.id])


def add_cdr(tmp_path: Path, name: str, kind: str, rows: pd.DataFrame) -> int:
    repository = core.repository
    source = tmp_path / name
    source.write_text('test source', encoding='utf-8')
    dataset_id, _created = repository.add_dataset(name, str(source), 'super')
    repository.replace_dataset_rows(dataset_id, rows)
    repository.update_dataset_profile(
        dataset_id, status='ready', progress=100, dataset_kind=kind, nr_mode='NSA',
        row_count=len(rows), column_count=len(rows.columns), processed_at=local_now_iso(),
    )
    return dataset_id


def voice_rows() -> pd.DataFrame:
    return pd.DataFrame({
        'Operator': ['EE', 'EE', 'Vodafone UK'],
        'Campaign': ['UK_Q1_2026'] * 3,
        'City': ['Leeds', 'York', 'Leeds'],
        'Session_ID_A': [1001, 1002, 1003],
        'Call_Status': ['Completed', 'Failed', 'Dropped'],
        'status': ['Completed', 'Failed', 'Dropped'],
        'event_start_time': ['2026-02-24 10:00:00', '2026-02-24 11:00:00', None],
        'Call_Start_Time': ['2026-02-24 10:00:00', '2026-02-24 11:00:00', '2026-02-24 12:00:00'],
        'Failure_Classification': ['', 'RF Problems', 'VoLTE Problems'],
        'Failure_Category': ['', 'Coverage problem', 'IMS'],
        'Call_Start_Latitude_A': [53.8, 53.9, 53.7],
        'Call_Start_Longitude_A': [-1.5, -1.0, -1.6],
    })


def data_rows() -> pd.DataFrame:
    return pd.DataFrame({
        'Operator': ['EE', 'Vodafone UK', 'Vodafone UK'],
        'Campaign': ['UK_Q1_2026'] * 3,
        'City': ['Leeds', 'Leeds', 'York'],
        'Test_ID': [5001, 5002, 5003],
        'Test_Name': ['HTTP DL', 'HTTP UL', 'Ping'],
        'Test_Result': ['Completed', 'Cutoff', 'Failed'],
        'status': ['Completed', 'Cutoff', 'Failed'],
        'event_start_time': ['2026-02-25 09:00:00', '2026-02-25 09:05:00', '2026-02-25 09:10:00'],
        'Failure_Classification': ['', 'RF Problems', ''],
        'Failure_Comment': ['', 'Interference on PCI 80', ''],
    })


def query(client, **payload):
    response = client.post('/api/non-qualified-calls/calls', json=payload)
    assert response.status_code == 200, response.text
    return response.json()


def test_module_is_active_by_default_and_marks_tabs_in_development(client):
    login(client)
    # Every module but PPT Reporting (old) is active for every user of a new deployment.
    assert client.get('/api/non-qualified-calls/state').status_code == 200
    assert 'href="/non-qualified-calls"' in client.get('/workspace').text
    settings = core.feature_activation_settings()
    settings['non-qualified-calls'] = {'default': 'none'}
    core.save_feature_activation_settings(settings)
    assert client.get('/api/non-qualified-calls/state').status_code == 403
    assert 'href="/non-qualified-calls"' not in client.get('/workspace').text
    # Once turned on, the module tabs show their stage: ALPHA, BETA or NEW; stable modules have no label.
    core.repository.set_application_state(core.MODULE_STAGE_LABELS_STATE_KEY, '1')
    page = client.get('/workspace').text
    assert '<span class="module-tab-label-mobile">Network</span><svg class="module-tab-new module-tab-new-network-insights"' in page
    assert '<span class="module-tab-label-mobile">Scoring</span><svg class="module-tab-new module-tab-new-scoring"' in page
    assert 'module-tab-new-workspace' not in page and 'module-tab-new-datasets-analysis' not in page
    assert '>STABLE</text>' not in page and '>ALPHA</text>' in page and '>BETA</text>' in page and '>NEW</text>' in page
    enable_module()
    page = client.get('/non-qualified-calls')
    assert page.status_code == 200
    assert 'module-tab-non-qualified-calls active' in page.text
    assert '<span class="module-tab-label-mobile">NQ Calls</span><svg class="module-tab-new module-tab-new-non-qualified-calls"' in page.text
    assert 'id="nq-table"' in page.text and 'non_qualified_calls.js' in page.text


def test_nq_calls_are_indexed_tracked_and_commented(client, tmp_path):
    enable_module()
    workspace_user('editor', 'user-editor')
    workspace_user('viewer', 'user-viewer')
    voice_id = add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)

    state = client.get('/api/non-qualified-calls/state').json()
    assert state['sync']['calls'] == 4 and state['sync']['datasets'] == 2
    assert [item['name'] for item in state['options']['statuses']][:2] == ['Open', 'Under Investigation']
    assert state['filter_options']['result'] == ['Cutoff', 'Dropped', 'Failed']
    assert {'super', 'editor', 'viewer'} <= set(state['users'])

    result = query(client)
    assert result['total'] == 4 and result['summary'] == {
        'total': 4, 'closed': 0, 'attended': 0, 'commented': 0, 'with_team': 0, 'assigned': 0, 'open': 4,
    }
    by_service = next(item for item in result['breakdowns'] if item['field'] == 'service')['items']
    assert {item['value']: item['count'] for item in by_service} == {'voice': 2, 'data': 2}
    dropped = next(call for call in result['calls'] if call['result'] == 'Dropped')
    # Empty values fall back to the next CDR column with a value.
    assert dropped['start_time'] == '2026-02-24 12:00:00' and dropped['status'] == 'Open'
    assert dropped['latitude'] == 53.7 and dropped['failure_classification'] == 'VoLTE Problems'
    assert query(client, filters={'service': ['voice']})['total'] == 2
    assert query(client, filters={'search': 'pci 80'})['total'] == 1

    key = dropped['call_key']
    url = f'/api/non-qualified-calls/calls/{key}'
    changed = client.patch(url, json={'changes': {'status': 'Under Investigation', 'team': 'Core Network'}, 'version': 0})
    assert changed.status_code == 200, changed.text
    detail = changed.json()
    assert detail['call']['status'] == 'Under Investigation' and detail['call']['version'] == 1
    assert {(entry['field'], entry['new_value']) for entry in detail['history']} == {
        ('status', 'Under Investigation'), ('team', 'Core Network')}
    assert ['Call_Status', 'Dropped'] in detail['fields'] and detail['call']['dataset_name'].endswith('Voice_2026_Q1.xlsx')
    # A stale version means another user changed the call first.
    assert client.patch(url, json={'changes': {'status': 'Resolved'}, 'version': 0}).status_code == 409
    assert client.patch(url, json={'changes': {'status': 'Unknown'}}).status_code == 400
    assert client.patch(url, json={'changes': {'assignee': 'nobody'}}).status_code == 400

    comment = client.post(f'{url}/comments', json={'body': 'Waiting for E2E traces from the vendor.'}).json()['comment']
    assert client.post(f'{url}/comments', json={'body': '   '}).status_code == 400
    listed = query(client, filters={'search': 'e2e traces'})
    assert listed['total'] == 1 and listed['calls'][0]['comment_count'] == 1
    assert listed['calls'][0]['last_comment']['created_by'] == 'super'
    assert query(client, filters={'without_comments': True})['total'] == 3
    assert query(client, filters={'team': [nq.UNASSIGNED]})['total'] == 3

    login(client, 'editor', 'editor123')
    assert client.patch(f"/api/non-qualified-calls/comments/{comment['id']}", json={'body': 'Changed'}).status_code == 403
    own = client.post(f'{url}/comments', json={'body': 'Traces requested.'}).json()['comment']
    edited = client.patch(f"/api/non-qualified-calls/comments/{own['id']}", json={'body': 'Traces requested twice.'})
    assert edited.status_code == 200 and edited.json()['comment']['edited_by'] == 'editor'
    bulk = client.post('/api/non-qualified-calls/calls/bulk', json={
        'call_keys': [call['call_key'] for call in result['calls']], 'changes': {'assignee': 'editor'}})
    assert bulk.json() == {'changed': 4}
    assert query(client, filters={'mine': True})['total'] == 4

    login(client, 'viewer', 'viewer123')
    assert query(client)['total'] == 4
    assert client.patch(url, json={'changes': {'status': 'Resolved'}}).status_code == 403
    assert client.post(f'{url}/comments', json={'body': 'Hi'}).status_code == 403

    login(client)  # A super-admin moderates any comment; deleted comments keep their history.
    deleted = client.delete(f"/api/non-qualified-calls/comments/{own['id']}").json()['comment']
    assert deleted['body'] == '' and deleted['deleted_by'] == 'super'
    detail = client.get(url).json()
    assert [entry['field'] for entry in detail['history']][-3:] == ['comment_edit', 'assignee', 'comment_delete']
    assert detail['history'][-1]['old_value'] == 'Traces requested twice.'

    # Processing the same CDR again keeps the follow-up of its calls.
    core.repository.replace_dataset_rows(voice_id, voice_rows())
    core.repository.update_dataset_profile(voice_id, processed_at=local_now_iso(), updated_at=local_now_iso())
    again = query(client, filters={'service': ['voice'], 'status': ['Under Investigation']})
    assert again['sync']['reindexed'] == 1 and again['total'] == 1 and again['calls'][0]['call_key'] == key
    assert query(client, filters={'open_only': True})['total'] == 4
    client.patch(url, json={'changes': {'status': 'Resolved'}})
    assert query(client, filters={'open_only': True})['total'] == 3
    assert query(client)['summary']['closed'] == 1


def test_statuses_and_teams_can_be_renamed_but_not_removed_while_in_use(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    login(client)
    key = query(client)['calls'][0]['call_key']
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'team': 'Transport'}})
    options = client.get('/api/non-qualified-calls/state').json()['options']
    teams = [{'name': 'Backhaul' if team['name'] == 'Transport' else team['name'], 'color': team['color'],
              'previous': team['name']} for team in options['teams']]
    saved = client.put('/api/non-qualified-calls/options', json={'statuses': options['statuses'], 'teams': teams})
    assert saved.status_code == 200, saved.text
    assert client.get(f'/api/non-qualified-calls/calls/{key}').json()['call']['team'] == 'Backhaul'
    removed = client.put('/api/non-qualified-calls/options', json={
        'statuses': options['statuses'], 'teams': [team for team in teams if team['name'] != 'Backhaul']})
    assert removed.status_code == 400 and 'in use' in removed.json()['detail']
    assert client.put('/api/non-qualified-calls/options', json={'statuses': [], 'teams': teams}).status_code == 400
    repeated = [{'name': 'Open'}, {'name': 'open'}]
    assert client.put('/api/non-qualified-calls/options', json={'statuses': repeated, 'teams': teams}).status_code == 400


def test_excel_export_lists_calls_comments_and_history(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    key = query(client, filters={'result': ['Cutoff']})['calls'][0]['call_key']
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'status': 'Pending Information'}})
    client.post(f'/api/non-qualified-calls/calls/{key}/comments', json={'body': 'Check the interference.'})
    response = client.post('/api/non-qualified-calls/export', json={'filters': {'result': ['Cutoff']}})
    assert response.status_code == 200
    assert response.headers['content-disposition'].endswith('.xlsx"')
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ['NQ Calls', 'Comments', 'History']
    calls = list(workbook['NQ Calls'].values)
    header = calls[0]
    assert len(calls) == 2
    row = dict(zip(header, calls[1]))
    assert (row['Result'], row['Status'], row['Comments'], row['Last Comment']) == (
        'Cutoff', 'Pending Information', 1, 'Check the interference.')
    assert list(workbook['Comments'].values)[1][4] == 'Check the interference.'
    assert list(workbook['History'].values)[1][4:7] == ('Status', 'Open', 'Pending Information')


def test_tracking_travels_with_workspace_packages_and_imports_once(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    login(client)
    key = query(client)['calls'][0]['call_key']
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'status': 'Resolved'}})
    client.post(f'/api/non-qualified-calls/calls/{key}/comments', json={'body': 'Fixed by the RAN team.'})
    payload = core._nq_call_tracking_payload(core.active_workspace)
    document = json.loads(payload)
    assert document['format'] == 'nq-call-tracking' and len(document['comments']) == 1
    assert 'nq_call_tracking' in core.full_workspace_archive_components()
    assert core.archive_workspace_components_for_target('nq-call-tracking') == ['nq_call_tracking']

    with core.repository.connection() as connection:
        for table in (nq.NQ_CALL_TRACKING_TABLE, nq.NQ_CALL_COMMENTS_TABLE, nq.NQ_CALL_HISTORY_TABLE):
            connection.execute(f'DELETE FROM {table}')
    assert core._restore_workspace_nq_call_tracking(core.active_workspace, payload) == 1
    detail = client.get(f'/api/non-qualified-calls/calls/{key}').json()
    assert detail['call']['status'] == 'Resolved' and detail['comments'][0]['body'] == 'Fixed by the RAN team.'
    assert len(detail['history']) == 1
    # Importing the same follow-up again changes nothing.
    assert core._restore_workspace_nq_call_tracking(core.active_workspace, payload) == 0
    assert len(client.get(f'/api/non-qualified-calls/calls/{key}').json()['comments']) == 1

    # Admin → Import/Export: the NQ Call Tracking package imports into a workspace.
    exported = client.get('/admin/import-export/export?export_target=nq-call-tracking')
    assert exported.status_code == 200, exported.text
    with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        assert json.loads(archive.read(manifest['archive_path']))['format'] == 'nq-call-tracking'
    assert manifest['kind'] == 'nq-call-tracking' and core.archive_workspace_components(manifest) == ['nq_call_tracking']
    with core.repository.connection() as connection:
        connection.execute(f'DELETE FROM {nq.NQ_CALL_COMMENTS_TABLE}')
    inspection = client.post('/admin/import-export/inspect',
                             files={'package': ('nq.zip', io.BytesIO(exported.content), 'application/zip')})
    assert inspection.status_code == 200, inspection.text
    job = client.post('/admin/import-export/import/jobs', data={
        'upload_id': inspection.headers['X-Import-Upload-Id'], 'confirmed_import': 'true', 'workspace_ids': 'default'})
    assert job.status_code == 200, job.text
    for _attempt in range(200):
        status = client.get(job.json()['status_url']).json()
        if status['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.01)
    assert status['status'] == 'ready', status
    assert len(client.get(f'/api/non-qualified-calls/calls/{key}').json()['comments']) == 1


def test_nq_tables_appear_in_database_management(client, tmp_path):
    enable_module()
    login(client)
    client.get('/api/non-qualified-calls/state')
    page = client.get('/admin').text
    for title in ('NQ Calls', 'NQ Call Tracking', 'NQ Call Comments', 'NQ Call History', 'NQ Call Options', 'NQ Call Sources'):
        assert f'>{title}<' in page or f'"{title}"' in page, title


def test_teams_members_progress_shared_filters_and_reporting_artifact(client, tmp_path):
    enable_module()
    workspace_user('editor', 'user-editor')
    workspace_user('viewer', 'user-viewer')
    voice_id = add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)

    # Teams are edited in Statuses & Teams with their members; a user can be in several teams.
    page_state = client.get('/api/non-qualified-calls/state').json()
    assert {'editor', 'viewer'} <= set(page_state['users'])
    statuses = [{**status, 'previous': status['name']} for status in page_state['options']['statuses']]
    payload = [{**team, 'previous': team['name'], 'members': ['editor'] if team['name'] == 'Core Network' else []}
               for team in page_state['options']['teams']]
    payload.append({'name': 'Field Ops', 'color': '#336699', 'previous': '', 'members': ['editor', 'viewer']})
    saved = client.put('/api/non-qualified-calls/options', json={'statuses': statuses, 'teams': payload})
    assert saved.status_code == 200, saved.text
    members = {team['name']: team['members'] for team in saved.json()['options']['teams']}
    assert members['Core Network'] == ['editor'] and members['Field Ops'] == ['editor', 'viewer']
    ghosts = client.put('/api/non-qualified-calls/options', json={'statuses': statuses, 'teams': [{'name': 'Ghosts', 'members': ['nobody']}]})
    assert ghosts.status_code == 400
    assert client.get('/api/non-qualified-calls/teams').status_code in {404, 405}
    page = client.get('/non-qualified-calls').text
    assert 'id="nq-members-dialog"' in page and 'Teams Users Assignments' not in client.get('/workspace-config').text

    # A call of a team with members only accepts its members as assignee.
    call = next(item for item in query(client)['calls'] if item['result'] == 'Dropped')
    url = f"/api/non-qualified-calls/calls/{call['call_key']}"
    assert client.patch(url, json={'changes': {'team': 'Core Network', 'assignee': 'viewer'}}).status_code == 400
    assigned = client.patch(url, json={'changes': {'team': 'Field Ops', 'assignee': 'viewer'}}).json()['call']
    assert assigned['assignee'] == 'viewer'
    # Moving it to a team the assignee does not belong to clears the assignee.
    moved = client.patch(url, json={'changes': {'team': 'Core Network'}}).json()['call']
    assert moved['team'] == 'Core Network' and moved['assignee'] == ''
    assert client.post(f'{url}/comments', json={'body': 'Checked the drive test.'}).status_code == 200

    # Summary breakdowns, in the order of the page and of the one-slide Executive Summary.
    result = query(client)
    assert [item['field'] for item in result['breakdowns']] == [
        'service', 'result', 'status', 'team', 'root_domain', 'failure_classification', 'campaign', 'operator', 'vendor',
        'region', 'cluster', 'city']
    by_campaign = next(item for item in result['breakdowns'] if item['field'] == 'campaign')
    assert by_campaign['label'] == 'By Campaign' and by_campaign['items'] == [{'value': 'UK_Q1_2026', 'count': 4}]

    # Progress View: distributions, attended calls and periods.
    progress = client.post('/api/non-qualified-calls/progress', json={'filters': {}, 'granularity': 'week'})
    assert progress.status_code == 200, progress.text
    progress = progress.json()
    assert progress['summary']['total'] == 4 and progress['summary']['attended'] == 1
    assert {item['field'] for item in progress['distributions']} >= {'status', 'team', 'assignee'}
    assert progress['periods'] and sum(period['detected'] for period in progress['periods']) == 4

    # The Filters panel is shared by every user and session of the workspace.
    filters = {'service': ['voice'], 'city': ['Leeds'], 'open_only': True, 'search': 'drop', 'unknown': ['x']}
    assert client.put('/api/non-qualified-calls/filters', json={'filters': filters}).status_code == 200
    login(client, 'viewer', 'viewer123')
    assert client.get('/api/non-qualified-calls/state').json()['saved_filters'] == {
        'city': ['Leeds'], 'open_only': True, 'search': 'drop', 'service': ['voice']}

    # Reporting Jobs: the Executive Summary and Progress Status artifact with its filters.
    from src.modules.report_tasks import ARTIFACT_PROVIDERS

    login(client)
    provider = ARTIFACT_PROVIDERS['non_qualified_calls']
    user = core.SessionUser(username='super', role='super-admin')
    artifacts = provider['generate'](
        {'formats': ['powerpoint', 'word', 'excel'], 'options': {'filters': {'datasets': [str(voice_id)]}, 'granularity': 'month'}},
        tmp_path, '20261006_120000', user)
    assert [item['file_name'] for item in artifacts] == [
        '20261006_120000 - Non-Qualified Calls - Executive Summary and Progress Status.pptx',
        '20261006_120000 - Non-Qualified Calls - Executive Summary and Progress Status.docx',
        '20261006_120000 - Non-Qualified Calls - Executive Summary and Progress Status.xlsx',
    ]
    assert all((tmp_path / item['file_name']).stat().st_size > 0 for item in artifacts)
    assert '2 Non-Qualified Calls' in artifacts[0]['details'][-1]


def test_page_exports_the_executive_summary_to_powerpoint_and_word(client, tmp_path):
    from docx import Document
    from pptx import Presentation

    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    page = client.get('/non-qualified-calls').text
    assert 'data-nq-document-export="powerpoint"' in page and 'data-nq-document-export="word"' in page
    body = {'filters': {'result': ['Cutoff']}, 'granularity': 'month'}
    deck = client.post('/api/non-qualified-calls/export/powerpoint', json=body)
    assert deck.status_code == 200 and deck.headers['content-disposition'].endswith('.pptx"')
    assert len(Presentation(io.BytesIO(deck.content)).slides) > 1
    titles = [slide.shapes.title.text_frame.text for slide in Presentation(io.BytesIO(deck.content)).slides
              if slide.shapes.title is not None]
    # The indicators and every breakdown of the Executive Summary share one slide.
    assert sum(title.startswith('Executive Summary') for title in titles) == 1
    document = client.post('/api/non-qualified-calls/export/word', json=body)
    assert document.status_code == 200 and document.headers['content-disposition'].endswith('.docx"')
    assert Document(io.BytesIO(document.content)).paragraphs
    assert client.post('/api/non-qualified-calls/export/pdf', json=body).status_code == 404


def test_root_causes_are_suggested_labelled_required_and_analysed(client, tmp_path):
    enable_module()
    voice = voice_rows()
    voice['Test_Name'] = ['CALL', 'WhatsApp CALL', 'CALL']
    voice['Cell_ID_A'] = ['[134445192]', '[2769016]->[134445192]', '[24219]']
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice)
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)

    state = client.get('/api/non-qualified-calls/state').json()
    domains = {item['name']: item for item in state['root_causes']['domains']}
    assert list(domains)[:2] == ['RF', 'RAN'] and state['root_causes']['require_to_close'] is False
    calls = {(call['service'], call['result']): call for call in query(client)['calls']}
    # The CDR failure classification suggests the domain, and its category or comment the cause.
    assert calls['voice', 'Failed']['suggested_root_cause'] == {'domain': 'RF', 'cause': 'Coverage', 'source': 'cdr'}
    # "Interference on PCI 80" does not say whether it is DL or UL: only the RF domain is suggested.
    assert calls['data', 'Cutoff']['suggested_root_cause'] == {'domain': 'RF', 'cause': '', 'source': 'cdr'}
    assert calls['voice', 'Dropped']['suggested_root_cause'] == {'domain': 'IMS/E2E', 'cause': '', 'source': 'cdr'}
    assert calls['data', 'Failed']['suggested_root_cause'] is None

    key = calls['voice', 'Dropped']['call_key']
    response = client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'root_domain': 'RF', 'root_cause': 'TAU reject'}})
    assert response.status_code == 400 and 'causes of the selected domain' in response.text
    detail = client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'root_domain': 'IMS/E2E', 'root_cause': 'No QCI1 established'}}).json()
    assert (detail['call']['root_domain'], detail['call']['root_cause']) == ('IMS/E2E', 'No QCI1 established')
    assert {entry['field'] for entry in detail['history']} == {'root_domain', 'root_cause'}
    assert query(client, filters={'root_domain': ['IMS/E2E']})['total'] == 1
    assert query(client, filters={'root_domain': [nq.UNASSIGNED]})['total'] == 3

    # Bulk: the suggestions label the calls without a root cause and keep the labelled ones.
    keys = [call['call_key'] for call in calls.values()]
    changed = client.post('/api/non-qualified-calls/calls/bulk', json={'call_keys': keys, 'changes': {'apply_suggestion': True}})
    assert changed.json()['changed'] == 2
    labelled = labelled_calls(client)
    assert labelled['voice', 'Dropped'] == ('IMS/E2E', 'No QCI1 established')
    assert labelled['data', 'Cutoff'] == ('RF', '') and labelled['voice', 'Failed'] == ('RF', 'Coverage')
    assert labelled['data', 'Failed'] == ('', '')

    # Renaming follows on every call; a cause in use cannot be removed.
    taxonomy = client.get('/api/non-qualified-calls/root-causes').json()
    rf = next(item for item in taxonomy['domains'] if item['name'] == 'RF')
    rf['previous'], rf['name'] = 'RF', 'Radio'
    for cause in rf['causes']:
        cause['previous'] = cause['name']
    saved = client.put('/api/non-qualified-calls/root-causes', json={'domains': taxonomy['domains'], 'require_to_close': True})
    assert saved.status_code == 200, saved.text
    assert labelled_calls(client)['voice', 'Failed'] == ('Radio', 'Coverage')
    current = saved.json()['domains']
    radio = next(item for item in current if item['name'] == 'Radio')
    radio['causes'] = [cause for cause in radio['causes'] if cause['name'] != 'Coverage']
    refused = client.put('/api/non-qualified-calls/root-causes', json={'domains': current, 'require_to_close': True})
    assert refused.status_code == 400 and 'in use' in refused.text

    # With the root cause required, a call cannot be closed without one.
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'root_domain': '', 'root_cause': ''}})
    blocked = client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'status': 'Resolved'}})
    assert blocked.status_code == 400 and 'root cause before closing' in blocked.text
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'root_domain': 'AAA'}})
    assert client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'status': 'Resolved'}}).status_code == 200

    stats = client.post('/api/non-qualified-calls/root-causes/stats', json={'filters': {}, 'include_suggestions': False}).json()
    assert stats['summary'] == {'total': 4, 'labelled': 3, 'suggested': 0, 'unclassified': 1}
    assert {item['value']: item['count'] for item in stats['domains']} == {'Radio': 2, 'AAA': 1, 'Not classified': 1}
    assert {row['name']: row['total'] for row in stats['call_types']} == {'Classic': 1, 'WhatsApp': 1, 'Data': 2}
    assert {row['name'] for row in stats['nr_modes']} == {'NSA'}
    # The last cell of the chain gives the eNB (ECI / 256); small identities are 2G/3G cells.
    nodes = {row['node']: row['calls'] for row in stats['nodes']}
    assert nodes == {'eNB 525176': 1, 'Cell 24219': 1}

    tracking = json.loads(nq.export_tracking_document(core.repository))
    assert tracking['root_causes']['require_to_close'] is True
    assert {(item['root_domain'], item['root_cause']) for item in tracking['tracking']} >= {('Radio', 'Coverage'), ('AAA', '')}
    page = client.get('/non-qualified-calls').text
    assert 'id="nq-root-dialog"' in page and 'data-nq-filter="root_domain"' in page and 'id="nq-root-nodes"' in page


def labelled_calls(client):
    return {(call['service'], call['result']): (call['root_domain'], call['root_cause']) for call in query(client)['calls']}


def test_root_cause_taxonomy_imports_with_the_tracking_and_shows_in_exports(client, tmp_path):
    from pptx import Presentation

    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    key = query(client)['calls'][0]['call_key']
    taxonomy = client.get('/api/non-qualified-calls/root-causes').json()
    taxonomy['domains'].append({'name': 'Transport', 'color': '#245a96', 'keywords': 'backhaul',
                                'causes': [{'name': 'Microwave fading', 'keywords': 'fading'}]})
    assert client.put('/api/non-qualified-calls/root-causes', json={'domains': taxonomy['domains']}).status_code == 200
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'root_domain': 'Transport', 'root_cause': 'Microwave fading'}})
    document = nq.export_tracking_document(core.repository)
    with core.repository.connection() as connection:
        connection.execute(f'DELETE FROM {nq.NQ_CALL_TRACKING_TABLE}')
        connection.execute(f"DELETE FROM {nq.NQ_ROOT_CAUSES_TABLE} WHERE domain = 'Transport'")
    nq.import_tracking_document(core.repository, document)
    restored = client.get('/api/non-qualified-calls/root-causes').json()
    transport = next(item for item in restored['domains'] if item['name'] == 'Transport')
    assert transport['causes'] == [{'name': 'Microwave fading', 'keywords': ['fading']}]
    detail = client.get(f'/api/non-qualified-calls/calls/{key}').json()['call']
    assert (detail['root_domain'], detail['root_cause']) == ('Transport', 'Microwave fading')

    exported = client.post('/api/non-qualified-calls/export', json={'filters': {}})
    sheet = load_workbook(io.BytesIO(exported.content))['NQ Calls']
    headers = [cell.value for cell in sheet[1]]
    assert {'Root Domain', 'Root Cause', 'Suggested Root Domain', 'Suggested Root Cause'} <= set(headers)
    deck = client.post('/api/non-qualified-calls/export/powerpoint', json={'filters': {}, 'granularity': 'month'})
    titles = [slide.shapes.title.text_frame.text for slide in Presentation(io.BytesIO(deck.content)).slides
              if slide.shapes.title is not None]
    assert any(title.startswith('Root Cause Analysis') for title in titles)
    assert any('eNB / gNB' in title for title in titles)
    page = client.get('/admin').text
    assert '>NQ Root Causes<' in page or '"NQ Root Causes"' in page


def test_admins_delete_single_history_entries_or_the_history_of_calls(client, tmp_path):
    enable_module()
    workspace_user('editor', 'user-editor')
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    keys = [call['call_key'] for call in query(client)['calls']]
    for key in keys:
        client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'status': 'Under Investigation', 'team': 'Transport'}})
    client.post(f'/api/non-qualified-calls/calls/{keys[0]}/comments', json={'body': 'Checked.'})
    detail = client.get(f'/api/non-qualified-calls/calls/{keys[0]}').json()
    assert len(detail['history']) == 2 and all(entry['id'] for entry in detail['history'])

    # Editors cannot delete the history.
    login(client, 'editor', 'editor123')
    assert client.delete(f"/api/non-qualified-calls/history/{detail['history'][0]['id']}").status_code == 403
    assert client.post('/api/non-qualified-calls/history/clear', json={'call_keys': keys}).status_code == 403

    login(client)
    assert client.delete(f"/api/non-qualified-calls/history/{detail['history'][0]['id']}").json() == {'deleted': 1}
    assert client.delete(f"/api/non-qualified-calls/history/{detail['history'][0]['id']}").status_code == 404
    after = client.get(f'/api/non-qualified-calls/calls/{keys[0]}').json()
    assert [entry['field'] for entry in after['history']] == ['team']
    # The follow-up and the comments stay.
    assert after['call']['status'] == 'Under Investigation' and len(after['comments']) == 1

    cleared = client.post('/api/non-qualified-calls/history/clear', json={'call_keys': keys[:2]})
    assert cleared.json() == {'deleted': 1 + 2}
    assert client.get(f'/api/non-qualified-calls/calls/{keys[1]}').json()['history'] == []
    assert len(client.get(f'/api/non-qualified-calls/calls/{keys[2]}').json()['history']) == 2
    assert client.post('/api/non-qualified-calls/history/clear', json={'call_keys': []}).status_code == 400
    page = client.get('/non-qualified-calls').text
    assert 'id="nq-bulk-history"' in page and 'id="nq-history-clear"' in page


def test_root_cause_suggestions_fall_back_to_comments_and_match_whole_words(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    calls = {call['result']: call for call in query(client)['calls']}
    # The CDR gives RF but no DL or UL: a comment of the engineers resolves the cause.
    cutoff = calls['Cutoff']['call_key']
    client.post(f'/api/non-qualified-calls/calls/{cutoff}/comments', json={'body': 'Strong UL interference seen on the eNB counters.'})
    assert query(client, filters={'result': ['Cutoff']})['calls'][0]['suggested_root_cause'] == {
        'domain': 'RF', 'cause': 'UL interference', 'source': 'comments'}
    # Without any CDR classification, a comment can choose the domain; keywords are whole words.
    failed = calls['Failed']['call_key']
    client.post(f'/api/non-qualified-calls/calls/{failed}/comments', json={'body': 'Poor performance, grant issue.'})
    assert client.get(f'/api/non-qualified-calls/calls/{failed}').json()['call']['suggested_root_cause'] is None
    client.post(f'/api/non-qualified-calls/calls/{failed}/comments', json={'body': 'Looks like an AAA rejection.'})
    assert client.get(f'/api/non-qualified-calls/calls/{failed}').json()['call']['suggested_root_cause'] == {
        'domain': 'AAA', 'cause': '', 'source': 'comments'}
    stats = client.post('/api/non-qualified-calls/root-causes/stats', json={'filters': {}}).json()
    assert {(item['domain'], item['cause']) for item in stats['causes']} >= {('RF', 'UL interference')}
    assert client.get('/api/non-qualified-calls/state').json()['root_cause_defaults'][0]['causes'][1]['name'] == 'UL interference'


def test_first_default_root_causes_are_upgraded_once(client, tmp_path):
    enable_module()
    login(client)
    client.get('/api/non-qualified-calls/state')
    repository = core.repository
    # A workspace seeded with the first defaults.
    with repository.connection() as connection:
        connection.execute(f"DELETE FROM {nq.NQ_ROOT_CAUSES_TABLE} WHERE domain = 'RF' AND cause = 'UL interference'")
        connection.execute(f"UPDATE {nq.NQ_ROOT_CAUSES_TABLE} SET keywords = '[\"interference\"]' WHERE domain = 'RF' AND cause = 'DL interference'")
        connection.execute(f"UPDATE {nq.NQ_ROOT_CAUSES_TABLE} SET keywords = '[\"volte\", \"ims\", \"e2e\"]' WHERE domain = 'IMS/E2E' AND cause = ''")
    repository.set_workspace_state(nq.ROOT_CAUSE_DEFAULTS_STATE_KEY, '1')
    taxonomy = client.get('/api/non-qualified-calls/root-causes').json()
    rf = next(item for item in taxonomy['domains'] if item['name'] == 'RF')
    assert [cause['name'] for cause in rf['causes']][:3] == ['DL interference', 'UL interference', 'Coverage']
    assert rf['causes'][0]['keywords'] == ['dl interference', 'downlink interference']
    assert next(item for item in taxonomy['domains'] if item['name'] == 'IMS/E2E')['keywords'] == ['volte', 'ims']
    # Edited values are kept, and the upgrade does not run again.
    rf['causes'] = [cause for cause in rf['causes'] if cause['name'] != 'UL interference']
    rf['causes'][0]['keywords'] = 'interference'
    assert client.put('/api/non-qualified-calls/root-causes', json={'domains': taxonomy['domains']}).status_code == 200
    again = client.get('/api/non-qualified-calls/root-causes').json()
    rf = next(item for item in again['domains'] if item['name'] == 'RF')
    assert rf['causes'][0]['keywords'] == ['interference'] and 'UL interference' not in [cause['name'] for cause in rf['causes']]


def test_everybody_sees_the_suggestion_rule_and_admins_edit_it(client, tmp_path):
    enable_module()
    workspace_user('editor', 'user-editor')
    workspace_user('manager', 'admin')
    workspace_user('viewer', 'user-viewer')
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    state = client.get('/api/non-qualified-calls/state').json()
    assert state['root_causes']['rule'] == nq.DEFAULT_ROOT_CAUSE_RULE and state['user']['can_configure_rule'] is True
    assert state['root_cause_rule_fields']['failure_classification'] == 'Failure Classification'

    def suggestion(service, result):
        return next(call for call in query(client)['calls'] if call['service'] == service and call['result'] == result)['suggested_root_cause']

    assert suggestion('voice', 'Failed') == {'domain': 'RF', 'cause': 'Coverage', 'source': 'cdr'}
    # Everybody sees the rule (the dialog is on the page); editors and viewers cannot change it.
    for username, password in (('editor', 'editor123'), ('viewer', 'viewer123')):
        login(client, username, password)
        assert client.get('/api/non-qualified-calls/state').json()['user']['can_configure_rule'] is False
        assert client.get('/api/non-qualified-calls/root-causes').json()['rule'] == nq.DEFAULT_ROOT_CAUSE_RULE
        assert 'id="nq-rule-dialog"' in client.get('/non-qualified-calls').text
        assert client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': nq.DEFAULT_ROOT_CAUSE_RULE}).status_code == 403
    login(client, 'manager', 'manager123')
    assert client.get('/api/non-qualified-calls/state').json()['user']['can_configure_rule'] is True
    assert client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': nq.DEFAULT_ROOT_CAUSE_RULE}).status_code == 200

    login(client)
    # Without the category fields, the CDR only gives the domain.
    rule = {**nq.DEFAULT_ROOT_CAUSE_RULE, 'cause_fields': ['failure_comment']}
    saved = client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': rule})
    assert saved.status_code == 200 and saved.json()['rule']['cause_fields'] == ['failure_comment']
    assert suggestion('voice', 'Failed') == {'domain': 'RF', 'cause': '', 'source': 'cdr'}
    assert client.get('/api/non-qualified-calls/root-causes').json()['rule']['cause_fields'] == ['failure_comment']
    # The comments resolve the cause the fields left open, unless the rule turns them off.
    key = next(call['call_key'] for call in query(client)['calls'] if call['result'] == 'Dropped')
    client.post(f'/api/non-qualified-calls/calls/{key}/comments', json={'body': 'qci1 bearer missing'})
    assert suggestion('voice', 'Dropped') == {'domain': 'IMS/E2E', 'cause': 'No QCI1 established', 'source': 'comments'}
    client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': {**rule, 'comments': False}})
    assert suggestion('voice', 'Dropped') == {'domain': 'IMS/E2E', 'cause': '', 'source': 'cdr'}
    # Partial matching finds keywords inside words: "rf" in "performance".
    data_failed = next(call['call_key'] for call in query(client)['calls'] if call['service'] == 'data' and call['result'] == 'Failed')
    client.post(f'/api/non-qualified-calls/calls/{data_failed}/comments', json={'body': 'Poor performance.'})
    assert suggestion('data', 'Failed') is None
    client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': {**rule, 'match': 'text'}})
    assert suggestion('data', 'Failed') == {'domain': 'RF', 'cause': '', 'source': 'comments'}
    client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': {**rule, 'comments': False}})
    bad = client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': {**rule, 'cause_fields': ['Call_Status']}})
    assert bad.status_code == 400 and 'Unknown root cause rule fields' in bad.text
    # The rule travels with the tracking and applies where no rule was chosen; no rule restores the default.
    document = nq.export_tracking_document(core.repository)
    assert json.loads(document)['root_causes']['rule']['comments'] is False
    assert client.put('/api/non-qualified-calls/root-causes/rule', json={'rule': None}).json()['rule'] == nq.DEFAULT_ROOT_CAUSE_RULE
    nq.import_tracking_document(core.repository, document)
    assert client.get('/api/non-qualified-calls/root-causes').json()['rule']['comments'] is False
    page = client.get('/non-qualified-calls').text
    assert 'id="nq-rule-dialog"' in page and 'id="nq-rule-open"' in page


def test_every_header_name_sorts_and_filters_the_calls(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    keys = [call['call_key'] for call in query(client)['calls']]
    client.patch(f'/api/non-qualified-calls/calls/{keys[0]}', json={'changes': {'root_domain': 'RF', 'root_cause': 'Coverage'}})
    for sort in ('cause', 'vendor', 'technology', 'campaign', 'service', 'result', 'assignee'):
        result = query(client, sort=sort, direction='asc')
        assert result['sort'] == sort and result['total'] == 4, sort
    assert query(client, sort='cause', direction='desc')['calls'][0]['root_cause'] == 'Coverage'
    page = client.get('/non-qualified-calls').text
    for field in ('service', 'vendor', 'campaign', 'technology', 'result', 'failure_classification', 'team', 'assignee',
                  'root_domain', 'root_cause'):
        assert f'data-filter="{field}"' in page and f'data-nq-filter="{field}"' in page, field


def test_progress_and_root_cause_values_filter_the_calls(client, tmp_path):
    enable_module()
    voice = voice_rows()
    voice['Test_Name'] = ['CALL', 'WhatsApp CALL', 'CALL']
    voice['Cell_ID_A'] = ['[134445192]', '[2769016]->[134445192]', '[24219]']
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice)
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)

    def total(**filters):
        return query(client, filters=filters)['total']

    assert total() == 4
    assert total(call_type=['WhatsApp']) == 1 and total(call_type=['Data']) == 2 and total(call_type=['Classic']) == 1
    assert total(nr_mode=['NSA']) == 4 and total(nr_mode=['SA']) == 0
    assert total(period=['month:2026-02']) == 4 and total(period=['week:2026-W09']) == 4 and total(period=['month:2026-03']) == 0
    assert total(age=['Over 90 days']) == 4 and total(age=['Under 7 days']) == 0
    # The root domain counting the suggestions, as the Root Cause Analysis does; the labelled pair.
    assert total(effective_domain=['RF']) == 2 and total(effective_domain=['IMS/E2E']) == 1
    assert total(effective_domain=[nq.UNASSIGNED]) == 1
    assert total(effective_cause=['RF||Coverage']) == 1 and total(root_pair=['RF||Coverage']) == 0
    key = next(call['call_key'] for call in query(client)['calls'] if call['result'] == 'Failed' and call['service'] == 'voice')
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'root_domain': 'RF', 'root_cause': 'Coverage'}})
    assert total(root_pair=['RF||Coverage']) == 1
    # The eNB of the last cell (ECI / 256) of each operator.
    # Nodes carry the mapped Operator: Vodafone UK is VF in the Operator Maps.
    assert total(node=['EE||eNB 525176']) == 1 and total(node=['VF||Cell 24219']) == 1
    # They combine with the other filters, travel with the progress and analysis, and are remembered.
    assert total(call_type=['Data'], effective_domain=['RF']) == 1
    progress = client.post('/api/non-qualified-calls/progress', json={'filters': {'call_type': ['Data']}}).json()
    assert progress['summary']['total'] == 2
    stats = client.post('/api/non-qualified-calls/root-causes/stats', json={'filters': {'node': ['EE||eNB 525176']}}).json()
    assert stats['summary']['total'] == 1
    assert client.put('/api/non-qualified-calls/filters', json={'filters': {'period': ['month:2026-02']}}).json()['filters'] == {
        'period': ['month:2026-02']}


def test_campaigns_read_year_quarter_and_mode_everywhere():
    from src.modules.column_names import campaign_sort_key, compact_campaign_value
    from src.modules.non_qualified_calls_visuals import _items, campaign_label

    raw = ['UK_Q2_SA_2026', 'UK_Q2_2026', 'UK_Q2_NSA_2026', 'UK_Q1_2026', '2026-Q2_SA', 'Special campaign']
    assert [campaign_label(value) for value in raw] == [
        '2026-Q2-SA', '2026-Q2', '2026-Q2-NSA', '2026-Q1', '2026-Q2-SA', 'Special campaign']
    # In a quarter: the plain campaign, then NSA, then SA.
    assert [compact_campaign_value(value) for value in sorted(raw[:4], key=campaign_sort_key)] == [
        '2026-Q1', '2026-Q2', '2026-Q2-NSA', '2026-Q2-SA']
    items = _items({'field': 'campaign', 'items': [{'value': 'UK_Q1_2026', 'count': 3}]}, {})
    assert items == [('2026-Q1', 3, '#b0234f')]

def test_indicator_cards_filter_by_follow_up_state(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    calls = query(client)['calls']
    first, second = calls[0]['call_key'], calls[1]['call_key']
    client.patch(f'/api/non-qualified-calls/calls/{first}', json={'changes': {'status': 'Resolved', 'team': 'Transport'}})
    client.post(f'/api/non-qualified-calls/calls/{second}/comments', json={'body': 'Checked.'})
    client.patch(f'/api/non-qualified-calls/calls/{second}', json={'changes': {'root_domain': 'RF'}})

    def total(**filters):
        return query(client, filters=filters)['total']

    summary = query(client)['summary']
    assert summary['attended'] == 2 and summary['closed'] == 1
    assert total(state=['closed']) == 1 and total(state=['attended']) == 2 and total(state=['not_attended']) == 2
    assert total(state=['with_team']) == 1 and total(state=['commented']) == 1 and total(state=['labelled']) == 1
    # Suggested: calls without a root cause that have a suggestion (as the Root Cause Analysis counts them).
    stats = client.post('/api/non-qualified-calls/root-causes/stats', json={'filters': {}}).json()
    assert total(rca_state=['suggested']) == stats['summary']['suggested']
    assert total(effective_domain=[nq.UNASSIGNED]) == stats['summary']['unclassified']


def test_operators_and_vendors_are_shown_and_filtered_with_their_mapped_labels(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    login(client)
    options = client.get('/api/non-qualified-calls/state').json()['filter_options']['operator']
    # The CDRs say Vodafone UK, an alias of VF in the Operator Maps.
    assert 'VF' in options and 'Vodafone UK' not in options
    calls = query(client, filters={'operator': ['VF']})['calls']
    assert calls and {call['operator'] for call in calls} == {'VF'}
    # A source spelling saved in the shared filters before selects its mapped label.
    assert query(client, filters={'operator': ['Vodafone UK']})['total'] == len(calls)


def add_stage_cdr(tmp_path: Path, name: str, kind: str, rows: pd.DataFrame, stage: str) -> int:
    dataset_id = add_cdr(tmp_path, name, kind, rows)
    core.repository.update_dataset_profile(dataset_id, cdr_stage=stage)
    return dataset_id


def joined_voice(rows: list[tuple[str, int, str, str]], campaign: str = 'UK_Q3_SA_2026') -> pd.DataFrame:
    """Voice calls as (JOIN_ID, Session_ID_A, result, start time)."""
    return pd.DataFrame({
        'JOIN_ID': [row[0] for row in rows], 'Session_ID_A': [row[1] for row in rows],
        'Operator': ['Vodafone UK'] * len(rows), 'Campaign': [campaign] * len(rows),
        'Session_Type': ['WhatsApp CALL'] * len(rows), 'Call_Status': [row[2] for row in rows],
        'status': [row[2] for row in rows], 'Call_Start_Time': [row[3] for row in rows],
        'Failure_Classification': ['RF Problems' if row[2] != 'Completed' else '' for row in rows],
    })


def test_join_id_identifies_calls_even_when_their_session_repeats(client, tmp_path):
    enable_module()
    # Two WhatsApp attempts share their A-side session but each has its own JOIN_ID.
    add_cdr(tmp_path, 'UK_Voice_Q3.xlsx', 'voice', joined_voice([
        ('0xAAA', 1043677052930, 'Failed', '2026-09-03 13:24:13'), ('0xBBB', 1043677052930, 'Failed', '2026-09-03 13:24:45'),
        ('0xCCC', 1043677052931, 'Completed', '2026-09-03 14:00:00'),
    ]))
    login(client)
    result = query(client)
    assert result['total'] == 2
    assert {call['join_id'] for call in result['calls']} == {'0xAAA', '0xBBB'}
    assert nq.call_key_for('voice', {'operator': 'EE', 'campaign': 'Other'}, '1', join_id='0xAAA') == \
        nq.call_key_for('voice', {'operator': 'Vodafone UK', 'campaign': 'UK_Q3_SA_2026'}, '2', join_id='0xaaa')


def test_daily_follow_up_moves_to_the_final_cdr_and_versions_are_told(client, tmp_path):
    enable_module()
    daily = add_stage_cdr(tmp_path, 'UK_Voice_20260921.xlsx', 'voice', joined_voice([
        ('0x1', 1, 'Failed', '2026-09-21 10:00:00'), ('0x2', 2, 'Dropped', '2026-09-21 11:00:00'),
        ('0x3', 3, 'Failed', '2026-09-21 12:00:00'), ('0x4', 4, 'Completed', '2026-09-21 13:00:00'),
    ]), 'daily')
    login(client)
    calls = {call['join_id']: call for call in query(client)['calls']}
    assert set(calls) == {'0x1', '0x2', '0x3'}
    key = calls['0x1']['call_key']
    client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'status': 'Under Investigation'}})
    client.post(f'/api/non-qualified-calls/calls/{key}/comments', json={'body': 'Analysed on the Daily CDR.'})

    # The Final CDR: 0x1 changed its failure, 0x2 is Completed, 0x3 dropped out and 0x5 is new.
    final_rows = joined_voice([
        ('0x1', 1, 'Dropped', '2026-09-21 10:00:00'), ('0x2', 2, 'Completed', '2026-09-21 11:00:00'),
        ('0x4', 4, 'Completed', '2026-09-21 13:00:00'), ('0x5', 5, 'Failed', '2026-09-22 09:00:00'),
    ])
    final = add_stage_cdr(tmp_path, 'UK_Voice_Q3_Final.xlsx', 'voice', final_rows, 'final')
    calls = {call['join_id']: call for call in query(client)['calls']}
    # The analysis follows the call to the Final CDR, which is its shown version.
    assert calls['0x1']['call_key'] == key and calls['0x1']['status'] == 'Under Investigation'
    assert calls['0x1']['dataset_id'] == final and calls['0x1']['result'] == 'Dropped'
    assert calls['0x1']['version_state'] == 'changed' and calls['0x1']['comment_count'] == 1
    assert calls['0x3']['version_state'] == 'not_in_final'
    assert '0x2' not in calls and calls['0x5']['version_state'] == ''
    # Completed in the Final CDR: listed on request, or while only the Daily CDR is chosen.
    assert [call['join_id'] for call in query(client, filters={'version': ['qualified']})['calls']] == ['0x2']
    only_daily = {call['join_id']: call for call in query(client, filters={'datasets': [str(daily)]})['calls']}
    assert set(only_daily) == {'0x1', '0x2', '0x3'}
    assert only_daily['0x1']['version_state'] == 'newer' and only_daily['0x1']['status'] == 'Under Investigation'
    assert only_daily['0x2']['version_state'] == 'qualified'
    detail = client.get(f'/api/non-qualified-calls/calls/{key}').json()
    assert [(version['stage'], version['latest']) for version in detail['versions']] == [('Final', True), ('Daily', False)]
    # Deleting the Daily CDR keeps the analysis of the calls the Final CDR still has.
    client.post(f'/datasets-analysis/delete/{daily}', follow_redirects=False)
    calls = {call['join_id']: call for call in query(client)['calls']}
    assert set(calls) == {'0x1', '0x5'} and calls['0x1']['status'] == 'Under Investigation'


def test_speech_calls_are_listed_once_and_their_samples_keep_their_own_follow_up(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'UK_Speech_Q3.xlsx', 'speech', pd.DataFrame({
        'JOIN_ID': ['0xS1', '0xS1', '0xS1', '0xS2'], 'Session_ID_A': [10, 10, 10, 11], 'Test_ID': [1, 2, 3, 4],
        'Operator': ['EE'] * 4, 'Campaign': ['UK_Q3_2026'] * 4, 'Test_Status': ['Failed', 'Completed', 'Failed', 'Completed'],
        'status': ['Failed', 'Completed', 'Failed', 'Completed'],
        'Test_Start_Time': ['2026-09-03 10:00:01', '2026-09-03 10:00:10', '2026-09-03 10:00:20', '2026-09-03 11:00:00'],
    }))
    login(client)
    result = query(client)
    assert result['total'] == 1 and result['calls'][0]['nq_samples'] == 2
    call_key = result['calls'][0]['call_key']
    detail = client.get(f'/api/non-qualified-calls/calls/{call_key}').json()
    assert len(detail['samples']) == 2
    sample_key = detail['samples'][1]['call_key']
    changed = client.patch(f'/api/non-qualified-calls/calls/{sample_key}', json={'changes': {'status': 'Resolved'}})
    assert changed.status_code == 200 and changed.json()['call']['parent_key'] == call_key
    client.post(f'/api/non-qualified-calls/calls/{sample_key}/comments', json={'body': 'Sample with garbled audio.'})
    assert query(client)['calls'][0]['status'] == 'Open'
    detail = client.get(f'/api/non-qualified-calls/calls/{call_key}').json()
    assert [sample['status'] for sample in detail['samples']] == ['Open', 'Resolved']
    workbook = load_workbook(io.BytesIO(client.post('/api/non-qualified-calls/export', json={'filters': {}}).content))
    assert 'Speech Samples' in workbook.sheetnames and len(list(workbook['Speech Samples'].values)) == 3


def test_nq_rate_of_each_campaign_and_operator(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Data_2026_Q1.xlsx', 'data', data_rows())
    login(client)
    rates = client.post('/api/non-qualified-calls/rates', json={'filters': {}}).json()
    voice = next(matrix for matrix in rates['matrices'] if matrix['service'] == 'voice')
    cells = voice['cells']['UK_Q1_2026']
    assert cells['EE'] == {'total': 2, 'nq': 1, 'rate': 50.0}
    assert voice['total'] == {'total': 3, 'nq': 2, 'rate': 66.67}
    assert [matrix['service'] for matrix in rates['matrices']] == ['voice', 'data', 'all']
    only_data = client.post('/api/non-qualified-calls/rates', json={'filters': {'service': ['data']}}).json()
    assert [matrix['service'] for matrix in only_data['matrices']] == ['data']
    # Each operator comes with its Operator Maps colour.
    assert voice['operator_colors']['EE'] == '#76B7B2'
    # The counts are reused until the indexed CDRs change: a changed CDR is counted again.
    assert nq._rate_counts_cache
    with core.repository.connection() as connection:
        connection.execute(f"UPDATE {nq.NQ_CALL_POPULATION_TABLE} SET is_nq = 0")
    assert client.post('/api/non-qualified-calls/rates', json={'filters': {}}).json()['matrices'][0]['total']['nq'] == 2
    with core.repository.connection() as connection:
        connection.execute(f"UPDATE {nq.NQ_CALL_SOURCES_TABLE} SET synced_at = 'later'")
    assert client.post('/api/non-qualified-calls/rates', json={'filters': {}}).json()['matrices'][0]['total']['nq'] == 0
    document = client.post('/api/non-qualified-calls/export/powerpoint', json={'filters': {}})
    assert document.status_code == 200


def test_analysis_fields_are_configured_filled_filtered_required_and_exported(client, tmp_path):
    enable_module()
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', voice_rows())
    login(client)
    saved = client.put('/api/non-qualified-calls/fields', json={'fields': [
        {'label': 'Final Category', 'type': 'list', 'in_table': True, 'in_summary': True, 'required_to_close': True,
         'options': [{'name': 'Poor Coverage LTE', 'color': '#c8102e'}, {'name': 'E2E'}]},
        {'label': 'Findings', 'type': 'long_text'}, {'label': 'Notice to VF3', 'type': 'yes_no'},
        {'label': 'Planned Date', 'type': 'date'},
    ]})
    assert saved.status_code == 200, saved.text
    keys = {field['label']: field['key'] for field in saved.json()['fields']}
    key = query(client)['calls'][0]['call_key']
    url = f'/api/non-qualified-calls/calls/{key}'
    assert client.patch(url, json={'changes': {'fields': {keys['Final Category']: 'Unknown'}}}).status_code == 400
    assert client.patch(url, json={'changes': {'fields': {keys['Planned Date']: '21/09/2026'}}}).status_code == 400
    # A required field must be filled before closing.
    refused = client.patch(url, json={'changes': {'status': 'Resolved'}})
    assert refused.status_code == 400 and 'Final Category' in refused.json()['detail']
    filled = client.patch(url, json={'changes': {'fields': {keys['Final Category']: 'poor coverage lte', keys['Findings']: 'Weak signal\nin the tunnel',
                                                          keys['Notice to VF3']: 'yes', keys['Planned Date']: '2026-10-15'}}})
    assert filled.status_code == 200, filled.text
    call = filled.json()['call']
    assert call['fields'][keys['Final Category']] == 'Poor Coverage LTE' and call['fields'][keys['Notice to VF3']] == 'Yes'
    assert call['version'] == 1
    history = filled.json()['history']
    assert ('field:' + keys['Final Category'], 'Poor Coverage LTE') in {(entry['field'], entry['new_value']) for entry in history}
    assert client.patch(url, json={'changes': {'status': 'Resolved'}}).status_code == 200
    assert query(client, filters={'fields': {keys['Final Category']: ['Poor Coverage LTE']}})['total'] == 1
    assert query(client, filters={'fields': {keys['Final Category']: [nq.UNASSIGNED]}})['total'] == 1
    assert query(client, filters={'search': 'in the tunnel'})['total'] == 1
    breakdown = next(item for item in query(client)['breakdowns'] if item['field'] == f"field:{keys['Final Category']}")
    assert {item['value']: item['count'] for item in breakdown['items']} == {'Poor Coverage LTE': 1, '': 1}
    sorted_calls = query(client, sort=f"field:{keys['Final Category']}", direction='desc')['calls']
    assert sorted_calls[0]['call_key'] == key
    # Renaming a value follows on every call; a value in use cannot be removed.
    fields = client.get('/api/non-qualified-calls/state').json()['fields']
    fields[0]['options'][0] = {'name': 'Poor Coverage 4G', 'previous': 'Poor Coverage LTE'}
    assert client.put('/api/non-qualified-calls/fields', json={'fields': fields}).status_code == 200
    assert client.get(url).json()['call']['fields'][keys['Final Category']] == 'Poor Coverage 4G'
    fields[0]['options'] = [{'name': 'E2E'}]
    assert client.put('/api/non-qualified-calls/fields', json={'fields': fields}).status_code == 400
    workbook = load_workbook(io.BytesIO(client.post('/api/non-qualified-calls/export', json={'filters': {}}).content))
    header = list(workbook['NQ Calls'].values)[0]
    assert 'Final Category' in header and header.index('Final Category') > header.index('Root Cause')
    # The analysis travels with the NQ Call Tracking package.
    document = json.loads(core._nq_call_tracking_payload(core.active_workspace))
    assert {field['label'] for field in document['fields']} >= {'Final Category', 'Findings'}
    with core.repository.connection() as connection:
        connection.execute(f'DELETE FROM {nq.NQ_FIELD_VALUES_TABLE}')
        connection.execute(f'DELETE FROM {nq.NQ_FIELDS_TABLE}')
    core._restore_workspace_nq_call_tracking(core.active_workspace, json.dumps(document).encode('utf-8'))
    assert client.get(url).json()['call']['fields'][keys['Final Category']] == 'Poor Coverage 4G'


def test_analysis_fields_are_proposed_from_a_user_input_workbook(client, tmp_path):
    from openpyxl import Workbook

    enable_module()
    login(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Failure Type', 'Auto_RCA_Category_A', 'Tunnel Failure', 'Findings', 'Planned Date', 'Netcheck Category'])
    sheet.append(['CALL', 'Will Come from Python Script- Yuriy', 'Yes', 'User Define', 'User Define', 'Value to be Taken from CDR'])
    sheet.append(['MRAB', None, 'No', None, None, None])
    output = io.BytesIO()
    workbook.save(output)
    response = client.post('/api/non-qualified-calls/fields/from-excel',
                           files={'workbook': ('User_Input_List.xlsx', output.getvalue(), 'application/vnd.ms-excel')})
    assert response.status_code == 200, response.text
    proposed = {field['label']: field for field in response.json()['fields']}
    assert proposed['Failure Type']['type'] == 'list' and [item['name'] for item in proposed['Failure Type']['options']] == ['CALL', 'MRAB']
    assert proposed['Tunnel Failure']['type'] == 'yes_no' and proposed['Findings']['type'] == 'long_text'
    assert proposed['Planned Date']['type'] == 'date'
    assert response.json()['skipped'] == ['Auto_RCA_Category_A', 'Netcheck Category']


def test_cdr_columns_and_optional_columns_of_the_calls_table(client, tmp_path):
    enable_module()
    rows = voice_rows()
    rows['Cellname_A'] = ['LEEDS_1', 'YORK_2', 'LEEDS_3']
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', rows)
    login(client)
    state = client.get('/api/non-qualified-calls/state').json()
    assert 'Cellname_A' in state['cdr_columns'] and 'join_id' in state['optional_columns']
    saved = client.put('/api/non-qualified-calls/table-columns', json={'builtin': ['join_id', 'nr_mode'], 'cdr': ['Cellname_A']})
    assert saved.status_code == 200 and saved.json()['table_columns'] == {'builtin': ['join_id', 'nr_mode'], 'cdr': ['Cellname_A'], 'export_cdr': []}
    # Region and Cluster are always in the table, in one column after Operator / Vendor.
    page = client.get('/non-qualified-calls').text
    assert page.index('data-sort="vendor"') < page.index('data-sort="region"') < page.index('data-sort="cluster"') < page.index('data-sort="city"')
    assert 'region' not in client.get('/api/non-qualified-calls/state').json()['optional_columns']
    calls = query(client, sort='cdr:Cellname_A', direction='asc')['calls']
    assert [call['extra']['Cellname_A'] for call in calls] == ['LEEDS_3', 'YORK_2']
    # Every optional and CDR column can be filtered by its values, like the other columns.
    options = client.get('/api/non-qualified-calls/state').json()['filter_options']
    assert options['column:cdr:Cellname_A'] == ['LEEDS_3', 'YORK_2'] and 'column:join_id' in options
    filtered = query(client, filters={'columns': {'cdr:Cellname_A': ['YORK_2']}})['calls']
    assert [call['extra']['Cellname_A'] for call in filtered] == ['YORK_2']
    assert query(client, filters={'columns': {'join_id': [nq.UNASSIGNED]}})['summary']['total'] >= 0
    saved_filters = client.put('/api/non-qualified-calls/filters', json={'filters': {'columns': {'cdr:Cellname_A': ['YORK_2'], 'bad key': ['x']}}})
    assert saved_filters.status_code == 200 and nq.saved_filters(core.repository)['columns'] == {'cdr:Cellname_A': ['YORK_2']}
    # More CDR columns can be exported to Excel only; the table's are exported too, with Cluster and both vendor identities.
    saved = client.put('/api/non-qualified-calls/table-columns', json={
        'builtin': ['join_id'], 'cdr': ['Cellname_A'], 'export_cdr': ['Cellname_A', 'Call_Status']})
    assert saved.json()['table_columns'] == {'builtin': ['join_id'], 'cdr': ['Cellname_A'], 'export_cdr': ['Call_Status']}
    sheet = load_workbook(io.BytesIO(client.post('/api/non-qualified-calls/export', json={'filters': {}}).content))['NQ Calls']
    headers = [cell.value for cell in sheet[1]]
    assert {'Cluster', 'Operator_Vendor', 'Vendor_Operator', 'CDR', 'Cellname_A', 'Call_Status'} <= set(headers)
    assert headers.index('Call_Status') == headers.index('Cellname_A') + 1 and headers[-1] == 'Call Key'
    rows = [dict(zip(headers, (cell.value for cell in row))) for row in sheet.iter_rows(min_row=2)]
    assert {row['Cellname_A'] for row in rows} == {'LEEDS_3', 'YORK_2'} and all(row['Call_Status'] for row in rows)
    assert client.put('/api/non-qualified-calls/table-columns', json={'builtin': ['unknown'], 'cdr': []}).status_code == 400


def test_session_type_is_a_call_column_with_its_own_filter(client, tmp_path):
    enable_module()
    rows = voice_rows()
    rows['Session_Type'] = ['MO Call', 'MT Call', 'MO Call']
    add_cdr(tmp_path, 'NetCheck_UK_CDR_Voice_2026_Q1.xlsx', 'voice', rows)
    login(client)
    assert client.get('/api/non-qualified-calls/state').json()['optional_columns']['session_type'] == 'Session Type'
    saved = client.put('/api/non-qualified-calls/table-columns', json={'builtin': ['session_type'], 'cdr': []})
    assert saved.status_code == 200
    calls = query(client, sort='session_type', direction='asc')['calls']
    assert [call['session_type'] for call in calls] == ['MO Call', 'MT Call']
    assert client.get('/api/non-qualified-calls/state').json()['filter_options']['column:session_type'] == ['MO Call', 'MT Call']
    filtered = query(client, filters={'columns': {'session_type': ['MT Call']}})['calls']
    assert [call['session_type'] for call in filtered] == ['MT Call']
    sheet = load_workbook(io.BytesIO(client.post('/api/non-qualified-calls/export', json={'filters': {}}).content))['NQ Calls']
    headers = [cell.value for cell in sheet[1]]
    assert headers.index('Session Type') == headers.index('Test Name') + 1


def test_follow_up_moves_once_to_the_join_id_call_keys(client, tmp_path):
    enable_module()
    rows = joined_voice([('0x9', 77, 'Failed', '2026-09-03 10:00:00')], campaign='UK_Q3_2026')
    dataset_id = add_cdr(tmp_path, 'UK_Voice_Q3.xlsx', 'voice', rows)
    login(client)
    new_key = query(client)['calls'][0]['call_key']
    # Recreate the index of the previous version: the call keyed by its session, with its follow-up.
    old_key = nq.call_key_for('voice', {'operator': 'Vodafone UK', 'campaign': 'UK_Q3_2026'}, '77')
    with core.repository.connection() as connection:
        connection.execute(f'UPDATE {nq.NQ_CALLS_TABLE} SET call_key = ?', (old_key,))
        connection.execute(f"UPDATE {nq.NQ_CALL_SOURCES_TABLE} SET revision = 'old'")
        connection.execute(f"INSERT INTO {nq.NQ_CALL_TRACKING_TABLE} (call_key, status, version, updated_by, updated_at) "
                           "VALUES (?, 'Under Investigation', 1, 'super', '2026-09-04T10:00:00')", (old_key,))
        connection.execute(f"INSERT INTO {nq.NQ_CALL_COMMENTS_TABLE} (uid, call_key, body, created_by, created_at) "
                           "VALUES ('c1', ?, 'Before the upgrade.', 'super', '2026-09-04T10:00:00')", (old_key,))
    core.repository.set_workspace_state(nq.KEY_SCHEME_STATE_KEY, '')
    call = query(client)['calls'][0]
    assert call['call_key'] == new_key and call['status'] == 'Under Investigation' and call['comment_count'] == 1
    assert core.repository.get_workspace_state(nq.KEY_SCHEME_STATE_KEY) == nq.KEY_SCHEME
    assert dataset_id
