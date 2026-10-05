from __future__ import annotations

from io import BytesIO
import re
from urllib.parse import unquote
import sqlite3

from pptx import Presentation

import src.DashboardAnalytic as app_module
from src.modules import scoring_jobs
from src.modules.repository import Repository
from tests.test_scoring_api import scoring_api


def test_scoring_page_reports_incomplete_catalogues_without_backfilling(scoring_api, monkeypatch):
    repository = scoring_api['repository']
    repository.set_main_cities(['Leeds', 'London'])

    def fail_if_scanned(*_args, **_kwargs):
        raise AssertionError('Opening Scoring must not scan materialized CDR rows.')

    monkeypatch.setattr(app_module, '_distinct_cdr_row_values', fail_if_scanned)
    monkeypatch.setattr(Repository, 'list_dataset_row_columns', fail_if_scanned)
    first = scoring_api['client'].get('/scoring')
    assert first.status_code == 200
    assert 'Some CDR filter catalogues are incomplete.' in first.text
    assert 'Main Cities' in first.text
    cached = repository.cdr_catalogues_by_dataset(scoring_api['complete_dataset_ids'])
    assert all(item['operators'] == [] for item in cached.values())
    assert repository.cdr_catalogues_by_dataset([]) == {}
    assert scoring_api['client'].get('/scoring').status_code == 200


def test_scoring_vendor_catalogue_uses_pure_labels_without_scanning_rows(scoring_api, monkeypatch):
    repository = scoring_api['repository']
    for dataset_id in scoring_api['complete_dataset_ids']:
        repository.replace_cdr_catalogue(
            dataset_id,
            vendors=['3_Ericsson', 'EE_Ericsson', 'VF_UK_Huawei', 'Mixed_Vendor', 'EE', 'O2', 'VF_SA'],
            vendors_only=['Ericsson', 'Huawei', 'Mixed_Vendor', 'EE - All', 'O2 - All', 'VF_SA - All'],
            regions=['North'], cities=['Leeds'], campaigns=['2026-Q2'],
            operators=['3', 'EE', 'O2', 'VF_UK', 'VF_SA'],
        )

    def fail_if_scanned(*_args, **_kwargs):
        raise AssertionError('Opening Scoring must use cached CDR catalogue values.')

    monkeypatch.setattr(app_module, '_distinct_cdr_row_values', fail_if_scanned)
    monkeypatch.setattr(Repository, 'list_dataset_row_columns', fail_if_scanned)
    captured = {}
    original_render = app_module.render_template

    def capture_context(request, template_name, context):
        captured.update(context)
        return original_render(request, template_name, context)

    monkeypatch.setattr(app_module, 'render_template', capture_context)
    page = scoring_api['client'].get('/scoring')

    assert page.status_code == 200
    catalogue = next(
        item['catalogue'] for item in captured['scoring_datasets']
        if item['file_name'] == 'UK_Q2_2026_NSA_Data.csv'
    )
    assert catalogue['operators'] == ['3', 'EE', 'O2', 'VF_SA', 'VF_UK']
    assert catalogue['vendors'] == ['EE - All', 'Ericsson', 'Huawei', 'Mixed_Vendor', 'O2 - All', 'VF_SA - All']


def test_operator_catalogue_replacement_preserves_unspecified_values_and_database_copy(scoring_api, tmp_path):
    repository = scoring_api['repository']
    dataset_id = scoring_api['dataset_id']
    app_module.cache_cdr_catalogue(
        dataset_id, repository.load_dataset_rows(dataset_id, ['Operator', 'City', 'Campaign'], {}), repository,
    )
    repository.replace_cdr_catalogue(
        dataset_id, vendors=['Ericsson'], vendors_only=['Ericsson'], regions=['West'], cities=['York'],
    )
    catalogue = repository.cdr_catalogues_by_dataset([dataset_id])[dataset_id]
    assert catalogue['operators'] == ['EE', 'O2']
    assert catalogue['campaigns'] == ['2026-Q2']
    job, _ = scoring_jobs.create_scoring_job(
        repository, scoring_api['complete_dataset_ids'], ['Operator'], 'NSA',
        context_filters={'City': ['York']},
    )
    destination = tmp_path / 'restored.db'
    with sqlite3.connect(repository.db_path) as source, sqlite3.connect(destination) as target:
        source.backup(target)
    restored = Repository(destination)
    restored.initialize()
    assert restored.cdr_catalogues_by_dataset([dataset_id])[dataset_id] == catalogue
    restored_job = scoring_jobs.get_scoring_job(restored, job['id'])
    assert restored_job['context_filters']['City'] == ['York']
    assert restored_job['configuration'] == job['configuration']


def test_scoring_export_labels_collapse_full_catalogue_filters_and_preserve_incomplete_jobs(scoring_api):
    repository = scoring_api['repository']
    repository.replace_cdr_catalogue(
        scoring_api['dataset_id'], vendors=['Nokia'], regions=['North'], cities=['Leeds'],
        vendors_only=['Nokia'],
        campaigns=['2026-Q2'], operators=['O2'],
    )
    complete = {'dataset_ids': [scoring_api['dataset_id']], 'context_filters': {
        'Region': ['North'], 'City': ['Leeds'], 'Operator': ['O2'],
        'Vendor': ['Nokia'], 'Campaign': ['2026-Q2'],
    }}
    normalized = app_module._scoring_export_job_with_catalogue_defaults(repository, complete)
    assert normalized['context_filters'] == {
        'Region': [], 'City': [], 'Operator': [], 'Vendor': [], 'Campaign': [],
    }
    assert complete['context_filters']['City'] == ['Leeds']

    explicitly_all = {'dataset_ids': [scoring_api['dataset_id']], 'context_filters': {
        'Region': [], 'City': [], 'Operator': [], 'Vendor': [], 'Campaign': [],
    }}
    assert app_module._scoring_export_job_with_catalogue_defaults(repository, explicitly_all) == explicitly_all

    partial = {'dataset_ids': [scoring_api['dataset_id']], 'context_filters': {'City': []}}
    assert app_module._scoring_export_job_with_catalogue_defaults(repository, partial) == partial
    partial_with_complete_city = {
        'dataset_ids': [scoring_api['dataset_id']], 'context_filters': {'City': ['Leeds']},
    }
    assert app_module._scoring_export_job_with_catalogue_defaults(
        repository, partial_with_complete_city,
    )['context_filters'] == {'City': []}

    missing = {
        'dataset_ids': [999999],
        'context_filters': {'Region': ['North'], 'City': ['Leeds'], 'Operator': ['O2'],
                            'Vendor': ['Nokia'], 'Campaign': ['2026-Q2']},
    }
    assert app_module._scoring_export_job_with_catalogue_defaults(repository, missing)['context_filters'] == missing['context_filters']


def test_scoring_api_filters_before_calculation_and_ppt_preserves_scope(scoring_api):
    response = scoring_api['client'].post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'nr_mode': 'NSA',
        'aggregation_levels': ['Operator'],
        'context_filters': {'Region': ['North'], 'City': ['Leeds'], 'Operator': ['O2'],
                            'Vendor': ['Nokia'], 'Campaign': ['2026-Q2']},
    })
    assert response.status_code == 200, response.text
    job = response.json()['job']
    assert job['context_filters']['City'] == ['Leeds']
    completed = scoring_jobs.run_scoring_job(scoring_api['repository'], job['id'])
    assert completed['status'] == 'completed', completed.get('error')
    frames = scoring_api['calls'][0]['frames']
    source_frames = frames.values() if isinstance(frames, dict) else [item[-1] for item in frames]
    assert all(set(frame['Operator']) == {'O2'} for frame in source_frames)
    exported = scoring_api['client'].get(f'/scoring/jobs/{job["id"]}/export/ppt')
    assert exported.status_code == 200, exported.text
    assert re.search(
        r'\d{8}_\d{6} - Scoring & GAP Analysis - NSA - North - All Clusters - Leeds - O2 - Nokia - 2026-Q2\.pptx',
        unquote(exported.headers['content-disposition']),
    )
    deck = Presentation(BytesIO(exported.content))
    for slide in list(deck.slides)[:2]:
        text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert 'Leeds' in text and 'North' in text and 'Nokia' in text and 'O2' in text
    saved = scoring_jobs.get_scoring_job(scoring_api['repository'], job['id'])
    assert saved['context_filters'] == job['context_filters']


def test_scoring_ppt_export_labels_full_selected_catalogues_as_all(scoring_api):
    repository = scoring_api['repository']
    selected_ids = scoring_api['complete_dataset_ids']
    repository.replace_cdr_catalogue(
        selected_ids[0], vendors=['Nokia'], vendors_only=['Nokia'], regions=['North'], cities=['Leeds'],
        campaigns=['2026-Q2'], operators=['EE', 'O2'],
    )
    for dataset_id in selected_ids[1:]:
        repository.replace_cdr_catalogue(
            dataset_id, vendors=['Nokia'], vendors_only=['Nokia'], regions=['North'], cities=['Leeds'],
            campaigns=['2026-Q2'], operators=['EE', 'O2'],
        )
    unrelated_id = scoring_api['add_ready_cdr'](
        name='Other_Q3_2026_NSA_Data.csv', campaign='2026-Q3',
    )
    repository.replace_cdr_catalogue(
        unrelated_id, vendors=['Huawei'], vendors_only=['Huawei'], regions=['South'], cities=['London'],
        campaigns=['2026-Q3'], operators=['Three UK'],
    )
    selected_catalogues = repository.cdr_catalogues_by_dataset(selected_ids)
    selected_values = {
        field: sorted({value for catalogue in selected_catalogues.values()
                       for value in catalogue[catalogue_field]}, key=str.casefold)
        for field, catalogue_field in {
            'Region': 'regions', 'City': 'cities', 'Operator': 'operators',
            'Vendor': 'vendors_only', 'Campaign': 'campaigns',
        }.items()
    }
    response = scoring_api['client'].post('/api/scoring/jobs', json={
        'dataset_ids': selected_ids, 'nr_mode': 'NSA', 'aggregation_levels': ['Operator'],
        'context_filters': selected_values,
    })
    assert response.status_code == 200, response.text
    job = response.json()['job']
    scoring_jobs.run_scoring_job(repository, job['id'])

    exported = scoring_api['client'].get(f'/scoring/jobs/{job["id"]}/export/ppt')
    assert exported.status_code == 200, exported.text
    filename = unquote(exported.headers['content-disposition'])
    assert 'All Regions' in filename and 'All Cities' in filename
    assert 'All Operators' in filename and 'All Vendors' in filename and '2026-Q2' in filename
    assert 'South' not in filename and 'London' not in filename and '2026-Q3' not in filename

    deck = Presentation(BytesIO(exported.content))
    cover_text = '\n'.join(shape.text for slide in list(deck.slides)[:2]
                            for shape in slide.shapes if shape.has_text_frame)
    assert 'Region: All Regions' in cover_text
    assert 'City: All Cities' in cover_text
    assert 'Vendor: All Vendors' in cover_text
    assert 'Operator: All Operators' in cover_text
    assert 'Campaigns: 2026-Q2' in cover_text
    assert 'South' not in cover_text and 'London' not in cover_text and 'Huawei' not in cover_text

    saved = scoring_jobs.get_scoring_job(repository, job['id'])
    assert saved['context_filters'] == selected_values
