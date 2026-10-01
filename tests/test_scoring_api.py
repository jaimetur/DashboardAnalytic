from __future__ import annotations

import csv
from copy import deepcopy
from io import BytesIO
from io import StringIO
import json
from pathlib import Path
import sqlite3
import time
from types import ModuleType

import pandas as pd
import pytest
from pptx import Presentation

import src.DashboardAnalytic as app_module
from src.modules import scoring_jobs
from src.modules.repository import Repository, local_now_iso
from tests.scoring_fixtures import scoring_configuration


def _login(client) -> None:
    response = client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=False)
    assert response.status_code == 303


def _mark_scoring_result_as_legacy(repository: Repository, job_id: int) -> None:
    result = scoring_jobs.get_scoring_job(repository, job_id, include_result=True)['result']
    result.pop('gap_direction', None)
    with sqlite3.connect(repository.db_path) as connection:
        connection.execute(
            'UPDATE scoring_jobs SET result_json = ? WHERE id = ?',
            (json.dumps(result, ensure_ascii=False), job_id),
        )


def _add_ready_cdr(
    repository: Repository,
    tmp_path: Path,
    *,
    name: str,
    nr_mode: str = 'NSA',
    kind: str = 'data',
    campaign: str = '2026-Q2',
) -> int:
    source = tmp_path / name
    source.write_text('test source', encoding='utf-8')
    dataset_id, _created = repository.add_dataset(name, str(source), 'admin')
    rows = pd.DataFrame({
        'Operator': ['EE', 'O2'],
        'Region': ['North', 'North'],
        'City': ['Leeds', 'Leeds'],
        'Vendor': ['Nokia', 'Nokia'],
        'Dataset_Kind': [kind, kind],
        'Environment': ['City', 'City'],
        'Campaign': [campaign, campaign],
        'score': [4.2, 3.6],
    })
    repository.replace_dataset_rows(dataset_id, rows)
    repository.update_dataset_profile(
        dataset_id,
        status='ready', progress=100, dataset_kind=kind, nr_mode=nr_mode,
        row_count=len(rows), column_count=len(rows.columns), processed_at=local_now_iso(),
    )
    repository.replace_cdr_catalogue(
        dataset_id, vendors=['Nokia'], regions=['North'], cities=['Leeds'], campaigns=[campaign],
    )
    return dataset_id


@pytest.fixture()
def scoring_api(client, tmp_path, monkeypatch):
    _login(client)
    workspace = app_module.active_workspace
    repository = Repository(
        workspace.database_path,
        global_db_path=app_module.repository.db_path,
        workspace_registry_db_path=app_module.workspace_registry.registry_path,
    )
    repository.initialize()
    repository.replace_scoring_configuration(scoring_configuration())

    calls = []
    engine = ModuleType('src.modules.scoring')
    engine.METHOD_VERSION = 'scoring-api-test-v1'
    engine.required_input_columns = lambda kind, levels: [*levels, 'score']

    def calculate_scoring(frames, levels, *, baseline_operator, configuration=None):
        calls.append({'frames': frames, 'levels': levels, 'baseline_operator': baseline_operator})
        scoring = [{
            'campaign': '2026-Q2', 'region': 'North', 'city': 'Leeds', 'vendor': 'Nokia',
            'dataset_type': 'data', 'environment': 'City', 'operator': operator,
            'kpi': 'Download Throughput', 'category': 'Data', 'value': 4.2,
            'score': 4.2, 'weighted_points': 4.2, 'max_points': 5.0, 'sample_count': 1,
        } for operator in ('EE', 'O2', 'Three UK', 'Vodafone UK')]
        charts = [{
            'campaign': '2026-Q2', 'region': 'North', 'city': 'Leeds', 'vendor': 'Nokia',
            'dataset_type': 'data', 'environment': 'City', 'operator': 'EE',
            'category': 'Overall', 'weighted_points': 4.2, 'max_points': 5.0, 'complete': True,
        }]
        gap = [{'operator': 'O2', 'baseline_operator': baseline_operator, 'gap_points': 0.6}]
        return {'scoring': scoring, 'charts': charts, 'gap': gap, 'warnings': []}

    engine.calculate_scoring = calculate_scoring
    monkeypatch.setattr(scoring_jobs, '_scoring_engine', lambda: engine)
    submitted = []

    def record_submission(task_repository, callback, *args, **kwargs):
        submitted.append((task_repository, callback, args, kwargs))
        return None

    monkeypatch.setattr(app_module, '_submit_workspace_job', record_submission)
    dataset_id = _add_ready_cdr(repository, tmp_path, name='UK_Q2_2026_NSA_Data.csv')
    voice_id = _add_ready_cdr(repository, tmp_path, name='UK_Q2_2026_NSA_Voice.csv', kind='voice')
    speech_id = _add_ready_cdr(repository, tmp_path, name='UK_Q2_2026_NSA_Speech.csv', kind='speech')
    return {
        'client': client,
        'repository': repository,
        'calls': calls,
        'submitted': submitted,
        'dataset_id': dataset_id,
        'complete_dataset_ids': [dataset_id, voice_id, speech_id],
        'add_ready_cdr': lambda **kwargs: _add_ready_cdr(repository, tmp_path, **kwargs),
        'tmp_path': tmp_path,
    }


def test_scoring_page_requires_login_and_renders_workspace_controls(client, scoring_api):
    client.get('/logout', follow_redirects=False)
    assert client.get('/scoring', follow_redirects=False).status_code == 303
    _login(client)
    page = client.get('/scoring')

    assert page.status_code == 200
    assert 'Scoring &amp; GAP Analysis' in page.text
    assert 'href="/scoring"' in page.text
    assert 'data-scoring-workspace' in page.text
    assert 'UK_Q2_2026_NSA_Data.csv' in page.text
    assert 'value="EE"' in page.text and 'value="O2"' in page.text
    assert '2026-Q2' in page.text
    assert 'Operator is required' in page.text
    assert 'data-calculate-scoring' in page.text
    assert 'data-recalculate-scoring' in page.text
    assert 'data-result-tab="scoring">Scoring Tables</button>' in page.text
    assert 'data-result-tab="charts">Scoring Charts</button>' in page.text
    assert 'GAP Analysis' in page.text
    assert 'Best Network Chart' in page.text


def test_scoring_page_exposes_profile_choices_and_active_profile(scoring_api, monkeypatch):
    captured = {}
    original_render = app_module.render_template

    def capture_context(request, template_name, context):
        captured.update(context)
        return original_render(request, template_name, context)

    monkeypatch.setattr(app_module, 'render_template', capture_context)
    page = scoring_api['client'].get('/scoring')

    assert page.status_code == 200
    profiles = captured['scoring_profiles']
    assert profiles == [{
        'id': 'netcheck-2026',
        'name': 'NetCheck 2026',
        'aggregation_hierarchy': scoring_api['repository'].get_scoring_configuration()['aggregation_hierarchy'],
    }]
    assert captured['scoring_active_profile_id'] == 'netcheck-2026'
    assert captured['aggregation_levels'] == profiles[0]['aggregation_hierarchy']


def test_scoring_baseline_options_follow_mapping_without_scanning_cdr_rows(scoring_api, monkeypatch):
    repository = scoring_api['repository']
    repository.replace_operator_mapping_groups([
        {'canonical': 'O2', 'aliases': ['O2 UK'], 'color': '#112233'},
        {'canonical': 'EE', 'aliases': ['EE UK'], 'color': '#445566'},
    ])

    def fail_if_cdr_rows_are_scanned(*_args, **_kwargs):
        raise AssertionError('Opening Scoring must not query dataset rows.')

    monkeypatch.setattr(Repository, 'list_distinct_dataset_row_values', fail_if_cdr_rows_are_scanned)
    page = scoring_api['client'].get('/scoring')

    assert page.status_code == 200
    baseline_select = page.text.split('<select data-baseline-operator', 1)[1].split('</select>', 1)[0]
    assert baseline_select.index('value="O2"') < baseline_select.index('value="EE"')
    assert 'value="O2" data-operator-color="#112233"' in baseline_select
    assert 'value="EE" data-operator-color="#445566"' in baseline_select


def test_scoring_selection_defaults_use_catalogue_metadata_without_reading_cdr_rows(scoring_api, monkeypatch):
    def fail_if_cdr_rows_are_read(*_args, **_kwargs):
        raise AssertionError('Loading the saved scoring selection must not read CDR rows.')

    monkeypatch.setattr(Repository, 'list_distinct_dataset_row_values', fail_if_cdr_rows_are_read)
    monkeypatch.setattr(Repository, 'dataset_row_count', fail_if_cdr_rows_are_read)
    monkeypatch.setattr(Repository, 'list_dataset_row_columns', fail_if_cdr_rows_are_read)

    response = scoring_api['client'].get('/api/scoring/selection')

    assert response.status_code == 200
    assert response.json() == {
        'selection': {
            'dataset_ids': scoring_api['complete_dataset_ids'],
            'aggregation_levels': ['Operator'],
            'nr_mode': 'NSA',
            'baseline_operator': 'EE',
            'scoring_profile_id': 'netcheck-2026',
            'context_filters': {
                'Region': [], 'City': [], 'Operator': [], 'Vendor': [], 'Campaign': [],
            },
        },
        'warnings': [],
        'persisted': False,
    }


def test_scoring_selection_persists_for_workspace_and_rejects_delayed_tab_updates(scoring_api):
    client = scoring_api['client']
    common_selection = {
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'aggregation_levels': ['Operator', 'Region'],
        'nr_mode': 'NSA',
        'baseline_operator': 'O2',
        'scoring_profile_id': 'netcheck-2026',
        'context_filters': {'Region': ['North'], 'City': [], 'Operator': [], 'Vendor': [], 'Campaign': []},
    }
    saved = client.put('/api/scoring/selection', json={
        'selection': common_selection, 'client_id': 'tab-a', 'client_revision': 1,
    })

    assert saved.status_code == 200
    assert saved.json()['accepted'] is True
    assert saved.json()['selection'] == common_selection
    assert client.get('/api/scoring/selection').json()['persisted'] is True

    newer_selection = {**common_selection, 'baseline_operator': 'EE'}
    newer = client.put('/api/scoring/selection', json={
        'selection': newer_selection, 'client_id': 'tab-a', 'client_revision': 2,
    })
    delayed = client.put('/api/scoring/selection', json={
        'selection': common_selection, 'client_id': 'tab-a', 'client_revision': 1,
    })

    assert newer.status_code == 200 and newer.json()['accepted'] is True
    assert delayed.status_code == 200 and delayed.json()['accepted'] is False
    assert delayed.json()['selection'] == newer_selection
    assert client.get('/api/scoring/selection').json()['selection'] == newer_selection

    other_tab = client.put('/api/scoring/selection', json={
        'selection': common_selection, 'client_id': 'tab-b', 'client_revision': 1,
    })
    assert other_tab.status_code == 200 and other_tab.json()['accepted'] is True
    assert client.get('/api/scoring/selection').json()['selection'] == common_selection

    empty_selection = {**common_selection, 'dataset_ids': []}
    cleared = client.put('/api/scoring/selection', json={'selection': empty_selection})
    assert cleared.status_code == 200 and cleared.json()['accepted'] is True
    assert client.get('/api/scoring/selection').json()['selection']['dataset_ids'] == []


def test_scoring_selection_recovers_from_deleted_cdr_and_methodology(scoring_api):
    client = scoring_api['client']
    repository = scoring_api['repository']
    selection = {
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'aggregation_levels': ['Operator', 'Region'],
        'nr_mode': 'NSA',
        'baseline_operator': 'EE',
        'scoring_profile_id': 'netcheck-2026',
        'context_filters': {'Region': ['North']},
    }
    saved = client.put('/api/scoring/selection', json={'selection': selection})
    assert saved.status_code == 200

    stale_dataset_id = scoring_api['complete_dataset_ids'][1]
    repository.delete_dataset(stale_dataset_id)
    configuration = repository.get_scoring_configuration()
    repository.replace_scoring_profiles({
        'active_profile_id': 'alternate-method',
        'profiles': [{
            'id': 'alternate-method', 'name': 'Alternate Method', 'configuration': configuration,
        }],
    })

    recovered = client.get('/api/scoring/selection')

    assert recovered.status_code == 200
    assert recovered.json()['selection']['dataset_ids'] == [
        scoring_api['complete_dataset_ids'][0], scoring_api['complete_dataset_ids'][2],
    ]
    assert recovered.json()['selection']['scoring_profile_id'] == 'alternate-method'
    assert recovered.json()['selection']['aggregation_levels'] == ['Operator', 'Region']
    assert any('CDR datasets' in warning for warning in recovered.json()['warnings'])
    assert any('methodology' in warning for warning in recovered.json()['warnings'])


def test_scoring_selection_survives_workspace_database_backup_restore(scoring_api, tmp_path):
    client = scoring_api['client']
    repository = scoring_api['repository']
    workspace = app_module.active_workspace
    selection = {
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'aggregation_levels': ['Operator', 'Region'],
        'nr_mode': 'NSA',
        'baseline_operator': 'EE',
        'scoring_profile_id': 'netcheck-2026',
        'context_filters': {'Region': ['North']},
    }
    saved = client.put('/api/scoring/selection', json={'selection': selection})
    assert saved.status_code == 200
    persisted_selection = saved.json()['selection']

    archive_path = app_module.create_recurring_database_backup({
        'components': ['workspace_database'],
        'backup_path': str(tmp_path / 'backups'),
        'max_backups': 5,
        'workspace_ids': [workspace.id],
    })
    changed = {**persisted_selection, 'baseline_operator': 'O2'}
    assert client.put('/api/scoring/selection', json={'selection': changed}).status_code == 200

    app_module.restore_database_backup(archive_path, ['workspace_database'])

    assert client.get('/api/scoring/selection').json()['selection'] == persisted_selection
    assert repository.get_workspace_state('scoring_calculation_selection_v1')


def test_scoring_api_validates_auth_mode_and_missing_jobs(client, scoring_api):
    assert client.get('/api/scoring/jobs').status_code == 200
    assert client.get('/api/scoring/jobs/999999').status_code == 404
    assert client.get('/scoring/jobs/999999/export/scoring').status_code == 404

    sa_id = _add_ready_cdr(
        scoring_api['repository'], scoring_api['tmp_path'], name='UK_Q2_2026_SA_Data.csv', nr_mode='SA',
    )
    response = client.post('/api/scoring/jobs', json={
        'dataset_ids': [scoring_api['dataset_id'], sa_id],
        'aggregation_levels': ['Operator'], 'nr_mode': 'NSA',
    })
    assert response.status_code == 400
    assert 'does not match the selected NSA NR Mode' in response.json()['detail']

    client.get('/logout', follow_redirects=False)
    assert client.get('/scoring', follow_redirects=False).status_code == 303
    assert client.get('/api/scoring/jobs').status_code == 401


def test_scoring_configuration_api_round_trips_complete_validated_document(scoring_api):
    client = scoring_api['client']
    repository = scoring_api['repository']
    original = repository.get_scoring_configuration()

    fetched = client.get('/api/workspace-config/scoring-configuration')
    assert fetched.status_code == 200
    assert fetched.json() == original
    assert 'configuration' not in fetched.json()

    updated = deepcopy(original)
    updated['gap_priority'] = list(reversed(updated['gap_priority']))
    saved = client.put('/api/workspace-config/scoring-configuration', json=updated)
    assert saved.status_code == 200, saved.text
    assert saved.json() == repository.get_scoring_configuration()
    assert saved.json()['gap_priority'] == updated['gap_priority']

    invalid = deepcopy(updated)
    invalid['gap_priority'].append(invalid['gap_priority'][0])
    rejected = client.put('/api/workspace-config/scoring-configuration', json=invalid)
    assert rejected.status_code == 400
    assert repository.get_scoring_configuration() == saved.json()


def test_scoring_profiles_api_edits_inactive_profile_and_switches_active(scoring_api):
    client = scoring_api['client']
    repository = scoring_api['repository']
    active_configuration = repository.get_scoring_configuration()
    profiles = client.get('/api/workspace-config/scoring-profiles').json()
    second_configuration = deepcopy(active_configuration)
    second_configuration['version'] = 'NetCheck 2025'
    profiles['profiles'].append({
        'id': 'netcheck-2025',
        'name': 'NetCheck 2025',
        'configuration': second_configuration,
    })

    saved = client.put('/api/workspace-config/scoring-profiles', json=profiles)
    assert saved.status_code == 200, saved.text
    assert client.get('/api/workspace-config/scoring-configuration').json() == active_configuration

    inactive_edit = deepcopy(saved.json())
    inactive_edit['profiles'][1]['configuration']['metrics'][0]['contexts']['DriveCity']['max_points'] = 81
    edited = client.put('/api/workspace-config/scoring-profiles', json=inactive_edit)
    assert edited.status_code == 200, edited.text
    assert client.get('/api/workspace-config/scoring-configuration').json() == active_configuration

    selected = edited.json()
    selected['active_profile_id'] = 'netcheck-2025'
    switched = client.put('/api/workspace-config/scoring-profiles', json=selected)
    assert switched.status_code == 200
    assert client.get('/api/workspace-config/scoring-configuration').json() == selected['profiles'][1]['configuration']


def test_scoring_job_can_select_inactive_profile_without_changing_workspace_default(scoring_api):
    client = scoring_api['client']
    repository = scoring_api['repository']
    active_id = repository.get_scoring_profiles()['active_profile_id']
    active_configuration = deepcopy(repository.get_scoring_configuration())
    profiles = repository.get_scoring_profiles()
    inactive = deepcopy(profiles['profiles'][0])
    inactive.update({'id': 'netcheck-2025', 'name': 'NetCheck 2025'})
    inactive['configuration']['version'] = 'NetCheck 2025'
    inactive['configuration']['aggregation_hierarchy'] = [
        'Region', 'Operator', 'Vendor', 'City', 'Campaign',
    ]
    inactive['configuration']['metrics'][0]['contexts']['DriveCity']['max_points'] = 88
    profiles['profiles'].append(inactive)
    repository.replace_scoring_profiles(profiles)

    response = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'nr_mode': 'NSA',
        'aggregation_levels': ['Region', 'City'],
        'scoring_profile_id': 'netcheck-2025',
    })

    assert response.status_code == 200, response.text
    job = response.json()['job']
    assert job['scoring_profile_id'] == 'netcheck-2025'
    assert job['scoring_profile_name'] == 'NetCheck 2025'
    assert job['configuration']['metrics'][0]['contexts']['DriveCity']['max_points'] == 88
    assert job['aggregation_hierarchy'] == inactive['configuration']['aggregation_hierarchy']
    assert job['aggregation_levels'] == ['Region', 'Operator', 'City']
    assert repository.get_scoring_profiles()['active_profile_id'] == active_id
    assert repository.get_scoring_configuration() == active_configuration

    defaulted = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'nr_mode': 'NSA',
        'aggregation_levels': ['Operator'],
        'scoring_profile_id': '',
    })
    assert defaulted.status_code == 200, defaulted.text
    assert defaulted.json()['job']['scoring_profile_id'] == active_id

    unknown = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'nr_mode': 'NSA',
        'aggregation_levels': ['Operator'],
        'scoring_profile_id': 'missing-profile',
    })
    assert unknown.status_code == 400
    assert 'was not found' in unknown.json()['detail']


def test_scoring_exports_accept_expanded_and_summary_table_modes(scoring_api):
    client = scoring_api['client']
    repository = scoring_api['repository']
    response = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'nr_mode': 'NSA',
        'aggregation_levels': ['Operator'],
    })
    assert response.status_code == 200, response.text
    job = scoring_jobs.run_scoring_job(repository, response.json()['job']['id'])
    assert job['status'] == 'completed'

    legacy_csv = client.get(f"/scoring/jobs/{job['id']}/export/scoring")
    city_raw_csv = client.get(f"/scoring/jobs/{job['id']}/export/scoring?environment=DriveCity")
    summary_csv = client.get(f"/scoring/jobs/{job['id']}/export/scoring?table_mode=summary")
    city_summary_csv = client.get(
        f"/scoring/jobs/{job['id']}/export/scoring?table_mode=summary&environment=DriveCity",
    )
    summary_ppt = client.get(f"/scoring/jobs/{job['id']}/export/ppt?table_mode=summary&gap_layout=adjacent")
    city_ppt = client.get(f"/scoring/jobs/{job['id']}/export/ppt?environment=DriveCity")
    invalid_environment = client.get(f"/scoring/jobs/{job['id']}/export/ppt?environment=Unknown")
    invalid_mode = client.get(f"/scoring/jobs/{job['id']}/export/gap?table_mode=compact")
    invalid_gap_layout = client.get(f"/scoring/jobs/{job['id']}/export/ppt?gap_layout=side")

    assert legacy_csv.status_code == 200
    assert city_raw_csv.status_code == 200
    raw_city_rows = list(csv.DictReader(StringIO(city_raw_csv.content.decode('utf-8-sig'))))
    assert raw_city_rows and {row['environment'] for row in raw_city_rows} == {'City'}
    assert summary_csv.status_code == 200
    assert summary_csv.headers['content-type'].startswith('text/csv')
    assert city_summary_csv.status_code == 200
    city_summary_rows = list(csv.DictReader(StringIO(city_summary_csv.content.decode('utf-8-sig'))))
    assert city_summary_rows and {row['environment'] for row in city_summary_rows} == {'DriveCity'}
    assert summary_ppt.status_code == 200
    assert summary_ppt.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.presentationml.presentation')
    assert city_ppt.status_code == 200
    assert invalid_environment.status_code == 400
    assert invalid_mode.status_code == 400
    assert invalid_gap_layout.status_code == 400


def test_workspace_config_renders_kpi_and_gap_priority_panels(scoring_api):
    page = scoring_api['client'].get('/workspace-config')

    assert page.status_code == 200
    assert '<h2>Scoring KPI Configuration</h2>' in page.text
    assert '<h2>GAP KPI Priority</h2>' in page.text
    assert 'data-scoring-environment' in page.text
    assert 'data-scoring-kpi-rows' in page.text
    assert 'data-scoring-priority-rows' in page.text
    assert '/static/js/scoring_config.js?v=' in page.text
    assert 'data-no-persist' in page.text
    assert 'Import or export this configuration in Admin' in page.text


def test_scoring_job_results_cache_force_and_exports(scoring_api):
    client = scoring_api['client']
    dataset_id = scoring_api['dataset_id']
    payload = {
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'aggregation_levels': ['Operator', 'Region'],
        'nr_mode': 'NSA',
        'baseline_operator': 'EE',
    }
    created = client.post('/api/scoring/jobs', json=payload)
    assert created.status_code == 200, created.text
    job_id = created.json()['job']['id']
    assert created.json()['job']['status'] == 'queued'
    assert len(scoring_api['submitted']) == 1

    unfinished_export = client.get(f'/scoring/jobs/{job_id}/export/scoring')
    assert unfinished_export.status_code == 409

    completed = scoring_jobs.run_scoring_job(scoring_api['repository'], job_id)
    assert completed['status'] == 'completed'
    assert scoring_api['calls'][0]['levels'] == ['Operator', 'Region']
    assert scoring_api['calls'][0]['baseline_operator'] == 'EE'

    cached = client.post('/api/scoring/jobs', json=payload)
    assert cached.status_code == 200
    assert cached.json()['cached'] is True
    assert cached.json()['job']['id'] == job_id
    assert len(scoring_api['submitted']) == 1

    _mark_scoring_result_as_legacy(scoring_api['repository'], job_id)

    scoring_api['repository'].replace_operator_mapping_groups([
        {'canonical': 'Three UK', 'aliases': ['3'], 'color': '#AABBCC'},
        {'canonical': 'O2', 'aliases': ['O2 UK'], 'color': '#112233'},
        {'canonical': 'Vodafone UK', 'aliases': ['VF'], 'color': '#445566'},
        {'canonical': 'EE', 'aliases': ['EE (UK)'], 'color': '#778899'},
    ])

    result = client.get(f'/api/scoring/jobs/{job_id}')
    assert result.status_code == 200
    assert result.json()['job']['status'] == 'completed'
    assert result.json()['scoring'][0]['weighted_points'] == 4.2
    assert result.json()['charts'][0]['category'] == 'Overall'
    assert result.json()['gap'][0]['gap_points'] == -0.6
    assert result.json()['gap_direction'] == 'operator_minus_reference'
    stored_result = scoring_jobs.get_scoring_job(scoring_api['repository'], job_id, include_result=True)['result']
    assert stored_result['gap'][0]['gap_points'] == 0.6
    assert 'gap_direction' not in stored_result
    assert [row['operator'] for row in result.json()['scoring']] == ['EE', 'O2', 'Three UK', 'Vodafone UK']
    score_table = result.json()['views']['score_tables'][0]
    assert score_table['operators'] == ['Three UK', 'O2', 'Vodafone UK', 'EE']
    assert score_table['operator_styles']['Three UK']['color'] == '#AABBCC'
    assert score_table['operator_styles']['EE']['color'] == '#778899'

    csv_response = client.get(f'/scoring/jobs/{job_id}/export/scoring')
    assert csv_response.status_code == 200
    assert csv_response.headers['content-type'].startswith('text/csv')
    assert 'Download Throughput' in csv_response.content.decode('utf-8-sig')
    gap_csv = client.get(f'/scoring/jobs/{job_id}/export/gap')
    assert gap_csv.status_code == 200
    assert 'gap_points' in gap_csv.content.decode('utf-8-sig')
    gap_csv_rows = pd.read_csv(StringIO(gap_csv.content.decode('utf-8-sig')))
    assert gap_csv_rows.loc[0, 'gap_points'] == -0.6

    ppt_response = client.get(f'/scoring/jobs/{job_id}/export/ppt')
    assert ppt_response.status_code == 200, ppt_response.text
    presentation = Presentation(BytesIO(ppt_response.content))
    slide_text = '\n'.join(
        shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert 'Scoring & GAP Analysis' in slide_text
    assert '2026-Q2' in slide_text
    scoring_slide = next(
        slide for slide in presentation.slides
        if any(shape.has_text_frame and 'Scoring Table' in shape.text for shape in slide.shapes)
    )
    scoring_table = next(shape.table for shape in scoring_slide.shapes if shape.has_table)
    # The first row groups Score and GAP; the next rows show operator and region hierarchy.
    assert scoring_table.cell(0, 5).text == 'Score'
    assert scoring_table.cell(0, 9).text == 'GAP'
    assert [scoring_table.cell(1, column).text for column in range(5, 9)] == [
        'Three UK', 'O2', 'Vodafone UK', 'EE',
    ]
    assert all(scoring_table.cell(2, column).text == 'North' for column in range(5, 9))
    assert str(scoring_table.cell(1, 5).fill.fore_color.rgb) == 'AABBCC'

    forced = client.post('/api/scoring/jobs', json={**payload, 'force': True})
    assert forced.status_code == 200
    assert forced.json()['job']['id'] != job_id
    assert forced.json()['cached'] is False
    assert len(scoring_api['submitted']) == 2


def test_global_raw_kpis_survive_job_reload_api_get_and_csv_export(scoring_api, monkeypatch):
    client = scoring_api['client']
    repository = scoring_api['repository']
    engine = scoring_jobs._scoring_engine()
    original_calculate = engine.calculate_scoring
    metric = repository.get_scoring_configuration()['metrics'][0]
    global_kpi = {
        'campaign': '2026-Q2', 'region': 'North', 'city': 'Leeds', 'vendor': 'Nokia',
        'dataset_type': 'voice', 'operator': 'EE', 'kpi_code': metric['code'],
        'kpi': metric['kpi'], 'category': metric['category'],
        'kpi_type': metric['kpi_type'], 'unit': '%', 'value': 98.75, 'sample_count': 12,
        'environment': 'All Environments', 'complete_coverage': True, 'missing_environments': [],
    }

    def calculate_with_global_kpis(frames, levels, *, baseline_operator, configuration=None):
        result = original_calculate(
            frames, levels, baseline_operator=baseline_operator, configuration=configuration,
        )
        result['global_kpis'] = [global_kpi]
        return result

    monkeypatch.setattr(engine, 'calculate_scoring', calculate_with_global_kpis)
    payload = {
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'aggregation_levels': ['Operator', 'Region'],
        'nr_mode': 'NSA', 'baseline_operator': 'EE',
    }
    created = client.post('/api/scoring/jobs', json=payload)
    assert created.status_code == 200, created.text
    job_id = created.json()['job']['id']
    completed = scoring_jobs.run_scoring_job(repository, job_id)
    assert completed['status'] == 'completed'
    persisted = scoring_jobs.get_scoring_job(repository, job_id, include_result=True)
    assert persisted['result']['global_kpis'] == [global_kpi]

    response = client.get(f'/api/scoring/jobs/{job_id}')
    assert response.status_code == 200, response.text
    api_result = response.json()
    assert api_result['global_kpis'] == [global_kpi]
    assert api_result['scoring'] == persisted['result']['scoring']
    assert api_result['charts'] == persisted['result']['charts']

    csv_response = client.get(f'/scoring/jobs/{job_id}/export/scoring?table_mode=expanded&environment=all')
    assert csv_response.status_code == 200, csv_response.text
    exported_rows = list(csv.DictReader(StringIO(csv_response.content.decode('utf-8-sig'))))
    global_row = next((row for row in exported_rows if (
        row['operator'] == 'EE' and row['kpi_code'] == metric['code'] and row['kpi_value'] == '98.75'
    )), None)
    assert global_row, [row for row in exported_rows if row['kpi_code'] == metric['code']][:8]
    assert global_row['kpi_value'] == '98.75'

    from src.modules.scoring_exports import export_scoring_csv

    without_global_kpis = deepcopy(persisted['result'])
    without_global_kpis.pop('global_kpis')
    job_metadata = scoring_jobs.get_scoring_job(repository, job_id)
    baseline_rows = list(csv.DictReader(StringIO(export_scoring_csv(
        job_metadata, without_global_kpis, 'scoring', 'expanded', environment='all',
    ))))
    score_signature = lambda rows: sorted(
        (row['environment'], row['region'], row['city'], row['campaign'],
         row['kpi_code'], row['row_type'], row['score_points'])
        for row in rows
    )
    assert score_signature(exported_rows) == score_signature(baseline_rows)


def test_scoring_ppt_export_normalizes_legacy_gap_without_mutating_saved_result(scoring_api, monkeypatch):
    client = scoring_api['client']
    response = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'],
        'aggregation_levels': ['Operator'],
        'nr_mode': 'NSA',
        'baseline_operator': 'EE',
    })
    assert response.status_code == 200, response.text
    job_id = response.json()['job']['id']
    scoring_jobs.run_scoring_job(scoring_api['repository'], job_id)
    _mark_scoring_result_as_legacy(scoring_api['repository'], job_id)

    from src.modules import scoring_exports
    captured = {}

    def capture_export(job, result, *_args, **kwargs):
        captured['result'] = deepcopy(result)
        captured['gap_layout'] = kwargs.get('gap_layout')
        captured['show_gap_values'] = kwargs.get('show_gap_values')
        return b'ppt'

    monkeypatch.setattr(scoring_exports, 'export_scoring_powerpoint', capture_export)
    exported = client.get(f'/scoring/jobs/{job_id}/export/ppt')
    assert exported.status_code == 200
    assert exported.content == b'ppt'
    assert captured['gap_layout'] == 'end'
    assert captured['show_gap_values'] is True
    assert captured['result']['gap'][0]['gap_points'] == -0.6
    assert captured['result']['gap_direction'] == 'operator_minus_reference'

    adjacent = client.get(f'/scoring/jobs/{job_id}/export/ppt?gap_layout=adjacent')
    assert adjacent.status_code == 200
    assert captured['gap_layout'] == 'adjacent'

    hidden_gaps = client.get(f'/scoring/jobs/{job_id}/export/ppt?show_gap_values=false')
    assert hidden_gaps.status_code == 200
    assert captured['show_gap_values'] is False

    stored = scoring_jobs.get_scoring_job(scoring_api['repository'], job_id, include_result=True)['result']
    assert stored['gap'][0]['gap_points'] == 0.6
    assert 'gap_direction' not in stored


def test_scoring_dataset_recalculate_and_automatic_queue_hook(scoring_api):
    client = scoring_api['client']
    dataset_id = scoring_api['dataset_id']

    response = client.post(f'/scoring/datasets/{dataset_id}/recalculate', follow_redirects=False)
    assert response.status_code == 303
    job_id = response.headers['location'].split('job_id=', 1)[1]
    job = scoring_jobs.get_scoring_job(scoring_api['repository'], int(job_id))
    assert job['status'] == 'queued'
    assert job['aggregation_levels'] == ['Operator']
    assert len(scoring_api['submitted']) == 1

    scoring_api['submitted'].clear()
    second_dataset_id = _add_ready_cdr(
        scoring_api['repository'], scoring_api['tmp_path'], name='UK_Q2_2026_NSA_Voice.csv', kind='voice',
    )
    app_module.queue_dataset_scoring(scoring_api['repository'], second_dataset_id, 'admin')
    queued = scoring_jobs.list_scoring_jobs(scoring_api['repository'])[0]
    assert queued['status'] == 'queued'
    assert second_dataset_id in queued['dataset_ids']
    assert len(queued['dataset_ids']) == 3
    assert queued['aggregation_levels'] == ['Operator']
    assert queued['baseline_operator'] == 'EE'
    assert len(scoring_api['submitted']) == 1
    assert scoring_api['submitted'][0][1] is scoring_jobs.run_scoring_job


def test_real_cdr_ingestion_automatically_scores_kpi_and_reuses_cache(client):
    _login(client)
    workspace = app_module.active_workspace
    configured_repository = Repository(workspace.database_path, app_module.repository.global_db_path,
                                       app_module.workspace_registry.registry_path)
    configured_repository.replace_scoring_configuration(scoring_configuration())
    uploads = [
        (
            'tiny_2026-Q2_NSA_data.csv',
            'data',
            'Campaign,G_Level_1,G_Level_2,Operator,Test_Name,Test_Result,Transfer_Duration\n'
            'UK_Q2_2026,Drive,City,EE,FDFS DL,Completed,10\n'
            'UK_Q2_2026,Drive,City,O2,FDFS DL,Completed,20\n'
            'UK_Q2_2026,Drive,City,O2,FDFS DL,Failed,30\n',
        ),
        (
            'tiny_2026-Q2_NSA_voice.csv',
            'voice',
            'Campaign,G_Level_1,G_Level_2,Operator,Session_Type,Call_Status,Call_Setup_Time\n'
            'UK_Q2_2026,Drive,City,EE,Voice,Completed,2\n',
        ),
        (
            'tiny_2026-Q2_NSA_speech.csv',
            'speech',
            'Campaign,G_Level_1,G_Level_2,Operator,Session_Type,LQ\n'
            'UK_Q2_2026,Drive,City,EE,Voice,4.2\n',
        ),
    ]
    for filename, kind, content in uploads:
        upload = client.post(
            '/datasets-analysis/upload',
            data={'dataset_kinds': kind, 'nr_modes': 'NSA'},
            files={'dataset_files': (filename, BytesIO(content.encode('utf-8')), 'text/csv')},
            follow_redirects=False,
        )
        assert upload.status_code == 303, upload.text

    workspace = app_module.active_workspace
    repository = Repository(
        workspace.database_path,
        global_db_path=app_module.repository.db_path,
        workspace_registry_db_path=app_module.workspace_registry.registry_path,
    )
    deadline = time.monotonic() + 10
    datasets_by_name = {}
    while time.monotonic() < deadline:
        datasets_by_name = {row['file_name']: row for row in repository.list_datasets()}
        if all(
            filename in datasets_by_name and datasets_by_name[filename]['status'] == 'ready'
            for filename, _kind, _content in uploads
        ):
            break
        time.sleep(0.05)
    selected_datasets = [datasets_by_name[filename] for filename, _kind, _content in uploads]
    assert all(dataset['status'] == 'ready' for dataset in selected_datasets)
    dataset_id = int(datasets_by_name[uploads[0][0]]['id'])
    expected_dataset_ids = [int(dataset['id']) for dataset in selected_datasets]

    deadline = time.monotonic() + 10
    automatic_job = None
    while time.monotonic() < deadline:
        response = client.get('/api/scoring/jobs')
        assert response.status_code == 200, response.text
        automatic_job = next((
            job for job in response.json()['jobs']
            if set(job['dataset_ids']) == set(expected_dataset_ids)
            and job['aggregation_levels'] == ['Operator']
        ), None)
        if automatic_job and automatic_job['status'] in {'completed', 'failed'}:
            break
        time.sleep(0.05)

    assert automatic_job is not None, 'The completed CDR did not queue its default scoring job.'
    assert automatic_job['status'] == 'completed', automatic_job
    result_response = client.get(f"/api/scoring/jobs/{automatic_job['id']}")
    assert result_response.status_code == 200, result_response.text
    result = result_response.json()

    c17_rows = [row for row in result['scoring'] if row['kpi_code'] == 'K12']
    assert {row['operator']: row['value'] for row in c17_rows} == {'EE': 100.0, 'O2': 50.0}
    assert any('Incomplete KPI or environment coverage' in warning for warning in result['warnings'])

    csv_response = client.get(f"/scoring/jobs/{automatic_job['id']}/export/scoring")
    assert csv_response.status_code == 200, csv_response.text
    exported_rows = list(csv.DictReader(StringIO(csv_response.content.decode('utf-8-sig'))))
    exported_c17 = {row['operator']: float(row['value']) for row in exported_rows if row['kpi_code'] == 'K12'}
    assert exported_c17 == {'EE': 100.0, 'O2': 50.0}

    cached = client.post('/api/scoring/jobs', json={
        'dataset_ids': expected_dataset_ids, 'aggregation_levels': ['Operator'], 'nr_mode': 'NSA',
        'baseline_operator': 'EE',
    })
    assert cached.status_code == 200, cached.text
    assert cached.json()['cached'] is True
    assert cached.json()['job']['id'] == automatic_job['id']


@pytest.mark.parametrize('job_status', ['queued', 'completed', 'failed'])
def test_scoring_delete_api_removes_job_and_recompute_uses_new_id(scoring_api, job_status):
    client = scoring_api['client']
    payload = {'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator'], 'nr_mode': 'NSA'}
    response = client.post('/api/scoring/jobs', json=payload)
    assert response.status_code == 200, response.text
    job_id = response.json()['job']['id']
    repository = scoring_api['repository']
    if job_status == 'completed':
        assert scoring_jobs.run_scoring_job(repository, job_id)['status'] == 'completed'
        assert client.get(f'/api/scoring/jobs/{job_id}').json()['scoring']
    elif job_status == 'failed':
        scoring_jobs._update_scoring_job(repository, job_id, status='failed', last_error='Test failure')
    response = client.delete(f'/api/scoring/jobs/{job_id}')
    assert response.status_code == 200, response.text
    assert response.json() == {'deleted': job_id}
    assert client.get(f'/api/scoring/jobs/{job_id}').status_code == 404
    assert client.get(f'/scoring/jobs/{job_id}/export/scoring').status_code == 404
    assert client.delete(f'/api/scoring/jobs/{job_id}').status_code == 404
    assert client.get('/api/scoring/jobs').json()['jobs'] == []
    assert repository.get_dataset(scoring_api['dataset_id']) is not None
    recreated = client.post('/api/scoring/jobs', json=payload)
    assert recreated.status_code == 200, recreated.text
    assert recreated.json()['cached'] is False
    assert recreated.json()['job']['id'] != job_id


def test_scoring_delete_requires_login_and_workspace_access(scoring_api, monkeypatch):
    client = scoring_api['client']
    created = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator'], 'nr_mode': 'NSA',
    })
    job_id = created.json()['job']['id']
    client.get('/logout', follow_redirects=False)
    assert client.delete(f'/api/scoring/jobs/{job_id}').status_code == 401
    assert scoring_jobs.get_scoring_job(scoring_api['repository'], job_id) is not None
    _login(client)
    monkeypatch.setattr(app_module.repository, 'user_has_workspace_access', lambda username, workspace_id: False)
    denied = client.delete(f'/api/scoring/jobs/{job_id}')
    assert denied.status_code == 403
    assert scoring_jobs.get_scoring_job(scoring_api['repository'], job_id) is not None
