"""GAP comparisons require a reference in the same selected hierarchy context."""
import csv
from copy import deepcopy
from io import BytesIO, StringIO

import pandas as pd
import pytest
from pptx import Presentation

from src.modules.scoring import calculate_scoring
from src.modules.scoring_exports import export_scoring_csv, export_scoring_powerpoint
from src.modules.scoring_views import build_scoring_views
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_exports import TEMPLATE


LEVELS = ['Operator', 'Vendor', 'Campaign']


def _configuration():
    configuration = scoring_configuration()
    configuration['metrics'] = [
        deepcopy(metric) for metric in configuration['metrics'] if metric['code'] == 'K1'
    ]
    configuration['gap_priority'] = ['K1']
    return configuration


def _source_rows(environments):
    rows = []
    observations = [
        ('2026-Q2', 'EE', 'Nokia', 'Completed'),
        ('2026-Q2', 'O2 UK', 'Nokia', 'Failed'),
        ('2026-Q2', 'O2 UK', 'Ericsson', 'Failed'),
        ('2026-Q3', 'EE', 'Nokia', 'Completed'),
        ('2026-Q3', 'O2 UK', 'Nokia', 'Failed'),
        ('2026-Q4', 'O2 UK', 'Nokia', 'Failed'),
    ]
    for environment in environments:
        for campaign, operator, vendor, status in observations:
            rows.append({
                'Campaign': campaign,
                'Operator': operator,
                'Vendor': vendor,
                'G_Level_1': 'Drive',
                'G_Level_2': 'City' if environment == 'City' else 'Connectionroad',
                'Session_Type': 'CALL',
                'Call_Status': status,
            })
    return rows


def _calculate(rows, *, hierarchical=True):
    configuration = _configuration()
    result = calculate_scoring(
        {'voice': pd.DataFrame(rows)}, LEVELS, baseline_operator='EE', configuration=configuration,
    )
    job = {
        'levels': LEVELS if hierarchical else ['Operator'],
        'baseline_operator': 'EE',
        'configuration': configuration,
    }
    if hierarchical:
        job.update({'aggregation_contract_version': 2, 'aggregation_levels': LEVELS})
    views = build_scoring_views(job, result)
    return result, views


def _gap_matrix(views, environment):
    return next(
        matrix for matrix in views['hierarchy_gap_tables']
        if matrix['context']['environment'] == environment
    )


def _leaf(matrix, operator, vendor, campaign):
    for column in matrix['hierarchy_columns']:
        path = {part['level']: part.get('value') for part in column['path']}
        if (path.get('Operator'), path.get('Vendor'), path.get('Campaign')) == (operator, vendor, campaign):
            return column['id']
    raise AssertionError(f'Missing hierarchy leaf for {operator}/{vendor}/{campaign}.')


def _kpi_row(matrix, code='K1'):
    return next(row for row in matrix['rows'] if row['kpi_code'] == code)


def _summary_for(views, environment, vendor='Nokia', campaign='2026-Q2'):
    return next(
        table for table in views['gap_summary_tables']
        if table['context']['environment'] == environment
        and table['context'].get('vendor') == vendor
        and table['context'].get('campaign') == campaign
    )


def test_vendor_and_campaign_levels_match_reference_only_within_the_same_context():
    result, views = _calculate(_source_rows(('City', 'Connectionroad')))

    matched_engine_gaps = [
        row for row in result['gap']
        if row['kpi_code'] == 'K1' and row['operator'] == 'O2 UK' and row['vendor'] == 'Nokia'
    ]
    assert {row['campaign'] for row in matched_engine_gaps} == {'2026-Q2', '2026-Q3'}
    assert all(row['gap_points'] < 0 for row in matched_engine_gaps)
    assert not any(
        row['kpi_code'] == 'K1' and row['operator'] == 'O2 UK'
        and (row['vendor'], row['campaign']) in {('Ericsson', '2026-Q2'), ('Nokia', '2026-Q4')}
        for row in result['gap']
    )

    combined = _gap_matrix(views, 'Combined')
    k1 = _kpi_row(combined)
    for campaign in ('2026-Q2', '2026-Q3'):
        same_vendor_leaf = _leaf(combined, 'O2 UK', 'Nokia', campaign)
        assert k1['gaps'][same_vendor_leaf] < 0
        assert k1['gap_partial'][same_vendor_leaf] is False
        assert k1['gap_environments'][same_vendor_leaf] == ['DriveCity', 'DriveConnectionroad']
    different_vendor_leaf = _leaf(combined, 'O2 UK', 'Ericsson', '2026-Q2')
    missing_reference_leaf = _leaf(combined, 'O2 UK', 'Nokia', '2026-Q4')
    assert k1['gaps'][different_vendor_leaf] is None
    assert k1['gaps'][missing_reference_leaf] is None


def test_hierarchy_partial_gap_uses_only_the_intersecting_environment():
    rows = [
        row for row in _source_rows(('City', 'Connectionroad'))
        if row['Campaign'] == '2026-Q2' and row['Vendor'] == 'Nokia'
        and (row['Operator'] == 'EE' or (row['Operator'] == 'O2 UK' and row['G_Level_2'] == 'City'))
    ]
    _, views = _calculate(rows)
    city_matrix = _gap_matrix(views, 'DriveCity')
    combined = _gap_matrix(views, 'Combined')
    city_leaf = _leaf(city_matrix, 'O2 UK', 'Nokia', '2026-Q2')
    combined_leaf = _leaf(combined, 'O2 UK', 'Nokia', '2026-Q2')
    row = _kpi_row(combined)
    assert row['gaps'][combined_leaf] == pytest.approx(_kpi_row(city_matrix)['gaps'][city_leaf])
    assert row['gap_partial'][combined_leaf] is True
    assert row['gap_environments'][combined_leaf] == ['DriveCity']


def _city_rows_with_one_road_series():
    """City rows, plus road rows for one other series: the road has results in some series, so nothing is scaled."""
    road = [row for row in _source_rows(('Connectionroad',))
            if row['Campaign'] == '2026-Q4' and row['Operator'] == 'O2 UK']
    return _source_rows(('City',)) + road


def test_missing_road_keeps_city_gap_and_combined_marks_common_environment_gap_partial():
    city_result, city_views = _calculate(_city_rows_with_one_road_series())
    city_matrix = _gap_matrix(city_views, 'DriveCity')
    city_leaf = _leaf(city_matrix, 'O2 UK', 'Nokia', '2026-Q2')
    assert _kpi_row(city_matrix)['gaps'][city_leaf] < 0

    combined = _gap_matrix(city_views, 'Combined')
    combined_leaf = _leaf(combined, 'O2 UK', 'Nokia', '2026-Q2')
    combined_row = _kpi_row(combined)
    assert combined_row['gaps'][combined_leaf] == pytest.approx(
        _kpi_row(city_matrix)['gaps'][city_leaf],
    )
    assert combined_row['gap_partial'][combined_leaf] is True
    assert combined_row['gap_environments'][combined_leaf] == ['DriveCity']
    assert '* marks a partial GAP calculated only from common available environment contributions (DriveCity).' in combined['note']

    scalar_result, scalar_views = _calculate(_city_rows_with_one_road_series(), hierarchical=False)
    scalar = _summary_for(scalar_views, 'Combined')
    scalar_row = _kpi_row(scalar)
    scalar_city = _kpi_row(_summary_for(scalar_views, 'DriveCity'))
    assert scalar_row['gaps']['O2 UK'] == pytest.approx(scalar_city['gaps']['O2 UK'])
    assert scalar_row['gap_partial']['O2 UK'] is True
    assert scalar_row['gap_environments']['O2 UK'] == ['DriveCity']

    scalar_export_job = {
        'aggregation_contract_version': 1,
        'aggregation_levels': ['Operator'], 'baseline_operator': 'EE',
        'configuration': scalar_result['configuration'],
    }
    csv_rows = list(csv.DictReader(StringIO(export_scoring_csv(
        scalar_export_job, scalar_result, 'gap', 'expanded', environment='all',
    ))))
    csv_gap = next(
        row for row in csv_rows
        if row['environment'] == 'Combined' and row['kpi_code'] == 'K1'
        and row['vendor'] == 'Nokia' and row['campaign'] == '2026-Q2'
    )
    assert float(csv_gap['gap_points']) == pytest.approx(scalar_row['gaps']['O2 UK'])
    assert csv_gap['gap_partial'] == 'True'
    assert csv_gap['gap_environments'] == '["DriveCity"]'

    presentation = Presentation(BytesIO(export_scoring_powerpoint(
        scalar_export_job, scalar_result, TEMPLATE, environment='all',
    )))
    combined_gap_tables = [
        shape.table
        for slide in presentation.slides
        if slide.shapes.title.text == 'GAP Analysis — All vs EE\nAll Environments'
        for shape in slide.shapes if shape.has_table
    ]
    assert combined_gap_tables
    expected_partial_gap = f"{scalar_row['gaps']['O2 UK']:.2f}*"
    assert any(
        cell.text == expected_partial_gap
        for table in combined_gap_tables for row in table.rows for cell in row.cells
    )

    without_reference = [row for row in _source_rows(('City', 'Connectionroad')) if row['Operator'] != 'EE']
    filtered_result, filtered_views = _calculate(without_reference)
    assert not any(row['kpi_code'] == 'K1' for row in filtered_result['gap'])
    filtered_combined = _gap_matrix(filtered_views, 'Combined')
    filtered_leaf = _leaf(filtered_combined, 'O2 UK', 'Nokia', '2026-Q2')
    assert _kpi_row(filtered_combined)['gaps'][filtered_leaf] is None
    assert _kpi_row(filtered_combined)['gap_partial'][filtered_leaf] is False
    assert _kpi_row(filtered_combined)['gap_environments'][filtered_leaf] == []


def test_scalar_combined_gap_requires_common_environment_coverage():
    rows = [
        row for row in _source_rows(('City', 'Connectionroad'))
        if row['Campaign'] == '2026-Q2' and row['Vendor'] == 'Nokia'
        and (row['Operator'] == 'EE' or (row['Operator'] == 'O2 UK' and row['G_Level_2'] == 'City'))
    ]
    _, views = _calculate(rows, hierarchical=False)
    combined = _kpi_row(_summary_for(views, 'Combined'))
    city = _kpi_row(_summary_for(views, 'DriveCity'))
    assert combined['gaps']['O2 UK'] == pytest.approx(city['gaps']['O2 UK'])
    assert combined['gap_partial']['O2 UK'] is True
    assert combined['gap_environments']['O2 UK'] == ['DriveCity']

    rows = _source_rows(('City', 'Connectionroad'))
    rows = [
        row for row in rows
        if row['Campaign'] == '2026-Q2' and row['Vendor'] == 'Nokia'
        and ((row['Operator'] == 'EE' and row['G_Level_2'] == 'City')
             or (row['Operator'] == 'O2 UK' and row['G_Level_2'] == 'Connectionroad'))
    ]
    _, views = _calculate(rows, hierarchical=False)
    combined = _summary_for(views, 'Combined')
    row = _kpi_row(combined)
    assert row['gaps']['O2 UK'] is None
    assert row['gap_partial']['O2 UK'] is False
    assert row['gap_environments']['O2 UK'] == []


def test_road_without_results_in_any_series_scales_the_combined_gap():
    result, views = _calculate(_source_rows(('City',)))
    assert result['environment_scaling']['scaled_environments'] == ['DriveConnectionroad']
    city = _gap_matrix(views, 'DriveCity')
    combined = _gap_matrix(views, 'Combined')
    city_leaf = _leaf(city, 'O2 UK', 'Nokia', '2026-Q2')
    combined_leaf = _leaf(combined, 'O2 UK', 'Nokia', '2026-Q2')
    metric = next(item for item in result['configuration']['metrics'] if item['code'] == 'K1')
    factor = sum(context['max_points'] for name, context in metric['contexts'].items()
                 if name in {'DriveCity', 'DriveConnectionroad'}) / metric['contexts']['DriveCity']['max_points']
    # The Combined GAP is the City GAP scaled to the full maximum, and it is not partial.
    assert _kpi_row(combined)['gaps'][combined_leaf] == pytest.approx(_kpi_row(city)['gaps'][city_leaf] * factor)
    assert _kpi_row(combined)['gap_partial'][combined_leaf] is False
