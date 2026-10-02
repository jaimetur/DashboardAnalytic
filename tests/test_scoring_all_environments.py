"""All Environments exports keep the full configured maximum when data is incomplete."""
from io import BytesIO

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from src.modules.scoring_exports import export_scoring_powerpoint
from src.modules.scoring_views import build_scoring_views
from tests.test_scoring_exports import (
    TEMPLATE,
    _comparison_matrices,
    _nested_shapes,
    _result,
    _slide_text,
)
from tests.test_scoring_table_exports import _job, _multi_environment_result


def _presentation(result):
    return Presentation(BytesIO(export_scoring_powerpoint(_job(), result, TEMPLATE)))


def _environment_matrix_slide(presentation, environment):
    for slide in presentation.slides:
        title = slide.shapes.title.text.split('\n')[0] if slide.shapes.title else ''
        if title != 'Scoring Tables — Breakdown' or environment not in _slide_text(slide):
            continue
        matrix = next(
            shape.table for shape in slide.shapes
            if shape.has_table and shape.table.cell(0, 0).text.strip() in {'CATEGORY', 'NETCHECK KPIs'}
        )
        return slide, matrix
    raise AssertionError(f'No Scoring Table slide found for {environment}.')


def _chart_for_environment(presentation, title, environment, chart_type):
    for slide in presentation.slides:
        slide_title = slide.shapes.title.text.split('\n')[0] if slide.shapes.title else ''
        if slide_title != title or environment not in _slide_text(slide):
            continue
        for shape in _nested_shapes(slide.shapes):
            if shape.has_chart and shape.chart.chart_type == chart_type:
                return shape.chart
    raise AssertionError(f'No {title} chart found for {environment}.')


def _total_row(table):
    return len(table.rows) - 1


def _assert_combined_kpi_allocations(combined, configuration):
    environments = configuration['scope']['environments']
    active_environments = [
        name for name, details in environments.items()
        if details['total_points'] > 0
    ]
    expected_by_code = {
        metric['code']: sum(
            metric['contexts'][environment]['max_points'] for environment in active_environments
        )
        for metric in configuration['metrics']
    }
    actual_by_code = {row['kpi_code']: row['max_points'] for row in combined['rows']}
    assert set(actual_by_code) == set(expected_by_code)
    for code, maximum in expected_by_code.items():
        assert actual_by_code[code] == pytest.approx(maximum)

    expected_by_category = {}
    for metric in configuration['metrics']:
        expected_by_category[metric['category']] = (
            expected_by_category.get(metric['category'], 0.0) + expected_by_code[metric['code']]
        )
    category_rows = [row for row in combined['expanded_rows'] if row['row_type'] == 'category']
    assert {row['category'] for row in category_rows} == set(expected_by_category)
    for row in category_rows:
        assert row['max_points'] == pytest.approx(expected_by_category[row['category']])


def test_all_environments_ppt_uses_full_allocation_in_tables_and_charts():
    result = _multi_environment_result()
    job = _job()
    views = build_scoring_views(job, result)
    combined = next(table for table in views['score_tables'] if table['context']['environment'] == 'Combined')
    city = next(table for table in views['score_tables'] if table['context']['environment'] == 'DriveCity')
    road = next(table for table in views['score_tables'] if table['context']['environment'] == 'DriveConnectionroad')

    assert combined['total']['max_points'] == pytest.approx(1000)
    assert city['total']['max_points'] == pytest.approx(650)
    assert road['total']['max_points'] == pytest.approx(350)
    assert combined['total']['weight_percent'] == pytest.approx(100)
    _assert_combined_kpi_allocations(combined, result['configuration'])
    for operator in combined['operators']:
        assert combined['total']['values'][operator]['points'] == pytest.approx(
            city['total']['values'][operator]['points'] + road['total']['values'][operator]['points'],
        )

    presentation = _presentation(result)
    matrices = _comparison_matrices(presentation)
    assert len(matrices) == 6
    matrices = [matrices[index] for index in (1, 3, 5)]
    assert [matrices[index].cell(_total_row(matrices[index]), 4).text for index in range(3)] == [
        '1000.00', '650.00', '350.00',
    ]

    donut = _chart_for_environment(
        presentation, 'Best Network Scoring per Service', 'All Environments', XL_CHART_TYPE.DOUGHNUT,
    )
    assert sum(float(value) for value in donut.series[0].values) == pytest.approx(1000)

    best_network = _chart_for_environment(
        presentation, 'Best Network Scoring per Service', 'All Environments', XL_CHART_TYPE.COLUMN_STACKED,
    )
    total_series = next(series for series in best_network.series if series.name == 'Total')
    for operator, points in zip(combined['operators'], total_series.values):
        assert float(points) == pytest.approx(combined['total']['values'][operator]['points'])

    category_chart = _chart_for_environment(
        presentation, 'Scoring per Category', 'All Environments', XL_CHART_TYPE.COLUMN_CLUSTERED,
    )
    for series in category_chart.series:
        assert sum(float(value) for value in series.values if value is not None) == pytest.approx(
            combined['total']['values'][series.name]['points'],
        )


def test_all_environments_ppt_keeps_full_maximum_and_marks_missing_road_incomplete():
    result = _result(operators=('EE', 'O2 UK'), warnings=[])
    job = _job()
    views = build_scoring_views(job, result)
    combined = next(table for table in views['score_tables'] if table['context']['environment'] == 'Combined')
    city = next(table for table in views['score_tables'] if table['context']['environment'] == 'DriveCity')

    assert combined['total']['max_points'] == pytest.approx(1000)
    assert combined['total']['weight_percent'] == pytest.approx(100)
    _assert_combined_kpi_allocations(combined, result['configuration'])
    for operator in combined['operators']:
        assert combined['total']['values'][operator]['complete'] is False
        assert combined['total']['values'][operator]['points'] == pytest.approx(city['total']['values'][operator]['points'])

    presentation = _presentation(result)
    matrices = _comparison_matrices(presentation)
    assert len(matrices) == 4
    combined_slide, combined_matrix = _environment_matrix_slide(presentation, 'All Environments')
    assert combined_matrix.cell(_total_row(combined_matrix), 4).text == '1000.00'
    assert all(
        combined_matrix.cell(_total_row(combined_matrix), column).text.endswith('*')
        for column in range(5, 5 + len(combined['operators']))
    )
    assert 'All Environments is incomplete because weighted environments are missing: DriveConnectionroad.' in _slide_text(combined_slide)

    donut = _chart_for_environment(
        presentation, 'Best Network Scoring per Service', 'All Environments', XL_CHART_TYPE.DOUGHNUT,
    )
    assert sum(float(value) for value in donut.series[0].values) == pytest.approx(1000)
