"""Scoring reports: scenarios with their own filters, aggregation and content options."""
from __future__ import annotations

from io import BytesIO
import json

import pytest
from pptx import Presentation

from src.modules.scoring_exports import export_scoring_report
from src.modules.scoring_reports import (
    default_compared_operators, default_report_configuration, default_scenario, delete_named_configuration,
    import_report_configurations, load_report_state, normalize_report_configuration, remember_last_configuration,
    report_configurations_document, save_named_configuration,
)
from tests.scoring_synthetic import synthetic_result
from tests.test_scoring_api import scoring_api  # noqa: F401  (fixture)
from tests.test_scoring_exports import TEMPLATE

OPERATORS = ['EE', 'O2 UK', 'Three UK', 'Vodafone UK']


class _StateRepository:
    def __init__(self):
        self.state = {}

    def get_workspace_state(self, key):
        return self.state.get(key)

    def set_workspace_state(self, key, value):
        self.state[key] = value


def _titles(content):
    presentation = Presentation(BytesIO(content))
    return [slide.shapes.title.text_frame.text.replace('\n', ' | ') if slide.shapes.title is not None else ''
            for slide in presentation.slides]


def test_default_scenario_follows_the_requested_defaults():
    scenario = default_scenario(context_filters={'City': ['London']}, aggregation_levels=['City', 'Operator'],
                                operators=['VF_UK', '3', 'EE', 'O2'])
    # Campaign is part of every default scenario.
    assert scenario['aggregation_levels'] == ['Operator', 'City', 'Campaign']
    for kind in ('best_network', 'most_reliable'):
        options = scenario['scorings'][kind]
        assert options['enabled'] is True and options['environments'] == 'all'
        assert options['charts'] == {'service': True, 'category': True, 'breakdown': True,
                                     'location_cards': True, 'trend': True}
        assert options['tables']['kpi_values'] is False and options['tables']['gap_values'] is False
        assert options['gap']['operators'] == ['VF_UK', '3']
    assert default_compared_operators(['Vodafone UK', 'Three UK', 'EE']) == ['Vodafone UK', 'Three UK']


def test_report_configuration_is_normalized():
    configuration = normalize_report_configuration({'scenarios': [
        {'name': ' London ', 'context_filters': {'City': ['London', 'London'], 'Unknown': ['x']},
         'aggregation_levels': ['Campaign', 'Bogus'], 'scorings': {'most_reliable': {'enabled': False}}},
        {},
    ]})
    first, second = configuration['scenarios']
    assert first['name'] == 'London' and first['context_filters'] == {'City': ['London']}
    assert first['aggregation_levels'] == ['Operator', 'Campaign']
    assert first['scorings']['most_reliable']['enabled'] is False
    assert first['scorings']['best_network']['tables']['summary'] is True
    assert second['name'] == 'Scenario 2'
    with pytest.raises(ValueError):
        normalize_report_configuration({'scenarios': []})
    with pytest.raises(ValueError, match='at least one scenario'):
        normalize_report_configuration({'scenarios': [{'scorings': {
            'best_network': {'enabled': False}, 'most_reliable': {'enabled': False}}}]})


def test_configurations_are_saved_remembered_and_exported():
    repository = _StateRepository()
    configuration = default_report_configuration(operators=OPERATORS)
    save_named_configuration(repository, 'National', configuration)
    save_named_configuration(repository, 'national', configuration)  # same name, any case: replaced
    remember_last_configuration(repository, {'scenarios': [default_scenario(name='London')]})
    state = load_report_state(repository)
    assert [item['name'] for item in state['configurations']] == ['national']
    assert state['last']['scenarios'][0]['name'] == 'London'
    document = report_configurations_document(state)
    other = _StateRepository()
    import_report_configurations(other, json.loads(json.dumps(document)))
    assert load_report_state(other)['configurations'][0]['configuration'] == state['configurations'][0]['configuration']
    delete_named_configuration(repository, 'NATIONAL')
    assert load_report_state(repository)['configurations'] == []
    with pytest.raises(ValueError):
        import_report_configurations(other, {'format': 'something-else'})


def test_report_document_has_one_section_per_scenario_and_scoring():
    national_job, national = synthetic_result(levels=('Operator',))
    cities_job, cities = synthetic_result(levels=('Operator', 'City'))
    national_scenario = default_scenario(name='National', operators=OPERATORS)
    cities_scenario = default_scenario(name='Main Cities', aggregation_levels=['Operator', 'City'], operators=OPERATORS)
    cities_scenario['scorings']['most_reliable']['enabled'] = False
    cities_scenario['scorings']['best_network']['charts']['breakdown'] = False
    cities_scenario['scorings']['best_network']['gap']['individual'] = False
    content = export_scoring_report([
        {'scenario': normalize_report_configuration({'scenarios': [national_scenario]})['scenarios'][0],
         'job': national_job, 'result': national},
        {'scenario': normalize_report_configuration({'scenarios': [cities_scenario]})['scenarios'][0],
         'job': cities_job, 'result': cities},
    ], TEMPLATE)
    titles = _titles(content)
    first = [title.split(' | ')[0] for title in titles]
    assert first[0] == 'Scoring & GAP Analysis'
    cover = Presentation(BytesIO(content)).slides[0]
    assert 'Scenarios: National, Main Cities' in ' '.join(shape.text_frame.text for shape in cover.shapes
                                                          if shape.has_text_frame)
    assert first.count('National — Best Network Scoring') == 1
    assert first.count('National — Most Reliable Network Scoring') == 1
    assert first.count('Main Cities — Best Network Scoring') == 1
    assert 'Main Cities — Most Reliable Network Scoring' not in first
    main_cities = titles[first.index('Main Cities — Best Network Scoring'):]
    assert not any('per Category (Breakdown)' in title for title in main_cities)
    assert not any(title.startswith('GAP Analysis — Three UK vs EE') for title in main_cities)
    assert any(title.startswith('Best Network Scoring per City | Main Cities · Best Network') for title in main_cities)
    # All Environments only: no environment blocks, and no "All Environments" in the subtitles.
    assert not any('Drive - City' in title for title in titles)
    assert not any('All Environments' in title for title in titles)
    # Only Vodafone and Three are compared with EE.
    assert not any('O2 UK vs EE' in title for title in titles)
    assert any('Three UK vs EE' in title for title in titles)


def test_split_environments_and_kpi_values():
    job, result = synthetic_result(levels=('Operator',))
    scenario = default_scenario(name='National', operators=OPERATORS)
    options = scenario['scorings']['best_network']
    options['environments'] = 'split'
    options['tables']['kpi_values'] = True
    scenario['scorings']['most_reliable']['enabled'] = False
    content = export_scoring_report([{'scenario': normalize_report_configuration({'scenarios': [scenario]})['scenarios'][0],
                                      'job': job, 'result': result}], TEMPLATE)
    presentation = Presentation(BytesIO(content))
    titles = [slide.shapes.title.text_frame.text.replace('\n', ' | ') for slide in presentation.slides]
    assert any(title.startswith('Best Network — Drive - City') for title in titles)
    breakdown = next(slide for slide, title in zip(presentation.slides, titles)
                     if title.startswith('Scoring Tables — Breakdown'))
    table = next(shape.table for shape in breakdown.shapes if shape.has_table)
    headers = {table.cell(row, column).text for row in range(4) for column in range(len(table.columns))}
    assert 'Value' in headers and 'Value / Score' in headers


def test_reporting_job_scoring_entry_keeps_its_report():
    from src.modules.report_tasks import normalize_definition
    report = default_report_configuration(operators=OPERATORS)
    definition = normalize_definition({'scoring': [{'nr_mode': 'NSA', 'report': report}]})
    assert definition['scoring'][0]['report']['scenarios'][0]['name'] == 'National'
    # The report content is required: there is no implicit default report.
    with pytest.raises(ValueError, match='Configure the report content'):
        normalize_definition({'scoring': [{'nr_mode': 'NSA'}]})


def test_report_api_generates_scenarios_and_remembers_the_configuration(scoring_api, monkeypatch):
    from types import ModuleType

    import pandas as pd

    from src.modules import scoring as scoring_engine
    from src.modules import scoring_jobs
    from tests.test_scoring_most_reliable import _reliable_configuration

    client = scoring_api['client']
    repository = scoring_api['repository']
    repository.replace_scoring_configuration(_reliable_configuration())
    rows = [{'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': 'City',
             'Session_Type': 'CALL', 'Call_Status': status}
            for operator, status in (('EE', 'Completed'), ('O2 UK', 'Completed'), ('Three UK', 'Failed'))]
    engine = ModuleType('src.modules.scoring')
    engine.METHOD_VERSION = 'scoring-report-api-test'
    engine.required_input_columns = lambda kind, levels: [*levels, 'score']
    engine.calculate_scoring = lambda frames, levels, *, baseline_operator, configuration=None: (
        scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, levels, baseline_operator=baseline_operator,
                                         configuration=configuration))
    monkeypatch.setattr(scoring_jobs, '_scoring_engine', lambda: engine)
    created = client.post('/api/scoring/jobs', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator'],
        'nr_mode': 'NSA', 'baseline_operator': 'EE',
    })
    job_id = created.json()['job']['id']
    scoring_jobs.run_scoring_job(repository, job_id)

    national = default_scenario(name='National', operators=['Three UK'])
    leeds = default_scenario(name='Leeds', context_filters={'City': ['Leeds']}, operators=['Three UK'])
    leeds['scorings']['most_reliable']['enabled'] = False
    configuration = {'scenarios': [national, leeds]}
    saved = client.post('/api/scoring/report-configurations', json={'name': 'Weekly', 'configuration': configuration})
    assert saved.status_code == 200, saved.text
    assert [item['name'] for item in client.get('/api/scoring/report-configurations').json()['configurations']] == ['Weekly']

    # The report uses the CDRs, NR Mode, methodology and GAP reference of the Calculation panel.
    job = scoring_jobs.get_scoring_job(repository, job_id)
    selection = {key: job[key] for key in ('dataset_ids', 'nr_mode', 'baseline_operator', 'scoring_profile_id')}
    assert client.post('/api/scoring/report/ppt', json={'configuration': configuration}).status_code == 400
    response = client.post('/api/scoring/report/ppt', json={'configuration': configuration, **selection})
    assert response.status_code == 200, response.text
    assert 'National' in response.headers['content-disposition'] and 'Leeds' in response.headers['content-disposition']
    titles = [title.split(' | ')[0] for title in _titles(response.content)]
    assert 'National — Best Network Scoring' in titles and 'Leeds — Best Network Scoring' in titles
    assert 'Leeds — Most Reliable Network Scoring' not in titles
    assert client.get('/api/scoring/report-configurations').json()['last']['scenarios'][1]['name'] == 'Leeds'
    # The Leeds scenario has its own scoring job with its filter.
    jobs = scoring_jobs.list_scoring_jobs(repository)
    assert any((job.get('context_filters') or {}).get('City') == ['Leeds'] for job in jobs)

    exported = client.get('/api/scoring/report-configurations/export').json()
    assert exported['format'] == 'drivetest-analyzer-scoring-report-configurations'
    assert client.delete('/api/scoring/report-configurations?name=Weekly').json()['configurations'] == []
    imported = client.post('/api/scoring/report-configurations/import',
                           files={'package': ('reports.json', json.dumps(exported), 'application/json')})
    assert [item['name'] for item in imported.json()['configurations']] == ['Weekly']
    bad = client.post('/api/scoring/report/ppt', json={'configuration': {'scenarios': []}, **selection})
    assert bad.status_code == 400


def test_levels_with_a_single_value_are_left_out_of_report_scenarios(monkeypatch):
    import src.DriveTestAnalyzer as core
    from src.modules.scoring_views import single_value_levels

    one_campaign = {'scoring': [
        {'operator': 'EE', 'campaign': 'UK_Q2_2026', 'city': 'London'},
        {'operator': 'O2 UK', 'campaign': 'UK_Q2_2026', 'city': 'Leeds'},
    ]}
    assert single_value_levels(['Operator', 'Campaign', 'City'], one_campaign) == ['Campaign']

    calculated = []

    def create(_repository, _selected, levels, *_args, **_kwargs):
        calculated.append(list(levels))
        return {'id': len(calculated), 'status': 'completed', 'levels': list(levels)}, True

    monkeypatch.setattr(core, 'validate_complete_scoring_cdr_selection', lambda *_args, **_kwargs: [1])
    monkeypatch.setattr(core, 'create_scoring_job', create)
    monkeypatch.setattr(core, 'get_scoring_job', lambda _repository, job_id, include_result=False: {
        'id': job_id, 'status': 'completed', 'aggregation_levels': calculated[job_id - 1], 'result': one_campaign,
    })
    scenario = default_scenario(aggregation_levels=['Operator', 'City', 'Campaign'])
    job = core._scenario_scoring_job(None, {'dataset_ids': [1], 'nr_mode': 'NSA'}, scenario, 'super')
    # The Campaign level splits nothing: the scenario is calculated again without it.
    assert calculated == [['Operator', 'City', 'Campaign'], ['Operator', 'City']]
    assert job['aggregation_levels'] == ['Operator', 'City']


def test_report_configurations_remember_the_configuration_they_were_chosen_from():
    scenario = {'name': 'National', 'scorings': {'best_network': {'enabled': True}}}
    assert normalize_report_configuration({'scenarios': [scenario], 'name': ' Weekly '})['name'] == 'Weekly'
    assert 'name' not in normalize_report_configuration({'scenarios': [scenario]})


def test_report_api_remembers_the_configuration_when_the_document_is_requested(scoring_api, monkeypatch):
    import src.DriveTestAnalyzer as core

    def failing_build(*_args, **_kwargs):
        raise ValueError('The scoring of the scenario could not be calculated.')

    monkeypatch.setattr(core, 'build_scoring_report_document', failing_build)
    client = scoring_api['client']
    configuration = {'name': 'Weekly', 'scenarios': [default_scenario(name='London', context_filters={'City': ['London']})]}
    response = client.post('/api/scoring/report/ppt', json={
        'configuration': configuration, 'dataset_ids': scoring_api['complete_dataset_ids'], 'nr_mode': 'NSA',
    })
    assert response.status_code == 400
    # Remembered before the document is built, so also when it is not ready (or never is).
    last = client.get('/api/scoring/report-configurations').json()['last']
    assert last['name'] == 'Weekly' and last['scenarios'][0]['name'] == 'London'
