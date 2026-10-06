"""Non-Qualified Calls report: Executive Summary and Progress Status in PowerPoint and Word."""

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches as DocxInches
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from src.modules.network_insights_export import _content_frame, _docx_table, _slide, _table_slides

FILTER_LABELS = (
    ('datasets', 'CDRs'), ('service', 'Service'), ('campaign', 'Campaign'), ('operator', 'Operator'),
    ('operator_vendor', 'Operator_Vendor'), ('vendor', 'Vendor'), ('region', 'Region'), ('cluster', 'Cluster'),
    ('city', 'City'), ('technology', 'Technology'), ('test_name', 'Test Name'), ('result', 'Result'),
    ('failure_classification', 'Failure Classification'), ('failure_category', 'Failure Category'),
    ('status', 'Status'), ('team', 'Team'), ('assignee', 'Assignee'),
)
FLAG_LABELS = {'open_only': 'Open calls only', 'without_comments': 'Without comments'}
GRANULARITY_LABELS = {'week': 'Week', 'month': 'Month', 'quarter': 'Quarter', 'year': 'Year'}
RASPBERRY = RGBColor(0xB0, 0x23, 0x4F)


def selection_lines(filters: dict[str, Any], dataset_names: dict[str, str], granularity: str) -> list[str]:
    """Readable filters of the report; empty filters include every value."""
    lines = []
    for key, label in FILTER_LABELS:
        values = [str(value) for value in filters.get(key) or [] if str(value).strip()]
        if key == 'datasets':
            values = [dataset_names.get(value, value) for value in values]
        values = ['Unassigned' if value == '__unassigned__' else value for value in values]
        if values:
            lines.append(f'{label}: {", ".join(values)}')
    lines.extend(text for key, text in FLAG_LABELS.items() if filters.get(key))
    if filters.get('search'):
        lines.append(f"Search: {filters['search']}")
    if not lines:
        lines.append('Filters: every Non-Qualified Call of the workspace')
    lines.append(f'Progress periods: {GRANULARITY_LABELS.get(granularity, "Month")}')
    lines.append(f'Generated: {datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")}')
    return lines


def _share(count: int, total: int) -> str:
    return f'{count * 100 / total:.1f}%' if total else '0.0%'


def _days(value: Any) -> str:
    return '—' if value is None else f'{value:.1f}'


def summary_table(summary: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    total = summary['total']
    return ['Indicator', 'Calls', 'Share'], [
        ['Non-Qualified Calls', f'{total:,}', '100%' if total else '0%'],
        ['Open', f"{summary['open']:,}", _share(summary['open'], total)],
        ['Closed', f"{summary['closed']:,}", _share(summary['closed'], total)],
        ['With Team', f"{summary['with_team']:,}", _share(summary['with_team'], total)],
        ['Assigned', f"{summary['assigned']:,}", _share(summary['assigned'], total)],
        ['Commented', f"{summary['commented']:,}", _share(summary['commented'], total)],
    ]


def breakdown_table(breakdown: dict[str, Any], total: int) -> tuple[list[str], list[list[str]]]:
    empty = 'Unassigned' if breakdown['field'] == 'team' else 'Not classified'
    return [breakdown['label'].removeprefix('By ').strip(), 'Calls', 'Share'], [
        [item.get('label') or item['value'] or empty, f"{item['count']:,}", _share(item['count'], total)]
        for item in breakdown['items']
    ]


def progress_summary_table(summary: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    return ['Indicator', 'Value'], [
        ['Attended (followed up or commented)', f"{summary['attended']:,}"],
        ['Not attended yet', f"{summary['not_attended']:,}"],
        ['Open', f"{summary['open']:,}"],
        ['Closed', f"{summary['closed']:,}"],
        ['Closure rate', f"{summary['closure_rate']:.1f}%"],
        ['Comments', f"{summary['comments']:,}"],
        ['Follow-up changes', f"{summary['changes']:,}"],
        ['Contributors', f"{summary['contributors']:,}"],
        ['Average days from the call to its first follow-up', _days(summary['avg_days_to_first_follow_up'])],
        ['Average days from the first follow-up to closing', _days(summary['avg_days_to_close'])],
    ]


def timeline_table(progress: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    statuses = progress['statuses']
    columns = ['Period', 'Detected', 'Attended', 'Comments', *[f'→ {status}' for status in statuses],
               'Closed', 'Reopened', 'Open backlog', 'Avg days to close']
    rows = [[period['period'], str(period['detected']), str(period['attended']), str(period['comments']),
             *[str(period['statuses'].get(status, 0)) for status in statuses],
             str(period['closed']), str(period['reopened']), str(period['open_backlog']), _days(period['avg_days_to_close'])]
            for period in progress['periods']]
    return columns, rows


def workload_table(title: str, rows: list[dict[str, Any]]) -> tuple[list[str], list[list[str]]]:
    return [title, 'Open', 'Closed', 'Total'], [[row['name'], str(row['open']), str(row['closed']), str(row['total'])] for row in rows]


def activity_table(rows: list[dict[str, Any]]) -> tuple[list[str], list[list[str]]]:
    return ['User', 'Changes', 'Comments', 'Calls closed'], [
        [row['name'], str(row['changes']), str(row['comments']), str(row['closed'])] for row in rows
    ]


def report_tables(summary: dict[str, Any], breakdowns: list[dict[str, Any]], progress: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    """Every table of the report, in order: Executive Summary, then Progress Status."""
    tables = [('Executive Summary', *summary_table(summary))]
    tables += [(breakdown['label'], *breakdown_table(breakdown, summary['total'])) for breakdown in breakdowns]
    tables.append(('Progress Status', *progress_summary_table(progress['summary'])))
    tables.append((f"Progress per {GRANULARITY_LABELS.get(progress['granularity'], 'Month')}", *timeline_table(progress)))
    tables.append(('Age of Open Calls', ['Age since the call', 'Open calls'], [[item['value'], str(item['count'])] for item in progress['aging']]))
    tables.append(('Team Workload', *workload_table('Team', progress['teams'])))
    tables.append(('Assignee Workload', *workload_table('Assignee', progress['assignees'])))
    tables.append(('User Activity', *activity_table(progress['activity'])))
    return tables


def _text_slide(presentation: Presentation, title: str, lines: list[str]) -> None:
    slide = _slide(presentation, title)
    left, top, width, height = _content_frame(slide)
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.word_wrap = True
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, line in enumerate(lines):
        paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(15)


def _pie_slide(presentation: Presentation, title: str, distributions: list[dict[str, Any]]) -> None:
    """Up to three doughnut charts side by side, each with its own legend."""
    slide = _slide(presentation, title)
    left, top, width, height = _content_frame(slide)
    shown = [item for item in distributions if item['items']][:3]
    if not shown:
        return
    column = width / len(shown)
    for index, distribution in enumerate(shown):
        data = CategoryChartData()
        items = distribution['items'][:10]
        data.categories = [item['value'] for item in items]
        data.add_series('Calls', [item['count'] for item in items])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(left + column * index), Inches(top),
                                       Inches(column - 0.15), Inches(height), data)
        chart = frame.chart
        chart.has_title = True
        chart.chart_title.text_frame.text = distribution['label']
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(13)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(9)
        plot.data_labels.number_format = '0'
        plot.data_labels.number_format_is_linked = False
        for point_index, item in enumerate(items):
            color = str(item.get('color') or '').lstrip('#')
            if len(color) == 6:
                fill = plot.series[0].points[point_index].format.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor.from_string(color.upper())


def _timeline_chart_slide(presentation: Presentation, progress: dict[str, Any]) -> None:
    periods = progress['periods'][-24:]
    if not periods:
        return
    slide = _slide(presentation, f"Attended and Closed per {GRANULARITY_LABELS.get(progress['granularity'], 'Month')}",
                   'The latest 24 periods')
    left, top, width, height = _content_frame(slide)
    data = CategoryChartData()
    data.categories = [period['period'] for period in periods]
    for key, label in (('detected', 'Detected'), ('attended', 'Attended'), ('closed', 'Closed')):
        data.add_series(label, [period[key] for period in periods])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(left), Inches(top), Inches(width), Inches(height), data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    for series, color in zip(chart.series, ('8A6B76', 'E08A1E', '2E8B57'), strict=False):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.size = Pt(9)


def with_colors(distributions: list[dict[str, Any]], options: dict[str, Any]) -> list[dict[str, Any]]:
    colors = {('status', item['name']): item['color'] for item in options['statuses']}
    colors.update({('team', item['name']): item['color'] for item in options['teams']})
    return [{**distribution, 'items': [{**item, 'color': colors.get((distribution['field'], item['value']), '')}
                                       for item in distribution['items']]} for distribution in distributions]


def _cover_selection(cover, lines: list[str]) -> None:
    """The report selection (filters, progress periods and generation time) below the cover's line."""
    from src.modules.exports import cover_details, remove_empty_placeholders

    cover_details(cover, lines)
    remove_empty_placeholders(cover)


def report_panels(summary: dict[str, Any], breakdowns: list[dict[str, Any]], progress: dict[str, Any],
                  options: dict[str, Any]) -> list[tuple[str, Any]]:
    """The Executive Summary and Progress View drawn like the page, as (title, PNG) panels."""
    from src.modules.non_qualified_calls_visuals import executive_summary_panels, progress_view_panels

    from src.modules.non_qualified_calls_visuals import table_panels

    label = GRANULARITY_LABELS.get(progress['granularity'], 'Month')
    tables = [panel for title, columns, rows in detail_tables(progress) for panel in table_panels(title, columns, rows)]
    return [*executive_summary_panels(summary, breakdowns, options), *progress_view_panels(progress, options, label), *tables]


def detail_tables(progress: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    """The figures the panels summarise: every period of the timeline and the activity of each user."""
    label = GRANULARITY_LABELS.get(progress['granularity'], 'Month')
    return [(f'Progress View · Progress per {label} · Details', *timeline_table(progress)),
            ('Progress View · User Activity', *activity_table(progress['activity']))]


def export_powerpoint(destination: Path, lines: list[str], summary: dict[str, Any], breakdowns: list[dict[str, Any]],
                      progress: dict[str, Any], options: dict[str, Any], template: Path | None = None) -> Path:
    from PIL import Image

    from src.config import settings
    from src.modules.cdr_reporting import _named_slide_layout, _remove_all_slides, _set_structural_slide_text

    presentation = Presentation(template or settings.ppt_templates_dir / 'Template_CDR_analysis.pptx')
    _remove_all_slides(presentation)
    cover = presentation.slides.add_slide(_named_slide_layout(presentation, 'Title Page'))
    _set_structural_slide_text(cover, 'Non-Qualified Calls', 'Executive Summary and Progress Status')
    _cover_selection(cover, lines)
    for title, image in report_panels(summary, breakdowns, progress, options):
        slide = _slide(presentation, title)
        left, top, width, height = _content_frame(slide)
        with Image.open(image) as raster:
            ratio = raster.width / raster.height
        image.seek(0)
        picture_width = min(width, height * ratio)
        slide.shapes.add_picture(image, Inches(left + (width - picture_width) / 2), Inches(top),
                                 width=Inches(picture_width), height=Inches(picture_width / ratio))
    closing = _named_slide_layout(presentation, 'Black logo end slide')
    if closing is not None:
        _set_structural_slide_text(presentation.slides.add_slide(closing), '', '')
    presentation.save(destination)
    return destination


def export_word(destination: Path, lines: list[str], summary: dict[str, Any], breakdowns: list[dict[str, Any]],
                progress: dict[str, Any], options: dict[str, Any] | None = None) -> Path:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = DocxInches(11.69), DocxInches(8.27)
    section.left_margin = section.right_margin = DocxInches(0.6)
    section.top_margin = section.bottom_margin = DocxInches(0.6)
    document.add_heading('Dashboard Analytic · Non-Qualified Calls', level=0)
    document.add_paragraph('Executive Summary and Progress Status')
    for line in lines:
        document.add_paragraph(line, style='List Bullet')
    usable = (section.page_width - section.left_margin - section.right_margin) / 914400
    for title, image in report_panels(summary, breakdowns, progress, options or {'statuses': [], 'teams': []}):
        heading = document.add_heading(title, level=1)
        # Every section starts on its own page, as each one has its own slide in PowerPoint.
        heading.paragraph_format.page_break_before = True
        document.add_picture(image, width=DocxInches(usable))
    document.save(destination)
    return destination
