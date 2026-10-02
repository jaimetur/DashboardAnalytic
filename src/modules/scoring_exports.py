"""Reference-style editable scoring matrices using the existing CDR master."""
from __future__ import annotations

from copy import deepcopy
import csv
import json
from io import BytesIO, StringIO
from math import ceil, isclose, isfinite
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from src.modules.cdr_reporting import (
    _named_slide_layout, _remove_all_slides, _set_slide_header,
    _set_structural_slide_text,
)
from src.modules.scoring_views import (
    THRESHOLD_COLORS, _gap_order_key, _hierarchy_display_value, build_scoring_views,
)
from src.modules.scoring_pptx_allocation import (
    _environment_display_label, add_maximum_allocation_donut, category_maximum_allocations,
    maximum_allocations_from_configuration,
)

_FONT = 'Ericsson Hilda'
_WHITE = '#FFFFFF'
_NEUTRAL = '#ECEFF1'
_GAP_SUMMARY_HEADER = '#455B65'
_GAP_SUMMARY_COMPARISON_HEADER = '#DDEBE6'
_GAP_SUMMARY_CATEGORY = '#E6F0F7'
_SCORING_EXPANDED_CATEGORY = '#DCE5E9'
_KPI_TYPE_COLORS = {'Reliable': '#D8EFCA', 'Diff': '#FFF2CC'}
_LEGACY_CAMPAIGN_WARNING = 'Campaigns are scored separately; the supplied Tableau Prep flow pools campaigns.'
_SCOPE_FILTER_FIELDS = ('Operator', 'Vendor', 'Region', 'City', 'Campaign')


def _table_for_mode(matrix: dict, table_mode: str) -> dict:
    """Project saved KPI rows without changing chart inputs or weighted totals."""
    if table_mode not in {'expanded', 'summary'}:
        raise ValueError('Table mode must be expanded or summary.')
    projected = deepcopy(matrix)
    row_key = 'category_rows' if table_mode == 'summary' else 'expanded_rows'
    total_key = 'category_total' if table_mode == 'summary' else 'expanded_total'
    projected['rows'] = deepcopy(matrix.get(row_key, matrix.get('rows', [])))
    projected['total'] = deepcopy(matrix.get(total_key, matrix.get('total', {})))
    projected['table_mode'] = table_mode
    return projected


def export_scoring_csv(job: dict[str, Any], result: dict[str, Any], kind: str,
                       table_mode: str, operator_mapping_groups: list[dict[str, Any]] | None = None,
                       *, environment: str = 'all') -> str:
    """Export the displayed table mode in a flat, explicit comparison format.

    Legacy CSV URLs without table_mode keep their original raw-result schema.
    Category scores are sums; category and final GAPs average available KPI gaps.
    """
    if kind not in {'scoring', 'gap'}:
        raise ValueError('CSV table kind must be scoring or gap.')
    views = build_scoring_views(job, result, operator_mapping_groups)
    matrix_key = 'hierarchy_score_tables' if kind == 'scoring' else 'hierarchy_gap_tables'
    matrices = views.get(matrix_key) or views.get('score_tables' if kind == 'scoring' else 'gap_summary_tables', [])
    _validate_environment(views, environment)
    if environment != 'all':
        matrices = [matrix for matrix in matrices if matrix.get('context', {}).get('environment') == environment]
    fields = ['table_type', 'row_type', 'environment', 'operator', 'vendor', 'region', 'city', 'campaign',
              'reference_operator', 'category', 'kpi_code', 'kpi', 'kpi_type', 'kpi_value',
              'score_points', 'max_points', 'score_weight_percent', 'gap_points', 'complete',
              'gap_partial', 'gap_environments']
    output = StringIO(newline='')
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    for original in matrices:
        matrix = _table_for_mode(original, table_mode)
        columns = {column['id']: column for column in matrix.get('hierarchy_columns', [])}
        operators = matrix['operators']
        if kind == 'gap':
            operators = [operator for operator in operators
                         if not columns.get(operator, {}).get('is_reference')
                         and operator != matrix['baseline_operator']]
        rows = matrix['rows'] + [{
            'row_type': 'total', 'category': 'Total', 'kpi': 'Weighted score / Average KPI GAP',
            'kpi_code': '', 'kpi_type': '', **matrix.get('total', {}),
        }]
        for row in rows:
            for operator in operators:
                column = columns.get(operator, {})
                context = {**matrix.get('context', {}), **column.get('context', {})}
                value = row.get('values', {}).get(operator, {})
                writer.writerow({
                    'table_type': kind, 'row_type': row.get('row_type', 'kpi'),
                    **{field: context.get(field) for field in ('environment', 'vendor', 'region', 'city', 'campaign')},
                    'operator': column.get('operator', operator), 'reference_operator': matrix['baseline_operator'],
                    **{field: row.get(field, '') for field in ('category', 'kpi_code', 'kpi', 'kpi_type')},
                    'kpi_value': value.get('value'), 'score_points': value.get('points'),
                    'max_points': row.get('max_points'), 'score_weight_percent': row.get('weight_percent'),
                    'gap_points': row.get('gaps', {}).get(operator), 'complete': value.get('complete', ''),
                    'gap_partial': row.get('gap_partial', {}).get(operator, False),
                    'gap_environments': json.dumps(row.get('gap_environments', {}).get(operator, [])),
                })
    return output.getvalue()


def _validate_environment(views: dict, environment: str) -> None:
    """Validate against the saved result, including custom methodology environments."""
    available = {matrix.get('context', {}).get('environment')
                 for key in ('score_tables', 'hierarchy_score_tables')
                 for matrix in views.get(key, [])}
    if environment != 'all' and environment not in available:
        raise ValueError('The selected environment is not available in this scoring job.')


def _export_environment_views(views: dict, environment: str) -> dict:
    """Put the existing aggregate first; never synthesize missing environment points."""
    _validate_environment(views, environment)
    projected = dict(views)
    for key in ('score_tables', 'gap_summary_tables', 'gap_tables',
                'hierarchy_score_tables', 'hierarchy_gap_tables'):
        matrices = views.get(key, [])
        if environment != 'all':
            matrices = [matrix for matrix in matrices
                        if matrix.get('context', {}).get('environment') == environment]
        matrices = sorted(matrices, key=lambda matrix: matrix.get('context', {}).get('environment') != 'Combined')
        projected[key] = []
        for matrix in matrices:
            if matrix.get('context', {}).get('environment') == 'Combined':
                matrix = {**matrix, 'context': {**matrix['context'], 'environment': 'All Environments'}}
            projected[key].append(matrix)
    return projected


def _number(value: Any) -> str:
    if value is None:
        return 'N/A'
    formatted = f'{float(value):.2f}'
    return '0.00' if formatted == '-0.00' else formatted


def _gap_number(row: dict, operator: str | None = None) -> str:
    """Mark comparisons based on only the common available environments."""
    value = row.get('gaps', {}).get(operator) if operator is not None else row.get('gap_points')
    partial = row.get('gap_partial', {})
    partial = partial.get(operator, False) if isinstance(partial, dict) else partial
    return _number(value) + ('*' if value is not None and partial else '')


def _score_total_label(rows: list[dict], score_ids: list[str], gap_ids: list[str], *, show_gap_values: bool) -> str:
    """Include the partial-coverage legend only when a displayed value is marked."""
    label = 'Weighted score / Average KPI GAP' if show_gap_values else 'Weighted score'
    partial = any(
        (value.get('points') is not None and not value.get('complete'))
        for row in rows for key in score_ids
        for value in [row.get('values', {}).get(key, {})]
    ) or any('*' in _gap_number(row, key) for row in rows for key in gap_ids)
    return label + ('; * incomplete' if partial else '')


def _scope(context: dict[str, Any], *, environment_label: bool = False) -> str:
    return ' · '.join(
        str(value) if key == 'environment' and not environment_label
        else f'{key.replace("_", " ").title()}: {value}'
        for key, value in context.items() if value is not None and str(value).strip()
    )


def _chart_subtitle(matrix: dict[str, Any]) -> str:
    environment = str(matrix.get('context', {}).get('environment') or '').strip()
    return matrix.get('environment_subtitle', _environment_display_label(environment))


def _slide(presentation, title: str, subtitle: str):
    layout = _named_slide_layout(presentation, 'Title Only')
    if layout is None:
        raise ValueError("The PowerPoint template needs a 'Title Only' layout.")
    slide = presentation.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    _set_slide_header(slide, title, subtitle)
    title_shape = slide.shapes.title
    if title_shape is not None:
        paragraphs = title_shape.text_frame.paragraphs
        if paragraphs:
            title_paragraph = paragraphs[0]
            title_paragraph.font.color.rgb = RGBColor.from_string('17232D')
            for run in title_paragraph.runs:
                run.font.color.rgb = RGBColor.from_string('17232D')
        for subtitle_paragraph in paragraphs[1:]:
            subtitle_paragraph.font.size = Pt(14)
            subtitle_paragraph.font.color.rgb = RGBColor.from_string('245A96')
            for run in subtitle_paragraph.runs:
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor.from_string('245A96')
    return slide


def _accent_gap_comparison_title(slide) -> None:
    paragraph = slide.shapes.title.text_frame.paragraphs[0]
    heading, separator, comparison = paragraph.text.partition(' — ')
    if not separator:
        return
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    paragraph.clear()
    for text, color in ((heading + separator, '17232D'), (comparison, 'A34E16')):
        run = paragraph.add_run()
        if properties is not None:
            run._r.insert(0, deepcopy(properties))
        run.text = text
        run.font.color.rgb = RGBColor.from_string(color)


def _text(slide, text: str, top: float, *, left: float = .55, width: float = 12.2,
          height: float = .4, size: float = 10, color: str = '#263746', align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    box.text = text
    for paragraph in box.text_frame.paragraphs:
        paragraph.font.name = _FONT
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
        paragraph.alignment = align
    return box


def _add_individual_gap_notes(slide, matrix: dict, rows: list[dict]) -> None:
    """Show the arithmetic mean of available KPI comparisons beside the table."""
    if matrix.get('hierarchy_columns'):
        comparisons = [
            (row.get('gaps', {}).get(column['id']),
             row.get('gap_partial', {}).get(column['id'], False))
            for row in rows for column in matrix['hierarchy_columns']
        ]
    else:
        comparisons = [(row.get('gap_points'), row.get('gap_partial', False)) for row in rows]
    valid = [(float(value), partial) for value, partial in comparisons
             if value is not None and isfinite(float(value))]
    average = sum(value for value, _ in valid) / len(valid) if valid else None
    value_text = _number(average) + ('*' if any(partial for _, partial in valid) else '')
    box = _text(
        slide, f'KPIs ordered by GAP\n\nAverage KPI GAP: {value_text} points\n\n'
        f'Operator − reference\nGreen: positive\nRed: negative\n\n{matrix["note"]}',
        1.8, left=10.2, width=2.5, height=4.65, size=13,
    )
    box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = box.text_frame.paragraphs[2]
    paragraph.clear()
    paragraph.add_run().text = 'Average KPI GAP: '
    value_run = paragraph.add_run()
    value_run.text = value_text
    value_run.font.bold = True
    value_run.font.color.rgb = RGBColor.from_string(
        '228B22' if average is not None and average > 0 else
        'CC2424' if average is not None and average < 0 else '263746'
    )
    paragraph.add_run().text = ' points'


def _add_gap_priority_arrow(slide, table, *, header_rows: int = 1) -> None:
    """Span the table from just below the Priority label to its bottom edge."""
    frame = table._graphic_frame
    label = _text(slide, 'Priority', frame.top.inches, left=frame.left.inches - .59,
                  width=.52, height=.28, size=7, color='#245A96', align=PP_ALIGN.CENTER)
    label.name = 'GAP KPI Priority Label'
    label.fill.solid()
    label.fill.fore_color.rgb = RGBColor(255, 255, 255)
    label.line.color.rgb = RGBColor.from_string('245A96')
    label.line.width = Pt(1)
    label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    label.text_frame.paragraphs[0].font.name = 'Arial'
    label.text_frame.paragraphs[0].font.bold = True
    body_top = label.top + label.height + Inches(.06)
    body_bottom = frame.top + frame.height
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, frame.left - Inches(.38), body_top,
        Inches(.12), max(Inches(.2), body_bottom - body_top),
    )
    arrow.name = 'GAP KPI Priority Arrow'
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor.from_string('245A96')
    arrow.line.fill.background()
    arrow._element.spPr.append(OxmlElement('a:effectLst'))
    for reference in arrow._element.xpath('./p:style/a:effectRef'):
        reference.set('idx', '0')


def _add_gap_color_scale(slide, matrix: dict, *, left: float, width: float, top: float = 6.57) -> None:
    """Add an editable red-white-green bar with this comparison's actual GAP range."""
    colors = matrix.get('gap_scale_colors') or {
        'negative': '#E57373', 'zero': '#FFFFFF', 'positive': '#70AD47',
    }
    maximum = float(matrix.get('gap_scale_max') or 0.0)

    def blend(start: str, end: str, amount: float) -> str:
        channels = [
            round(int(start[index:index + 2], 16) * (1 - amount)
                  + int(end[index:index + 2], 16) * amount)
            for index in (1, 3, 5)
        ]
        return '#' + ''.join(f'{channel:02X}' for channel in channels)

    heading = _text(slide, 'GAP color scale', top, left=left, width=width, size=7.5, height=.14)
    heading.text_frame.paragraphs[0].font.bold = True
    bar_left, bar_width, bar_top, bar_height = left + width * .25, width * .5, top + .16, .12
    segments = 24
    for index in range(segments):
        midpoint = (index + .5) / segments
        if midpoint <= .5:
            color = blend(colors['negative'], colors['zero'], midpoint * 2)
        else:
            color = blend(colors['zero'], colors['positive'], (midpoint - .5) * 2)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(bar_left + bar_width * index / segments), Inches(bar_top),
            Inches(bar_width / segments + .005), Inches(bar_height),
        )
        shape.name = f'GAP Color Scale Segment {index + 1}'
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color.lstrip('#'))
        shape.line.fill.background()
    _text(slide, f'−{maximum:.2f} (operator trails)', top + .28, left=left,
          width=width * .43, size=6.5, height=.13)
    _text(slide, '· 0 ·', top + .28, left=left + width * .44,
          width=width * .12, size=6.5, height=.13, align=PP_ALIGN.CENTER)
    _text(slide, f'+{maximum:.2f} (operator leads)', top + .28, left=left + width * .57,
          width=width * .43, size=6.5, height=.13, align=PP_ALIGN.RIGHT)


def _cell(cell, text: str, *, color: str = _WHITE, foreground: str = '#17232D',
          size: float = 9, bold: bool = False, left: bool = False) -> None:
    cell.text = text
    cell.margin_left = cell.margin_right = Inches(.025)
    cell.margin_top = cell.margin_bottom = 0
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(color.lstrip('#'))
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.name = _FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(foreground.lstrip('#'))
        paragraph.alignment = PP_ALIGN.LEFT if left else PP_ALIGN.CENTER
        paragraph.line_spacing = 1.0
        paragraph.space_before = paragraph.space_after = 0
    properties = cell._tc.get_or_add_tcPr()
    border_namespace = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    edges = ('lnL', 'lnR', 'lnT', 'lnB')
    for edge_index, edge in enumerate(edges):
        existing = properties.find(f'{border_namespace}{edge}')
        if existing is not None:
            properties.remove(existing)
        line = OxmlElement(f'a:{edge}')
        line.set('w', '4500')
        fill = OxmlElement('a:solidFill')
        rgb = OxmlElement('a:srgbClr')
        rgb.set('val', '65717A')
        fill.append(rgb)
        line.append(fill)
        properties.insert(edge_index, line)


def _operator_color(matrix: dict, operator: str) -> str:
    color = matrix.get('operator_styles', {}).get(operator, {}).get('color')
    if color:
        return color
    for column in matrix.get('hierarchy_columns', []):
        if column.get('operator') == operator and column.get('color'):
            return column['color']
    return '#365F91'


def _operator_label(matrix: dict, operator: str) -> str:
    return matrix.get('operator_styles', {}).get(operator, {}).get('label', operator)


def _campaigns_for_export(job: dict[str, Any], result: dict[str, Any]) -> list[str]:
    """Prefer campaign values in scored contexts over the selected-source union."""
    campaigns: list[str] = []

    def collect(value: Any) -> None:
        if value is None:
            return
        label = str(value).strip()
        if label and label.casefold() not in {item.casefold() for item in campaigns}:
            campaigns.append(label)

    for key in ('scoring', 'score_rows', 'totals', 'charts'):
        rows = result.get(key)
        if not isinstance(rows, list):
            continue
        for row in rows:
            if not isinstance(row, dict):
                continue
            campaign = next((value for field, value in row.items()
                             if str(field).casefold() in {'campaign', 'period', 'quarter'} and value is not None), None)
            collect(campaign)
    if campaigns:
        return sorted(campaigns, key=str.casefold)

    source_metadata = job.get('source_metadata', [])
    if isinstance(source_metadata, dict):
        source_metadata = source_metadata.get('sources', [])
    if isinstance(source_metadata, list):
        for source in source_metadata:
            if not isinstance(source, dict):
                continue
            source_campaigns = source.get('campaigns', [])
            for campaign in source_campaigns if isinstance(source_campaigns, list) else [source_campaigns]:
                collect(campaign)
    if not campaigns:
        fallback = job.get('campaigns', [])
        for campaign in fallback if isinstance(fallback, list) else [fallback]:
            collect(campaign)
    return sorted(campaigns, key=str.casefold)


def _filter_value(value: Any) -> str:
    if isinstance(value, (list, tuple, set)):
        return ', '.join(str(item).strip() for item in value if item is not None and str(item).strip())
    return str(value).strip() if value is not None else ''


def _scope_filter_fields(job: dict[str, Any]) -> tuple[str, ...]:
    configured = job.get('aggregation_hierarchy')
    if not isinstance(configured, list):
        configuration = job.get('configuration')
        configured = configuration.get('aggregation_hierarchy') if isinstance(configuration, dict) else None
    if not isinstance(configured, list):
        return _SCOPE_FILTER_FIELDS
    canonical = {field.casefold(): field for field in _SCOPE_FILTER_FIELDS}
    fields = []
    for value in configured:
        field = canonical.get(str(value).strip().casefold())
        if field and field not in fields:
            fields.append(field)
    fields.extend(field for field in _SCOPE_FILTER_FIELDS if field not in fields)
    return tuple(fields)


def _canonical_scope_filter_labels(payload: Any, fields: tuple[str, ...] = _SCOPE_FILTER_FIELDS) -> list[str] | None:
    if not isinstance(payload, dict) or any(not isinstance(value, list) for value in payload.values()):
        return None
    normalized = {str(key).strip().casefold(): value for key, value in payload.items()}
    if payload and not any(field.casefold() in normalized for field in fields):
        return None
    return [
        f'{field}: {_filter_value(normalized.get(field.casefold(), [])) or "All"}'
        for field in fields
    ]


def _context_filter_labels(job: dict[str, Any]) -> list[str]:
    labels: list[str] = []
    seen: set[str] = set()
    reserved = {'period', 'quarter', 'nr mode', 'aggregation', 'aggregation level', 'baseline'}

    def append(label: Any, value: Any, operator: Any = None) -> None:
        name = str(label or '').strip().replace('_', ' ')
        rendered_value = _filter_value(value)
        if not name or not rendered_value or name.casefold() in reserved:
            return
        formatted_name = ' '.join(
            part.upper() if part.casefold() == 'ran' else part.capitalize() for part in name.split()
        )
        description = f'{formatted_name} {operator} {rendered_value}' if operator else f'{formatted_name}: {rendered_value}'
        identity = description.casefold()
        if identity not in seen:
            seen.add(identity)
            labels.append(description)

    for payload_name in ('context_filters', 'filters'):
        payload = job.get(payload_name)
        if isinstance(payload, dict):
            canonical_labels = (
                _canonical_scope_filter_labels(payload, _scope_filter_fields(job))
                if payload_name == 'context_filters' else None
            )
            if canonical_labels is not None:
                for label in canonical_labels:
                    append(*label.split(': ', 1))
                continue
            if any(key in payload for key in ('column', 'field', 'name', 'key')):
                field = payload.get('column') or payload.get('field') or payload.get('name') or payload.get('key')
                append(field, payload.get('value', payload.get('values')), payload.get('operator'))
            else:
                for name, value in payload.items():
                    if isinstance(value, dict):
                        append(name, value.get('value', value.get('values')), value.get('operator'))
                    else:
                        append(name, value)
        elif isinstance(payload, list):
            for item in payload:
                if isinstance(item, dict):
                    field = item.get('column') or item.get('field') or item.get('name') or item.get('key')
                    append(field, item.get('value', item.get('values')), item.get('operator'))
                elif isinstance(item, str) and item.strip() and item.casefold() not in seen:
                    seen.add(item.casefold())
                    labels.append(item.strip())
        elif isinstance(payload, str) and payload.strip() and payload.casefold() not in seen:
            seen.add(payload.casefold())
            labels.append(payload.strip())
    return labels


def _scoring_filter_subtitle(job: dict[str, Any], environment: str | None = None) -> str:
    nr_mode = str(job.get('nr_mode') or 'NSA').strip()
    mode_label = {'nsa': 'Non-Standalone', 'sa': 'Standalone'}.get(nr_mode.casefold(), nr_mode)
    levels = job.get('aggregation_levels') or job.get('levels') or []
    if isinstance(levels, str):
        levels = [levels]
    level_labels = [str(level).strip() for level in levels if str(level).strip()] if isinstance(levels, list) else []
    plan = job.get('_scoring_display_selections')
    if plan is None:
        from src.modules.cdr_report_filenames import scoring_display_selections
        plan = scoring_display_selections(job.get('context_filters'), aggregation_values=level_labels,
                                          nr_mode_label=nr_mode)
    def display(field):
        selected = ', '.join(plan['values'][field])
        return selected + (', ...' if selected else '...') if plan['omitted'][field] else selected
    return '\n'.join([mode_label, *[
        f'{field.title()}: {display(field)}'
        for field in ('aggregation', 'operator', 'vendor', 'region', 'city')
    ]])


def prepare_scoring_display_selections(job: dict[str, Any], result: dict[str, Any],
                                       template_path: Path) -> dict:
    """Compute shared presentation and download-name selections without changing filters."""
    from src.modules.cdr_report_filenames import scoring_display_selections
    return scoring_display_selections(
        job.get('context_filters'),
        campaign_values=_campaigns_for_export(job, result),
        aggregation_values=job.get('aggregation_levels') or job.get('levels'),
        nr_mode_label=job.get('nr_mode') or 'NSA',
    )


def _fit_scoring_intro_subtitle(presentation, slide, layout_name: str, subtitle: str) -> None:
    title_shape = next(
        (shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3}),
        None,
    )
    if title_shape is None:
        return
    subtitle_shape = next(
        (shape for shape in slide.placeholders if shape.placeholder_format.type == 4),
        None,
    )
    if subtitle_shape is None:
        # Transition layouts may omit the subtitle; clone the template's real placeholder.
        title_layout = _named_slide_layout(presentation, 'Title Page')
        source = next((shape for shape in title_layout.placeholders
                       if shape.placeholder_format.type == 4), None) if title_layout else None
        if source is None:
            raise ValueError("The PowerPoint template needs a subtitle placeholder for scoring exports.")
        slide.shapes.clone_placeholder(source)
        subtitle_shape = next(shape for shape in slide.placeholders
                              if shape.placeholder_format.type == 4)

    title_text = title_shape.text_frame.text.split('\x0b', 1)[0].split('\n', 1)[0]
    title_shape.text_frame.text = title_text
    _set_shape_geometry(title_shape, Inches(.7))
    title_shape.top = Inches(1.8)
    for paragraph in title_shape.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.font.size = Pt(44)
        paragraph.font.color.rgb = RGBColor.from_string('FFFFFF')

    lines = subtitle.splitlines()
    subtitle_shape.height = Inches(.45)
    subtitle_shape.left = title_shape.left
    subtitle_shape.width = title_shape.width
    subtitle_shape.top = title_shape.top + title_shape.height + Inches(.05)
    subtitle_shape.text_frame.text = lines[0] if lines else ''
    subtitle_shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    mode = subtitle_shape.text_frame.paragraphs[0]
    mode.alignment = PP_ALIGN.LEFT
    mode.font.size = Pt(18)
    mode.font.bold = True
    mode.font.color.rgb = RGBColor.from_string('FFC700')
    mode.line_spacing = 1.0
    mode.space_before = mode.space_after = Pt(0)
    properties = mode._p.get_or_add_pPr()
    properties.set('marL', '0')
    properties.set('indent', '0')
    for bullet in properties.xpath('./a:buNone'):
        properties.remove(bullet)
    properties.insert_element_before(OxmlElement('a:buNone'), 'a:tabLst', 'a:defRPr', 'a:extLst')

    filters_shape = _text(
        slide, '', 3.6, left=title_shape.left / Inches(1),
        width=title_shape.width / Inches(1), height=1.85,
        size=16, color='#CCEEF4',
    )
    filters_shape.name = 'Scoring Aggregations and Filters'
    frame = filters_shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, line in enumerate(['Aggregations & Filters:', *lines[1:]]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(16 if index == 0 else _hierarchy_content_font(
            [line], filters_shape.width / Inches(1) - .72, maximum=16,
        ))
        paragraph.font.bold = index == 0
        properties = paragraph._p.get_or_add_pPr()
        properties.set('marL', str(Inches(.22) if index else 0))
        properties.set('indent', '0')
        properties.insert(0, OxmlElement('a:buNone'))
        paragraph.font.color.rgb = RGBColor.from_string('CCEEF4')
        paragraph.line_spacing = 1.0
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(3)
    for shape in (title_shape, subtitle_shape, filters_shape):
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.name = 'Aptos'
            paragraph.font._rPr.set('spc', '0')
            paragraph.font._rPr.set('kern', '0')
            for run in paragraph.runs:
                run.font.name = 'Aptos'
                run.font._rPr.set('spc', '0')
                run.font._rPr.set('kern', '0')


def _set_shape_geometry(shape, height: int) -> None:
    """Materialize inherited placeholder geometry before resizing its height."""
    left, top, width = shape.left, shape.top, shape.width
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height


def _scoring_environment_title(environment: str, configuration: dict) -> str:
    scope = configuration.get('scope', {}).get('environments', {}).get(environment, {})
    if scope.get('display_name'):
        return str(scope['display_name'])
    filters = [str(scope.get(key) or '').strip() for key in ('g_level_1', 'g_level_2')]
    filters = [value for value in filters if value]
    if filters:
        filters = ['Connecting Roads' if value.casefold() in {'connectionroad', 'connectingroads'}
                   else value for value in filters]
        return ' - '.join(filters)
    return _environment_display_label(environment)


def _add_scoring_intro_slides(
    presentation, job: dict[str, Any], result: dict[str, Any], *, environment: str | None = None,
) -> None:
    configuration = job.get('configuration') or result.get('configuration') or {}
    title = (_scoring_environment_title(environment, configuration) if environment
             else 'Scoring & GAP Analysis')
    subtitle = _scoring_filter_subtitle(job, environment)
    campaigns = _campaigns_for_export(job, result)
    plan = job.get('_scoring_display_selections')
    campaign_value = ', '.join(plan['values']['campaign']) if plan else ', '.join(campaigns) or 'All Campaigns'
    if plan and plan['omitted']['campaign']:
        campaign_value += ', ...' if campaign_value else '...'
    campaign_text = f'Campaigns: {campaign_value}'
    layout_name = 'Title Page'
    layout = _named_slide_layout(presentation, layout_name)
    if layout is None:
        raise ValueError(f"The PowerPoint template needs a '{layout_name}' layout for scoring exports.")
    slide = presentation.slides.add_slide(layout)
    _set_structural_slide_text(slide, title, subtitle)
    _fit_scoring_intro_subtitle(presentation, slide, layout_name, subtitle)
    campaign_top = 5.9
    campaign_shape = _text(slide, campaign_text, campaign_top, left=.52, width=10.68, height=.32,
                           size=16, color=_WHITE)
    campaign_shape.name = 'Scoring Campaigns'
    campaign_shape.text_frame.word_wrap = False
    for paragraph in campaign_shape.text_frame.paragraphs:
        paragraph.font.size = Pt(_hierarchy_content_font(
            [campaign_text], campaign_shape.width / Inches(1) - .5, maximum=16,
        ))
        paragraph.font.name = 'Aptos'
        for run in paragraph.runs:
            run.font.name = 'Aptos'


def _header_foreground(color: str) -> str:
    red, green, blue = (int(color.lstrip('#')[index:index + 2], 16) for index in (0, 2, 4))
    return '#17232D' if .299 * red + .587 * green + .114 * blue > 155 else _WHITE


def _operator_pages(table: dict) -> list[list[str]]:
    operators = table['operators']
    baseline = table['baseline_operator']
    if len(operators) <= 5:
        return [operators]
    if baseline in operators:
        compared = [operator for operator in operators if operator != baseline]
        return [[operator for operator in operators if operator in compared[start:start + 4] or operator == baseline]
                for start in range(0, len(compared), 4)]
    return [operators[start:start + 5] for start in range(0, len(operators), 5)]


def _scalar_matrix_as_hierarchy(matrix: dict) -> dict:
    converted = deepcopy(matrix)
    converted['hierarchy_levels'] = ['Operator']
    converted['hierarchy_columns'] = [
        {'id': operator, 'operator': operator, 'path': [{'level': 'Operator', 'value': _operator_label(matrix, operator)}],
         'color': _operator_color(matrix, operator), 'is_reference': operator == matrix['baseline_operator']}
        for operator in matrix['operators']
    ]
    return converted


def _kpi_type_color(metric: dict) -> str:
    """Keep the same KPI classification colors across Score and GAP tables."""
    return _KPI_TYPE_COLORS.get(str(metric.get('kpi_type', '')), _NEUTRAL)


_SUMMARY_BEST_COLOR = '#C6EFCE'
_SUMMARY_WORST_COLOR = '#FFC7CE'


def _add_summary_operator_legend(slide) -> None:
    """Explain the best and worst score highlights above Summary tables on the right."""
    for index, (label, color) in enumerate((('Best operator', _SUMMARY_BEST_COLOR),
                                            ('Worst operator', _SUMMARY_WORST_COLOR))):
        table = slide.shapes.add_table(
            1, 1, Inches(10.1 + index * 1.35), Inches(1.25), Inches(1.3), Inches(.18),
        ).table
        _cell(table.cell(0, 0), label, color=color, size=8)


def _summary_extreme_operators(values: dict[str, dict]) -> tuple[set[str], set[str]]:
    """Return tied best and worst operators when the values establish a ranking."""
    available = {}
    for operator, value in values.items():
        points = value.get('points') if isinstance(value, dict) else value
        try:
            points = float(points)
        except (TypeError, ValueError):
            continue
        if isfinite(points):
            available[operator] = points
    if len(available) < 2:
        return set(), set()
    best = max(available.values())
    worst = min(available.values())
    if isclose(best, worst, rel_tol=1e-9, abs_tol=1e-9):
        return set(), set()
    return (
        {operator for operator, points in available.items()
         if isclose(points, best, rel_tol=1e-9, abs_tol=1e-9)},
        {operator for operator, points in available.items()
         if isclose(points, worst, rel_tol=1e-9, abs_tol=1e-9)},
    )


def _summary_extreme_columns(values: dict[str, dict], columns: list[dict]) -> tuple[set[str], set[str]]:
    """Rank hierarchy leaves only against operators in the same non-operator path."""
    groups: dict[tuple[tuple[str, Any], ...], dict[str, dict]] = {}
    for column in columns:
        context = tuple(
            (entry['level'], entry.get('value'))
            for entry in column.get('path', []) if entry.get('level') != 'Operator'
        )
        groups.setdefault(context, {})[column['id']] = values.get(column['id'], {})
    best_columns: set[str] = set()
    worst_columns: set[str] = set()
    for group in groups.values():
        best, worst = _summary_extreme_operators(group)
        best_columns.update(best)
        worst_columns.update(worst)
    return best_columns, worst_columns


def _score_tables(presentation, matrices: list[dict], legend: list[dict], *, gap_layout: str = 'end',
                  show_gap_values: bool = True, title: str = 'Scoring Tables — Breakdown') -> None:
    if gap_layout == 'adjacent':
        matrices = [_scalar_matrix_as_hierarchy(matrix) for matrix in matrices]
        _hierarchy_score_tables(presentation, matrices, legend, gap_layout=gap_layout,
                                show_gap_values=show_gap_values, title=title)
        return

    for matrix in matrices:
        metric_pages = [matrix['rows']]
        is_summary = matrix.get('table_mode') == 'summary'
        fixed_count = 4 if is_summary else 5
        weight_index = 2 if is_summary else 3
        maximum_index = weight_index + 1
        for operators in [matrix['operators']]:
            baseline = matrix['baseline_operator']
            comparisons = ([name for name in operators if name != baseline]
                           if show_gap_values and baseline in matrix['operators'] else [])
            for page_index, metrics in enumerate(metric_pages):
                total_label = _score_total_label(
                    [*metrics, matrix['total']], operators, comparisons, show_gap_values=show_gap_values,
                )
                page_label = f' · KPIs {page_index + 1}/{len(metric_pages)}' if len(metric_pages) > 1 else ''
                slide = _slide(presentation, title, _chart_subtitle(matrix) + page_label)
                include_total = page_index == len(metric_pages) - 1
                headers = ['CATEGORY', 'CATEGORY' if is_summary else 'KPI',
                           *([] if is_summary else ['Type of KPI']), 'Score weight\n(%)', 'Max score',
                           *[_operator_label(matrix, name) for name in operators],
                           *[f'GAP {_operator_label(matrix, name)}\n− {_operator_label(matrix, baseline)}' for name in comparisons]]
                count = len(metrics) + 1 + int(include_total)
                table = slide.shapes.add_table(count, len(headers), Inches(.55), Inches(1.55),
                                               presentation.slide_width - Inches(1.1), Inches(5.2)).table
                widths = [1.22, 2.60, *([] if is_summary else [.75]), .68, .79]
                remaining = (presentation.slide_width / 914400 - 1.1 - sum(widths)) / max(1, len(headers) - fixed_count)
                for index, column in enumerate(table.columns):
                    column.width = Inches(widths[index] if index < fixed_count else remaining)
                table.rows[0].height = Inches(.46)
                metric_font, body_heights = _hierarchy_metric_layout(
                    metrics, total_label,
                    available_height=5.2 - .46, kpi_width=widths[1],
                )
                for row, height in zip(list(table.rows)[1:], body_heights):
                    row.height = Inches(height)
                numeric_texts = [
                    _number(value.get('points')) + ('*' if value.get('points') is not None and not value.get('complete') else '')
                    for metric in metrics + [matrix['total']] for value in metric.get('values', {}).values()
                ] + [_gap_number(metric, operator) for metric in metrics + [matrix['total']] for operator in comparisons]
                numeric_font = _hierarchy_content_font(
                    numeric_texts, remaining, maximum=min(11, min(body_heights) * 72 * .82),
                )
                colors = [_GAP_SUMMARY_HEADER, _GAP_SUMMARY_HEADER,
                          *([] if is_summary else [_GAP_SUMMARY_HEADER]), '#4EA72E', '#FF0000',
                          *[_operator_color(matrix, name) for name in operators], *['#FFFF00'] * len(comparisons)]
                for index, header in enumerate(headers):
                    _cell(table.cell(0, index), header, color=colors[index],
                          foreground=_header_foreground(colors[index]), size=8.5, bold=True)
                category_color = _GAP_SUMMARY_CATEGORY if is_summary else _SCORING_EXPANDED_CATEGORY
                for row_index, metric in enumerate(metrics, 1):
                    subtotal = metric.get('row_type') == 'category'
                    row_color = ('#E3E6E7' if is_summary else category_color) if subtotal else None
                    best, worst = (_summary_extreme_operators(metric.get('values', {}))
                                   if matrix.get('table_mode') == 'summary' and subtotal else (set(), set()))
                    _cell(table.cell(row_index, 0), metric['category'], color=category_color,
                          size=metric_font, left=True, bold=True)
                    _cell(table.cell(row_index, 1), metric['kpi'], color=row_color or '#E7E8E9',
                          bold=True, left=True, size=metric_font)
                    if not is_summary:
                        _cell(table.cell(row_index, 2), metric.get('kpi_type', ''),
                              color=row_color or _kpi_type_color(metric), size=metric_font, bold=subtotal)
                    _cell(table.cell(row_index, weight_index), _number(metric['weight_percent']) + '%', color=row_color or _WHITE,
                          size=metric_font, bold=subtotal)
                    _cell(table.cell(row_index, maximum_index), _number(metric['max_points']), color=row_color or _WHITE,
                          size=metric_font, bold=subtotal)
                    for offset, operator in enumerate(operators, fixed_count):
                        value = metric['values'][operator]
                        partial = value['points'] is not None and not value['complete']
                        score_color = (_SUMMARY_BEST_COLOR if operator in best else
                                       _SUMMARY_WORST_COLOR if operator in worst else
                                       row_color or value.get('color', _WHITE if value['complete'] else _NEUTRAL))
                        _cell(table.cell(row_index, offset), _number(value['points']) + ('*' if partial else ''),
                              color=score_color,
                              size=numeric_font, bold=subtotal)
                    for offset, operator in enumerate(comparisons, fixed_count + len(operators)):
                        _cell(table.cell(row_index, offset), _gap_number(metric, operator),
                              color=row_color or metric.get('gap_colors', {}).get(operator, '#FFF0D8'),
                              size=numeric_font, bold=True)
                first = 0
                while first < len(metrics):
                    last = first
                    while last + 1 < len(metrics) and metrics[last + 1]['category'] == metrics[first]['category']:
                        last += 1
                    if last > first:
                        cell = table.cell(first + 1, 0)
                        cell.merge(table.cell(last + 1, 0))
                        _cell(cell, metrics[first]['category'], color=category_color,
                              size=metric_font, left=True, bold=True)
                    first = last + 1
                if include_total:
                    total = matrix['total']
                    index = count - 1
                    total_color = '#D8DFE4' if is_summary else category_color
                    best, worst = (_summary_extreme_operators(total.get('values', {}))
                                   if is_summary else (set(), set()))
                    _cell(table.cell(index, 0), 'TOTAL', color=total_color, bold=True, left=True, size=metric_font)
                    _cell(table.cell(index, 1), total_label, color=total_color, size=metric_font, left=True)
                    if not is_summary:
                        _cell(table.cell(index, 2), '', color=total_color, size=metric_font)
                    for column, value in ((weight_index, _number(total['weight_percent']) + '%'),
                                          (maximum_index, _number(total['max_points']))):
                        _cell(table.cell(index, column), value, color=total_color, bold=True, size=numeric_font)
                    for column, operator in enumerate(operators, fixed_count):
                        value = total['values'][operator]
                        score_color = (_SUMMARY_BEST_COLOR if operator in best else
                                       _SUMMARY_WORST_COLOR if operator in worst else total_color)
                        _cell(table.cell(index, column), _number(value['points']) + ('*' if value['points'] is not None and not value['complete'] else ''),
                              color=score_color, bold=True, size=numeric_font)
                    for column, operator in enumerate(comparisons, fixed_count + len(operators)):
                        _cell(table.cell(index, column), _gap_number(total, operator), color=total_color, bold=True, size=numeric_font)
                if is_summary:
                    _remove_summary_category_column(table)
                    _add_summary_operator_legend(slide)
                else:
                    for index, item in enumerate(legend):
                        legend_table = slide.shapes.add_table(1, 1, Inches(.55 + index * 1.15), Inches(7.02), Inches(1.1), Inches(.18)).table
                        _cell(legend_table.cell(0, 0), item.get('band', ''), color=item['color'], size=8)
                _text(slide, f'GAP = operator − reference; ±{matrix.get("gap_scale_max", 0):.2f} points; green + / red −',
                      7.03, left=6.5, width=6.2, size=8.5, height=.18)
                _text(slide, matrix['coverage_note'], 7.27, size=8, height=.18)


def _format_chart(chart, *, maximum: float, labels=XL_DATA_LABEL_POSITION.OUTSIDE_END) -> None:
    # Axis identifiers are unsigned integers in DrawingML, including cross references.
    axis_references = chart._chartSpace.xpath('.//c:axId | .//c:crossAx')
    axis_ids = {value: str(index + 1) for index, value in enumerate(
        dict.fromkeys(reference.get('val') for reference in axis_references),
    )}
    for axis_reference in axis_references:
        axis_reference.set('val', axis_ids[axis_reference.get('val')])
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.name = 'Arial'
    chart.legend.font.size = Pt(11)
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = 'Arial'
        axis.tick_labels.font.size = Pt(10)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = max(1, maximum) * 1.12
    chart.plots[0].has_data_labels = True
    data_labels = chart.plots[0].data_labels
    data_labels.position = labels
    data_labels.number_format = '0.0'
    data_labels.font.name = 'Arial'
    data_labels.font.size = Pt(9)
    # Single-level categories need no native hierarchy label dividers.
    for flag in chart.category_axis._element.xpath('./c:noMultiLvlLbl'):
        flag.set('val', '1' if chart.plots[0].categories.depth <= 1 else '0')
    for gridline in chart.category_axis._element.xpath('./c:majorGridlines | ./c:minorGridlines'):
        gridline.getparent().remove(gridline)
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.minor_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.format.line.fill.background()
    if labels == XL_DATA_LABEL_POSITION.CENTER:
        data_labels.font.color.rgb = RGBColor.from_string('FFFFFF')


def _hide_chart_legend_entries(chart, indexes: set[int]) -> None:
    """Hide selected native series legend entries without changing the chart data."""
    if not indexes:
        return
    legend = chart.legend._element
    position = 1 if len(legend) and legend[0].tag.endswith('}legendPos') else 0
    for index in sorted(indexes):
        entry = OxmlElement('c:legendEntry')
        entry_index = OxmlElement('c:idx')
        entry_index.set('val', str(index))
        delete = OxmlElement('c:delete')
        delete.set('val', '1')
        entry.extend((entry_index, delete))
        legend.insert(position, entry)
        position += 1


def _fill_series(series, color: str) -> None:
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor.from_string(color.lstrip('#'))
    series.format.line.color.rgb = RGBColor.from_string('FFFFFF')
    series.format.line.width = Pt(.5)


def _charts(presentation, matrices: list[dict]) -> None:
    for matrix in matrices:
        slide = _slide(presentation, 'Scoring per Category', _chart_subtitle(matrix))
        categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
        data = CategoryChartData()
        data.categories = categories
        maximum = 0.0
        for operator in matrix['operators']:
            values = []
            for category in categories:
                cells = [row['values'][operator] for row in matrix['rows'] if row['category'] == category]
                available = [cell['points'] for cell in cells if cell['points'] is not None]
                value = sum(available) if available else None
                values.append(value)
                maximum = max(maximum, value or 0)
            data.add_series(_operator_label(matrix, operator), values)
        chart_width = presentation.slide_width - Inches(1.4)
        chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(1.65),
                                             chart_width, Inches(4.65), data)
        chart = chart_shape.chart
        _format_chart(chart, maximum=maximum)
        chart.plots[0].gap_width = 140
        chart.plots[0].overlap = -20
        chart.plots[0].data_labels.font.size = Pt(7)
        chart.legend.font.size = Pt(10)
        for series, operator in zip(chart.series, matrix['operators']):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(_operator_color(matrix, operator).lstrip('#'))
        slide.shapes.add_group_shape([chart_shape]).name = 'Scoring Category Comparison Chart'
        _text(slide, matrix['coverage_note'], 6.95, size=9)


def _add_family_allocation_donut(slide, matrix: dict) -> None:
    voice_categories = {'CLASSIC CALLS', 'WHATSAPP CALLS', 'MULTI RAB'}
    maximums = [
        sum(row.get('max_points') or 0 for row in matrix['rows']
            if (row.get('source_kind') in {'voice', 'speech'}
                if row.get('source_kind') else row['category'] in voice_categories) == is_voice)
        for is_voice in (True, False)
    ]
    allocations = matrix.get('environment_allocations') or {
        str(matrix['context'].get('environment') or 'All Environments'): {
            'voice': maximums[0], 'data': maximums[1],
        },
    }
    return add_maximum_allocation_donut(
        slide, matrix, allocations, left=9.7, top=1.75, width=3.0, height=5.0,
    )


def _add_category_allocation_donut(slide, matrix: dict) -> None:
    voice_categories = {'CLASSIC CALLS', 'WHATSAPP CALLS', 'MULTI RAB'}
    maximums = [
        sum(row.get('max_points') or 0 for row in matrix['rows']
            if (row.get('source_kind') in {'voice', 'speech'}
                if row.get('source_kind') else row['category'] in voice_categories) == is_voice)
        for is_voice in (True, False)
    ]
    allocations = matrix.get('environment_allocations') or {
        str(matrix['context'].get('environment') or 'All Environments'): {
            'voice': maximums[0], 'data': maximums[1],
        },
    }
    return add_maximum_allocation_donut(
        slide, matrix, allocations, left=9.7, top=1.75, width=3.0, height=5.0,
        category_allocations=category_maximum_allocations(matrix),
    )


def _stacked_category_chart(presentation, matrix: dict) -> None:
    """Compare KPI category contributions as one stacked bar per operator."""
    pages = _stacked_chart_pages(matrix)
    if len(pages) > 1:
        for page in pages:
            _stacked_category_chart(presentation, page)
        return
    slide = _slide(presentation, 'Best Network Scoring per Category', _chart_subtitle(matrix))
    categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
    operators = list(matrix['operators'])
    _stacked_chart_page_note(slide, matrix)
    data = CategoryChartData()
    data.categories = [_operator_label(matrix, operator) for operator in operators]
    series_specs = []
    maximum = 0.0
    for operator_index, operator in enumerate(operators):
        for category_index, category in enumerate(categories):
            cells = [row['values'][operator] for row in matrix['rows'] if row['category'] == category]
            available = [cell['points'] for cell in cells if cell['points'] is not None]
            value = sum(available) if available else None
            values = [value if index == operator_index else None for index in range(len(operators))]
            label = _operator_label(matrix, operator) if category_index == 0 else f'{operator} · {category}'
            data.add_series(label, values)
            series_specs.append((operator, category_index, category_index > 0))
    totals = [matrix['total']['values'][operator]['points'] for operator in operators]
    data.add_series('TOTAL SCORE', totals)
    chart_width = Inches(8.9)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Inches(.55), Inches(1.65), chart_width, Inches(4.65), data,
    )
    chart = chart_shape.chart
    maximum = max(
        (sum(row['values'][operator]['points'] or 0 for row in matrix['rows']
             if row['values'][operator]['points'] is not None) for operator in operators),
        default=0.0,
    )
    _format_chart(chart, maximum=matrix.get('_stacked_chart_maximum', maximum), labels=XL_DATA_LABEL_POSITION.CENTER)
    chart.legend.font.size = Pt(9)
    hidden_legend_entries = set()
    for series_index, (operator, category_index, hide_from_legend) in enumerate(series_specs):
        series = chart.series[series_index]
        _fill_series(
            series, _hierarchy_chart_color(_operator_color(matrix, operator), category_index, len(categories)),
        )
        if hide_from_legend:
            hidden_legend_entries.add(series_index)
    _hide_chart_legend_entries(chart, hidden_legend_entries | {len(series_specs)})
    _format_stacked_segment_labels(chart)
    _add_total_labels(chart)
    category_key = _add_category_shade_bar(
        slide, categories, left=Inches(.55), top=Inches(6.38), width=chart_width, height=Inches(.4),
    )
    hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, [
        {'id': operator, 'operator': _operator_label(matrix, operator)} for operator in operators
    ])
    chart_group = [chart_shape] + ([category_key] if category_key is not None else []) + ([hierarchy_grid] if hierarchy_grid is not None else [])
    slide.shapes.add_group_shape(chart_group).name = 'Scoring Stacked Operator Chart'
    _add_category_allocation_donut(slide, matrix)
    _text(slide, matrix['coverage_note'], 6.95, size=9)


def _stacked_chart_pages(matrix: dict, *, limit: int = 12) -> list[dict]:
    """Keep preferred operators together and preserve hierarchy order within each page."""
    columns = matrix.get('hierarchy_columns') or [
        {'id': operator, 'operator': operator} for operator in matrix['operators']
    ]
    if len(columns) <= limit:
        return [matrix]
    grouped = {}
    for column in columns:
        grouped.setdefault(column['operator'], []).append(column)
    preferred = {'3', 'three', 'threeuk', 'ee', 'vfuk', 'vodafoneuk'}
    groups = sorted(grouped.items(), key=lambda group: (
        ''.join(character for character in _operator_label(matrix, group[0]).lower() if character.isalnum()) not in preferred
        and ''.join(character for character in str(group[0]).lower() if character.isalnum()) not in preferred,
        list(grouped).index(group[0]),
    ))
    pages = []
    current = []
    for _, group in groups:
        for start in range(0, len(group), limit):
            chunk = group[start:start + limit]
            if current and len(current) + len(chunk) > limit:
                pages.append(current)
                current = []
            current.extend(chunk)
    if current:
        pages.append(current)
    projected = []
    positions = {column['id']: index for index, column in enumerate(columns)}
    for index, page in enumerate(pages, 1):
        page = sorted(page, key=lambda column: positions[column['id']])
        item = dict(matrix)
        item['operators'] = list(dict.fromkeys(column['operator'] for column in page))
        if matrix.get('hierarchy_columns'):
            item['hierarchy_columns'] = page
        item['_stacked_chart_page'] = f'Page {index} of {len(pages)}'
        item['_stacked_chart_maximum'] = max(
            (value['points'] for value in matrix['total']['values'].values() if value['points'] is not None),
            default=matrix['total']['max_points'],
        )
        projected.append(item)
    return projected


def _stacked_chart_page_note(slide, matrix: dict) -> None:
    if matrix.get('_stacked_chart_page'):
        _text(slide, matrix['_stacked_chart_page'], 1.42, left=.65, width=8.65, height=.2, size=9)


def _add_chart_hierarchy_grid(slide, chart_shape, columns: list[dict]) -> Any:
    """Align an editable aggregation grid with the chart's explicit plot area."""
    if not columns:
        return None
    if not columns[0].get('path'):
        columns = [dict(column, path=[{'level': 'Operator', 'value': column['operator']}])
                   for column in columns]
    depth = len(columns[0]['path'])
    chart = chart_shape.chart
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    plot = chart._chartSpace.chart.plotArea
    layout = plot.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}layout')
    if layout is None:
        layout = OxmlElement('c:layout')
        plot.insert(0, layout)
    for child in list(layout):
        layout.remove(child)
    narrow_leaves = chart_shape.width.inches * .89 / len(columns) < .65
    row_heights = [.30 if narrow_leaves and index == depth - 1 else .22 for index in range(depth)]
    grid_height = sum(row_heights)
    chart_shape.height -= Inches(grid_height + .1)
    plot_left = .55 / chart_shape.width.inches
    plot_width = 1 - .75 / chart_shape.width.inches
    manual = OxmlElement('c:manualLayout')
    for name, value in (('layoutTarget', 'inner'), ('xMode', 'edge'), ('yMode', 'edge'),
                        ('wMode', 'factor'), ('hMode', 'factor'),
                        ('x', str(plot_left)), ('y', '.13'), ('w', str(plot_width)), ('h', '.87')):
        element = OxmlElement('c:' + name)
        element.set('val', value)
        manual.append(element)
    layout.append(manual)
    left = chart_shape.left + Inches(.55)
    width = chart_shape.width - Inches(.75)
    top = chart_shape.top + chart_shape.height
    grid = slide.shapes.add_table(depth, len(columns), left, top, width, Inches(grid_height))
    grid.name = 'Scoring Chart Aggregation Grid'
    row_top = top / 914400
    for row_index, depth_index in enumerate(reversed(range(depth))):
        row = grid.table.rows[row_index]
        row.height = Inches(row_heights[depth_index])
        start = 0
        while start < len(columns):
            prefix = columns[start]['path'][:depth_index + 1]
            end = start
            while end + 1 < len(columns) and columns[end + 1]['path'][:depth_index + 1] == prefix:
                end += 1
            cell = grid.table.cell(row_index, start)
            if end > start:
                cell.merge(grid.table.cell(row_index, end))
            value = str(_hierarchy_display_value(columns[start]['path'][depth_index]))
            if narrow_leaves and depth_index == depth - 1 and len(value) == 7 and value[4:6] == '-Q':
                value = value[:5] + '\n' + value[5:]
            _cell(cell, value,
                  color='#EDF3F7' if row_index % 2 == 0 else '#E1EBF1',
                  foreground='#344858', size=8, bold=depth_index == 0)
            start = end + 1
        level = columns[0]['path'][depth_index]['level']
        label = _text(slide, 'Category (0)' if level == 'Category' else level, row_top,
                      left=chart_shape.left.inches, width=.51,
                      height=row_heights[depth_index], size=7, color='#344858', align=PP_ALIGN.RIGHT)
        label.name = 'Scoring Chart Aggregation Level ' + level
        label.text_frame.margin_left = label.text_frame.margin_right = 0
        label.text_frame.word_wrap = False
        row_top += row_heights[depth_index]
    return grid


def _format_stacked_segment_labels(chart) -> None:
    """Use contrasting labels and omit numbers that cannot fit thin segments."""
    chart.plots[0].gap_width = 45
    plot = chart._chartSpace.xpath('.//c:barChart')[0]
    defaults = plot.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}dLbls')
    for series in list(chart.series)[:-1]:
        labels = deepcopy(defaults)
        color = _header_foreground(str(series.format.fill.fore_color.rgb)).lstrip('#')
        for properties in labels.xpath('.//a:defRPr | .//a:rPr'):
            for fill in properties.xpath('./a:solidFill'):
                properties.remove(fill)
            fill = OxmlElement('a:solidFill')
            rgb = OxmlElement('a:srgbClr')
            rgb.set('val', color)
            fill.append(rgb)
            properties.insert(0, fill)
        for index, value in enumerate(series.values):
            if value is not None and value < chart.value_axis.maximum_scale * .04:
                point = OxmlElement('c:dLbl')
                point_index = OxmlElement('c:idx')
                point_index.set('val', str(index))
                delete = OxmlElement('c:delete')
                delete.set('val', '1')
                point.extend((point_index, delete))
                labels.insert(0, point)
        category = series._element.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}cat')
        series._element.insert(series._element.index(category), labels)


def _best_network(presentation, matrices: list[dict]) -> None:
    voice_categories = {'CLASSIC CALLS', 'WHATSAPP CALLS', 'MULTI RAB'}
    for matrix in [page for source in matrices for page in _stacked_chart_pages(source)]:
        slide = _slide(presentation, 'Best Network Scoring per Service', _chart_subtitle(matrix))
        _stacked_chart_page_note(slide, matrix)
        hierarchy_columns = matrix.get('hierarchy_columns', [])
        if hierarchy_columns:
            chart_columns = hierarchy_columns
            operators = list(dict.fromkeys(column['operator'] for column in chart_columns))
        else:
            operators = list(matrix['operators'])
            chart_columns = [
                {'id': operator, 'operator': operator}
                for operator in operators
            ]
        data = CategoryChartData()
        if hierarchy_columns:
            _add_hierarchy_chart_categories(data, hierarchy_columns)
        else:
            data.categories = [_operator_label(matrix, operator) for operator in operators]
        maximums = []
        family_rows = {}
        for family in ('Voice', 'Data'):
            metrics = [row for row in matrix['rows']
                       if (row.get('source_kind') in {'voice', 'speech'}
                           if row.get('source_kind') else row['category'] in voice_categories) == (family == 'Voice')]
            family_rows[family] = metrics
            maximums.append(sum(row['max_points'] or 0 for row in metrics))

        voice_series_count = len(operators)
        for family in ('Voice', 'Data'):
            for operator in operators:
                values = []
                for column in chart_columns:
                    if column['operator'] != operator:
                        values.append(None)
                        continue
                    available = [row['values'][column['id']]['points'] for row in family_rows[family]
                                 if row['values'][column['id']]['points'] is not None]
                    values.append(sum(available) if available else None)
                label = _operator_label(matrix, operator)
                name = f'Voice · {label}' if family == 'Voice' else label
                data.add_series(name, values)
        totals = [matrix['total']['values'][column['id']]['points'] for column in chart_columns]
        data.add_series('Total', totals)
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_STACKED, Inches(.65), Inches(1.75),
            Inches(8.65), Inches(4.25), data,
        )
        chart = chart_shape.chart
        _format_chart(chart, maximum=matrix.get('_stacked_chart_maximum')
                      or max((value for value in totals if value is not None), default=0)
                      or matrix['total']['max_points'], labels=XL_DATA_LABEL_POSITION.CENTER)
        chart.legend.font.size = Pt(8.5)
        if hierarchy_columns:
            chart.category_axis.tick_labels.font.size = Pt(9)
        legend_hidden = set(range(voice_series_count)) | {2 * voice_series_count}
        _hide_chart_legend_entries(chart, legend_hidden)
        for operator_index, operator in enumerate(operators):
            color = _operator_color(matrix, operator).lstrip('#')
            channels = [int(color[index:index + 2], 16) for index in (0, 2, 4)]
            voice_color = '#{:02X}{:02X}{:02X}'.format(
                *(round(channel + (255 - channel) * .45) for channel in channels)
            )
            _fill_series(chart.series[operator_index], voice_color)
            _fill_series(chart.series[voice_series_count + operator_index], color)
        _format_stacked_segment_labels(chart)
        _add_total_labels(chart)
        service_heading = _text(slide, 'Service types', 6.22, left=.65, width=8.65,
                                height=.22, size=9, color='#4A5B65')
        service_heading.name = 'Scoring Chart Service Key Heading'
        service_heading.text_frame.paragraphs[0].font.bold = True
        service_key = _add_category_shade_bar(
            slide, ['Data', 'Voice'], left=Inches(.65), top=Inches(6.48),
            width=Inches(8.65), height=Inches(.35),
        )
        service_key.name = 'Scoring Chart Service Shade Key'
        for index, color in enumerate(('555555', 'C4C4C4')):
            cell = service_key.table.cell(0, index)
            cell.fill.fore_color.rgb = RGBColor.from_string(color)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)
                paragraph.font.color.rgb = RGBColor.from_string(
                    _header_foreground(color).lstrip('#'))
        grid_columns = chart_columns if hierarchy_columns else [
            dict(column, operator=_operator_label(matrix, column['operator'])) for column in chart_columns
        ]
        hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, grid_columns)
        slide.shapes.add_group_shape(
            [chart_shape, service_heading, service_key] + ([hierarchy_grid] if hierarchy_grid is not None else []),
        ).name = 'Scoring Best Network Service Chart'
        _add_family_allocation_donut(slide, matrix)
        _text(slide, matrix['coverage_note'], 7.03, size=9, height=.35)


def _add_total_labels(chart) -> None:
    """Overlay an invisible native line series with labels at each stacked total."""
    plot_area = chart._chartSpace.chart.plotArea
    bar_plot = plot_area.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}barChart')
    series = bar_plot.findall('{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')[-1]
    bar_plot.remove(series)
    for element in list(series):
        if element.tag.endswith('}invertIfNegative'):
            series.remove(element)
    shape_properties = series.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr')
    if shape_properties is not None:
        series.remove(shape_properties)
    shape_properties = OxmlElement('c:spPr')
    line = OxmlElement('a:ln')
    line.append(OxmlElement('a:noFill'))
    shape_properties.append(line)
    series.insert(3, shape_properties)
    marker = series.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}marker')
    if marker is not None:
        series.remove(marker)
    marker = OxmlElement('c:marker')
    symbol = OxmlElement('c:symbol')
    symbol.set('val', 'none')
    marker.append(symbol)
    series.insert(4, marker)
    line_plot = OxmlElement('c:lineChart')
    grouping = OxmlElement('c:grouping')
    grouping.set('val', 'standard')
    line_plot.extend([grouping, series])
    labels = deepcopy(bar_plot.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}dLbls'))
    labels.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}dLblPos').set('val', 't')
    for run_properties in labels.xpath('.//a:defRPr | .//a:rPr'):
        run_properties.set('b', '1')
        run_properties.set('sz', '1200')
        for fill in run_properties.xpath('./a:solidFill'):
            run_properties.remove(fill)
        fill = OxmlElement('a:solidFill')
        color = OxmlElement('a:srgbClr')
        color.set('val', '000000')
        fill.append(color)
        run_properties.insert(0, fill)
    line_plot.append(labels)
    for axis_id in bar_plot.findall('{http://schemas.openxmlformats.org/drawingml/2006/chart}axId'):
        line_plot.append(deepcopy(axis_id))
    plot_area.insert(plot_area.index(bar_plot) + 1, line_plot)


def _hierarchy_chart(presentation, matrix: dict) -> None:
    pages = _stacked_chart_pages(matrix)
    if len(pages) > 1:
        for page in pages:
            _hierarchy_chart(presentation, page)
        return
    columns = matrix.get('hierarchy_columns', [])
    slide = _slide(presentation, 'Best Network Scoring per Category', _chart_subtitle(matrix))
    _stacked_chart_page_note(slide, matrix)
    data = CategoryChartData()
    _add_hierarchy_chart_categories(data, columns)
    categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
    operators = list(dict.fromkeys(column['operator'] for column in columns))
    category_rows = {
        category: [row for row in matrix['rows'] if row['category'] == category]
        for category in categories
    }
    series_specs = []
    for operator in operators:
        for category_index, category in enumerate(categories):
            values = []
            for column in columns:
                if column['operator'] != operator:
                    values.append(None)
                    continue
                cells = [row['values'][column['id']] for row in category_rows[category]]
                available = [cell['points'] for cell in cells if cell['points'] is not None]
                values.append(sum(available) if available else None)
            operator_label = _operator_label(matrix, operator)
            name = operator_label if category_index == 0 else f'{operator_label} · {category}'
            data.add_series(name, values)
            series_specs.append((operator, category_index, category_index > 0))

    data.add_series('TOTAL SCORE', [matrix['total']['values'][column['id']]['points'] for column in columns])

    maximum = max((sum(row['values'][column['id']]['points'] or 0
                       for row in matrix['rows']
                       if row['values'][column['id']]['points'] is not None)
                   for column in columns), default=0.0)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Inches(.55), Inches(1.92), Inches(8.9), Inches(4.48), data,
    )
    chart = chart_shape.chart
    _format_chart(chart, maximum=matrix.get('_stacked_chart_maximum', maximum), labels=XL_DATA_LABEL_POSITION.CENTER)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(8)
    chart.category_axis.tick_labels.font.size = Pt(9)
    hidden_legend_entries = set()
    for series_index, (operator, category_index, hide_from_legend) in enumerate(series_specs):
        series = chart.series[series_index]
        color = _operator_color(matrix, operator)
        _fill_series(series, _hierarchy_chart_color(color, category_index, len(categories)))
        if hide_from_legend:
            hidden_legend_entries.add(series_index)
    _hide_chart_legend_entries(chart, hidden_legend_entries | {len(series_specs)})
    _format_stacked_segment_labels(chart)
    _add_total_labels(chart)

    category_key = _add_category_shade_bar(
        slide, categories, left=Inches(.55), top=Inches(6.46),
        width=Inches(8.9), height=Inches(.4),
    )
    hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, columns)
    chart_group = [chart_shape] + ([category_key] if category_key is not None else []) + ([hierarchy_grid] if hierarchy_grid is not None else [])
    slide.shapes.add_group_shape(chart_group).name = 'Scoring Chart With Category Key'
    _add_category_allocation_donut(slide, matrix)
    _text(slide, matrix['coverage_note'], 7.06, size=8, height=.2)


def _hierarchy_category_comparison_chart(presentation, matrix: dict) -> None:
    """Show each category with readable hierarchy labels and one series per operator."""
    categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
    for category in categories:
        metrics = [row for row in matrix['rows'] if row['category'] == category]
        maximum = max((sum(row['values'][column['id']]['points'] or 0 for row in metrics)
                   for column in matrix['hierarchy_columns']), default=0.0)
        columns = [dict(column, path=[{'level': 'Category', 'value': category}] + column['path'])
                   for column in matrix['hierarchy_columns']]
        operators = list(dict.fromkeys(column['operator'] for column in columns))
        slide = _slide(presentation, 'Scoring per Category', _chart_subtitle(matrix))
        heading = category
        _text(slide, heading, 1.42, left=.7, width=11.9, height=.25, size=12, color='#4A5B65')
        data = CategoryChartData()
        _add_hierarchy_chart_categories(data, columns)
        for operator in operators:
            values = []
            for column in columns:
                points = [row['values'][column['id']]['points'] for row in metrics
                          if row['values'][column['id']]['points'] is not None]
                values.append(sum(points) if points and column['operator'] == operator else None)
            data.add_series(_operator_label(matrix, operator), values)
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(1.8),
            presentation.slide_width - Inches(1.4), Inches(4.85), data,
        )
        chart = chart_shape.chart
        _format_chart(chart, maximum=maximum)
        chart.plots[0].gap_width = 120
        chart.plots[0].overlap = 100
        chart.plots[0].data_labels.font.size = Pt(min(10, max(6, chart_shape.width.inches * .89 * 72 / max(1, len(columns)) / 3.2)))
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.legend.font.size = Pt(11)
        for series, operator in zip(chart.series, operators):
            _fill_series(series, _operator_color(matrix, operator))
        hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, columns)
        slide.shapes.add_group_shape([chart_shape, hierarchy_grid]).name = 'Scoring Hierarchy Category Comparison Chart'
        _text(slide, matrix['coverage_note'], 6.95, size=9)


def _add_category_shade_bar(slide, categories: list[str], *, left: int, top: int, width: int, height: int):
    """Create a native continuous grayscale key bar grouped with its editable chart."""
    if not categories:
        return None
    key_shape = slide.shapes.add_table(1, len(categories), left, top, width, height)
    key_shape.name = 'Scoring Chart Category Shade Key'
    table = key_shape.table
    for index, category in enumerate(categories):
        color = _hierarchy_chart_color('#606060', index, len(categories)).lstrip('#')
        cell = table.cell(0, index)
        cell.text = str(category)
        cell.margin_left = cell.margin_right = Inches(.035)
        cell.margin_top = cell.margin_bottom = 0
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(color)
        cell.text_frame.word_wrap = True
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = _FONT
            paragraph.font.size = Pt(8)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor.from_string(_header_foreground(color).lstrip('#'))
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.line_spacing = 1.0
            paragraph.space_before = paragraph.space_after = 0
    return key_shape


def _hierarchy_chart_color(color: str, category_index: int, category_count: int = 7) -> str:
    """Use lighter shades of each operator color to distinguish stacked KPI groups."""
    channels = [int(color.lstrip('#')[index:index + 2], 16) for index in (0, 2, 4)]
    tint = .72 * category_index / max(1, category_count - 1)
    return '#{:02X}{:02X}{:02X}'.format(*(round(channel + (255 - channel) * tint) for channel in channels))


def _wrapped_line_count(text: Any, width_inches: float, font_size: float) -> int:
    capacity = max(1, int((width_inches * 72 - 3.6) / max(.1, font_size * .65)))
    lines = 1
    used = 0
    for word in str(text or '').split():
        word_length = len(word)
        if used and used + 1 + word_length > capacity:
            lines += 1
            used = 0
        if word_length > capacity:
            extra_lines = ceil(word_length / capacity) - 1
            lines += extra_lines
            used = word_length % capacity
        else:
            used += (1 if used else 0) + word_length
    return lines


def _hierarchy_metric_layout(metrics: list[dict], total_label: str | None, *, available_height: float,
                             kpi_width: float) -> tuple[float, list[float]]:
    texts = [str(row.get('kpi', '')) for row in metrics]
    if total_label is not None:
        texts.append(total_label)
    font_size = min(7.5, max(2.5, available_height * 72 / max(1, len(texts)) * .82))
    line_counts = [_wrapped_line_count(text, kpi_width, font_size) for text in texts]
    for _ in range(3):
        base_height = available_height / max(1, sum(line_counts))
        font_size = min(7.5, max(2.5, base_height * 72 * .82))
        updated = [_wrapped_line_count(text, kpi_width, font_size) for text in texts]
        if updated == line_counts:
            break
        line_counts = updated
    base_height = available_height / max(1, sum(line_counts))
    return font_size, [base_height * lines for lines in line_counts]


def _hierarchy_content_font(texts: list[str], width_inches: float, *, maximum: float = 7.5) -> float:
    glyph_units = {
        '.': .28, '-': .4, '+': .48, '*': .36,
    }
    widest = max(
        (sum(glyph_units.get(char, .56 if char.isdigit() else .52) for char in str(text)) for text in texts),
        default=1.0,
    )
    usable_width = max(1.0, width_inches * 72 - 3.6)
    return min(maximum, max(2.5, usable_width / max(1.0, widest)))


def _add_hierarchy_chart_categories(data: CategoryChartData, columns: list[dict]) -> None:
    """Build native multi-level chart categories from hierarchy values."""
    nodes = {}
    for column in columns:
        parent = None
        prefix = ()
        for item in column.get('path', []):
            value = _hierarchy_display_value(item)
            prefix += (item.get('value'),)
            if prefix in nodes:
                parent = nodes[prefix]
                continue
            node = data.add_category(value) if parent is None else parent.add_sub_category(value)
            nodes[prefix] = node
            parent = node


def _hierarchy_header_groups(table, columns: list[dict], levels: list[str], *, start_col: int,
                             leaf_width: int, header_rows: int, physical_column_width: float,
                             leaf_label: str = 'Score', header_color: str | None = None,
                             gap_reference: str | None = None) -> None:
    for level_index, level in enumerate(levels):
        start = 0
        while start < len(columns):
            representative = columns[start]
            prefix = tuple((entry['level'], entry.get('value')) for entry in representative['path'][:level_index + 1])
            end = start
            while end + 1 < len(columns):
                next_prefix = tuple(
                    (entry['level'], entry.get('value')) for entry in columns[end + 1]['path'][:level_index + 1]
                )
                if next_prefix != prefix:
                    break
                end += 1
            left_column = start_col + start * leaf_width
            right_column = start_col + (end + 1) * leaf_width - 1
            cell = table.cell(level_index, left_column)
            if right_column > left_column:
                cell.merge(table.cell(level_index, right_column))
            label = _hierarchy_display_value(representative['path'][level_index])
            if level == 'Operator' and gap_reference:
                label = f'{label} − {gap_reference}'
            color = representative['color'] if level == 'Operator' else (header_color or '#E6ECFA')
            header_size = min(7.5, max(5.0, physical_column_width * leaf_width
                                      * (end - start + 1) * 12))
            _cell(cell, str(label), color=color, foreground=_header_foreground(color), size=header_size, bold=True)
            start = end + 1
    subheader_row = len(levels)
    for column_index, column in enumerate(columns):
        color = column['color']
        leaf_font_size = min(7.5, max(5.0, physical_column_width * 12))
        _cell(table.cell(subheader_row, start_col + column_index * leaf_width), leaf_label,
              color=color, foreground=_header_foreground(color), size=leaf_font_size, bold=True)
        if leaf_width == 2:
            _cell(table.cell(subheader_row, start_col + column_index * leaf_width + 1), 'GAP',
                  color='#FFFF00', foreground='#17232D', size=leaf_font_size, bold=True)


def _score_column_plan(columns: list[dict], gap_layout: str, *, show_gap_values: bool = True) -> list[tuple[dict, str]]:
    scores = [(column, 'Score') for column in columns]
    if not show_gap_values:
        return scores
    gaps = [(column, 'GAP') for column in columns if not column.get('is_reference')]
    if gap_layout == 'end':
        return scores + gaps
    return [item for column in columns for item in (
        [(column, 'Score')] if column.get('is_reference') else [(column, 'Score'), (column, 'GAP')]
    )]


def _score_column_headers(table, plan: list[tuple[dict, str]], levels: list[str],
                          *, gap_layout: str, leaf_width: float, fixed_count: int = 5) -> None:
    groups = [('Score', 0, sum(kind == 'Score' for _, kind in plan)),
              ('GAP', sum(kind == 'Score' for _, kind in plan), len(plan))]
    if gap_layout == 'adjacent' and any(kind == 'GAP' for _, kind in plan):
        groups = [('Score / GAP', 0, len(plan))]
    for label, start, end in groups:
        if start == end:
            continue
        cell = table.cell(0, fixed_count + start)
        if end - start > 1:
            cell.merge(table.cell(0, fixed_count + end - 1))
        _cell(cell, label, color='#FFFF00' if label == 'GAP' else '#455B65',
              foreground='#17232D' if label == 'GAP' else _WHITE, size=7, bold=True)
    for depth, level in enumerate(levels):
        start = 0
        while start < len(plan):
            column, kind = plan[start]
            prefix = tuple((entry['level'], entry.get('value')) for entry in column['path'][:depth + 1])
            end = start
            while end + 1 < len(plan):
                next_column, next_kind = plan[end + 1]
                next_prefix = tuple((entry['level'], entry.get('value'))
                                    for entry in next_column['path'][:depth + 1])
                if next_prefix != prefix or (gap_layout == 'end' and next_kind != kind):
                    break
                end += 1
            cell = table.cell(depth + 1, fixed_count + start)
            if end > start:
                cell.merge(table.cell(depth + 1, fixed_count + end))
            value = _hierarchy_display_value(column['path'][depth])
            color = column['color'] if level == 'Operator' else '#E6ECFA'
            size = min(7, max(2.5, leaf_width * (end - start + 1) * 12))
            _cell(cell, str(value) if value is not None else 'Not specified', color=color,
                  foreground=_header_foreground(color), size=size, bold=True)
            start = end + 1
    for index, (column, kind) in enumerate(plan, fixed_count):
        color = '#FFFF00' if kind == 'GAP' else column['color']
        cell = table.cell(len(levels) + 1, index)
        _cell(cell, kind, color=color, foreground=_header_foreground(color),
              size=min(7, max(2.5, leaf_width * 12)), bold=True)


def _hierarchy_score_tables(presentation, matrices: list[dict], legend: list[dict], *, gap_layout: str = 'end',
                            show_gap_values: bool = True, title: str = 'Scoring Tables — Breakdown') -> None:
    for matrix in matrices:
        is_summary = matrix.get('table_mode') == 'summary'
        fixed_count = 4 if is_summary else 5
        weight_index = 2 if is_summary else 3
        maximum_index = weight_index + 1
        columns = matrix['hierarchy_columns']
        levels = matrix['hierarchy_levels']
        metrics = matrix['rows']
        plan = _score_column_plan(columns, gap_layout, show_gap_values=show_gap_values)
        total_label = _score_total_label(
            [*metrics, matrix['total']],
            [column['id'] for column, kind in plan if kind == 'Score'],
            [column['id'] for column, kind in plan if kind == 'GAP'],
            show_gap_values=show_gap_values,
        )
        slide = _slide(presentation, title, _chart_subtitle(matrix))
        header_rows = len(levels) + 2
        row_count = header_rows + len(metrics) + 1
        table_height = 5.0
        table_width = presentation.slide_width / 914400 - 1.1
        fixed_widths = [0.9, 2.15, *([] if is_summary else [0.65]), 0.47, 0.52]
        table = slide.shapes.add_table(row_count, fixed_count + len(plan), Inches(.55), Inches(1.5),
                                       Inches(table_width), Inches(table_height)).table
        leaf_width = (table_width - sum(fixed_widths)) / max(1, len(plan))
        for index, column in enumerate(table.columns):
            column.width = Inches(fixed_widths[index] if index < fixed_count else leaf_width)
        header_level_height = min(.23, 1.05 / max(1, len(levels)))
        table.rows[0].height = Inches(.22)
        for row_index in range(1, len(levels) + 1):
            table.rows[row_index].height = Inches(header_level_height)
        table.rows[header_rows - 1].height = Inches(.22)
        available_body_height = table_height - header_level_height * len(levels) - .44
        metric_font, body_heights = _hierarchy_metric_layout(
            metrics, total_label,
            available_height=available_body_height, kpi_width=fixed_widths[1],
        )
        for row_index, row_height in enumerate(body_heights, header_rows):
            table.rows[row_index].height = Inches(row_height)
        _score_column_headers(table, plan, levels, gap_layout=gap_layout, leaf_width=leaf_width,
                              fixed_count=fixed_count)
        static_headers = [('CATEGORY', _GAP_SUMMARY_HEADER),
                          ('CATEGORY' if is_summary else 'KPI', _GAP_SUMMARY_HEADER),
                          *([] if is_summary else [('Type of KPI', _GAP_SUMMARY_HEADER)]),
                          ('Score weight\n(%)', '#4EA72E'), ('Max score', '#FF0000')]
        for index, (label, color) in enumerate(static_headers):
            cell = table.cell(0, index)
            cell.merge(table.cell(header_rows - 1, index))
            _cell(cell, label, color=color, foreground=_header_foreground(color), size=7, bold=True)
        data_font = min(7, metric_font)
        total = matrix['total']
        numeric_texts = []
        for metric in metrics + [total]:
            for column, kind in plan:
                value = metric['values'][column['id']]
                numeric_texts.append(_gap_number(metric, column['id']) if kind == 'GAP'
                                     else _number(value['points']) + ('*' if value['points'] is not None and not value['complete'] else ''))
        numeric_font = _hierarchy_content_font(
            numeric_texts, leaf_width, maximum=min(11, min(body_heights) * 72 * .82),
        )
        category_color = _GAP_SUMMARY_CATEGORY if is_summary else _SCORING_EXPANDED_CATEGORY
        for row_offset, metric in enumerate(metrics, header_rows):
            subtotal = metric.get('row_type') == 'category'
            row_color = ('#E3E6E7' if is_summary else category_color) if subtotal else None
            best, worst = (_summary_extreme_columns(
                metric.get('values', {}), [column for column, kind in plan if kind == 'Score'],
            ) if matrix.get('table_mode') == 'summary' and subtotal else (set(), set()))
            _cell(table.cell(row_offset, 0), metric['category'], color=category_color,
                  size=data_font, left=True, bold=True)
            _cell(table.cell(row_offset, 1), metric['kpi'], color=row_color or '#E7E8E9',
                  bold=subtotal, left=True, size=metric_font)
            if not is_summary:
                _cell(table.cell(row_offset, 2), metric.get('kpi_type', ''),
                      color=row_color or _kpi_type_color(metric), size=data_font, bold=subtotal)
            _cell(table.cell(row_offset, weight_index), _number(metric['weight_percent']) + '%', color=row_color or _WHITE,
                  size=data_font, bold=subtotal)
            _cell(table.cell(row_offset, maximum_index), _number(metric['max_points']), color=row_color or _WHITE,
                  size=data_font, bold=subtotal)
            for index, (column, kind) in enumerate(plan, fixed_count):
                leaf_id = column['id']
                value = metric['values'][leaf_id]
                if kind == 'GAP':
                    _cell(table.cell(row_offset, index), _gap_number(metric, leaf_id),
                          color=row_color or metric.get('gap_colors', {}).get(leaf_id, THRESHOLD_COLORS['Unavailable']),
                          size=numeric_font, bold=True)
                else:
                    partial = value['points'] is not None and not value['complete']
                    score_color = (_SUMMARY_BEST_COLOR if leaf_id in best else
                                   _SUMMARY_WORST_COLOR if leaf_id in worst else
                                   row_color or value.get('color', _NEUTRAL))
                    _cell(table.cell(row_offset, index), _number(value['points']) + ('*' if partial else ''),
                          color=score_color, size=numeric_font, bold=subtotal)
        total_index = row_count - 1
        total_color = '#D8DFE4' if is_summary else category_color
        best, worst = (_summary_extreme_columns(
            total.get('values', {}), [column for column, kind in plan if kind == 'Score'],
        ) if is_summary else (set(), set()))
        for index, label in enumerate(('TOTAL', total_label, *([] if is_summary else ['']),
                                      _number(total['weight_percent']) + '%', _number(total['max_points']))):
            _cell(table.cell(total_index, index), label, color=total_color, bold=True, left=index < 2, size=data_font)
        for index, (column, kind) in enumerate(plan, fixed_count):
            value = total['values'][column['id']]
            text = (_gap_number(total, column['id']) if kind == 'GAP' else
                    _number(value['points']) + ('*' if value['points'] is not None and not value['complete'] else ''))
            color = (_SUMMARY_BEST_COLOR if kind == 'Score' and column['id'] in best else
                     _SUMMARY_WORST_COLOR if kind == 'Score' and column['id'] in worst else total_color)
            _cell(table.cell(total_index, index), text, color=color, size=numeric_font, bold=True)
        _merge_category_cells(table, metrics, header_rows,
                              color=category_color, include_subtotals=True, bold_categories=True)
        if is_summary:
            _remove_summary_category_column(table)
            _add_summary_operator_legend(slide)
        else:
            for index, item in enumerate(legend):
                legend_table = slide.shapes.add_table(
                    1, 1, Inches(6.8 + index * 1.15), Inches(1.25), Inches(1.1), Inches(.18),
                ).table
                _cell(legend_table.cell(0, 0), item.get('band', ''), color=item['color'], size=8)
        _text(slide, matrix['coverage_note'], 7.27, size=8, height=.18)


def _remove_summary_category_column(table) -> None:
    """Retain category total labels in a single summary identity column."""
    grid = table._tbl.tblGrid
    removed_width = int(grid[0].get('w'))
    grid.remove(grid[0])
    grid[0].set('w', str(int(grid[0].get('w')) + removed_width))
    for row in table._tbl.tr_lst:
        row.remove(row.tc_lst[0])


def _merge_category_cells(table, metrics: list[dict], start_row: int,
                          *, color: str = '#E6F0F7', include_subtotals: bool = False,
                          bold_categories: bool = False,
                          alternate_colors: tuple[str, str] | None = None) -> None:
    first = 0
    run_index = 0
    while first < len(metrics):
        last = first
        while (last + 1 < len(metrics) and metrics[last + 1]['category'] == metrics[first]['category']
               and (include_subtotals or metrics[last + 1].get('row_type') != 'category')):
            last += 1
        cell = table.cell(start_row + first, 0)
        run_color = alternate_colors[run_index % 2] if alternate_colors else color
        if alternate_colors:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(run_color.lstrip('#'))
        if last > first:
            cell.merge(table.cell(start_row + last, 0))
            subtotal = metrics[first].get('row_type') == 'category'
            _cell(cell, metrics[first]['category'], color=run_color,
                  size=7, left=True, bold=bold_categories or subtotal)
        run_index += 1
        first = last + 1


def _hierarchy_gap_projection(matrix: dict, columns: list[dict]) -> dict:
    """Keep a selected set of leaves in a hierarchy GAP view."""
    projected = deepcopy(matrix)
    leaf_ids = [column['id'] for column in columns]
    projected['hierarchy_columns'] = deepcopy(columns)
    projected['operators'] = leaf_ids
    projected['operator_styles'] = {
        leaf_id: projected['operator_styles'][leaf_id]
        for leaf_id in leaf_ids if leaf_id in projected.get('operator_styles', {})
    }
    projected['rows'] = [
        row | {
            'gaps': {leaf_id: row.get('gaps', {}).get(leaf_id) for leaf_id in leaf_ids},
            'gap_colors': {leaf_id: row.get('gap_colors', {}).get(leaf_id, THRESHOLD_COLORS['Unavailable'])
                           for leaf_id in leaf_ids},
        }
        for row in projected.get('rows', [])
    ]
    compared_ids = [column['id'] for column in columns if not column.get('is_reference')]
    if len({column['operator'] for column in columns if not column.get('is_reference')}) == 1:
        if projected.get('table_mode') == 'summary':
            projected['rows'].sort(key=lambda row: _gap_order_key(row, compared_ids))
        else:
            kpis = [row for row in projected['rows'] if row.get('row_type') != 'category']
            projected['rows'] = sorted(kpis, key=lambda row: _gap_order_key(row, compared_ids))
    projected['total'] = {
        'gaps': {leaf_id: matrix.get('total', {}).get('gaps', {}).get(leaf_id) for leaf_id in leaf_ids},
        'gap_partial': {leaf_id: matrix.get('total', {}).get('gap_partial', {}).get(leaf_id, False)
                        for leaf_id in leaf_ids},
        'gap_environments': {leaf_id: matrix.get('total', {}).get('gap_environments', {}).get(leaf_id, [])
                             for leaf_id in leaf_ids},
    }
    return projected


def _hierarchy_gap_tables(presentation, matrices: list[dict], *,
                          title: str = 'GAP Analysis — All vs reference',
                          show_priority: bool = False) -> None:
    for matrix in matrices:
        columns = matrix['hierarchy_columns']
        levels = matrix['hierarchy_levels']
        metrics = [row for row in matrix['rows'] if row.get('row_type') != 'category']
        slide = _slide(presentation, title, _chart_subtitle(matrix))
        if not title.startswith('GAP Analysis — All vs '):
            _accent_gap_comparison_title(slide)
        header_rows = len(levels) + 1
        row_count = header_rows + len(metrics)
        table_height = 4.65
        table_width = 9.3 if show_priority else 12.03
        fixed_widths = [1.0, 2.55, .7]
        table = slide.shapes.add_table(row_count, 3 + len(columns), Inches(.65), Inches(1.55),
                                       Inches(table_width), Inches(table_height)).table
        leaf_width = (table_width - sum(fixed_widths)) / max(1, len(columns))
        for index, column in enumerate(table.columns):
            column.width = Inches(fixed_widths[index] if index < 3 else leaf_width)
        header_level_height = min(.25, 1.05 / max(1, len(levels)))
        subheader_height = .25
        for row_index in range(len(levels)):
            table.rows[row_index].height = Inches(header_level_height)
        table.rows[len(levels)].height = Inches(subheader_height)
        available_body_height = table_height - header_level_height * len(levels) - subheader_height
        metric_font, body_heights = _hierarchy_metric_layout(
            metrics, None, available_height=available_body_height,
            kpi_width=fixed_widths[1],
        )
        for row_index, row_height in enumerate(body_heights, header_rows):
            table.rows[row_index].height = Inches(row_height)
        _hierarchy_header_groups(table, columns, levels, start_col=3, leaf_width=1,
                                 header_rows=header_rows, physical_column_width=leaf_width,
                                 leaf_label='GAP', gap_reference=_operator_label(matrix, matrix['baseline_operator']),
                                 header_color=_GAP_SUMMARY_COMPARISON_HEADER)
        data_font = min(7, metric_font)
        for index, label in enumerate(('Category', 'KPI', 'Type of KPI')):
            cell = table.cell(0, index)
            if header_rows > 1:
                cell.merge(table.cell(header_rows - 1, index))
            _cell(cell, label, color=_GAP_SUMMARY_HEADER,
                  foreground=_WHITE, size=7, bold=True)
        gap_texts = [
            _gap_number(row, column['id'])
            for row in metrics for column in columns
        ]
        leaf_font = _hierarchy_content_font(gap_texts, leaf_width, maximum=data_font)
        for row_offset, row in enumerate(metrics, header_rows):
            subtotal = row.get('row_type') == 'category'
            row_color = _GAP_SUMMARY_CATEGORY if subtotal else None
            _cell(table.cell(row_offset, 0), row['category'], color=row_color or _GAP_SUMMARY_CATEGORY,
                  foreground='#17232D', size=data_font, bold=True, left=True)
            _cell(table.cell(row_offset, 1), row['kpi'], color=row_color or '#E6ECFA', size=metric_font, left=True,
                  bold=subtotal)
            _cell(table.cell(row_offset, 2), row['kpi_type'],
                  color=row_color or _kpi_type_color(row),
                  size=data_font, bold=subtotal)
            for column_index, column in enumerate(columns, 3):
                leaf_id = column['id']
                _cell(table.cell(row_offset, column_index), _gap_number(row, leaf_id),
                      color=row_color or row['gap_colors'].get(leaf_id, THRESHOLD_COLORS['Unavailable']),
                      size=leaf_font, bold=True)
        _merge_category_cells(table, metrics, header_rows, include_subtotals=True,
                              bold_categories=True, alternate_colors=('#E6F0F7', '#D7E5EE'))
        _text(slide, matrix['note'], 7.05, height=.25, size=9)
        if show_priority:
            _add_individual_gap_notes(slide, matrix, metrics)
            _add_gap_priority_arrow(slide, table, header_rows=header_rows)
        _add_gap_color_scale(slide, matrix, left=.65, width=table_width)


def _gap_summary_tables(presentation, matrices: list[dict]) -> None:
    """Export the combined comparison before individual operator GAP slides."""
    for matrix in matrices:
        operators = matrix['operators']
        pages = [[row for row in matrix['rows'] if row.get('row_type') != 'category']]
        for page_index, rows in enumerate(pages):
            page_label = f' · Page {page_index + 1}/{len(pages)}' if len(pages) > 1 else ''
            reference = _operator_label(matrix, matrix['baseline_operator'])
            slide = _slide(presentation, 'GAP Analysis — All vs ' + reference,
                           _chart_subtitle(matrix) + page_label)
            table = slide.shapes.add_table(len(rows) + 1, len(operators) + 3, Inches(.65), Inches(1.65),
                                           Inches(12.03), Inches(4.55)).table
            widths = [1.45, 4.05, 1.1] + [5.43 / max(1, len(operators))] * len(operators)
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            headers = ['Category', 'NETCHECK KPIs', 'Type of KPI'] + [
                f'{_operator_label(matrix, operator)} − {reference}' for operator in operators]
            table.rows[0].height = Inches(.43)
            metric_font, body_heights = _hierarchy_metric_layout(
                rows, None, available_height=4.55 - .43, kpi_width=4.05,
            )
            for row, height in zip(list(table.rows)[1:], body_heights):
                row.height = Inches(height)
            for index, header in enumerate(headers):
                color = _operator_color(matrix, operators[index - 3]) if index >= 3 else _GAP_SUMMARY_HEADER
                _cell(table.cell(0, index), header, color=color,
                      foreground=_header_foreground(color), size=10, bold=True)
            for index, row in enumerate(rows, 1):
                subtotal = row.get('row_type') == 'category'
                row_color = _GAP_SUMMARY_CATEGORY if subtotal else None
                _cell(table.cell(index, 0), row['category'], color=row_color or _GAP_SUMMARY_CATEGORY,
                      foreground='#17232D', size=metric_font, bold=True, left=True)
                _cell(table.cell(index, 1), row['kpi'], color=row_color or '#E6ECFA', size=metric_font,
                      left=True, bold=subtotal)
                _cell(table.cell(index, 2), row['kpi_type'],
                      color=row_color or _kpi_type_color(row),
                      size=metric_font, bold=subtotal)
                for column, operator in enumerate(operators, 3):
                    _cell(table.cell(index, column), _gap_number(row, operator),
                          color=row_color or row['gap_colors'][operator], size=metric_font, bold=True)
            _text(slide, matrix['note'], 7.12, height=.25, size=9)
            _merge_category_cells(table, rows, 1, include_subtotals=True, bold_categories=True,
                                  alternate_colors=('#E6F0F7', '#D7E5EE'))
            _add_gap_color_scale(slide, matrix, left=.65, width=12.03)


def _gap_tables(presentation, matrices: list[dict]) -> None:
    for matrix in matrices:
        pages = [[row for row in matrix['rows'] if row.get('row_type') != 'category']]
        for page_index, rows in enumerate(pages):
            page_label = f' · Page {page_index + 1}/{len(pages)}' if len(pages) > 1 else ''
            comparison = f'{_operator_label(matrix, matrix["operator"])} vs {_operator_label(matrix, matrix["baseline_operator"])}'
            subtitle = _chart_subtitle(matrix) + page_label
            slide = _slide(presentation, 'GAP Analysis — ' + comparison, subtitle)
            _accent_gap_comparison_title(slide)
            if not rows:
                _text(slide, 'No comparable KPI gaps are available. See the scoring matrix for missing values.', 1.8, height=1, size=16)
                continue
            table = slide.shapes.add_table(len(rows) + 1, 4, Inches(.65), Inches(1.65), Inches(9.3),
                                           Inches(4.55)).table
            for column, width in zip(table.columns, (1.65, 4.45, 1.5, 1.7)):
                column.width = Inches(width)
            table.rows[0].height = Inches(.43)
            metric_font, body_heights = _hierarchy_metric_layout(
                rows, None, available_height=4.55 - .43, kpi_width=4.45,
            )
            for row, height in zip(list(table.rows)[1:], body_heights):
                row.height = Inches(height)
            headers = ['Category', 'NETCHECK KPIs', 'Type of KPI',
                       f'{_operator_label(matrix, matrix["operator"])} − {_operator_label(matrix, matrix["baseline_operator"])}']
            for index, header in enumerate(headers):
                color = _operator_color(matrix, matrix['operator']) if index == 3 else _GAP_SUMMARY_HEADER
                _cell(table.cell(0, index), header, color=color,
                      foreground=_header_foreground(color), size=11, bold=True)
            for index, row in enumerate(rows, 1):
                subtotal = row.get('row_type') == 'category'
                row_color = _GAP_SUMMARY_CATEGORY if subtotal else None
                _cell(table.cell(index, 0), row['category'], color=row_color or _GAP_SUMMARY_CATEGORY,
                      foreground='#17232D', size=metric_font, bold=True, left=True)
                _cell(table.cell(index, 1), row['kpi'], color=row_color or '#E6ECFA', size=metric_font,
                      left=True, bold=subtotal)
                _cell(table.cell(index, 2), row['kpi_type'],
                      color=row_color or _kpi_type_color(row),
                      size=metric_font, bold=subtotal)
                _cell(table.cell(index, 3), _gap_number(row), color=row_color or row.get('gap_color', '#FFFF80'),
                      size=metric_font, bold=True)
            _add_individual_gap_notes(slide, matrix, rows)
            _merge_category_cells(table, rows, 1, include_subtotals=True, bold_categories=True,
                                  alternate_colors=('#E6F0F7', '#D7E5EE'))
            _add_gap_priority_arrow(slide, table)
            _add_gap_color_scale(slide, matrix, left=.65, width=9.3)
            _text(slide, 'KPIs are ordered by GAP, from lowest to highest. GAP Priority does not affect this order.', 7.12, height=.25, size=9)


def export_scoring_powerpoint(job: dict[str, Any], result: dict[str, Any], template_path: Path,
                             operator_mapping_groups: list[dict[str, Any]] | None = None,
                             *, table_mode: str = 'expanded', gap_layout: str = 'end',
                             environment: str = 'all', show_gap_values: bool = True) -> bytes:
    """Export saved points as comparison matrices, charts and prioritized gaps."""
    if not template_path.is_file():
        raise ValueError('The configured CDR PowerPoint template is missing.')
    if table_mode not in {'expanded', 'summary'}:
        raise ValueError('Table mode must be expanded or summary.')
    if gap_layout not in {'end', 'adjacent'}:
        raise ValueError('GAP layout must be end or adjacent.')
    job = dict(job)
    if '_scoring_display_selections' not in job:
        job['_scoring_display_selections'] = prepare_scoring_display_selections(job, result, template_path)
    presentation = Presentation(template_path)
    _remove_all_slides(presentation)
    views = _export_environment_views(build_scoring_views(job, result, operator_mapping_groups), environment)
    configuration = job.get('configuration') or result.get('configuration') or {}
    for matrix_key in ('score_tables', 'gap_summary_tables', 'gap_tables',
                       'hierarchy_score_tables', 'hierarchy_gap_tables'):
        for matrix in views.get(matrix_key, []):
            matrix['environment_subtitle'] = _scoring_environment_title(
                matrix['context']['environment'], configuration,
            )
    environment_allocations = maximum_allocations_from_configuration(configuration)
    for matrix_key in ('score_tables', 'hierarchy_score_tables'):
        for matrix in views.get(matrix_key, []):
            matrix['environment_allocations'] = environment_allocations
            matrix['environment_labels'] = {
                name: str(scope.get('display_name') or _environment_display_label(name))
                for name, scope in configuration.get('scope', {}).get('environments', {}).items()
            }
    _add_scoring_intro_slides(presentation, job, result)
    subtitle = 'All Environments' if environment == 'all' else _scoring_environment_title(environment, configuration)
    matrices = views['score_tables']
    hierarchy_matrices = views.get('hierarchy_score_tables', [])
    use_hierarchy = bool(hierarchy_matrices) and len(hierarchy_matrices[0].get('hierarchy_levels', [])) > 1
    if use_hierarchy:
        gap_by_environment = {
            matrix['context'].get('environment'): matrix
            for matrix in views.get('hierarchy_gap_tables', [])
        }
        grouped_hierarchy: dict[str, list[dict[str, Any]]] = {}
        for matrix in hierarchy_matrices:
            grouped_hierarchy.setdefault(str(matrix['context'].get('environment') or 'Unspecified'), []).append(matrix)
        for current_environment, environment_matrices in grouped_hierarchy.items():
            _add_scoring_intro_slides(presentation, job, result, environment=current_environment)
            for matrix in environment_matrices:
                _best_network(presentation, [matrix])
                _hierarchy_chart(presentation, matrix)
                _hierarchy_category_comparison_chart(presentation, matrix)
                for mode, title in (('summary', 'Scoring Tables — Summary'),
                                    ('expanded', 'Scoring Tables — Breakdown')):
                    _hierarchy_score_tables(
                        presentation, [_table_for_mode(matrix, mode)], views.get('threshold_legend', []),
                        gap_layout=gap_layout, show_gap_values=show_gap_values, title=title,
                    )
                gap_matrix = gap_by_environment.get(matrix['context'].get('environment'))
                if gap_matrix is None:
                    continue
                # Finish the full comparison block before advancing to the next environment.
                reference = _operator_label(gap_matrix, gap_matrix['baseline_operator'])
                compared_columns = [column for column in gap_matrix['hierarchy_columns']
                                    if not column.get('is_reference')]
                if compared_columns:
                    all_matrix = _hierarchy_gap_projection(_table_for_mode(gap_matrix, 'expanded'), compared_columns)
                    _hierarchy_gap_tables(
                        presentation, [all_matrix], title=f'GAP Analysis — All vs {reference}',
                    )
                compared_operators = dict.fromkeys(
                    column['operator'] for column in gap_matrix['hierarchy_columns']
                    if not column.get('is_reference')
                )
                for operator in compared_operators:
                    operator_columns = [column for column in gap_matrix['hierarchy_columns']
                                        if column['operator'] == operator]
                    operator_matrix = _hierarchy_gap_projection(_table_for_mode(gap_matrix, 'expanded'), operator_columns)
                    _hierarchy_gap_tables(
                        presentation, [operator_matrix],
                        title=f'GAP Analysis — {operator} vs {reference}', show_priority=True,
                    )
    elif matrices:
        grouped_scores: dict[str, list[dict[str, Any]]] = {}
        for matrix in matrices:
            grouped_scores.setdefault(str(matrix['context'].get('environment') or 'Unspecified'), []).append(matrix)
        for current_environment, environment_matrices in grouped_scores.items():
            _add_scoring_intro_slides(presentation, job, result, environment=current_environment)
            for matrix in environment_matrices:
                _best_network(presentation, [matrix])
                _stacked_category_chart(presentation, matrix)
                _charts(presentation, [matrix])
                for mode, title in (('summary', 'Scoring Tables — Summary'),
                                    ('expanded', 'Scoring Tables — Breakdown')):
                    _score_tables(
                        presentation, [_table_for_mode(matrix, mode)], views.get('threshold_legend', []),
                        gap_layout=gap_layout, show_gap_values=show_gap_values, title=title,
                    )
                _gap_summary_tables(presentation, [_table_for_mode(table, 'expanded') for table in views['gap_summary_tables'] if table['context'] == matrix['context']])
                _gap_tables(presentation, [_table_for_mode(table, 'expanded') for table in views['gap_tables'] if table['context'] == matrix['context']])
    else:
        slide = _slide(presentation, 'Scoring Tables', subtitle)
        _text(slide, 'No scoring measurements are available for this saved job.', 1.8, height=1, size=16)
    slide_warnings = [str(warning) for warning in result.get('warnings', [])
                      if str(warning) != _LEGACY_CAMPAIGN_WARNING]
    for slide in presentation.slides:
        slide.notes_slide.notes_text_frame.text = '\n'.join(slide_warnings)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
