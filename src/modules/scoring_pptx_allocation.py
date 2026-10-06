"""Editable maximum-allocation doughnuts used in Scoring PowerPoint exports."""
from __future__ import annotations

import math
import re
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


_GLOBAL_COLORS = ('176E77', 'E6A81D', 'C55A11', '5B9BD5', 'A64D79')
_VOICE_COLOR = '4472C4'
_DATA_COLOR = '7030A0'
_CATEGORY_COLORS = ('4472C4', '7030A0', 'C55A11', '5B9BD5', 'A64D79', '548235', 'D65F8D', '8064A2')


def _number(value: Any) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0.0
    return number if math.isfinite(number) and number > 0 else 0.0


def _legend_points(value: float) -> str:
    """Match the web legend's maximum of three fractional digits."""
    return f'{value:.3f}'.rstrip('0').rstrip('.')


def _metric_source_kind(metric: dict[str, Any]) -> str:
    source_kind = metric.get('source_kind')
    if source_kind in (None, '') and isinstance(metric.get('calculation'), dict):
        source_kind = metric['calculation'].get('source_kind')
    return str(source_kind or '').strip().casefold()


def _add_percentage_labels(slide, values, colors, chart_left, chart_top, chart_size: float, hole: int) -> None:
    """Add editable native text labels over doughnut slices that have room."""
    total = sum(_number(value) for value in values)
    if total <= 0:
        return
    outer_radius = chart_size * .45
    inner_radius = outer_radius * hole / 100
    middle_radius = (outer_radius + inner_radius) / 2
    center_x = chart_left + chart_size / 2
    center_y = chart_top + chart_size / 2
    angle = 0.0
    for value, color in zip(values, colors):
        share = _number(value) / total
        label_text = f'{share * 100:.1f}%'
        span = share * 2 * math.pi
        midpoint = math.radians(angle + share * 180)
        center_dx = middle_radius * math.sin(midpoint)
        center_dy = -middle_radius * math.cos(midpoint)
        padding = 1 / 72
        fitted = None
        for font_size in (8.5, 8.0, 7.5, 7.0, 6.5):
            label_height = (font_size + 4) / 72
            label_width_inches = max(.3, (len(label_text) * font_size * .62 + 6) / 72)
            # Arial Bold advances differ substantially between digits, the
            # decimal point and the percent sign. A uniform character width
            # rejects labels that fit the narrow inner allocation band.
            glyph_width = sum(.556 if character.isdigit() else
                              .278 if character == '.' else .889
                              for character in label_text) * font_size / 72
            glyph_height = font_size * 1.2 / 72
            if middle_radius * span < glyph_width + 2 * padding:
                continue
            half_width, half_height = glyph_width / 2, glyph_height / 2
            closest_radius = math.hypot(max(abs(center_dx) - half_width, 0),
                                        max(abs(center_dy) - half_height, 0))
            if closest_radius < inner_radius + padding:
                continue
            corners = ((center_dx + x, center_dy + y)
                       for x in (-half_width, half_width) for y in (-half_height, half_height))
            if all(inner_radius + padding <= math.hypot(x, y) <= outer_radius - padding
                   and abs(math.atan2(math.sin(math.atan2(x, -y) - midpoint),
                                      math.cos(math.atan2(x, -y) - midpoint)))
                   <= span / 2 - padding / middle_radius
                   for x, y in corners):
                fitted = (font_size, label_height, label_width_inches)
                break
        if fitted is None:
            angle += share * 360
            continue
        font_size, label_height, label_width_inches = fitted
        x = center_x + center_dx - label_width_inches / 2
        y = center_y + center_dy - label_height / 2
        rgb = tuple(int(color[offset:offset + 2], 16) for offset in (0, 2, 4))
        luminance = .2126 * rgb[0] + .7152 * rgb[1] + .0722 * rgb[2]
        text_color = '333333' if luminance > 155 else 'FFFFFF'
        label = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(label_width_inches), Inches(label_height),
        )
        label.name = 'Maximum Allocation Percentage Label'
        text_frame = label.text_frame
        text_frame.clear()
        text_frame.margin_left = 0
        text_frame.margin_right = 0
        text_frame.margin_top = 0
        text_frame.margin_bottom = 0
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = label_text
        run.font.name = 'Arial'
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(text_color)
        angle += share * 360


def maximum_allocations_from_configuration(
    configuration: dict[str, Any] | None,
) -> dict[str, dict[str, float]]:
    """Return per-environment Voice/Data maximum points from configured KPI maxima.

    Each metric is counted once per environment. Leaf columns and result rows are
    deliberately ignored so the allocation cannot multiply with the hierarchy.
    """
    if not isinstance(configuration, dict):
        return {}
    scope = configuration.get('scope')
    environments = scope.get('environments') if isinstance(scope, dict) else None
    metrics = configuration.get('metrics')
    if not isinstance(environments, dict) or not isinstance(metrics, list):
        return {}

    allocations: dict[str, dict[str, float]] = {}
    for environment, details in environments.items():
        if not isinstance(details, dict) or _number(details.get('total_points')) <= 0:
            continue
        name = str(environment).strip()
        if name:
            allocations[name] = {'voice': 0.0, 'data': 0.0, 'max': 0.0}

    for metric in metrics:
        if not isinstance(metric, dict):
            continue
        kind = _metric_source_kind(metric)
        family = 'voice' if kind in {'voice', 'speech'} else 'data' if kind == 'data' else None
        if family is None:
            continue
        contexts = metric.get('contexts')
        if not isinstance(contexts, dict):
            continue
        for environment, allocation in allocations.items():
            context = contexts.get(environment)
            if isinstance(context, dict):
                allocation[family] += _number(context.get('max_points'))
    for allocation in allocations.values():
        allocation['max'] = allocation['voice'] + allocation['data']
    return allocations


def scaled_environment_allocations(
    allocations: dict[str, dict[str, float]], scaling: dict[str, Any] | None,
) -> dict[str, dict[str, float]]:
    """Allocations of the Combined scores when environments without results were scaled out.

    Those environments get no points and the others share theirs in proportion;
    each keeps its configured maximum as ``original_max``.
    """
    scaled = set((scaling or {}).get('scaled_environments') or [])
    if not scaled:
        return allocations
    total = sum(values.get('max', 0.0) for values in allocations.values())
    kept = sum(values.get('max', 0.0) for name, values in allocations.items() if name not in scaled)
    if kept <= 0:
        return allocations
    result = {}
    for name, values in allocations.items():
        factor = 0.0 if name in scaled else total / kept
        result[name] = {'voice': values.get('voice', 0.0) * factor, 'data': values.get('data', 0.0) * factor,
                        'max': values.get('max', 0.0) * factor, 'original_max': values.get('max', 0.0)}
    return result


def category_maximum_allocations(matrix: dict[str, Any]) -> dict[str, float]:
    """Return configured maximum points once per KPI category in a score matrix."""
    allocations: dict[str, float] = {}
    for row in matrix.get('rows', []):
        if not isinstance(row, dict) or row.get('row_type') == 'category':
            continue
        category = str(row.get('category') or '').strip()
        if category:
            allocations[category] = allocations.get(category, 0.0) + _number(row.get('max_points'))
    return allocations


def _selected_allocations(
    matrix: dict[str, Any], environment_allocations: dict[str, dict[str, Any]],
) -> list[tuple[str, float, float]]:
    context = matrix.get('context')
    environment = str(context.get('environment') or '').strip() if isinstance(context, dict) else ''
    if environment.casefold() in {'combined', 'all', 'all environments', 'allenvironments'}:
        selected = list(environment_allocations.items())
    else:
        match = next((name for name in environment_allocations
                      if name.casefold() == environment.casefold()), None)
        selected = [(match, environment_allocations[match])] if match is not None else []
    return [
        (str(name), _number(values.get('voice')), _number(values.get('data')))
        for name, values in selected if isinstance(values, dict)
    ]


def _environment_kind(environment: str) -> str:
    name = environment.casefold()
    return 'city' if 'city' in name else 'road' if 'road' in name else 'location'


def _environment_display_label(environment: str) -> str:
    key = ''.join(character for character in environment.casefold() if character.isalnum())
    if key == 'drivecity':
        return 'Drive - City'
    if key in {'driveconnectionroad', 'driveconnectingroads', 'driveroad'}:
        return 'Drive - Connecting Roads'
    return environment


def _category_icon_kind(category: str) -> str:
    key = ''.join(character for character in category.casefold() if character.isalnum())
    if 'whatsapp' in key:
        return 'whatsapp'
    if 'classic' in key or 'calls' in key or key == 'voice':
        return 'voice'
    if key == 'data':
        return 'data'
    if 'multirab' in key:
        return 'multirab'
    if 'transfer' in key:
        return 'transfer'
    if 'brows' in key:
        return 'browsing'
    if 'video' in key:
        return 'video'
    if 'interactivity' in key:
        return 'interactivity'
    return 'category' if key == 'category' else 'location'


def _category_display_label(category: str) -> str:
    return category.strip().title()


def _allocation_chart(slide, categories, values, colors, left, top, size, hole):
    data = CategoryChartData()
    data.categories = categories
    data.add_series('Allocation', values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(left), Inches(top), Inches(size), Inches(size), data,
    ).chart
    chart.has_title = False
    chart.has_legend = False
    # Use the same proportional plot frame for each overlay, avoiding automatic
    # pie margins that otherwise make the ring radii differ between charts.
    layout = chart._chartSpace.plotArea.find(qn('c:layout'))
    if layout is None:
        layout = OxmlElement('c:layout')
        chart._chartSpace.plotArea.insert(0, layout)
    manual = OxmlElement('c:manualLayout')
    for tag, value in (('layoutTarget', 'inner'), ('xMode', 'edge'), ('yMode', 'edge'),
                       ('wMode', 'factor'), ('hMode', 'factor'),
                       ('x', '.05'), ('y', '.05'), ('w', '.9'), ('h', '.9')):
        element = OxmlElement(f'c:{tag}')
        element.set('val', value)
        manual.append(element)
    layout.append(manual)
    plot = chart.plots[0]
    plot.vary_by_categories = True
    hole_size = plot._element.find(qn('c:holeSize'))
    if hole_size is None:
        hole_size = OxmlElement('c:holeSize')
        plot._element.append(hole_size)
    hole_size.set('val', str(hole))
    for point, color in zip(chart.series[0].points, colors):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(color)
    # No-fill chart and plot areas allow the inner native charts to remain
    # visible over the outer Global chart in Office and PDF renderers.
    chart_properties = OxmlElement('c:spPr')
    chart_properties.append(OxmlElement('a:noFill'))
    chart._chartSpace.insert_element_before(
        chart_properties, 'c:txPr', 'c:externalData', 'c:printSettings', 'c:userShapes', 'c:extLst',
    )
    plot_properties = OxmlElement('c:spPr')
    plot_properties.append(OxmlElement('a:noFill'))
    chart._chartSpace.plotArea.insert_element_before(plot_properties, 'c:extLst')
    return chart


# These paths match allocationIconPath in the web scoring interface.
_ICON_PATHS = {
    'whatsapp': 'M3 4h18v13H9l-6 4ZM8 7h3l1 3-2 1 3 3 1-2 3 1v3c-4 1-10-5-9-9Z',
    'voice': 'M6 3h4l2 5-3 2c2 4 3 5 7 7l2-3 5 2v4c-1 3-5 2-8 0C7 16 2 7 6 3Z',
    'multirab': 'M12 3v18M5 21l7-18 7 18M4 7a11 11 0 0 1 16 0M7 10a7 7 0 0 1 10 0',
    'transfer': 'M7 3v17m-4-4 4 4 4-4M17 21V4m-4 4 4-4 4 4',
    'browsing': 'M2 4h20v16H2ZM2 8h20M5 6h1m2 0h1m2 0h1M8 12l-3 2 3 2m8-4 3 2-3 2',
    'video': 'M3 4h18v16H3ZM9 8l7 4-7 4Z',
    'interactivity': 'M2 12h4l3-8 5 16 3-8h5',
    'category': 'M3 3h7v7H3ZM14 3h7v7h-7ZM3 14h7v7H3ZM14 14h7v7h-7Z',
    'data': 'M7 3h10v18H7ZM10 7h10m-3-3 3 3-3 3M14 16H4m3-3-3 3 3 3',
    'city': 'M3 21V7h7v14M10 21V3h10v18M1 21h22M6 10v2m0 3v2m8-10h3m-3 4h3m-3 4h3',
    'road': 'M5 21 9 3m10 18L15 3M12 3v3m0 3v3m0 3v3m0 2v1',
    'location': 'M12 22s8-8 8-13a8 8 0 0 0-16 0c0 5 8 13 8 13ZM12 6a3 3 0 1 0 0 6 3 3 0 0 0 0-6Z',
}


def _icon_geometry(kind: str):
    """Translate the fixed web icon paths into editable DrawingML geometry.

    Only the commands used by these paths are needed. Their circular arcs have
    no rotation, so DrawingML arcTo preserves them without rasterization.
    """
    geometry = OxmlElement('a:custGeom')
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        geometry.append(OxmlElement(f'a:{tag}'))
    rect = OxmlElement('a:rect')
    for attr, value in (('l', '0'), ('t', '0'), ('r', 'w'), ('b', 'h')):
        rect.set(attr, value)
    geometry.append(rect)
    paths = OxmlElement('a:pathLst')
    path = OxmlElement('a:path')
    path.set('w', '24000')
    path.set('h', '24000')
    path.set('fill', 'none')
    paths.append(path)
    geometry.append(paths)
    tokens = re.findall(r'[A-Za-z]|[-+]?(?:\d*\.\d+|\d+)', _ICON_PATHS.get(kind, _ICON_PATHS['location']))
    current = (0., 0.)
    subpath = current
    control = None
    command = ''
    index = 0

    def point(element, coordinates):
        node = OxmlElement('a:pt')
        node.set('x', str(round(coordinates[0] * 1000)))
        node.set('y', str(round(coordinates[1] * 1000)))
        element.append(node)

    while index < len(tokens):
        if tokens[index].isalpha():
            command = tokens[index]
            index += 1
        absolute = command.isupper()
        op = command.upper()
        if op == 'Z':
            path.append(OxmlElement('a:close'))
            current = subpath
            control = None
            continue
        count = {'M': 2, 'L': 2, 'H': 1, 'V': 1, 'C': 6, 'S': 4, 'A': 7}[op]
        values = [float(value) for value in tokens[index:index + count]]
        index += count
        origin = (0., 0.) if absolute else current
        pairs = [(values[i] + origin[0], values[i + 1] + origin[1])
                 for i in range(0, len(values) - 1, 2)] if op in {'M', 'L', 'C', 'S'} else []
        if op in {'M', 'L', 'H', 'V'}:
            endpoint = (pairs[-1] if pairs else
                        (values[0] + origin[0], current[1]) if op == 'H' else
                        (current[0], values[0] + origin[1]))
            element = OxmlElement('a:moveTo' if op == 'M' else 'a:lnTo')
            point(element, endpoint)
            if op == 'M':
                subpath = endpoint
                command = 'L' if absolute else 'l'
        elif op in {'C', 'S'}:
            if op == 'S':
                reflected = (2 * current[0] - control[0], 2 * current[1] - control[1]) if control else current
                pairs.insert(0, reflected)
            element = OxmlElement('a:cubicBezTo')
            for coordinates in pairs:
                point(element, coordinates)
            endpoint = pairs[-1]
        else:
            radius, _, _, large, sweep, dx, dy = values
            endpoint = (dx + origin[0], dy + origin[1])
            vx, vy = (current[0] - endpoint[0]) / 2, (current[1] - endpoint[1]) / 2
            factor = math.sqrt(max(0., (radius * radius - vx * vx - vy * vy) / (vx * vx + vy * vy)))
            if bool(large) == bool(sweep):
                factor = -factor
            cx = (current[0] + endpoint[0]) / 2 + factor * vy
            cy = (current[1] + endpoint[1]) / 2 - factor * vx
            start = math.atan2(current[1] - cy, current[0] - cx)
            finish = math.atan2(endpoint[1] - cy, endpoint[0] - cx)
            span = (finish - start) % (2 * math.pi) if sweep else -((start - finish) % (2 * math.pi))
            element = OxmlElement('a:arcTo')
            element.set('wR', str(round(radius * 1000)))
            element.set('hR', str(round(radius * 1000)))
            element.set('stAng', str(round(math.degrees(start) * 60000)))
            element.set('swAng', str(round(math.degrees(span) * 60000)))
        path.append(element)
        current = endpoint
        control = pairs[-2] if op in {'C', 'S'} else None
    return geometry


def _remove_shape_effects(shape) -> None:
    """Override theme effects so legend marks cannot inherit shadows."""
    properties = shape._element.spPr
    for tag in ('a:effectLst', 'a:effectDag'):
        existing = properties.find(qn(tag))
        if existing is not None:
            properties.remove(existing)
    properties.append(OxmlElement('a:effectLst'))
    for reference in shape._element.xpath('./p:style/a:effectRef'):
        reference.set('idx', '0')


def _legend_icon(slide, kind: str, x: float, y: float, size: float, color: str) -> None:
    """Draw the web icon as one editable, unfilled native vector path."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(size), Inches(size))
    shape.name = f'Maximum Allocation Icon {kind}'
    properties = shape._element.spPr
    preset = properties.find(qn('a:prstGeom'))
    properties.replace(preset, _icon_geometry(kind))
    shape.fill.background()
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(.85)
    shape.line._get_or_add_ln().set('cap', 'rnd')
    shape.line._get_or_add_ln().append(OxmlElement('a:round'))
    _remove_shape_effects(shape)


def _environment_segment_icons(slide, allocations, environment_allocations,
                               chart_left: float, chart_top: float, chart_size: float) -> float:
    total = sum(voice + data_points for _, voice, data_points in allocations)
    if total <= 0:
        return chart_top + chart_size
    size = .45
    radius = chart_size * .45 + size / 2 + .08
    angle = 0.0
    marker_bottom = chart_top + chart_size
    for environment, voice, data_points in allocations:
        share = (voice + data_points) / total
        midpoint = math.radians(angle + share * 180)
        x = chart_left + chart_size / 2 + radius * math.sin(midpoint) - size / 2
        y = chart_top + chart_size / 2 - radius * math.cos(midpoint) - size / 2
        first_shape = len(slide.shapes)
        _legend_icon(slide, _environment_kind(environment), x, y, size,
                     _global_color(environment, environment_allocations))
        parts = list(slide.shapes)[first_shape:]
        for part in parts:
            part.line.width = Pt(1.3)
        icon = parts[0] if len(parts) == 1 else slide.shapes.add_group_shape(parts)
        icon.name = f'Maximum Allocation Environment Segment Icon {environment}'
        marker_bottom = max(marker_bottom, y + size)
        angle += share * 360
    return marker_bottom


def _legend_row(slide, x, y, width, height, icon, color, label, detail, bold=False, wrap=False, struck=''):
    _legend_icon(slide, icon, x, y + max(0, (height - .18) / 2), .18, color)
    swatch = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x + .20), Inches(y + (height - .09) / 2), Inches(.09), Inches(.09),
    )
    swatch.name = f'Maximum Allocation Swatch {label.splitlines()[0]}'
    swatch.fill.solid()
    swatch.fill.fore_color.rgb = RGBColor.from_string(color)
    swatch.line.fill.background()
    _remove_shape_effects(swatch)
    textbox = slide.shapes.add_textbox(
        Inches(x + .32), Inches(y), Inches(max(.2, width - .32)), Inches(height),
    )
    textbox.text_frame.clear()
    textbox.text_frame.word_wrap = wrap
    textbox.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.font.name = 'Arial'
    paragraph.font.bold = bold
    textbox.text_frame.margin_left = textbox.text_frame.margin_right = 0
    textbox.text_frame.margin_top = textbox.text_frame.margin_bottom = 0
    font_size = min(8.5, max(6.5, (width - .32) * 72 /
                            (max(1, len(label) + len(detail) + len(struck) + 3) * .52)))
    paragraph.font.size = Pt(font_size)
    name_run = paragraph.add_run()
    name_run.text = f'{label}: '
    name_run.font.name = 'Arial'
    name_run.font.size = Pt(font_size)
    name_run.font.bold = bold
    name_run.font.color.rgb = RGBColor.from_string('465565')
    if struck:
        struck_run = paragraph.add_run()
        struck_run.text = f'{struck} '
        struck_run.font.name = 'Arial'
        struck_run.font.size = Pt(font_size)
        struck_run.font.color.rgb = RGBColor.from_string('7A8691')
        struck_run._r.get_or_add_rPr().set('strike', 'sngStrike')
    detail_run = paragraph.add_run()
    detail_run.text = detail
    detail_run.font.name = 'Arial'
    detail_run.font.size = Pt(font_size)
    detail_run.font.bold = bold
    detail_run.font.color.rgb = RGBColor.from_string('8A3D0A')


def _global_color(environment: str, environment_allocations: dict[str, Any]) -> str:
    original_order = list(environment_allocations)
    index = next((i for i, name in enumerate(original_order)
                  if str(name).casefold() == environment.casefold()), 0)
    return _GLOBAL_COLORS[index % len(_GLOBAL_COLORS)]


def _allocation_legend_heading(slide, x, y, width, height, label):
    textbox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    textbox.text_frame.clear()
    textbox.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    textbox.text_frame.word_wrap = False
    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.text = label
    paragraph.font.name = 'Arial'
    paragraph.font.bold = True
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = RGBColor.from_string('465565')


def add_maximum_allocation_donut(
    slide,
    matrix: dict[str, Any],
    environment_allocations: dict[str, dict[str, Any]],
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    category_allocations: dict[str, float] | None = None,
):
    """Add an editable environment doughnut with an inner allocation ring.

    Geometry is in inches and describes the whole widget: the chart sits above
    its legend, with a combined total followed by environment maxima. Both
    combined and single-environment matrices show environment totals on the
    outer ring and service or category totals on the inner ring. The returned
    value is the outer native chart.
    """
    allocations = _selected_allocations(matrix, environment_allocations)
    if not allocations:
        return None

    category_allocations = {
        str(category): _number(value) for category, value in (category_allocations or {}).items()
        if str(category).strip() and _number(value) > 0
    }
    category_mode = bool(category_allocations)
    inner_labels = [_category_display_label(category) for category in category_allocations] if category_mode else ['Voice', 'Data']
    inner_colors = ([_CATEGORY_COLORS[index % len(_CATEGORY_COLORS)]
                     for index in range(len(inner_labels))]
                    if category_mode else [_VOICE_COLOR, _DATA_COLOR])
    if category_mode:
        inner_totals = list(category_allocations.values())
    elif len(allocations) > 1:
        inner_totals = [sum(voice for _, voice, _ in allocations),
                        sum(data_points for _, _, data_points in allocations)]
    else:
        inner_totals = [allocations[0][1], allocations[0][2]]

    center_total = sum(voice + data_points for _, voice, data_points in allocations)
    legend_items = []
    if len(allocations) > 1:
        legend_items.append(('location', '465565', 'Total Points',
                             f'{_legend_points(center_total)} pts (100.0%)'))
    legend_items.append(('heading', 'Points per Environment:'))
    for environment, voice, data_points in allocations:
        value = voice + data_points
        percent = value * 100 / center_total if center_total else 0
        # Points moved by environment scaling show the configured points struck through.
        configured = environment_allocations.get(environment, {}).get('original_max')
        struck = f'{_legend_points(configured)} pts' if configured is not None and abs(configured - value) > .005 else ''
        legend_items.append((_environment_kind(environment), _global_color(environment, environment_allocations),
                             matrix.get('environment_labels', {}).get(environment, _environment_display_label(environment)),
                             f'{_legend_points(value)} pts ({percent:.1f}%)', struck))
    legend_items.append(('heading', 'Points per KPI Category:' if category_mode else 'Points per Service:'))
    if category_mode:
        for index, category in enumerate(inner_labels):
            value = inner_totals[index]
            percent = value * 100 / center_total if center_total else 0
            legend_items.append((_category_icon_kind(category), inner_colors[index], _category_display_label(category),
                                 f'{_legend_points(value)} pts ({percent:.1f}%)'))
    else:
        for index, (label, color, kind) in enumerate((
            ('Voice', _VOICE_COLOR, 'voice'), ('Data', _DATA_COLOR, 'data'),
        )):
            value = inner_totals[index]
            percent = value * 100 / center_total if center_total else 0
            legend_items.append((kind, color, label, f'{_legend_points(value)} pts ({percent:.1f}%)'))

    # Category legends have more rows; compact them to preserve enough chart
    # diameter for the environment and major category percentage labels.
    legend_heights = [.18 if category_mode else .22 for _ in legend_items]
    chart_size = min(width * .90, height - .35 - sum(legend_heights))
    chart_left = left + (width - chart_size) / 2
    chart_top = top
    totals = [voice + data_points for _, voice, data_points in allocations]
    chart = _allocation_chart(
        slide, [environment for environment, _, _ in allocations], totals,
        [_global_color(environment, environment_allocations) for environment, _, _ in allocations],
        chart_left, chart_top, chart_size, 72,
    )
    marker_bottom = _environment_segment_icons(slide, allocations, environment_allocations,
                                               chart_left, chart_top, chart_size)
    # Move the inner ring inward by the same subtle fraction used in the web chart.
    # The slightly smaller hole keeps its band width approximately unchanged.
    ring_size = chart_size * (.72 - 2 / 156)
    inner_hole = 62
    _allocation_chart(
        slide, inner_labels, inner_totals, inner_colors,
        chart_left + (chart_size - ring_size) / 2, chart_top + (chart_size - ring_size) / 2,
        ring_size, inner_hole,
    )
    _add_percentage_labels(slide, totals,
                           [_global_color(environment, environment_allocations)
                            for environment, _, _ in allocations],
                           chart_left, chart_top, chart_size, 72)
    _add_percentage_labels(slide, inner_totals, inner_colors,
                           chart_left + (chart_size - ring_size) / 2,
                           chart_top + (chart_size - ring_size) / 2, ring_size, inner_hole)

    center_width = ring_size * inner_hole / 100 * .9 - .06
    center_text = f'{center_total:.2f}'
    total_size = max(2, min(20, center_width * 72 / (len(center_text) * .65)))
    unit_size = max(1, min(8.5, total_size * 12 / 31))
    center_height = (total_size + unit_size) * 1.35 / 72
    center = slide.shapes.add_textbox(
        Inches(chart_left + (chart_size - center_width) / 2),
        Inches(chart_top + (chart_size - center_height) / 2),
        Inches(center_width), Inches(center_height),
    )
    center.name = 'Maximum Score Allocation Total'
    frame = center.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, (text, size, color) in enumerate(((center_text, total_size, '1C3745'),
                                                ('max points', unit_size, '5C707A'))):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.text = text
        paragraph.font.name = 'Arial'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor.from_string(color)

    label_top = max(top + chart_size + .25, marker_bottom + .08)
    label_width = width
    label_y = label_top
    for item, label_height in zip(legend_items, legend_heights):
        x = left
        y = label_y
        if item[0] == 'heading':
            _allocation_legend_heading(slide, x + .12, y, label_width - .12, label_height, item[1])
        else:
            kind, color, label, detail, *rest = item
            global_row = label == 'Total Points'
            indent = .12 if global_row else .26
            _legend_row(slide, x + indent, y, label_width - indent, label_height,
                        kind, color, label, detail, bold=global_row, wrap=False, struck=rest[0] if rest else '')
        label_y += label_height
    return chart
