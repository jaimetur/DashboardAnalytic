"""NetCheck-style insight slides of the Scoring & GAP Analysis exports.

Location cards, KPI GAP profiles, scoring trends and campaign comparisons use the
insights of the scoring views (see ``scoring_insights``).
"""
from __future__ import annotations

from math import ceil, floor
from typing import Any, Callable

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

VOICE_COLOR = '#F2A900'
DATA_COLOR = '#0E6B66'
PANEL_COLOR = '#0B2A5B'
TITLE_BLUE = '#1450A8'
CARD_COLOR = '#F2F3F5'
LOSS_COLOR = '#E8414F'
WIN_COLOR = '#4CA65A'
# KPI GAP profiles for every location the location cards show (one per City, Cluster or Region).
MAX_PROFILE_SCOPES = 12
PROFILE_TABLE_BOTTOM = 6.68
MAX_GROUP_SLIDES = 6
MAX_COMPARISONS = 8
_FONT = 'Arial'
FALLBACK_COLORS = ('#E60000', '#0B6E8F', '#7A3DB8', '#00A3AD', '#F2A900', '#4CA65A', '#8C564B', '#5B6770')


def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip('#'))


def _points(value: float | None, digits: int = 0) -> str:
    return 'N/A' if value is None else f'{value:,.{digits}f}'


def _box(slide, left, top, width, height, color, *, shape=MSO_SHAPE.RECTANGLE):
    box = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb(color)
    box.line.fill.background()
    box.shadow.inherit = False
    return box


def _lighten(color: str, amount: float) -> RGBColor:
    red, green, blue = (int(color.lstrip('#')[index:index + 2], 16) for index in (0, 2, 4))
    return RGBColor(*(round(value + (255 - value) * amount) for value in (red, green, blue)))


def _gradient_bar(slide, left, top, width, height, color: str):
    """A data bar like the Excel ones: the colour fading to white, with a border of the same colour."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    bar.fill.gradient()
    bar.fill.gradient_angle = 0
    stops = bar.fill.gradient_stops
    stops[0].position, stops[1].position = 0, 1
    stops[0].color.rgb = _rgb(color)
    stops[1].color.rgb = _lighten(color, .88)
    bar.line.color.rgb = _rgb(color)
    bar.line.width = Pt(.5)
    bar.shadow.inherit = False
    return bar


def _write(frame, paragraphs: list[tuple[str, float, bool, str]], *, align=PP_ALIGN.LEFT) -> None:
    """Fill a text frame with (text, size, bold, color) paragraphs."""
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.08)
    frame.margin_top = frame.margin_bottom = Inches(.04)
    frame.clear()
    for index, (text, size, bold, color) in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.name = _FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)


def _write_runs(frame, runs: list[tuple[str, bool, str | None]], *, size: float, color: str) -> None:
    """One paragraph of (text, bold, color) runs; a run without a colour uses ``color``."""
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.clear()
    paragraph = frame.paragraphs[0]
    for text, bold, run_color in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = _FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(run_color or color)


def _legend_key(slide, left: float, top: float, items: list[tuple[str, str]]) -> None:
    for label, color in items:
        _box(slide, left, top + .03, .16, .16, color)
        box = slide.shapes.add_textbox(Inches(left + .22), Inches(top), Inches(1.1), Inches(.24))
        _write(box.text_frame, [(label, 9, False, '#4A5B65')])
        box.text_frame.margin_left = 0
        left += 1.35


def _block_profiles(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep the GAP profiles of at most a few contexts: the latest campaign first."""
    contexts = list(dict.fromkeys(tuple(sorted(item['context'].items())) for item in items))
    if len(contexts) > MAX_PROFILE_SCOPES:
        from src.modules.scoring_insights import campaign_sort_key
        campaigns = sorted({item['context'].get('campaign') for item in items if item['context'].get('campaign')},
                           key=campaign_sort_key)
        if campaigns:
            items = [item for item in items if item['context'].get('campaign') == campaigns[-1]]
            contexts = list(dict.fromkeys(tuple(sorted(item['context'].items())) for item in items))
    if len(contexts) > MAX_PROFILE_SCOPES:
        return []
    return [item for item in items if item['operator'] != item['reference']]


def _context_text(context: dict[str, Any]) -> str:
    from src.modules.scoring_insights import _display
    names = {'vendor': 'Vendor', 'region': 'Region', 'cluster': 'Cluster', 'city': 'City', 'campaign': 'Campaign'}
    return ' · '.join(f'{label}: {_display(label, context[field])}' for field, label in names.items() if field in context)


def _latest_campaign_groups(groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    from src.modules.scoring_insights import campaign_sort_key
    campaigns = [group.get('campaign') for group in groups if group.get('campaign')]
    if len(set(campaigns)) <= 1:
        return groups
    latest = max(campaigns, key=campaign_sort_key)
    return [group for group in groups if group.get('campaign') in (None, latest)]


def add_location_card_slides(presentation, groups: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                             new_slide: Callable) -> None:
    """One slide per group: a card per City, Cluster or Region with Voice and Data points per operator.

    With several campaigns only the latest one gets cards; the scoring trend shows the others.
    """
    groups = _latest_campaign_groups(groups)
    for group in groups[:MAX_GROUP_SLIDES]:
        level = group['level']
        plural = {'City': 'cities', 'Cluster': 'clusters', 'Region': 'regions'}[level]
        slide = new_slide(presentation, f'{scoring_label} Scoring per {level}',
                          ' · '.join(part for part in (subtitle, group['title']) if part))
        cards = group['cards']
        maximum = group['maximum']
        voice_max, data_max = group.get('max_voice') or 0, group.get('max_data') or 0
        panel = _box(slide, .45, 1.55, 2.25, 5.35, PANEL_COLOR)
        panel.name = 'Scoring Location Cards Panel'
        paragraphs = [(f'The {scoring_label} scoring per {level.lower()} for the selected {plural}.', 11, False, '#FFFFFF'),
                      ('', 6, False, '#FFFFFF')]
        if group['scaled']:
            paragraphs.append((f'The scoring points are scaled to a maximum of {_points(maximum)} points, providing '
                               'comparability with the overall scoring.', 10, False, '#FFFFFF'))
            paragraphs.append(('', 6, False, '#FFFFFF'))
        paragraphs.append((f'Total: {_points(maximum)} points', 12, True, '#FFFFFF'))
        if voice_max or data_max:
            paragraphs.append((f'Voice: {_points(voice_max)} · Data: {_points(data_max)}', 10, False, '#FFFFFF'))
        _write(panel.text_frame, paragraphs, align=PP_ALIGN.CENTER)
        panel.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        count = len(cards)
        columns = count if count <= 4 else ceil(count / 2)
        rows = 1 if count <= 4 else 2
        gap = .12
        area_left, area_top, area_width, area_height = 2.85, 1.55, 10.05, 5.35
        width = (area_width - gap * (columns - 1)) / columns
        height = (area_height - gap * (rows - 1)) / rows
        label_size = 8 if columns <= 4 else 7
        for index, card in enumerate(cards):
            left = area_left + (index % columns) * (width + gap)
            top = area_top + (index // columns) * (height + gap)
            background = _box(slide, left, top, width, height, CARD_COLOR)
            background.name = f'Scoring Location Card {card["label"]}'
            heading = slide.shapes.add_textbox(Inches(left + .05), Inches(top + .05), Inches(width - .1), Inches(.32))
            _write(heading.text_frame, [(str(card['label']).upper(), 12 if columns <= 4 else 10, True, TITLE_BLUE)])
            data = CategoryChartData()
            data.categories = [str(item['label']) for item in card['operators']]
            data.add_series('Voice', [item['voice'] for item in card['operators']])
            data.add_series('Data', [item['data'] for item in card['operators']])
            data.add_series('Total', [item['total'] for item in card['operators']])
            chart_shape = slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_STACKED, Inches(left + .03), Inches(top + .38),
                Inches(width - .06), Inches(height - .42), data,
            )
            chart_shape.name = f'Scoring Location Chart {card["label"]}'
            chart = chart_shape.chart
            _format_card_chart(chart, maximum, label_size)
        _legend_key(slide, 7.0, 6.98, [('Voice', VOICE_COLOR), ('Data', DATA_COLOR)])
        note = slide.shapes.add_textbox(Inches(9.7), Inches(6.98), Inches(3.2), Inches(.24))
        _write(note.text_frame, [(f'Scoring points scaled to a maximum of {_points(maximum)} points' if group['scaled']
                                  else f'Maximum scoring: {_points(maximum)} points', 8, False, '#6B7A85')],
               align=PP_ALIGN.RIGHT)


def _format_card_chart(chart, maximum: float, label_size: float) -> None:
    from src.modules.scoring_exports import _add_total_labels, _fill_series
    chart.has_legend = False
    chart.font.name = _FONT
    chart.font.size = Pt(label_size)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = max(1, maximum) * 1.18
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.tick_labels.font.size = Pt(label_size)
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.format.line.color.rgb = _rgb('#BFC6CC')
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = XL_DATA_LABEL_POSITION.CENTER
    labels.number_format = '0'
    labels.number_format_is_linked = False
    labels.font.size = Pt(label_size)
    labels.font.color.rgb = _rgb('#FFFFFF')
    _fill_series(chart.series[0], VOICE_COLOR)
    _fill_series(chart.series[1], DATA_COLOR)
    _add_total_labels(chart)
    for run_properties in chart._chartSpace.xpath('.//c:lineChart//a:defRPr'):
        run_properties.set('sz', str(round((label_size + 1) * 100)))
    for number_format in chart._chartSpace.xpath('.//c:lineChart/c:dLbls/c:numFmt'):
        number_format.set('formatCode', '0')


def _table(slide, rows: int, columns: int, left: float, top: float, width: float, row_height: float, widths: list[float]):
    shape = slide.shapes.add_table(rows, columns, Inches(left), Inches(top), Inches(width), Inches(row_height * rows))
    table = shape.table
    for index, column_width in enumerate(widths):
        table.columns[index].width = Inches(column_width)
    for row in table.rows:
        row.height = Inches(row_height)
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set('firstRow', '0')
    tbl_pr.set('bandRow', '0')
    return shape, table


def _cell(cell, text: str, *, size: float, bold: bool = False, color: str = '#263746', fill: str | None = '#FFFFFF',
          align=PP_ALIGN.LEFT, underline: bool = False) -> None:
    cell.margin_left = cell.margin_right = Inches(.04)
    cell.margin_top = cell.margin_bottom = Inches(0)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
    frame = cell.text_frame
    frame.word_wrap = False
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = _FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    run.font.color.rgb = _rgb(color)


def _environment_label(name: str, labels: dict[str, str]) -> str:
    return labels.get(name) or name


def _profile_table(slide, profile: dict[str, Any], *, kind: str, left: float, width: float,
                   environment_labels: dict[str, str], header_color: str) -> None:
    rows = profile['to_maximum'] if kind == 'maximum' else profile['to_reference']
    columns = profile['columns']
    title = (f"{profile['label']}: Gap to Maximum" if kind == 'maximum'
             else f"{profile['label']}: Gap to {profile['reference_label']}")
    count = max(1, len(rows))
    top = 1.62
    # The banner and the column headers (up to three lines) have explicit heights, so the
    # bars drawn under the Total column line up with their rows whatever the KPI count.
    banner_height, header_height = .3, .46
    row_height = min(.22, (PROFILE_TABLE_BOTTOM - top - banner_height - header_height) / count)
    size = 7 if row_height >= .17 else 6
    environment_width = .62
    total_width = 1.0
    fixed = [.5, 1.3]
    kpi_width = width - sum(fixed) - environment_width * len(columns) - total_width
    widths = [*fixed, kpi_width, *([environment_width] * len(columns)), total_width]
    gap_key = 'gap_to_maximum' if kind == 'maximum' else 'gap_to_reference'
    total_key = 'total_gap_to_maximum' if kind == 'maximum' else 'total_gap_to_reference'
    # Data bars in every value column, each column with its own scale; the values are
    # written over them, right-aligned.
    bar_columns = [[row[gap_key].get(name) for row in rows] for name in columns]
    bar_columns.append([row[total_key] for row in rows])
    for offset, values in enumerate(bar_columns):
        column = 3 + offset
        scale = max((abs(value) for value in values if value is not None), default=0) or 1
        bar_left = left + sum(widths[:column]) + .03
        bar_width = widths[column] - .06
        for index, value in enumerate(values):
            if value is None or value == 0:
                continue
            color = LOSS_COLOR if kind == 'maximum' or value < 0 else WIN_COLOR
            bar = _gradient_bar(slide, bar_left, top + banner_height + header_height + index * row_height + .02,
                                max(.02, bar_width * abs(value) / scale), row_height - .04, color)
            bar.name = 'Scoring GAP Bar'
    shape, table = _table(slide, len(rows) + 2, len(widths), left, top, width, row_height, widths)
    table.rows[0].height = Inches(banner_height)
    table.rows[1].height = Inches(header_height)
    shape.name = f"Scoring KPI GAP {'Maximum' if kind == 'maximum' else 'Reference'} Table"
    table.cell(0, 0).merge(table.cell(0, len(widths) - 1))
    _cell(table.cell(0, 0), title, size=size + 3, bold=True, color='#FFFFFF', fill=header_color, align=PP_ALIGN.RIGHT)
    headers = ['Service', 'Area', 'KPIs', *[_environment_label(name, environment_labels) for name in columns], 'Total KPI']
    for column, text in enumerate(headers):
        _cell(table.cell(1, column), text, size=size, bold=True, fill='#E9ECEF',
              align=PP_ALIGN.LEFT if column < 3 else PP_ALIGN.CENTER)
        table.cell(1, column).text_frame.word_wrap = True
    for row_index, row in enumerate(rows, start=2):
        fill = '#FFFFFF' if row_index % 2 else '#F7F8F9'
        underline = bool(profile.get('underline_most_reliable') and row.get('most_reliable'))
        _cell(table.cell(row_index, 0), row['service'].upper(), size=size, fill=fill, align=PP_ALIGN.CENTER)
        _cell(table.cell(row_index, 1), str(row.get('category') or ''), size=size, fill=fill, underline=underline)
        _cell(table.cell(row_index, 2), str(row.get('kpi') or ''), size=size, fill=fill, underline=underline)
        for offset, name in enumerate(columns):
            value = row[gap_key].get(name)
            _cell(table.cell(row_index, 3 + offset), _points(value, 2), size=size, fill=None, align=PP_ALIGN.RIGHT)
        _cell(table.cell(row_index, len(widths) - 1), _points(row[total_key], 2), size=size, bold=True, fill=None,
              align=PP_ALIGN.RIGHT)


def add_kpi_gap_profile_slides(presentation, profiles: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                               environment_labels: dict[str, str], new_slide: Callable,
                               operator_colors: dict[str, str]) -> None:
    """Per compared operator: the KPIs sorted by the points lost against the maximum and the GAP to the reference."""
    for profile in _block_profiles(profiles):
        context = _context_text(profile['context'])
        slide = new_slide(presentation, f"KPI GAP Profile — {profile['label']} vs {profile['reference_label']}",
                          ' · '.join(part for part in (subtitle, context) if part))
        header = profile.get('color') or operator_colors.get(profile['operator']) or '#C8102E'
        _profile_table(slide, profile, kind='maximum', left=.35, width=6.2,
                       environment_labels=environment_labels, header_color=header)
        _profile_table(slide, profile, kind='reference', left=6.78, width=6.2,
                       environment_labels=environment_labels, header_color=header)
        lost = sum(row['total_gap_to_maximum'] for row in profile['to_maximum'])
        reference = sum(row['total_gap_to_reference'] for row in profile['to_reference'])
        reference_color = operator_colors.get(profile['reference']) or '#17232D'
        runs = [(profile['label'], True, header), (' loses ', False, None),
                (f'{lost:,.2f} {scoring_label} points', True, None), (' against the maximum; its GAP to ', False, None),
                (profile['reference_label'], True, reference_color), (' is ', False, None),
                (f'{reference:+,.2f} points', True, None),
                ('. Gap to Maximum is the KPI maximum minus the points scored; Gap to the reference is the operator '
                 'points minus the reference points.', False, None)]
        if profile.get('underline_most_reliable'):
            runs.append((' Underlined KPIs are used for the Most Reliable Network scoring.', False, None))
        note = slide.shapes.add_textbox(Inches(.35), Inches(PROFILE_TABLE_BOTTOM + .14), Inches(12.6), Inches(.36))
        note.name = 'Scoring KPI GAP Profile Note'
        _write_runs(note.text_frame, runs, size=8.5, color='#4A5B65')


def add_trend_slides(presentation, trends: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                     new_slide: Callable) -> None:
    """A line chart of each operator's scoring across the campaigns of every group."""
    from src.modules.scoring_exports import _fill_series  # noqa: F401  (keeps helpers in one module)
    for trend in trends[:MAX_GROUP_SLIDES]:
        slide = new_slide(presentation, f'{scoring_label} Scoring Trend',
                          ' · '.join(part for part in (subtitle, trend['title']) if part))
        data = CategoryChartData()
        data.categories = trend['campaigns']
        for series in trend['series']:
            data.add_series(series['label'], series['points'])
        chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.55), Inches(1.6), Inches(12.2),
                                             Inches(5.3), data)
        chart_shape.name = 'Scoring Trend Chart'
        chart = chart_shape.chart
        chart.font.name = _FONT
        chart.font.size = Pt(10)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        values = [value for series in trend['series'] for value in series['points'] if value is not None]
        low = min(values, default=0)
        chart.value_axis.minimum_scale = max(0, floor((low - 25) / 50) * 50)
        chart.value_axis.maximum_scale = trend['maximum']
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb('#E3E6E9')
        chart.value_axis.format.line.fill.background()
        chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
        for index, (plot_series, series) in enumerate(zip(chart.series, trend['series'])):
            color = series.get('color') or FALLBACK_COLORS[index % len(FALLBACK_COLORS)]
            plot_series.format.line.color.rgb = _rgb(color)
            plot_series.format.line.width = Pt(2.25)
            plot_series.smooth = False
            plot_series.marker.format.fill.solid()
            plot_series.marker.format.fill.fore_color.rgb = _rgb(color)
            plot_series.marker.format.line.color.rgb = _rgb(color)
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.number_format = '0'
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.position = XL_DATA_LABEL_POSITION.ABOVE
        plot.data_labels.font.size = Pt(8)


def _delta_color(value: float | None, scale: float) -> str:
    if value is None or scale <= 0 or value == 0:
        return '#FFFFFF'
    ratio = min(1.0, abs(value) / scale)
    start = (255, 255, 255)
    end = (232, 65, 79) if value < 0 else (76, 166, 90)
    return '#{:02X}{:02X}{:02X}'.format(*(round(a + (b - a) * ratio) for a, b in zip(start, end)))


def _comparison_table(slide, comparison: dict[str, Any], *, left: float, width: float) -> None:
    rows = comparison['rows']
    top = 2.2
    row_height = min(.2, (7.05 - top) / (len(rows) + 2.6))
    size = 7 if row_height >= .17 else 6
    widths = [.5, width - .5 - 4 * .62 - .8, .62, .62, .62, .62, .8]
    heading = slide.shapes.add_textbox(Inches(left), Inches(1.86), Inches(width), Inches(.32))
    title = f"{comparison['label']}" + (f" · {comparison['title']}" if comparison['title'] else '')
    _write(heading.text_frame, [(title, 12, True, '#17232D')], align=PP_ALIGN.CENTER)
    shape, table = _table(slide, len(rows) + 2, len(widths), left, top, width, row_height, widths)
    shape.name = 'Scoring Campaign Comparison Table'
    previous, latest = comparison['previous_campaign'], comparison['latest_campaign']
    headers = ['Service', 'KPI', f'{previous}\nKPI value', f'{previous}\nPoints', f'{latest}\nKPI value',
               f'{latest}\nPoints', 'Δ Points']
    table.rows[0].height = Inches(row_height * 1.6)
    for column, text in enumerate(headers):
        _cell(table.cell(0, column), text, size=size, bold=True, fill='#E9ECEF',
              align=PP_ALIGN.LEFT if column < 2 else PP_ALIGN.CENTER)
        table.cell(0, column).text_frame.word_wrap = True
    scale = max((abs(row['delta']) for row in rows if row['delta'] is not None), default=0)
    for index, row in enumerate(rows, start=1):
        fill = '#FFFFFF'
        _cell(table.cell(index, 0), row['service'].upper(), size=size, fill=fill)
        _cell(table.cell(index, 1), str(row.get('kpi') or ''), size=size, fill=fill)
        for column, key, digits in ((2, 'previous_value', 2), (3, 'previous_points', 2),
                                    (4, 'latest_value', 2), (5, 'latest_points', 2)):
            _cell(table.cell(index, column), _points(row[key], digits), size=size, fill=fill, align=PP_ALIGN.RIGHT)
        _cell(table.cell(index, 6), _points(row['delta'], 2), size=size, bold=True,
              fill=_delta_color(row['delta'], scale), align=PP_ALIGN.RIGHT)
    last = len(rows) + 1
    table.cell(last, 0).merge(table.cell(last, 5))
    _cell(table.cell(last, 0), 'Scoring Points Gap:', size=size + 2, bold=True, fill='#FFFFFF', align=PP_ALIGN.RIGHT)
    total = comparison['total_delta']
    _cell(table.cell(last, 6), f'{total:+,.2f}', size=size + 2, bold=True,
          color=LOSS_COLOR if total < 0 else WIN_COLOR, fill='#FFFFFF', align=PP_ALIGN.RIGHT)


def add_campaign_comparison_slides(presentation, comparisons: list[dict[str, Any]], *, scoring_label: str,
                                   subtitle: str, new_slide: Callable) -> None:
    """Each operator's KPI points in the two latest campaigns, two comparisons per slide."""
    selected = comparisons[:MAX_COMPARISONS]
    for start in range(0, len(selected), 2):
        pair = selected[start:start + 2]
        first = pair[0]
        slide = new_slide(presentation, f'{scoring_label} Campaign Comparison',
                          f"{subtitle} · {first['previous_campaign']} → {first['latest_campaign']}")
        intro = slide.shapes.add_textbox(Inches(.45), Inches(1.52), Inches(12.4), Inches(.3))
        _write(intro.text_frame, [(
            'KPI values and scoring points of the two latest campaigns; Δ Points is the latest minus the previous '
            'campaign, from the largest loss to the largest gain.', 9, False, '#4A5B65')])
        intro.text_frame.margin_left = 0
        width = 6.15 if len(pair) == 2 else 8.0
        lefts = [.45, 6.78] if len(pair) == 2 else [2.65]
        for comparison, left in zip(pair, lefts):
            _comparison_table(slide, comparison, left=left, width=width)


LOSS_LOW = (249, 214, 92)
LOSS_HIGH = (200, 16, 46)
MAX_LOSS_BARS = 30
MAX_LOSS_MAPS = 8


def _loss_color(ratio: float) -> tuple[int, int, int]:
    ratio = max(0.0, min(1.0, ratio))
    return tuple(round(low + (high - low) * ratio) for low, high in zip(LOSS_LOW, LOSS_HIGH))


def render_points_loss_map(areas: list[dict[str, Any]], background: list[list[float]], *,
                           width: int = 1500, height: int = 1300, labels: int = 12) -> bytes:
    """PNG map: the test locations in grey, each place as a bubble and each route as dots, coloured by points lost."""
    from io import BytesIO
    from math import cos, radians, sqrt
    from PIL import Image, ImageDraw, ImageFont

    coordinates = [(lat, lon) for lat, lon in background]
    for area in areas:
        if area.get('latitude') is not None:
            coordinates.append((area['latitude'], area['longitude']))
        coordinates.extend((lat, lon) for lat, lon in area.get('route') or [])
    image = Image.new('RGB', (width, height), (255, 255, 255))
    if not coordinates:
        output = BytesIO()
        image.save(output, format='PNG')
        return output.getvalue()
    latitudes = [lat for lat, _lon in coordinates]
    longitudes = [lon for _lat, lon in coordinates]
    scale_x = cos(radians(sum(latitudes) / len(latitudes)))
    min_x, max_x = min(longitudes) * scale_x, max(longitudes) * scale_x
    min_y, max_y = min(latitudes), max(latitudes)
    margin = 60
    span = max(max_x - min_x, 1e-6) / (width - 2 * margin), max(max_y - min_y, 1e-6) / (height - 2 * margin)
    unit = max(span)
    offset_x = (width - (max_x - min_x) / unit) / 2
    offset_y = (height - (max_y - min_y) / unit) / 2

    def project(lat: float, lon: float) -> tuple[float, float]:
        return offset_x + (lon * scale_x - min_x) / unit, height - offset_y - (lat - min_y) / unit

    draw = ImageDraw.Draw(image)
    for lat, lon in background:
        x, y = project(lat, lon)
        draw.ellipse((x - 2, y - 2, x + 2, y + 2), fill=(214, 220, 226))
    peak = max((area['points'] for area in areas), default=0) or 1
    for area in sorted(areas, key=lambda item: item['points']):
        color = _loss_color(area['points'] / peak)
        if area.get('kind') == 'route' and area.get('route'):
            for lat, lon in area['route']:
                x, y = project(lat, lon)
                draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=color)
        elif area.get('latitude') is not None:
            x, y = project(area['latitude'], area['longitude'])
            radius = 8 + 34 * sqrt(area['points'] / peak)
            draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color, outline=(255, 255, 255), width=3)
    try:
        font = ImageFont.truetype('Arial.ttf', 26)
    except OSError:
        font = ImageFont.load_default()
    places = [area for area in areas if area.get('kind') != 'route' and area.get('latitude') is not None]
    for area in places[:labels]:
        x, y = project(area['latitude'], area['longitude'])
        radius = 8 + 34 * sqrt(area['points'] / peak)
        draw.text((x + radius + 6, y - 14), area['name'], fill=(23, 35, 45), font=font,
                  stroke_width=3, stroke_fill=(255, 255, 255))
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def _ring_centroid(ring: list[list[float]]) -> tuple[float, float, float]:
    """Area and centroid (longitude, latitude) of one ring (shoelace formula)."""
    area = cx = cy = 0.0
    for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1]):
        cross = x1 * y2 - x2 * y1
        area += cross
        cx += (x1 + x2) * cross
        cy += (y1 + y2) * cross
    if abs(area) < 1e-12:
        return 0.0, ring[0][0], ring[0][1]
    return abs(area) / 2, cx / (3 * area), cy / (3 * area)


def render_points_loss_choropleth(boundaries: dict[str, list[list[list[float]]]], values: dict[str, float], *,
                                  height: int = 1300, labels: int = 8, label_room: int = 420,
                                  places: list[tuple[str, str]] | None = None,
                                  outlines: list[list[list[float]]] | None = None) -> bytes:
    """PNG map: every polygon of the layer, those with losses coloured by the points lost.

    The areas that lose most are labelled, or ``places`` (a label and the area it is drawn on).
    """
    from io import BytesIO
    from math import cos, radians
    from PIL import Image, ImageDraw, ImageFont

    rings = [(name, ring) for name, polygons in boundaries.items() for ring in polygons if len(ring) > 2]
    longitudes = [x for _name, ring in rings for x, _y in ring]
    latitudes = [y for _name, ring in rings for _x, y in ring]
    if not rings:
        image = Image.new('RGB', (height, height), (255, 255, 255))
        output = BytesIO()
        image.save(output, format='PNG')
        return output.getvalue()
    west, east, south, north = min(longitudes), max(longitudes), min(latitudes), max(latitudes)
    scale_x = cos(radians((south + north) / 2))
    margin = 30
    unit = (north - south) / (height - 2 * margin) or 1e-6
    # Room on the right for the labels of the areas near the east coast.
    width = max(int((east - west) * scale_x / unit) + 2 * margin + label_room, height // 3)
    image = Image.new('RGB', (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)

    def project(x: float, y: float) -> tuple[float, float]:
        return margin + (x - west) * scale_x / unit, height - margin - (y - south) / unit

    # The outlines of the countries, in grey under their areas.
    for ring in outlines or []:
        if len(ring) > 2:
            draw.polygon([project(x, y) for x, y in ring], fill=(241, 243, 245), outline=(225, 228, 232))
    peak = max(values.values(), default=0) or 1
    floor_value = min((value for value in values.values() if value > 0), default=0)
    span = (peak - floor_value) or 1
    for name, ring in rings:
        value = values.get(name, 0)
        fill = _loss_color((value - floor_value) / span) if value > 0 else (234, 236, 239)
        draw.polygon([project(x, y) for x, y in ring], fill=fill, outline=(255, 255, 255))
    try:
        font = ImageFont.truetype('Arial.ttf', 30)
    except OSError:
        font = ImageFont.load_default()
    named = places if places is not None else [
        (name, name) for name, value in sorted(values.items(), key=lambda item: -item[1]) if value > 0]
    for label, name in named[:labels]:
        polygons = boundaries.get(name) or []
        if not polygons:
            continue
        _area, x, y = max((_ring_centroid(ring) for ring in polygons), key=lambda item: item[0])
        px, py = project(x, y)
        draw.ellipse((px - 5, py - 5, px + 5, py + 5), fill=(23, 35, 45))
        draw.text((px + 10, py - 16), label, fill=(23, 35, 45), font=font, stroke_width=3, stroke_fill=(255, 255, 255))
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def _loss_scale(slide, left: float, top: float, width: float, low: float, high: float) -> None:
    """Colour scale under the map, from the smallest to the largest loss of an area."""
    scale = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + .55), Inches(top), Inches(width - 1.1), Inches(.13))
    scale.name = 'Scoring Points Lost Scale'
    scale.fill.gradient()
    scale.fill.gradient_angle = 0
    stops = scale.fill.gradient_stops
    stops[0].position, stops[1].position = 0, 1
    stops[0].color.rgb, stops[1].color.rgb = RGBColor(*LOSS_LOW), RGBColor(*LOSS_HIGH)
    scale.line.fill.background()
    scale.shadow.inherit = False
    for text, position, align in ((_points(low, 1), left, PP_ALIGN.RIGHT), (_points(high, 1), left + width - .5, PP_ALIGN.LEFT)):
        label = slide.shapes.add_textbox(Inches(position), Inches(top - .06), Inches(.5), Inches(.25))
        label.text_frame.margin_left = label.text_frame.margin_right = 0
        _write(label.text_frame, [(text, 8, False, '#6B7A85')], align=align)


def _loss_bars(slide, bars: list[dict[str, Any]], *, left: float, top: float, width: float, height: float) -> None:
    from pptx.chart.data import CategoryChartData

    data = CategoryChartData()
    data.categories = [area['name'] for area in reversed(bars)]
    data.add_series('Points lost', [round(area['points'], 2) for area in reversed(bars)])
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(left), Inches(top),
                                         Inches(width), Inches(height), data)
    chart_shape.name = 'Scoring Points Lost Bars'
    chart = chart_shape.chart
    chart.has_title = False
    chart.has_legend = False
    chart.font.name = _FONT
    chart.font.size = Pt(7 if len(bars) > 20 else 8)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = _rgb('#E3E6E9')
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    plot = chart.plots[0]
    plot.gap_width = 35
    plot.has_data_labels = True
    plot.data_labels.number_format = '0.0'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(7)
    plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    peak = max((area['points'] for area in bars), default=0) or 1
    low = min((area['points'] for area in bars), default=0)
    series = plot.series[0]
    for index, area in enumerate(reversed(bars)):
        point = series.points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor(*_loss_color((area['points'] - low) / ((peak - low) or 1)))


LOSS_FIELD_TITLES = {'Area': 'Map Area', 'City': 'City', 'Cluster': 'Cluster', 'Region': 'Region'}


def add_points_loss_slides(presentation, maps: list[dict[str, Any]], background: list[list[float]], *,
                           scoring_label: str, subtitle: str, new_slide: Callable,
                           boundaries: Callable[[str], dict] | None = None,
                           environment_labels: dict[str, str] | None = None,
                           outlines: Callable[[str], list] | None = None) -> None:
    """Points lost per area, like the NetCheck points-loss slides: a map and the areas that lose most.

    The main slide colours the map areas (the ITL3 areas in the UK, the workspace's Map Areas
    elsewhere) by the points of the city or Connecting Roads route that loses most among those
    measured in each one, as the web map does by default, and ranks the cities and routes; without
    map areas the cities are bubbles on the test locations. Clusters and Regions with polygons get
    their own slide.
    """
    from io import BytesIO

    boundaries = boundaries or (lambda _field: {})
    outlines = outlines or (lambda _field: [])
    by_series: dict[tuple, dict[str, dict[str, Any]]] = {}
    for item in maps:
        series = (item['operator'], tuple(sorted((key, str(value)) for key, value in item['context'].items())))
        by_series.setdefault(series, {})[item['field']] = item
    slides = []
    for fields in by_series.values():
        city = fields.get('City')
        if city is not None:
            layer = fields.get('Area')
            slides.append((city, 'City', layer if layer is not None and boundaries('Area') else None))
        for field in ('Cluster', 'Region'):
            if fields.get(field) is not None and boundaries(field):
                slides.append((fields[field], field, fields[field]))
    for ranking, field, layer in slides[:MAX_LOSS_MAPS]:
        context = _context_text(ranking['context'])
        level = LOSS_FIELD_TITLES.get(field, field)
        slide = new_slide(presentation, f"Points Lost per {level} — {ranking['label']}",
                          ' · '.join(part for part in (subtitle, context) if part))
        bars = [area for area in ranking['areas'] if area['name'] != 'Not specified'][:MAX_LOSS_BARS]
        # All environments: each area names the environments where it loses points.
        labels = environment_labels or {}
        bars = [{**area, 'name': f"{area['name']} ({', '.join(labels.get(name, name) for name in area['environments'])})"}
                if area.get('environments') else area for area in bars]
        share = sum(area['points'] for area in bars) / ranking['total'] if ranking['total'] else 0
        plural = {'City': 'cities and routes', 'Cluster': 'clusters', 'Region': 'regions'}.get(field, 'areas')
        by_rows = layer is not None and layer['field'] == 'Area' and field == 'City'
        area_label = (layer or {}).get('area_label') or LOSS_FIELD_TITLES.get((layer or {}).get('field'), 'area')
        map_text = f"The map colours each {area_label} by the points " if layer is not None else 'The map shows where '
        # The operator, large and in its colour, above the map and the ranking.
        heading = slide.shapes.add_textbox(Inches(.45), Inches(1.35), Inches(12.4), Inches(.5))
        heading.name = 'Scoring Points Lost Operator'
        _write(heading.text_frame, [(str(ranking['label']), 24, True, ranking.get('color') or '#17232D')])
        intro = slide.shapes.add_textbox(Inches(.45), Inches(1.85), Inches(12.4), Inches(.5))
        intro.text_frame.word_wrap = True
        _write_runs(intro.text_frame, [
            (map_text, False, None), (ranking['label'], True, ranking.get('color') or None),
            (f" loses in the city or route that loses most among those measured in it ({_points(ranking['total'], 1)} {scoring_label} points lost in total). "
             if by_rows else
             f" loses in it ({_points(ranking['total'], 1)} {scoring_label} points lost in total). " if layer is not None
             else f" loses its {scoring_label} points ({_points(ranking['total'], 1)} in total). ", False, None),
            (f'The {len(bars)} {plural} listed', True, None),
            (f' account for {share * 100:.0f}% of the points lost.', False, None),
        ], size=11, color='#17232D')
        map_top, map_height, map_width = 2.35, 4.5, 6.0
        if layer is not None:
            # As the web map does by default: the points of the city or route that loses most among those
            # measured in each area.
            from src.modules.scoring_points_loss import row_area_links, row_area_values
            values = (row_area_values(ranking['areas']) if by_rows else
                      {area['name']: area['points'] for area in layer['areas'] if area['name'] != 'Not specified'})
            # The cities and routes that lose most are labelled once, on their main area.
            if by_rows:
                links = row_area_links(ranking['areas'])
                places = [(area['name'], links[area['name']][0]) for area in ranking['areas'] if links.get(area['name'])]
            else:
                places = None
            image = render_points_loss_choropleth(boundaries(layer['field']), values, places=places,
                                                  outlines=outlines(layer['field']))
        else:
            image = render_points_loss_map(ranking['areas'], background)
        from PIL import Image
        with Image.open(BytesIO(image)) as opened:
            ratio = opened.width / opened.height
        width = min(map_width, map_height * ratio)
        height = width / ratio
        picture = slide.shapes.add_picture(BytesIO(image), Inches(.45 + (map_width - width) / 2), Inches(map_top),
                                           Inches(width), Inches(height))
        picture.name = 'Scoring Points Lost Map'
        _loss_bars(slide, bars, left=6.55, top=2.4, width=6.35, height=4.55)
        if layer is not None and values:
            positive = [value for value in values.values() if value > 0]
            _loss_scale(slide, .45, 6.95, map_width, min(positive), max(positive))
        else:
            legend = slide.shapes.add_textbox(Inches(.45), Inches(6.9), Inches(6.0), Inches(.3))
            legend.text_frame.margin_left = 0
            _write(legend.text_frame, [('Bubbles are cities and dots the Connecting Roads routes; colour and size grow '
                                        'with the points lost. Grey: test locations.', 8, False, '#6B7A85')])
