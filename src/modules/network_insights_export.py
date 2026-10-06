"""PowerPoint and Word summaries of a Network Insights analysis."""
from __future__ import annotations

from collections.abc import Callable
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx import Presentation
from pptx.util import Inches, Pt

from src.modules.exports import MUTED, ORANGE, _add_textbox
from src.utils.fonts import load_image_font

# Renders a Canvas chart model to PNG bytes at the requested pixel size.
ChartRenderer = Callable[[dict[str, Any], int, int], bytes]

TABLE_ROWS_PER_SLIDE = 12


def _number(value: Any, digits: int = 1, suffix: str = '') -> str:
    if value is None or value == '':
        return '—'
    if isinstance(value, (int, float)):
        return f'{value:,.{digits}f}{suffix}' if isinstance(value, float) else f'{value:,}{suffix}'
    return str(value)


def _signed(value: Any, suffix: str = '') -> str:
    return '—' if value is None else f'{value:+.1f}{suffix}'


def summary_selection_lines(selection: dict[str, Any]) -> list[str]:
    """Readable lines describing the analysed CDRs, grouping and filters."""
    lines = [
        *([f"NR Mode: {selection['nr_mode']}"] if selection.get('nr_mode') else []),
        f"Technology: {selection.get('technology_label') or selection.get('technology') or 'LTE'}",
        f"Grouping: {selection.get('group_label') or 'Operator → Campaign'}",
    ]
    technology = selection.get('technology') or 'lte'
    for radio, label, coverage, interference in (('lte', 'LTE', 'coverage_threshold', 'interference_threshold'),
                                                  ('nr', 'NR', 'nr_coverage_threshold', 'nr_interference_threshold')):
        if technology in {radio, 'lte_nr'} and selection.get(coverage) is not None:
            lines.append(f"{label} thresholds: low coverage below {selection.get(coverage)} dBm RSRP, "
                         f"high interference below {selection.get(interference)} dB SINR")
    for label, key in (('Operators', 'operators'), ('Operator_Vendor', 'operator_vendors'), ('Vendors', 'vendors'),
                       ('Campaigns', 'campaigns'), ('Regions', 'regions'), ('Clusters', 'clusters'), ('Cities', 'cities')):
        values = selection.get(key) or []
        lines.append(f"{label}: {', '.join(values) if values else 'All'}")
    datasets = selection.get('dataset_names') or []
    lines.append(f"CDRs ({len(datasets)}): {', '.join(datasets) if datasets else 'All ready CDRs'}")
    return lines


def technology_sections(analysis: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    """(title suffix, section) per technology; LTE+NR exports LTE and NR separately."""
    sections = analysis.get('sections') or [analysis]
    if len(sections) == 1:
        return [('', sections[0])]
    return [(f" · {section.get('technology_label') or section.get('technology')}", section) for section in sections]


def overview_table(analysis: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    comparison = analysis.get('comparison')
    columns = ['Group', 'Samples', 'RSRP mean', 'RSRP median', 'RSRP P10', 'Low cov. %',
               'SINR mean', 'SINR median', 'SINR P10', 'High interf. %', 'eNodeBs', 'Cells']
    if comparison:
        columns += [f'Δ RSRP med. ({comparison["latest"]} vs {comparison["previous"]})', 'Δ Low cov.', 'Δ SINR med.', 'Δ High interf.']
    rows = []
    for row in analysis.get('rf_rows') or analysis.get('overview') or []:
        values = [str(row.get('operator') or ''), _number(row.get('samples'), 0), _number(row.get('rsrp_mean')), _number(row.get('rsrp_median')),
                  _number(row.get('rsrp_p10')), _number(row.get('low_coverage_share'), suffix='%'), _number(row.get('sinr_mean')),
                  _number(row.get('sinr_median')), _number(row.get('sinr_p10')), _number(row.get('high_interference_share'), suffix='%'),
                  _number(row.get('observed_enodebs'), 0), _number(row.get('observed_cells'), 0)]
        if comparison:
            deltas = row.get('deltas') or {}
            values += [_signed(deltas.get('rsrp_median')), _signed(deltas.get('low_coverage_share'), ' pp'),
                       _signed(deltas.get('sinr_median')), _signed(deltas.get('high_interference_share'), ' pp')]
        rows.append(values)
    return columns, rows


def hotspot_table(rows: list[dict[str, Any]], unit: str, limit: int = 10) -> tuple[list[str], list[list[str]]]:
    columns = ['City', 'Region', 'Samples', f'Mean ({unit})', 'Share below threshold', 'Latitude', 'Longitude']
    return columns, [
        [row.get('city') or '—', row.get('region') or '—', _number(row.get('samples'), 0), _number(row.get('mean')),
         _number(row.get('bad_share'), suffix='%'), f"{row.get('latitude')}", f"{row.get('longitude')}"]
        for row in rows[:limit]
    ]


def spectrum_tables(analysis: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    spectrum = analysis.get('spectrum') or {}
    observed = [
        [row.get('operator') or '', row.get('technology') or '', row.get('band') or '', row.get('band_class') or '',
         _number(row.get('share'), suffix='%'), _number(row.get('samples'), 0), _number(row.get('typical_bandwidth_mhz'), 0, ' MHz')]
        for row in spectrum.get('observed') or []
    ]
    licensed = [
        [row.get('operator') or '', _number(row.get('Low')), _number(row.get('Mid')), _number(row.get('High (TDD)')), _number(row.get('total'))]
        for row in spectrum.get('licensed') or []
    ]
    tables = []
    if observed:
        tables.append(('Observed Spectrum', ['Group', 'Technology', 'Band', 'Class', 'Share', 'Samples', 'Typical bandwidth'], observed))
    if licensed:
        tables.append(('Licensed Spectrum (MHz)', ['Operator', 'Low', 'Mid', 'High (TDD)', 'Total'], licensed))
    return tables


def deployment_tables(deployment: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    if 'groupings' in deployment:
        return [table for grouping in deployment['groupings'] for table in deployment_tables(grouping)]
    tables = []
    for inventory in deployment.get('inventories') or []:
        rows = [[row['group'], _number(row['sites'], 0), _number(row['cells'], 0)] for row in inventory.get('rows') or []]
        totals = inventory.get('totals') or {}
        if totals:
            rows.append(['Total', _number(totals.get('sites'), 0), _number(totals.get('cells'), 0)])
        if rows:
            tables.append((f"Network Deployment · {deployment.get('group_label') or 'All'} · {inventory.get('operator')} ({inventory.get('file_name')})",
                           [deployment.get('group_label') or 'Group', 'Sites', 'Cells'], rows))
    return tables


def cluster_sites_tables(analysis: dict[str, Any], deployment: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    """Export both data sources independently of the panel's selected view."""
    observed = [[row.get('operator') or '', _number(row.get('observed_enodebs'), 0),
                 _number(row.get('observed_cells'), 0), _number(row.get('samples'), 0),
                 _number(row['samples'] / row['observed_enodebs']) if row.get('observed_enodebs') else '—']
                for _suffix, section in technology_sections(analysis) for row in _labelled_rows(analysis, section)]
    inventory = [[row['operator'], _number(row['sites'], 0), _number(row['cells'], 0)]
                 for row in (deployment.get('cluster_inventories') or {}).get('rows') or []]
    tables = []
    if observed:
        tables.append(('Cluster Sites Density · Observed Sites/Cells',
                       ['Group', 'Observed eNodeBs', 'Observed cells', 'Samples', 'Samples per eNodeB'], observed))
    if inventory:
        tables.append(('Cluster Sites Density · Inventory Sites/Cells',
                       ['Operator', 'Inventory sites', 'Inventory cells'], inventory))
    return tables


def _labelled_rows(analysis: dict[str, Any], section: dict[str, Any]) -> list[dict[str, Any]]:
    rows = section.get('rf_rows') or []
    if len(analysis.get('sections') or []) < 2:
        return rows
    return [{**row, 'operator': f"{row.get('operator') or ''} · {section.get('technology_label') or ''}"} for row in rows]


def _render_image(render: ChartRenderer, payload: Any, width: int = 1600, height: int = 900) -> tuple[BytesIO | None, str]:
    if not isinstance(payload, dict):
        return None, 'No chart for this selection.'
    try:
        return BytesIO(render(payload, width, height)), ''
    except Exception as exc:  # A chart that cannot be drawn must not drop the whole summary.
        return None, f'The chart could not be rendered: {exc}'


def _section_visuals(section: dict[str, Any], render: ChartRenderer) -> tuple[list[tuple[str, BytesIO | None, str, Any]], list[tuple[str, list[tuple[str, BytesIO | None, str, Any]]]]]:
    """The RF Quality CDFs (RSRP | SINR) and, for every group, its maps (Coverage | Interference) with the areas table
    shown under each map, as the page lays them out. Panels are (caption, PNG or None, note, table or None)."""
    charts = section.get('charts') or {}
    maps = section.get('maps') or {}
    label = section.get('technology_label') or ''
    cdfs = [('', *_render_image(render, payload), None) for payload in (charts.get('rsrp_cdf'), charts.get('sinr_cdf')) if isinstance(payload, dict)]
    groups = []
    for group in maps.get('groups') or [maps]:
        panels = [
            (caption, *_render_image(render, group.get(key)), area_table(group.get(hotspots) or [], unit))
            for caption, key, hotspots, unit in (
                (f'{label} Coverage (mean RSRP per grid square)'.strip(), 'coverage', 'coverage_hotspots', 'dBm'),
                (f'{label} Interference (mean SINR per grid square)'.strip(), 'interference', 'interference_hotspots', 'dB'),
            ) if isinstance(group.get(key), dict)
        ]
        if panels:
            groups.append((str(group.get('operator') or ''), panels))
    return cdfs, groups


def deployment_table_groups(deployment: dict[str, Any]) -> list[tuple[str, list[tuple[str, list[str], list[list[str]]]]]]:
    """One group per Network Deployment view with each operator's inventory side by side (Vodafone, then Three)."""
    if 'groupings' in deployment:
        return [group for grouping in deployment['groupings'] for group in deployment_table_groups(grouping)]
    panels = []
    for inventory in deployment.get('inventories') or []:
        rows = [[row['group'], _number(row['sites'], 0), _number(row['cells'], 0)] for row in inventory.get('rows') or []]
        totals = inventory.get('totals') or {}
        if totals:
            rows.append(['Total', _number(totals.get('sites'), 0), _number(totals.get('cells'), 0)])
        if rows:
            panels.append((f"{inventory.get('operator')} ({inventory.get('file_name')})", [deployment.get('group_label') or 'Group', 'Sites', 'Cells'], rows))
    return [(f"Network Deployment · {deployment.get('group_label') or 'All'}", panels)] if panels else []


CLASS_BAR_COLUMNS = ('RSRP classes', 'SINR classes')
OVERVIEW_CARD_FIELDS = (
    ('Median RSRP', 'rsrp_median', ' dBm', 'rsrp_median', ' dB'), ('Low coverage', 'low_coverage_share', '%', 'low_coverage_share', ' pp'),
    ('Median SINR', 'sinr_median', ' dB', 'sinr_median', ' dB'), ('High interference', 'high_interference_share', '%', 'high_interference_share', ' pp'),
    ('Observed eNodeBs', 'observed_enodebs', '', '', ''), ('Samples', 'samples', '', '', ''),
)


def overview_cards(analysis: dict[str, Any], section: dict[str, Any]) -> list[dict[str, Any]]:
    """The Overview cards of one technology: title, colour and the six main indicators."""
    colours = analysis.get('colours') or {}
    cards = []
    for row in section.get('overview') or section.get('rf_rows') or []:
        deltas = row.get('deltas') or {}
        values = []
        for label, key, unit, delta_key, delta_unit in OVERVIEW_CARD_FIELDS:
            digits = 0 if key in {'observed_enodebs', 'samples'} else 1
            text = _number(row.get(key), digits, unit)
            if delta_key and deltas.get(delta_key) is not None:
                text += f' ({_signed(deltas.get(delta_key), delta_unit)})'
            values.append((label, text))
        cards.append({'title': str(row.get('operator') or ''), 'colour': colours.get(row.get('operator')) or '#6D46A8', 'values': values})
    return cards


def rf_quality_table(section: dict[str, Any]) -> tuple[list[str], list[list[str]], list[tuple[list[dict[str, Any]], list[dict[str, Any]]]]]:
    """The table under the RF Quality CDFs; the class columns are drawn as coloured bars from the returned classes."""
    label = section.get('technology_label') or ''
    columns = ['Group', 'Samples', f'{label} RSRP samples'.strip(), 'Median RSRP', 'P10 RSRP', 'Low coverage', 'RSRP classes',
               'Median SINR', 'P10 SINR', 'High interference', 'SINR classes']
    rows, bars = [], []
    for row in section.get('rf_rows') or section.get('overview') or []:
        rows.append([str(row.get('operator') or ''), _number(row.get('samples'), 0), _number(row.get('rsrp_samples', row.get('samples')), 0),
                     _number(row.get('rsrp_median')), _number(row.get('rsrp_p10')), _number(row.get('low_coverage_share'), suffix='%'), '',
                     _number(row.get('sinr_median')), _number(row.get('sinr_p10')), _number(row.get('high_interference_share'), suffix='%'), ''])
        bars.append((row.get('rsrp_classes') or [], row.get('sinr_classes') or []))
    return columns, rows, bars


def area_table(rows: list[dict[str, Any]], unit: str, limit: int = 8) -> tuple[list[str], list[list[str]]]:
    """The compact areas table shown under each map: Area, Region, Samples, Mean and share Below the threshold."""
    return ['Area', 'Region', 'Samples', 'Mean', 'Below'], [
        [row.get('city') or '—', row.get('region') or '—', _number(row.get('samples'), 0), _number(row.get('mean'), suffix=f' {unit}'),
         _number(row.get('bad_share'), suffix='%')]
        for row in rows[:limit]
    ]


CLASS_BAR_BLOCKS = 16
CLASS_BAR_BLOCK = '\u2588'


def class_bar_segments(classes: list[dict[str, Any]], blocks: int = CLASS_BAR_BLOCKS) -> list[tuple[int, str]]:
    """Split a bar of ``blocks`` characters between the quality classes in proportion to their shares.

    The bar is written as coloured block characters inside the table cell, so it always stays in its row.
    """
    shares = [(float(item.get('share') or 0), str(item.get('colour') or '#999999')) for item in classes]
    total = sum(share for share, _colour in shares)
    if total <= 0:
        return []
    exact = [share / total * blocks for share, _colour in shares]
    counts = [int(value) for value in exact]
    for index in sorted(range(len(exact)), key=lambda item: exact[item] - counts[item], reverse=True)[:blocks - sum(counts)]:
        counts[index] += 1
    return [(count, colour) for count, (_share, colour) in zip(counts, shares, strict=False) if count]

def _single_line(value: Any) -> str:
    """Keep table values intact while removing explicit line-break characters."""
    return ' '.join(str(value).splitlines())


def _column_text_widths(columns: list[str], rows: list[list[str]]) -> list[float]:
    """Estimate text widths in points at 1 pt, reserving space for font substitution."""
    normal, bold = load_image_font(100), load_image_font(100, bold=True)
    return [max(float(bold.getlength(_single_line(label))),
                max((float(normal.getlength(_single_line(row[index]))) for row in rows), default=0)) / 100 * 1.15
            for index, label in enumerate(columns)]


def _table_geometry(columns: list[str], rows: list[list[str]], width: float, maximum_font: float = 10) -> tuple[list[float], float]:
    measured = _column_text_widths(columns, rows)
    padding = 0.14 * 72
    font = min(maximum_font, (width * 72 - padding * len(columns)) / max(sum(measured), 1))
    font = max(1, font)
    required = [length * font + padding for length in measured]
    scale = width / sum(required)
    return [value * scale for value in required], font


def _ppt_table_parts(columns: list[str], rows: list[list[str]], width: float):
    """Repeat the identity column if a wide comparison needs multiple metric panels."""
    measured = _column_text_widths(columns, rows)
    indices, parts = [0], []
    for index in range(1, len(columns)):
        proposed = [*indices, index]
        if len(indices) > 1 and sum(measured[item] * 8 + 0.14 * 72 for item in proposed) > width * 72:
            parts.append(indices)
            indices = [0]
        indices.append(index)
    parts.append(indices)
    return [([columns[index] for index in part], [[row[index] for index in part] for row in rows]) for part in parts]


def _add_table(slide, columns: list[str], rows: list[list[str]], left: float, top: float, width: float,
               *, row_height: float = 0.32, maximum_font: float = 10):
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(left), Inches(top), Inches(width), Inches(row_height * (len(rows) + 1)))
    table = shape.table
    widths, font_size = _table_geometry(columns, rows, width, maximum_font)
    for index, column_width in enumerate(widths):
        table.columns[index].width = Inches(column_width)
    for row_index, values in enumerate([columns, *rows]):
        for column, value in enumerate(values):
            cell = table.cell(row_index, column)
            cell.text = _single_line(value)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.text_frame.word_wrap = False
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_before = paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.bold = row_index == 0
    for row in table.rows:
        row.height = Inches(row_height)
    return table


def _content_frame(slide) -> tuple[float, float, float, float]:
    layout = slide.slide_layout
    shapes = [shape for shape in layout.placeholders if shape.placeholder_format.type not in {1, 3, 4}]
    shape = max(shapes, key=lambda item: item.width * item.height)
    return tuple(value / 914400 for value in (shape.left, shape.top, shape.width, shape.height))


def _slide(presentation: Presentation, title: str, subtitle: str = ''):
    from src.modules.cdr_reporting import _named_slide_layout, _remove_template_chart_placeholders, _set_slide_header

    layout = _named_slide_layout(presentation, 'Title + 1 rows + 1 columns')
    if layout is None:
        raise ValueError('The Dashboard template requires the Title + 1 rows + 1 columns layout.')
    slide = presentation.slides.add_slide(layout)
    _set_slide_header(slide, title, subtitle)
    _remove_template_chart_placeholders(slide)
    for placeholder in slide.placeholders:
        if placeholder.has_text_frame:
            placeholder.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return slide


def _table_slides(presentation: Presentation, title: str, columns: list[str], rows: list[list[str]], subtitle: str = '') -> None:
    from src.modules.cdr_reporting import _named_slide_layout

    layout = _named_slide_layout(presentation, 'Title + 1 rows + 1 columns')
    content = max((shape for shape in layout.placeholders if shape.placeholder_format.type not in {1, 3, 4}),
                  key=lambda shape: shape.width * shape.height)
    parts = _ppt_table_parts(columns, rows, content.width / 914400)
    for part_index, (part_columns, part_rows) in enumerate(parts, start=1):
        pages = [part_rows[index:index + TABLE_ROWS_PER_SLIDE] for index in range(0, len(part_rows), TABLE_ROWS_PER_SLIDE)] or [[]]
        for page_number, page in enumerate(pages, start=1):
            page_title = title + (f' · Part {part_index}/{len(parts)}' if len(parts) > 1 else '')
            if len(pages) > 1:
                page_title += f' · {page_number}/{len(pages)}'
            slide = _slide(presentation, page_title, subtitle)
            left, top, width, _height = _content_frame(slide)
            if page:
                _add_table(slide, part_columns, page, left, top, width)
            else:
                _add_textbox(slide, left, top, width, 0.3, 'No data for this selection.', size=12, color=MUTED)


PANEL_GAP = 0.3


def _panel_frames(left: float, width: float, count: int) -> list[tuple[float, float]]:
    panel_width = (width - PANEL_GAP * (count - 1)) / count
    return [(left + index * (panel_width + PANEL_GAP), panel_width) for index in range(count)]


def _paired_table_slides(presentation: Presentation, title: str, panels: list[tuple[str, list[str], list[list[str]]]], subtitle: str = '') -> None:
    """Tables side by side on the same slides, as the page shows them; long tables continue on more slides."""
    if len(panels) == 1 and not panels[0][0]:
        _table_slides(presentation, title, panels[0][1], panels[0][2], subtitle)
        return
    page_count = max(1, *((len(rows) + TABLE_ROWS_PER_SLIDE - 1) // TABLE_ROWS_PER_SLIDE for _caption, _columns, rows in panels))
    for page_number in range(page_count):
        slide = _slide(presentation, title + (f' · {page_number + 1}/{page_count}' if page_count > 1 else ''), subtitle)
        left, top, width, _height = _content_frame(slide)
        for (caption, columns, rows), (panel_left, panel_width) in zip(panels, _panel_frames(left, width, len(panels)), strict=False):
            page = rows[page_number * TABLE_ROWS_PER_SLIDE:(page_number + 1) * TABLE_ROWS_PER_SLIDE]
            _add_textbox(slide, panel_left, top, panel_width, 0.3, caption, size=12, bold=True)
            if page:
                _add_table(slide, columns, page, panel_left, top + 0.4, panel_width)
            elif page_number == 0:
                _add_textbox(slide, panel_left, top + 0.4, panel_width, 0.3, 'No data for this selection.', size=11, color=MUTED)


def _overview_slides(presentation: Presentation, title: str, cards: list[dict[str, Any]], per_row: int = 4, rows_per_slide: int = 2) -> None:
    """The Overview cards of the page: one card per group with its colour and six indicators."""
    per_slide = per_row * rows_per_slide
    pages = [cards[index:index + per_slide] for index in range(0, len(cards), per_slide)] or [[]]
    for page_number, page in enumerate(pages, start=1):
        slide = _slide(presentation, title + (f' · {page_number}/{len(pages)}' if len(pages) > 1 else ''))
        left, top, width, height = _content_frame(slide)
        card_width = (width - PANEL_GAP * (per_row - 1)) / per_row
        card_height = (height - PANEL_GAP * (rows_per_slide - 1)) / rows_per_slide
        for index, card in enumerate(page):
            card_left = left + (index % per_row) * (card_width + PANEL_GAP)
            card_top = top + (index // per_row) * (card_height + PANEL_GAP)
            colour = RGBColor.from_string(card['colour'].lstrip('#').upper())
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(card_left), Inches(card_top), Inches(card_width), Inches(card_height))
            box.adjustments[0] = 0.06
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            box.line.color.rgb = RGBColor(0xD9, 0xDE, 0xEA)
            strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(card_left + 0.08), Inches(card_top), Inches(card_width - 0.16), Inches(0.06))
            strip.fill.solid()
            strip.fill.fore_color.rgb = colour
            strip.line.fill.background()
            heading = slide.shapes.add_textbox(Inches(card_left + 0.15), Inches(card_top + 0.12), Inches(card_width - 0.3), Inches(0.35))
            run = heading.text_frame.paragraphs[0].add_run()
            run.text = card['title']
            run.font.size, run.font.bold, run.font.color.rgb = Pt(13), True, colour
            cell_width = (card_width - 0.3) / 2
            cell_height = (card_height - 0.55) / 3
            for value_index, (label, value) in enumerate(card['values']):
                box_left = card_left + 0.15 + (value_index % 2) * cell_width
                box_top = card_top + 0.5 + (value_index // 2) * cell_height
                frame = slide.shapes.add_textbox(Inches(box_left), Inches(box_top), Inches(cell_width), Inches(cell_height)).text_frame
                frame.word_wrap = True
                caption = frame.paragraphs[0].add_run()
                caption.text = label.upper()
                caption.font.size, caption.font.bold, caption.font.color.rgb = Pt(7), True, RGBColor(0x5B, 0x63, 0x77)
                figure = frame.add_paragraph().add_run()
                figure.text = value
                figure.font.size, figure.font.bold, figure.font.color.rgb = Pt(10), True, RGBColor(0x1F, 0x24, 0x33)


def _add_picture(slide, image: BytesIO, left: float, top: float, width: float, height: float) -> float:
    """Fit a PNG in a box, centred horizontally at its top; return the picture height."""
    with Image.open(image) as raster:
        ratio = raster.width / raster.height
    image.seek(0)
    picture_width = min(width, height * ratio)
    picture_height = picture_width / ratio
    slide.shapes.add_picture(image, Inches(left + (width - picture_width) / 2), Inches(top), width=Inches(picture_width), height=Inches(picture_height))
    return picture_height


def _add_bar_table(slide, columns: list[str], rows: list[list[str]], bars: list[tuple[list[dict[str, Any]], list[dict[str, Any]]]],
                   left: float, top: float, width: float, row_height: float) -> None:
    filler = CLASS_BAR_BLOCK * CLASS_BAR_BLOCKS
    sized = [[filler if columns[index] in CLASS_BAR_COLUMNS else value for index, value in enumerate(row)] for row in rows]
    table = _add_table(slide, columns, sized, left, top, width, row_height=row_height, maximum_font=9)
    for bar_index, column_name in enumerate(CLASS_BAR_COLUMNS):
        column = columns.index(column_name)
        for row_index, classes in enumerate(bars, start=1):
            paragraph = table.cell(row_index, column).text_frame.paragraphs[0]
            size = paragraph.runs[0].font.size if paragraph.runs else Pt(8)
            for run in list(paragraph.runs):
                run._r.getparent().remove(run._r)
            for count, colour in class_bar_segments(classes[bar_index]):
                run = paragraph.add_run()
                run.text = CLASS_BAR_BLOCK * count
                run.font.size = size
                run.font.color.rgb = RGBColor.from_string(colour.lstrip('#').upper())

def _visual_slides(presentation: Presentation, title: str, panels: list[tuple[str, BytesIO | None, str, Any]],
                   below: tuple[list[str], list[list[str]], list[Any]] | None = None) -> None:
    """Charts side by side as on the page, each with an optional table under it (the map areas), or one table
    under all of them (RF Quality with its class bars). Rows that do not fit continue on further slides."""
    row_height = 0.24
    slide = _slide(presentation, title)
    left, top, width, height = _content_frame(slide)
    has_tables = below is not None or any(table for *_rest, table in panels)
    image_height = height * (0.52 if has_tables else 1) - 0.35
    remaining_rows = []
    for (caption, image, note, table), (panel_left, panel_width) in zip(panels, _panel_frames(left, width, len(panels)), strict=False):
        cursor = top
        if caption:
            _add_textbox(slide, panel_left, cursor, panel_width, 0.3, caption, size=12, bold=True)
            cursor += 0.35
        if image is None:
            _add_textbox(slide, panel_left, cursor, panel_width, 0.6, note, size=12, color=ORANGE)
            cursor += image_height
        else:
            cursor += _add_picture(slide, image, panel_left, cursor, panel_width, image_height)
        if table:
            columns, rows = table
            if rows:
                _add_table(slide, columns, rows, panel_left, cursor + 0.1, panel_width, row_height=row_height, maximum_font=9)
            else:
                _add_textbox(slide, panel_left, cursor + 0.1, panel_width, 0.3, 'No area falls below the threshold.', size=11, color=MUTED)
    if below is None:
        return
    columns, rows, bars = below
    row_height = 0.3
    table_top = top + 0.35 + image_height + 0.1
    fit = max(1, int((top + height - table_top) / row_height) - 1)
    first, rest = (rows[:fit], bars[:fit]), (rows[fit:], bars[fit:])
    if first[0]:
        _add_bar_table(slide, columns, first[0], first[1], left, table_top, width, row_height)
    per_slide = max(1, int(height / row_height) - 1)
    pages = [(rest[0][index:index + per_slide], rest[1][index:index + per_slide]) for index in range(0, len(rest[0]), per_slide)]
    for page_number, (page_rows, page_bars) in enumerate(pages, start=2):
        continuation = _slide(presentation, f'{title} · {page_number}/{len(pages) + 1}')
        left, top, width, _height = _content_frame(continuation)
        _add_bar_table(continuation, columns, page_rows, page_bars, left, top, width, row_height)


def export_network_insights_powerpoint(destination: Path, analysis: dict[str, Any], deployment: dict[str, Any],
                                       selection: dict[str, Any], render: ChartRenderer, template: Path | None = None) -> Path:
    from src.config import settings
    from src.modules.cdr_reporting import _named_slide_layout, _remove_all_slides, _set_structural_slide_text

    presentation = Presentation(template or settings.ppt_templates_dir / 'Template_CDR_analysis.pptx')
    _remove_all_slides(presentation)
    cover = presentation.slides.add_slide(_named_slide_layout(presentation, 'Title Page'))
    _set_structural_slide_text(cover, 'Summary Network Insights', 'Dashboard Analytic')
    selection_slide = _slide(presentation, 'Analysis Selection')
    left, top, width, height = _content_frame(selection_slide)
    shape = selection_slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.word_wrap = True
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, line in enumerate(summary_selection_lines(selection)):
        paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(15)
    sections = technology_sections(analysis)
    visuals = [(suffix, section, *_section_visuals(section, render)) for suffix, section in sections]
    # Sections follow the order of the page: Overview, RF Quality, maps, Network Deployment, Cluster Sites Density, Spectrum.
    for suffix, section in sections:
        _overview_slides(presentation, f'Overview{suffix}', overview_cards(analysis, section))
        columns, rows = overview_table(section)
        _table_slides(presentation, f'RF Quality Overview{suffix}', columns, rows, selection.get('group_label') or '')
    for suffix, section, cdfs, _groups in visuals:
        if cdfs:
            _visual_slides(presentation, f'RF Quality{suffix}', cdfs, rf_quality_table(section))
    for suffix, _section, _cdfs, groups in visuals:
        for group, panels in groups:
            _visual_slides(presentation, f'Coverage & Interference Maps{suffix}' + (f' · {group}' if group else ''), panels)
    for title, panels in deployment_table_groups(deployment):
        _paired_table_slides(presentation, title, panels)
    for title, columns, rows in [*cluster_sites_tables(analysis, deployment), *spectrum_tables(analysis)]:
        _table_slides(presentation, title, columns, rows)
    warnings = analysis.get('warnings') or []
    if warnings:
        slide = _slide(presentation, 'Analysis Notes')
        left, top, width, height = _content_frame(slide)
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        shape.text_frame.text = '\n'.join(str(warning) for warning in warnings)
        shape.text_frame.word_wrap = True
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    closing = _named_slide_layout(presentation, 'Black logo end slide')
    if closing is not None:
        _set_structural_slide_text(presentation.slides.add_slide(closing), '', '')
    presentation.save(destination)
    return destination


def _docx_table(document: Document, columns: list[str], rows: list[list[str]], width: float | None = None):
    table = document.add_table(rows=1, cols=len(columns))
    table.style = 'Light Grid Accent 1'
    table.autofit = False
    section = document.sections[-1]
    if width is None:
        width = (section.page_width - section.left_margin - section.right_margin) / 914400
    widths, font = _table_geometry(columns, rows, width, maximum_font=9)
    for index, value in enumerate(widths):
        table.columns[index].width = DocxInches(value)
    for row_index, values in enumerate([columns, *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = DocxInches(widths[index])
            cell.text = _single_line(value)
            properties = cell._tc.get_or_add_tcPr()
            properties.append(OxmlElement('w:noWrap'))
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                item = OxmlElement(f'w:{side}')
                item.set(qn('w:w'), '72' if side in {'left', 'right'} else '36')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            properties.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = DocxPt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.name = 'Arial'
                    run.font.size = DocxPt(font)
                    run.bold = row_index == 0
    return table


DOCX_COLUMN_GAP = 0.3


def _docx_columns(document: Document, count: int) -> float:
    """Continue on a new continuous section with ``count`` columns; return each column's width in inches."""
    section = document.add_section(WD_SECTION.CONTINUOUS)
    columns = section._sectPr.find(qn('w:cols'))
    if columns is None:
        columns = OxmlElement('w:cols')
        section._sectPr.append(columns)
    columns.set(qn('w:num'), str(count))
    columns.set(qn('w:space'), str(int(DOCX_COLUMN_GAP * 1440)))
    usable = (section.page_width - section.left_margin - section.right_margin) / 914400
    return (usable - DOCX_COLUMN_GAP * (count - 1)) / count


def _docx_section_heading(document: Document, title: str) -> None:
    """Every section of the summary starts on a new page, as each one has its own slide in PowerPoint."""
    heading = document.add_heading(title, level=1)
    heading.paragraph_format.page_break_before = True


def _docx_side_by_side(document: Document, title: str, panels: list[tuple[str, Any]], add_panel: Callable[[Any, float], None],
                       *, new_page: bool = True) -> None:
    """A heading on a new page, then the panels side by side in two columns (Coverage | Interference, Vodafone | Three)."""
    _docx_section_heading(document, title)
    width = _docx_columns(document, len(panels)) if len(panels) > 1 else None
    for index, (caption, content) in enumerate(panels):
        paragraph = document.add_paragraph() if caption or index else None
        if index and width is not None:
            # The column break opens the caption, so both panels start at the same height.
            paragraph.add_run().add_break(WD_BREAK.COLUMN)
        if caption:
            paragraph.add_run(caption).bold = True
        add_panel(content, width)
    if width is not None:
        _docx_columns(document, 1)


def _docx_overview_cards(document: Document, cards: list[dict[str, Any]], per_row: int = 4) -> None:
    """The Overview cards as a grid: each cell holds a group's coloured title and its six indicators."""
    if not cards:
        document.add_paragraph('No data for this selection.')
        return
    table = document.add_table(rows=(len(cards) + per_row - 1) // per_row, cols=per_row)
    table.style = 'Table Grid'
    for index, card in enumerate(cards):
        cell = table.cell(index // per_row, index % per_row)
        title = cell.paragraphs[0].add_run(card['title'])
        title.bold = True
        title.font.size = DocxPt(11)
        title.font.color.rgb = DocxRGBColor.from_string(card['colour'].lstrip('#').upper())
        for label, value in card['values']:
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_after = DocxPt(0)
            caption = paragraph.add_run(f'{label}: ')
            caption.font.size = DocxPt(8)
            caption.font.color.rgb = DocxRGBColor(0x5B, 0x63, 0x77)
            figure = paragraph.add_run(value)
            figure.bold = True
            figure.font.size = DocxPt(9)


def _docx_bar_table(document: Document, columns: list[str], rows: list[list[str]], bars: list[tuple[list[dict[str, Any]], list[dict[str, Any]]]]) -> None:
    """The RF Quality table with the RSRP and SINR classes drawn as coloured bars in their cells."""
    filler = CLASS_BAR_BLOCK * CLASS_BAR_BLOCKS
    sized = [[filler if columns[index] in CLASS_BAR_COLUMNS else value for index, value in enumerate(row)] for row in rows]
    table = _docx_table(document, columns, sized)
    for bar_index, column_name in enumerate(CLASS_BAR_COLUMNS):
        column = columns.index(column_name)
        for row_index, classes in enumerate(bars, start=1):
            paragraph = table.cell(row_index, column).paragraphs[0]
            size = paragraph.runs[0].font.size if paragraph.runs else DocxPt(8)
            for run in list(paragraph.runs):
                run._r.getparent().remove(run._r)
            for count, colour in class_bar_segments(classes[bar_index]):
                run = paragraph.add_run(CLASS_BAR_BLOCK * count)
                run.font.size = size
                run.font.name = 'Arial'
                run.font.color.rgb = DocxRGBColor.from_string(colour.lstrip('#').upper())

def export_network_insights_word(destination: Path, analysis: dict[str, Any], deployment: dict[str, Any],
                                 selection: dict[str, Any], render: ChartRenderer) -> Path:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = DocxInches(11.69), DocxInches(8.27)
    section.left_margin = section.right_margin = DocxInches(0.6)
    section.top_margin = section.bottom_margin = DocxInches(0.6)
    tables = [*(overview_table(section) for _suffix, section in technology_sections(analysis)),
              *(hotspot_table((section.get('maps') or {}).get(key) or [], unit) for _suffix, section in technology_sections(analysis)
                for key, unit in [('coverage_hotspots', 'dBm'), ('interference_hotspots', 'dB')]),
              *((columns, rows) for _title, columns, rows in [*spectrum_tables(analysis), *deployment_tables(deployment), *cluster_sites_tables(analysis, deployment)])]
    needed = max((sum(length * 9 + 0.14 * 72 for length in _column_text_widths(columns, rows)) / 72
                  for columns, rows in tables), default=0)
    if needed > 10.49:
        section.page_width = DocxInches(max(16.54, min(22, needed + 1.2)))
        section.page_height = DocxInches(11.69)
    document.add_heading('Dashboard Analytic · Summary Network Insights', level=0)
    for line in summary_selection_lines(selection):
        document.add_paragraph(line, style='List Bullet')
    for warning in analysis.get('warnings') or []:
        document.add_paragraph(f'Warning: {warning}')
    def add_image(content: tuple[BytesIO | None, str, Any], width: float | None) -> None:
        image, note, table = content
        if image is None:
            document.add_paragraph(note)
        else:
            document.add_picture(image, width=DocxInches(width or 6.3))
        if table is not None:
            add_table(table, width)

    def add_table(content: tuple[list[str], list[list[str]]], width: float | None) -> None:
        columns, rows = content
        if rows:
            _docx_table(document, columns, rows, width)
        else:
            document.add_paragraph('No data for this selection.')

    sections = technology_sections(analysis)
    visuals = [(suffix, section, *_section_visuals(section, render)) for suffix, section in sections]
    # Sections follow the order of the page: Overview, RF Quality, maps, Network Deployment, Cluster Sites Density, Spectrum.
    for suffix, section in sections:
        _docx_section_heading(document, f'Overview{suffix}')
        _docx_overview_cards(document, overview_cards(analysis, section))
        _docx_section_heading(document, f'RF Quality Overview{suffix}')
        columns, rows = overview_table(section)
        _docx_table(document, columns, rows)
    for suffix, section, cdfs, _groups in visuals:
        if cdfs:
            _docx_side_by_side(document, f'RF Quality{suffix}', [(caption, (image, note, None)) for caption, image, note, _table in cdfs], add_image, new_page=True)
            columns, rows, bars = rf_quality_table(section)
            if rows:
                _docx_bar_table(document, columns, rows, bars)
    for suffix, _section, _cdfs, groups in visuals:
        for group, panels in groups:
            _docx_side_by_side(document, f'Coverage & Interference Maps{suffix}' + (f' · {group}' if group else ''),
                               [(caption, (image, note, table)) for caption, image, note, table in panels], add_image, new_page=True)
    for title, panels in deployment_table_groups(deployment):
        _docx_side_by_side(document, title, [(caption, (columns, rows)) for caption, columns, rows in panels], add_table)
    for title, columns, rows in [*cluster_sites_tables(analysis, deployment), *spectrum_tables(analysis)]:
        _docx_section_heading(document, title)
        _docx_table(document, columns, rows)
    document.save(destination)
    return destination

