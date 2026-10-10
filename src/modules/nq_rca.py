"""Root Cause Analysis of the NQ Analysis Center: the RCA script results and the recommendation.

Three sources describe the root cause of a Non-Qualified Call:

- the **RCA script** (a Python script of the team) writes, for each JOIN_ID, the categories and
  subcategories of both sides and a suggested category and cause; its CSV or Excel files are
  imported and linked to the calls by JOIN_ID;
- the **NetCheck RCA** of the CDR (Failure Classification, Category and Subcategory);
- the **analysts**, who choose the Selected Root Category and Root Cause.

Both RCAs keep their own category → cause hierarchy. The recommendation maps each one onto the
root cause catalog of the workspace (names first, then the keywords of the domains, categories
and causes), adds the keyword reading of the comments and what the analysts chose before for the
same RCA values, and proposes the category and cause most sources agree on.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from threading import Lock
from typing import Any

from src.modules.nq_catalog import RCA_REFS, keyword_found, normalized_text

RCA_FILE_FORMAT_HELP = (
    'A CSV (comma, semicolon or tab separated) or Excel file with a JOIN_ID column and the RCA columns '
    'Auto_RCA_Category_A, Auto_RCA_Subcategory_A, Auto_RCA_Category_B, Auto_RCA_Subcategory_B, '
    'RCA_Suggested_Category and RCA_Suggested_Cause; any other column is kept and shown with the call.'
)
RCA_TEMPLATE_COLUMNS = ('JOIN_ID', 'Auto_RCA_Category_A', 'Auto_RCA_Subcategory_A', 'Auto_RCA_Category_B',
                        'Auto_RCA_Subcategory_B', 'RCA_Suggested_Category', 'RCA_Suggested_Cause')
MAX_RCA_ROWS = 500_000
MAX_EXTRA_COLUMNS = 40
# Column identities (letters and digits only) of the known RCA columns.
_RCA_COLUMNS = {
    'joinid': 'join_id',
    'autorcacategorya': 'auto_rca_category_a', 'autorcasubcategorya': 'auto_rca_subcategory_a',
    'autorcacategoryb': 'auto_rca_category_b', 'autorcasubcategoryb': 'auto_rca_subcategory_b',
    'rcasuggestedcategory': 'rca_category', 'rcacategory': 'rca_category', 'suggestedcategory': 'rca_category',
    'rcasuggestedcause': 'rca_cause', 'rcacause': 'rca_cause', 'suggestedcause': 'rca_cause',
}
SOURCE_LABELS = {'script': 'RCA script', 'netcheck': 'NetCheck RCA', 'comments': 'Comments', 'learned': 'Previous decisions'}
SOURCE_WEIGHTS = {'script': 3, 'learned': 3, 'netcheck': 2, 'comments': 1}


def _identity(value: Any) -> str:
    return re.sub(r'[^a-z0-9]', '', str(value or '').casefold())


def _cell(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r'\s+', ' ', str(value)).strip()
    return '' if text.casefold() in {'nan', 'none', 'null', '<na>'} else text[:300]


def _table(content: bytes, file_name: str) -> list[list[Any]]:
    if file_name.casefold().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    text = content.decode('utf-8-sig', errors='replace')
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def read_rca_file(content: bytes, file_name: str) -> dict[str, Any]:
    """The RCA results of a file: one entry per JOIN_ID (the last row wins), the recognized columns and the rest."""
    rows = _table(content, file_name)
    if not rows:
        raise ValueError('The RCA file is empty.')
    headers = [_cell(value) for value in rows[0]]
    columns: dict[int, str] = {}
    extra: dict[int, str] = {}
    for index, header in enumerate(headers):
        known = _RCA_COLUMNS.get(_identity(header))
        if known and known not in columns.values():
            columns[index] = known
        elif header and len(extra) < MAX_EXTRA_COLUMNS:
            extra[index] = header
    if 'join_id' not in columns.values():
        raise ValueError('The RCA file needs a JOIN_ID column.')
    if not any(value in RCA_REFS for value in columns.values()):
        raise ValueError('The RCA file has none of the RCA columns (Auto_RCA_Category_A… RCA_Suggested_Cause).')
    join_index = next(index for index, value in columns.items() if value == 'join_id')
    results: dict[str, dict[str, Any]] = {}
    skipped = duplicates = 0
    for row in rows[1:MAX_RCA_ROWS + 1]:
        join_id = _cell(row[join_index] if join_index < len(row) else '')
        if not join_id:
            if any(_cell(value) for value in row):
                skipped += 1
            continue
        values = {key: _cell(row[index]) for index, key in columns.items() if key != 'join_id' and index < len(row)}
        values.update({f'extra:{name}': _cell(row[index]) for index, name in extra.items() if index < len(row) and _cell(row[index])})
        values = {key: value for key, value in values.items() if value}
        key = join_id.casefold()
        if key in results:
            duplicates += 1
        results[key] = {'join_id': join_id, 'values': values}
    return {'results': list(results.values()), 'columns': [RCA_REFS[value] for value in columns.values() if value in RCA_REFS],
            'extra': list(extra.values()), 'skipped': skipped, 'duplicates': duplicates}


def rca_template() -> bytes:
    """An empty RCA results file with its columns and one example row."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(RCA_TEMPLATE_COLUMNS)
    writer.writerow(['0x0123ABCD', 'RF Problems', 'Coverage problem', 'RF Problems', 'Coverage problem',
                     'Poor Coverage LTE', 'NLOS due to clutter or terrain profile'])
    return output.getvalue().encode('utf-8')


# -- mapping onto the root cause catalog --------------------------------------------------------
def _by_name(items: list[dict[str, Any]], value: str) -> dict[str, Any] | None:
    key = normalized_text(value)
    return next((item for item in items if normalized_text(item['name']) == key), None) if key else None


def _first_match(items: list[dict[str, Any]], text: str, match: str, domain: str = '') -> dict[str, Any] | None:
    for item in items:
        if domain and item.get('domain') and item['domain'] != domain:
            continue
        if keyword_found(item.get('keywords') or [], text, match):
            return item
    return None


# The results of the catalog and the rule of one request, by their text: many calls share the same NetCheck and
# script values, so each combination is matched once. The catalog and rule objects themselves are the key.
_MEMO_SIZE = 8
_memos: list[tuple[Any, Any, dict[Any, Any]]] = []
_memos_guard = Lock()


def _memo(root: Any, rule: Any) -> dict[Any, Any]:
    with _memos_guard:
        for entry in _memos:
            if entry[0] is root and entry[1] is rule:
                return entry[2]
        memo: dict[Any, Any] = {}
        _memos.insert(0, (root, rule, memo))
        del _memos[_MEMO_SIZE:]
        return memo


def map_to_catalog(category: str, cause: str, root: dict[str, list[dict[str, Any]]], match: str = 'words',
                   extra_text: str = '') -> dict[str, str]:
    """The catalog category and cause a pair of RCA values stands for: same names first, then keywords."""
    memo = _memo(root, None)
    key = ('map', category, cause, match, extra_text)
    if key not in memo:
        memo[key] = _map_to_catalog(category, cause, root, match, extra_text)
    return dict(memo[key])


def _map_to_catalog(category: str, cause: str, root: dict[str, list[dict[str, Any]]], match: str, extra_text: str) -> dict[str, str]:
    found_category = _by_name(root['categories'], category) or _first_match(
        root['categories'], ' | '.join(normalized_text(value) for value in (category, cause, extra_text) if value), match)
    found_cause = _by_name(root['causes'], cause) or _first_match(root['causes'], normalized_text(cause), match) \
        or _first_match(root['causes'], normalized_text(category), match)
    domain = (found_category or {}).get('domain') or (found_cause or {}).get('domain') or ''
    return {'domain': domain, 'category': (found_category or {}).get('name', ''), 'cause': (found_cause or {}).get('name', '')}


def suggest_from_text(call: Any, root: dict[str, list[dict[str, Any]]], rule: dict[str, Any],
                      comments: str = '') -> dict[str, str] | None:
    """The domain, category and cause the CDR fields (or, else, the comments) suggest with the workspace rule.

    1. A domain keyword in the rule's domain fields (Failure Classification) chooses the domain.
    2. Category keywords in those fields, the cause fields and the technology choose the category,
       and cause keywords in the cause fields (Failure Category, Subcategory, Comment) the cause:
       in the domain found, when the rule keeps them in it.
    3. Without a category and cause, the comments are read the same way.
    """
    def text(fields: list[str]) -> str:
        values = []
        for field in fields:
            try:
                values.append(normalized_text(call[field]))
            except (KeyError, IndexError, TypeError):
                continue
        return ' | '.join(value for value in values if value)

    domain_text = text(rule.get('domain_fields') or [])
    cause_text = text(rule.get('cause_fields') or [])
    technology = text(['technology', 'failure_technology'])
    memo = _memo(root, rule)
    key = ('suggest', domain_text, cause_text, technology, comments)
    if key not in memo:
        memo[key] = _suggest(domain_text, cause_text, technology, comments, root, rule)
    return dict(memo[key]) if memo[key] else None


def _suggest(domain_text: str, cause_text: str, technology: str, comments: str, root: dict[str, list[dict[str, Any]]],
             rule: dict[str, Any]) -> dict[str, str] | None:
    match = rule.get('match', 'words')
    domain = next((item for item in root['domains'] if keyword_found(item.get('keywords') or [], domain_text, match)), None)
    scope = domain['name'] if domain and rule.get('causes_in_domain', True) else ''
    category_text = ' | '.join(value for value in (domain_text, cause_text, technology) if value)
    category = _first_match(root['categories'], category_text, match, scope)
    cause = _first_match(root['causes'], cause_text, match, scope)
    if category or cause:
        found_domain = (domain or {}).get('name') or (category or {}).get('domain') or (cause or {}).get('domain') or ''
        return {'domain': found_domain, 'category': (category or {}).get('name', ''), 'cause': (cause or {}).get('name', ''),
                'source': 'netcheck'}
    notes = normalized_text(comments) if rule.get('comments', True) else ''
    if notes:
        category = _first_match(root['categories'], notes, match, scope)
        cause = _first_match(root['causes'], notes, match, scope)
        if category or cause:
            found_domain = (domain or {}).get('name') or (category or {}).get('domain') or (cause or {}).get('domain') or ''
            return {'domain': found_domain, 'category': (category or {}).get('name', ''), 'cause': (cause or {}).get('name', ''),
                    'source': 'comments'}
    if domain is not None:
        return {'domain': domain['name'], 'category': '', 'cause': '', 'source': 'netcheck'}
    if notes and rule.get('comment_domain', True):
        commented = next((item for item in root['domains'] if keyword_found(item.get('keywords') or [], notes, match)), None)
        if commented is not None:
            return {'domain': commented['name'], 'category': '', 'cause': '', 'source': 'comments'}
    return None


def learned_key(source: str, category: str, cause: str) -> tuple[str, str, str]:
    return source, normalized_text(category), normalized_text(cause)


def recommend(call: Any, rca: dict[str, str], root: dict[str, list[dict[str, Any]]], rule: dict[str, Any],
              comments: str = '', learned: dict[tuple[str, str, str], dict[tuple[str, str], set[str]]] | None = None) -> dict[str, Any]:
    """The root cause each source proposes for a call and the recommended Selected Root Category and Cause.

    ``rca`` holds the imported script values of the call; ``learned`` holds, for the RCA values of
    the calls already decided, the calls that chose each category and cause. Each other call with
    the same script or NetCheck values counts once.
    """
    match = rule.get('match', 'words')
    learned = learned or {}
    sources: list[dict[str, Any]] = []
    script_category = rca.get('rca_category') or rca.get('auto_rca_category_a') or ''
    script_cause = rca.get('rca_cause') or rca.get('auto_rca_subcategory_a') or ''
    if script_category or script_cause:
        mapped = map_to_catalog(script_category, script_cause, root, match)
        sources.append({'source': 'script', 'raw_category': script_category, 'raw_cause': script_cause, **mapped})
    netcheck_category = _value(call, 'failure_category') or _value(call, 'failure_classification')
    netcheck_cause = _value(call, 'failure_subcategory')
    suggestion = suggest_from_text(call, root, rule, comments)
    if netcheck_category or netcheck_cause or (suggestion and suggestion['source'] == 'netcheck'):
        mapped = suggestion if suggestion and suggestion['source'] == 'netcheck' else {'domain': '', 'category': '', 'cause': ''}
        sources.append({'source': 'netcheck', 'raw_category': netcheck_category, 'raw_cause': netcheck_cause,
                        'domain': mapped['domain'], 'category': mapped['category'], 'cause': mapped['cause']})
    if suggestion and suggestion['source'] == 'comments':
        sources.append({'source': 'comments', 'raw_category': '', 'raw_cause': '', 'domain': suggestion['domain'],
                        'category': suggestion['category'], 'cause': suggestion['cause']})
    voters: dict[tuple[str, str], set[str]] = {}
    own = _value(call, 'call_key')
    for source, category, cause in (('script', script_category, script_cause), ('netcheck', netcheck_category, netcheck_cause)):
        if category or cause:
            for choice, keys in learned.get(learned_key(source, category, cause), {}).items():
                voters.setdefault(choice, set()).update(keys - {own})
    voters = {choice: keys for choice, keys in voters.items() if keys}
    if voters:
        (category, cause), keys = max(sorted(voters.items()), key=lambda item: len(item[1]))
        count = len(keys)
        if count >= 2:
            item = _by_name(root['categories'], category)
            reason = _by_name(root['causes'], cause)
            sources.append({'source': 'learned', 'raw_category': '', 'raw_cause': '', 'count': count,
                            'domain': (item or reason or {}).get('domain', ''), 'category': (item or {}).get('name', ''),
                            'cause': (reason or {}).get('name', '')})
    for source in sources:
        source['label'] = SOURCE_LABELS[source['source']]
    return {'sources': sources, 'recommendation': _consensus(sources, root)}


def _value(call: Any, key: str) -> str:
    try:
        return str(call[key] or '').strip()
    except (KeyError, IndexError, TypeError):
        return ''


def _consensus(sources: list[dict[str, Any]], root: dict[str, list[dict[str, Any]]]) -> dict[str, Any] | None:
    """The category most sources (weighted) propose and, among the sources that agree on it, the best cause."""
    category_votes: Counter = Counter()
    for source in sources:
        if source['category']:
            category_votes[source['category']] += SOURCE_WEIGHTS[source['source']]
    cause_votes: Counter = Counter()
    best_category = category_votes.most_common(1)[0][0] if category_votes else ''
    for source in sources:
        if source['cause'] and (not best_category or source['category'] in {best_category, ''}):
            cause_votes[source['cause']] += SOURCE_WEIGHTS[source['source']]
    best_cause = cause_votes.most_common(1)[0][0] if cause_votes else ''
    if not best_category and not best_cause:
        return None
    agreeing = [source['source'] for source in sources
                if (best_category and source['category'] == best_category) or (not best_category and source['cause'] == best_cause)]
    if len(agreeing) >= 2:
        confidence = 'high'
    elif agreeing and agreeing[0] in {'script', 'learned'}:
        confidence = 'medium'
    else:
        confidence = 'low'
    category = _by_name(root['categories'], best_category)
    cause = _by_name(root['causes'], best_cause)
    return {'category': best_category, 'cause': best_cause,
            'domain': (category or {}).get('domain') or (cause or {}).get('domain') or '',
            'sources': agreeing, 'confidence': confidence}
